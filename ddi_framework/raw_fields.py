# -*- coding: utf-8 -*-
"""
动态原始字段契约与 Excel 辅助页工具。

本模块负责把不同分页、不同样本返回的字段合并为稳定并集，并保留字段来源和中文备注。
"""

# 导入数据类定义工具。
from dataclasses import dataclass

# 导入数据类字段工具。
from dataclasses import field

# 导入 Excel 工作簿工具。
from openpyxl import Workbook


# 定义动态字段收集结果。
@dataclass
class DynamicFieldTable:

    """
    保存字段并集及按并集对齐的数据行。

    :param columns: 字段描述列表。
    :type columns: list[dict]
    :param rows: 已按字段并集对齐的数据行。
    :type rows: list[dict]
    :param changes: 字段变化记录。
    :type changes: list[dict]
    """

    # 保存动态字段描述。
    columns: list = field(default_factory=list)

    # 保存对齐后的数据行。
    rows: list = field(default_factory=list)

    # 保存字段变化记录。
    changes: list = field(default_factory=list)


# 定义动态字段收集器。
class DynamicFieldCollector:

    """
    按首次出现顺序合并多页和多样本原始字段。

    :param configured_titles: 已知原始键到中文标题的映射。
    :type configured_titles: dict | None
    """

    # 初始化动态字段收集器。
    def __init__(self, configured_titles=None):

        """
        初始化字段标题配置。

        :param configured_titles: 已知原始键到中文标题的映射。
        :type configured_titles: dict | None
        :return: None。
        :rtype: None
        """

        # 保存标准化配置标题。
        self.configured_titles = {str(key): str(value) for key, value in (configured_titles or {}).items() if str(key).strip() and str(value).strip()}

    # 合并多页记录和字段元数据。
    def collect(self, pages):

        """
        合并多页字段并将每条记录对齐到字段并集。

        :param pages: 分页对象集合，每项可包含 rows、headers、page、metadata。
        :type pages: list[dict]
        :return: 动态字段表。
        :rtype: DynamicFieldTable
        """

        # 初始化字段索引和描述集合。
        columns = []
        column_index = {}

        # 初始化字段变化日志。
        changes = []

        # 初始化暂存原始记录集合。
        raw_rows = []

        # 遍历全部分页响应。
        for page_number, page in enumerate(pages or [], start=1):

            # 读取当前页码。
            current_page = int(page.get("page", page_number) or page_number) if isinstance(page, dict) else page_number

            # 读取当前页字段元数据。
            page_headers = page.get("headers", []) if isinstance(page, dict) else []

            # 读取当前页记录。
            page_rows = page.get("rows", []) if isinstance(page, dict) else []

            # 规范化当前页字段元数据。
            normalized_headers = self._normalize_headers(page_headers, page_rows)

            # 记录当前页新增字段。
            page_new_keys = []

            # 遍历当前页字段。
            for header in normalized_headers:

                # 读取原始字段键。
                key = str(header.get("key", "")).strip()

                # 跳过空字段键。
                if not key:
                    continue

                # 记录首次出现的新字段。
                if key not in column_index:
                    column_index[key] = len(columns)
                    columns.append({"key": key, "title": self._resolve_title(key, header), "source": "configured" if key in self.configured_titles else str(header.get("source", "response") or "response"), "ordinal": len(columns), "note": str(header.get("note", "") or ""), "first_page": current_page, "last_page": current_page, "dynamic": bool(page_number > 1)})
                    page_new_keys.append(key)
                else:
                    # 更新字段最后出现页码。
                    columns[column_index[key]]["last_page"] = current_page

            # 记录分页字段新增变化。
            if page_new_keys and page_number > 1:
                changes.append({"page": current_page, "added": list(page_new_keys), "missing": []})

            # 构建当前页记录并保留页码。
            for row in page_rows if isinstance(page_rows, list) else []:
                if isinstance(row, dict):
                    raw_rows.append((current_page, dict(row)))

        # 计算每页缺失字段并对齐所有记录。
        for page_number, row in raw_rows:

            # 计算当前行缺失字段。
            missing_keys = [column["key"] for column in columns if column["key"] not in row]

            # 记录分页缺失字段。
            if missing_keys:
                changes.append({"page": page_number, "added": [], "missing": missing_keys})

        # 生成按字段并集对齐的记录。
        aligned_rows = [{column["key"]: row.get(column["key"], "") for column in columns} for _, row in raw_rows]

        # 返回动态字段表。
        return DynamicFieldTable(columns, aligned_rows, self._deduplicate_changes(changes))

    # 规范化字段元数据。
    def _normalize_headers(self, headers, rows):

        """
        将字段列表、字典字段和记录键统一转换为字段元数据。

        :param headers: 原始字段元数据。
        :type headers: list
        :param rows: 当前页记录。
        :type rows: list
        :return: 规范化字段元数据列表。
        :rtype: list[dict]
        """

        # 初始化字段元数据结果。
        result = []

        # 解析显式字段定义。
        for value in headers if isinstance(headers, list) else []:
            if isinstance(value, dict):
                result.append({"key": value.get("key") or value.get("name") or value.get("field"), "title": value.get("title") or value.get("label") or value.get("caption") or "", "note": value.get("note") or value.get("remark") or value.get("备注") or "", "source": "metadata"})
            else:
                result.append({"key": value, "title": value, "note": "", "source": "metadata"})

        # 从记录键补齐没有字段元数据的动态字段。
        known_keys = {str(value.get("key", "")) for value in result}
        for row in rows if isinstance(rows, list) else []:
            if isinstance(row, dict):
                for key in row:
                    if str(key) not in known_keys:
                        result.append({"key": key, "title": key, "note": "", "source": "response"})
                        known_keys.add(str(key))

        # 返回规范化字段元数据。
        return result

    # 解析字段正式标题。
    def _resolve_title(self, key, header):

        """
        按配置标题、中文备注、接口标题和原始键优先级解析展示标题。

        :param key: 原始字段键。
        :type key: str
        :param header: 当前字段元数据。
        :type header: dict
        :return: 正式展示标题。
        :rtype: str
        """

        # 优先使用商业配置中的中文标题。
        if key in self.configured_titles:
            return self.configured_titles[key]

        # 读取接口中文备注。
        note = str(header.get("note", "")).strip()
        if note:
            return note

        # 读取接口字段标题。
        title = str(header.get("title", "")).strip()
        return title or key

    # 去除重复字段变化记录。
    def _deduplicate_changes(self, changes):

        """
        合并相同页码和字段集合的变化记录。

        :param changes: 原始变化记录。
        :type changes: list[dict]
        :return: 去重后的变化记录。
        :rtype: list[dict]
        """

        # 初始化去重映射。
        unique = {}

        # 遍历变化记录。
        for change in changes:
            key = (change.get("page"), tuple(change.get("added", [])), tuple(change.get("missing", [])))
            unique[key] = change

        # 返回保持插入顺序的变化记录。
        return list(unique.values())


# 定义动态原始表 Excel 写入函数。
def write_dynamic_raw_excel(file_path, field_table, sheet_name="原始数据"):

    """
    写入原始字段并集工作表及隐藏字段元数据工作表。

    :param file_path: 目标 Excel 文件路径。
    :type file_path: str
    :param field_table: 已完成字段合并的原始表数据。
    :type field_table: DynamicFieldTable
    :param sheet_name: 原始数据工作表名称。
    :type sheet_name: str
    :return: 已写入的 Excel 文件路径。
    :rtype: str
    """

    # 创建新的 Excel 工作簿。
    workbook = Workbook()

    # 获取默认原始数据工作表。
    data_sheet = workbook.active

    # 设置原始数据工作表名称。
    data_sheet.title = sheet_name

    # 写入正式原始字段标题。
    data_sheet.append([column.get("title", "") for column in field_table.columns])

    # 按字段并集顺序写入全部原始记录。
    for row in field_table.rows:

        # 写入缺失字段已补空值的原始数据行。
        data_sheet.append([row.get(column.get("key", ""), "") for column in field_table.columns])

    # 冻结用户熟悉的标题行。
    data_sheet.freeze_panes = "A2"

    # 在存在字段时设置筛选范围。
    if field_table.columns:

        # 将自动筛选仅应用于原始明细表。
        data_sheet.auto_filter.ref = data_sheet.dimensions

    # 创建不干扰明细阅读的字段元数据工作表。
    metadata_sheet = workbook.create_sheet("字段元数据")

    # 写入字段契约元数据标题。
    metadata_sheet.append(["正式标题", "原始键", "中文备注", "字段来源", "首次出现页码", "最后出现页码", "动态新增字段", "是否存在分页缺失", "商业配置映射状态"])

    # 汇总存在字段缺失的字段键。
    missing_keys = {key for change in field_table.changes for key in change.get("missing", [])}

    # 遍历字段并写入每个字段的元数据。
    for column in field_table.columns:

        # 读取当前字段原始键。
        key = column.get("key", "")

        # 计算商业配置标题是否命中。
        mapping_status = "已映射" if column.get("source") == "configured" else "未映射"

        # 写入字段的完整来源与分页信息。
        metadata_sheet.append([column.get("title", ""), key, column.get("note", ""), column.get("source", ""), column.get("first_page", ""), column.get("last_page", ""), "是" if column.get("dynamic") else "否", "是" if key in missing_keys else "否", mapping_status])

    # 隐藏辅助字段元数据工作表。
    metadata_sheet.sheet_state = "hidden"

    # 保存完整工作簿到目标路径。
    workbook.save(file_path)

    # 返回写入后的目标路径。
    return str(file_path)


# 定义正式原始 Excel 写入函数。
def write_formal_raw_excel(file_path, headers, rows, sheet_name="原始数据"):

    """
    按商业页面或接口提供的标题文本与字段顺序写入正式原始 Excel。

    :param file_path: 正式原始 Excel 的目标文件路径。
    :type file_path: str
    :param headers: 页面或接口返回的有序字段定义。
    :type headers: list[str] | list[dict]
    :param rows: 接口或页面解析得到的原始记录。
    :type rows: list[dict]
    :param sheet_name: 原始数据工作表名称。
    :type sheet_name: str
    :return: 已写入的正式原始 Excel 文件路径。
    :rtype: str
    """

    # 解析不改名且保持顺序的正式原始字段定义。
    columns = normalize_formal_raw_columns(headers, rows)

    # 创建正式原始 Excel 工作簿。
    workbook = Workbook()

    # 获取默认原始数据工作表。
    data_sheet = workbook.active

    # 设置受 Excel 长度限制的原始数据工作表名称。
    data_sheet.title = str(sheet_name or "原始数据")[:31]

    # 写入页面或接口返回的原始标题文本。
    data_sheet.append([column["title"] for column in columns])

    # 遍历全部原始记录。
    for row in rows or []:

        # 读取当前原始记录字典。
        record = row if isinstance(row, dict) else {}

        # 按原始字段键顺序写入当前记录值。
        data_sheet.append([record.get(column["key"], "") for column in columns])

    # 冻结原始数据标题行。
    data_sheet.freeze_panes = "A2"

    # 为原始数据工作表设置自动筛选。
    if columns:

        # 将自动筛选范围限制在正式原始明细。
        data_sheet.auto_filter.ref = data_sheet.dimensions

    # 创建不干扰业务明细的字段元数据工作表。
    metadata_sheet = workbook.create_sheet("字段元数据")

    # 写入字段元数据标题行。
    metadata_sheet.append(["正式标题", "原始键", "中文备注", "字段来源", "首次出现页码", "最后出现页码", "动态新增字段", "是否存在分页缺失", "商业配置映射状态"])

    # 遍历正式原始字段写入基础元数据。
    for column in columns:

        # 写入缺少额外接口元数据时的保守字段来源信息。
        metadata_sheet.append([column["title"], column["key"], "", "response", 1, 1, "否", "否", "未映射"])

    # 隐藏字段元数据工作表以保持原始明细视觉简洁。
    metadata_sheet.sheet_state = "hidden"

    # 保存完整的正式原始 Excel 工作簿。
    workbook.save(file_path)

    # 关闭已保存的 Excel 工作簿。
    workbook.close()

    # 返回正式原始 Excel 文件路径。
    return str(file_path)


# 定义正式原始字段规范化函数。
def normalize_formal_raw_columns(headers, rows):

    """
    将页面或接口字段定义规范化为不改名的有序原始列。

    :param headers: 页面或接口返回的有序字段定义。
    :type headers: list[str] | list[dict]
    :param rows: 接口或页面解析得到的原始记录。
    :type rows: list[dict]
    :return: 按页面或接口顺序排列的原始列定义。
    :rtype: list[dict]
    """

    # 初始化正式原始列定义集合。
    columns = []

    # 初始化已登记的原始字段键集合。
    known_keys = set()

    # 遍历页面或接口声明的字段定义。
    for header in headers or []:

        # 解析字典字段定义中的原始键。
        raw_key = header.get("key", header.get("name", header.get("field", ""))) if isinstance(header, dict) else header

        # 规范化原始字段键文本。
        key = str(raw_key or "")

        # 跳过无法从原始记录取值的空字段键。
        if not key:

            # 继续解析后续字段定义。
            continue

        # 跳过重复出现的同一原始字段键。
        if key in known_keys:

            # 继续解析后续字段定义。
            continue

        # 读取页面或接口提供的原始标题文本。
        raw_title = header.get("title", header.get("label", header.get("caption", key))) if isinstance(header, dict) else header

        # 保留原始标题文本且不追加前缀或序号。
        title = str(raw_title if raw_title is not None else key)

        # 记录当前原始字段键。
        known_keys.add(key)

        # 追加当前原始列定义。
        columns.append({"key": key, "title": title})

    # 遍历全部原始记录以补齐分页或样本中新增的字段。
    for row in rows or []:

        # 跳过非字典格式的原始记录。
        if not isinstance(row, dict):

            # 继续读取下一条原始记录。
            continue

        # 遍历当前原始记录保持的字段插入顺序。
        for raw_key in row:

            # 规范化当前新增字段键文本。
            key = str(raw_key or "")

            # 跳过空字段键或已登记字段。
            if not key or key in known_keys:

                # 继续读取下一字段。
                continue

            # 记录当前新增原始字段键。
            known_keys.add(key)

            # 追加使用原始键作为标题的动态新增字段。
            columns.append({"key": key, "title": key})

    # 返回保持真实标题与顺序的原始列定义。
    return columns
