# -*- coding: utf-8 -*-
"""
米伦厂商合作平台协议适配器。

本模块负责米伦软件厂商合作平台（百姓堂厂商合作平台同构站点）的会话登录、
流向清单过滤查询、分页读取和本地 Excel 生成。
"""

# 导入操作系统路径模块。
import os

# 导入 URL 表单编码函数。
from urllib.parse import urlencode

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义米伦平台数据类型与单据类型清单映射。
BILL_TYPE_FILTERS = {
    "购进流向": "('采购退货单','采购入库单')",
    "纯销流向": "('零售销售单')",
}


class MilunPlatform(BasePlatform):

    """
    米伦厂商合作平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 米伦平台适配器实例。
    :rtype: MilunPlatform
    """

    # 定义米伦平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化米伦平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://bxtyylsyf.hicp.net:1686/logon.htm")

        # 保存平台根地址。
        self.base_url = account.get("接口根地址", "http://bxtyylsyf.hicp.net:1686")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.base_url}/Submitlogon.htm")

        # 保存清单查询接口地址。
        self.query_api_url = account.get("查询接口", f"{self.base_url}/GetFlowQueryList.json")

        # 保存清单页地址。
        self.query_page_url = account.get("清单页", f"{self.base_url}/FlowQuery.html")

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 20))

        # 初始化会话 Cookie 串。
        self.session_cookie = ""

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录状态。
        self.logged_in = False

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验数据类型是否已接入。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否存在单据类型过滤定义。
        if data_type not in BILL_TYPE_FILTERS:

            # 抛出数据类型未接入异常。
            raise ValueError(f"米伦厂商合作平台暂未接入数据类型：{data_type}")

    # 定义平台登录方法。
    def login(self):

        """
        执行米伦厂商合作平台协议登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("米伦厂商合作平台登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 访问登录页以初始化站点状态。
        self.http.get(self.login_page_url, referer=self.base_url)

        # 构建登录请求体。
        payload = {"username": self.clean_login_name(), "password": self.clean_login_password()}

        # 构建登录请求头。
        headers = {"Content-Type": "application/x-www-form-urlencoded", "Referer": self.login_page_url, "Origin": self.base_url}

        # 编码表单请求体。
        data = urlencode(payload).encode("utf-8")

        # 提交登录请求。
        result = self.http.request("POST", self.login_api_url, headers=headers, data=data, timeout=60, allow_http_error=True)

        # 解析登录响应。
        login_result = self.decode_json_response(result)

        # 校验登录结果。
        if int(login_result.get("errcode", -1)) != 0:

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 登录失败：{login_result.get('errmsg') or login_result}")

        # 保存会话 Cookie 串（平台以响应 JSON 回传会话而非 Set-Cookie）。
        self.session_cookie = str(login_result.get("cookie", ""))

        # 标记当前登录状态。
        self.logged_in = True

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("登录状态", "成功", "green"), self.ph.kv("落地页", login_result.get("location", ""), "cyan")))

        # 返回登录响应数据。
        return login_result

    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行。
        :rtype: list[dict]
        """

        # 查询米伦平台指定类型完整数据。
        query_result = self.query_milun_data(data_type, start_date, end_date)

        # 返回查询结果行。
        return query_result["rows"]

    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定类型业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询米伦平台指定类型完整数据。
        query_result = self.query_milun_data(data_type, start_date, end_date)

        # 使用当前工作目录作为默认输出目录。
        output_dir = output_dir or os.getcwd()

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建输出文件名。
        file_name = self.build_output_file_name(data_type, query_result["start_date"], query_result["end_date"])

        # 构建输出文件路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan"), self.ph.kv("字段数", len(query_result["headers"]), "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义文本响应解析方法。
    def decode_text(self, result):

        """
        解析文本响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 响应文本。
        :rtype: str
        """

        # 解码 UTF-8 响应文本。
        return result.body.decode("utf-8", errors="replace").strip()

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 导入 JSON 编解码模块。
        import json

        # 解码响应文本。
        response_text = self.decode_text(result)

        # 处理空响应。
        if not response_text:

            # 返回空响应结构。
            return {}

        # 返回 JSON 解析结果。
        return json.loads(response_text)

    # 定义米伦平台数据查询方法。
    def query_milun_data(self, data_type, start_date, end_date):

        """
        查询米伦平台指定类型全部分页数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 查询结果、表头和日期范围。
        :rtype: dict
        """

        # 校验数据类型。
        self.validate_data_type(data_type)

        # 确保当前会话已经登录。
        self.ensure_login()

        # 缺省日期回退为当天。
        start_date = start_date or self.today_text()
        end_date = end_date or self.today_text()

        # 打印查询阶段标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date, "cyan"), self.ph.kv("结束日期", end_date, "cyan")))

        # 构建清单过滤串。
        filter_text = self.build_filter_text(data_type, start_date, end_date)

        # 初始化合并数据行集合。
        merged_rows = []

        # 初始化合并表头集合。
        merged_headers = []

        # 初始化总页数。
        total_pages = 1

        # 初始化当前页码。
        page_number = 1

        # 循环读取全部分页。
        while page_number <= total_pages:

            # 读取当前页响应。
            response_data = self.fetch_page(filter_text, page_number)

            # 校验业务错误码。
            errcode = int(response_data.get("errcode", -1))

            # 判断当前页是否返回无效清单业务码。
            if errcode == 2:

                # 无效清单为空结果业务应答，记录日志后终止读取。
                self.ph.info(self.ph.join(self.ph.kv("查询结果", "当前过滤条件下无有效清单", "yellow")))

                # 结束分页循环。
                break

            # 判断当前页是否返回其他业务错误。
            if errcode != 0:

                # 抛出查询业务异常。
                raise RuntimeError(f"{self.company_name} {data_type} 查询失败：{response_data.get('errmsg') or response_data}")

            # 更新总页数。
            total_pages = int(response_data.get("Pages", 1) or 1)

            # 解析嵌套 JSON 字符串数据块。
            current_rows = self.parse_nested_rows(response_data.get("data", ""))

            # 合并当前页数据行。
            merged_rows.extend(current_rows)

            # 按行序合并表头。
            merged_headers = self.merge_headers(merged_headers, current_rows)

            # 打印分页读取日志。
            self.ph.info(self.ph.join(self.ph.kv("当前页", f"{page_number}/{total_pages}", "cyan"), self.ph.kv("本页记录", len(current_rows), "cyan"), self.ph.kv("累计记录", len(merged_rows), "cyan")))

            # 页码递增。
            page_number += 1

        # 返回查询结果。
        return {"headers": merged_headers, "rows": merged_rows, "start_date": start_date, "end_date": end_date}

    # 定义当天日期文本读取方法。
    def today_text(self):

        """
        读取当天日期文本。

        :return: 当天日期文本。
        :rtype: str
        """

        # 导入日期时间对象。
        from datetime import datetime

        # 返回当天日期文本。
        return datetime.now().strftime("%Y-%m-%d")

    # 定义清单过滤串构建方法。
    def build_filter_text(self, data_type, start_date, end_date):

        """
        构建清单过滤串。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 清单过滤串。
        :rtype: str
        """

        # 返回记账日期区间与单据类型清单拼接的过滤串。
        return f"[记账日期] >= '{start_date}' And [记账日期] <= '{end_date}' And [单据类型] in {BILL_TYPE_FILTERS[data_type]}"

    # 定义分页读取方法。
    def fetch_page(self, filter_text, page_number):

        """
        读取指定页码的清单数据。

        :param filter_text: 清单过滤串。
        :type filter_text: str
        :param page_number: 当前页码。
        :type page_number: int
        :return: 页面响应数据。
        :rtype: dict
        """

        # 构建请求头（会话 Cookie 为登录响应回传串）。
        headers = {"Content-Type": "application/json; charset=UTF-8", "Cookie": self.session_cookie, "Referer": f"{self.query_page_url}?page={page_number}", "Origin": self.base_url}

        # 发送清单查询请求。
        result = self.http.request("POST", f"{self.query_api_url}?page={page_number}", headers=headers, data=filter_text.encode("utf-8"), timeout=180, allow_http_error=True)

        # 返回 JSON 响应数据。
        return self.decode_json_response(result)

    # 定义嵌套数据块解析方法。
    def parse_nested_rows(self, data_text):

        """
        解析响应中嵌套 JSON 字符串的数据块。

        :param data_text: 嵌套 JSON 字符串。
        :type data_text: str
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 导入 JSON 编解码模块。
        import json

        # 判断数据块是否为空。
        if not data_text:

            # 返回空数据行。
            return []

        # 解析嵌套 JSON 字符串。
        rows = json.loads(data_text)

        # 过滤非字典数据行。
        return [row for row in rows if isinstance(row, dict)]

    # 定义表头按行序合并方法。
    def merge_headers(self, merged_headers, rows):

        """
        按行序合并动态表头。

        :param merged_headers: 已合并表头。
        :type merged_headers: list[str]
        :param rows: 当前页数据行。
        :type rows: list[dict]
        :return: 合并后的表头。
        :rtype: list[str]
        """

        # 初始化当前合并表头副本。
        current_headers = list(merged_headers)

        # 遍历当前页数据行。
        for row in rows:

            # 遍历当前行字段名称。
            for header_name in row.keys():

                # 判断字段名称是否已存在。
                if header_name in current_headers:

                    # 进入下一字段。
                    continue

                # 追加新字段。
                current_headers.append(header_name)

        # 返回合并后的表头。
        return current_headers

    # 定义通用 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头列表。
        :type headers: list[str]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(header_name, "") if isinstance(row, dict) else "" for header_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error
