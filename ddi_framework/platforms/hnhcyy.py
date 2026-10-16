# -*- coding: utf-8 -*-
"""
河南省慧采医药有限公司流向管理系统协议适配器。

本模块负责河南省慧采医药有限公司流向管理系统的账号登录、分页查询和本地 Excel 生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期时间对象。
from datetime import datetime

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义河南慧采医药流向管理系统协议适配器。
class HnhcyyPlatform(BasePlatform):

    """
    河南慧采医药流向管理系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 河南慧采医药平台适配器实例。
    :rtype: HnhcyyPlatform
    """

    # 定义河南慧采医药适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化河南慧采医药平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://1.192.219.160:61334/index.html#/login")

        # 保存平台根地址。
        self.base_url = account.get("接口根地址") or self.build_origin(self.login_page_url)

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.base_url}/actionapi/Users/UserLogin")

        # 保存购进流向查询接口地址。
        self.purchase_query_url = account.get("购进流向接口", f"{self.base_url}/actionapi/Users/jhrkDataSearch")

        # 保存配送流向查询接口地址。
        self.distribution_query_url = account.get("配送流向接口", f"{self.base_url}/actionapi/Users/pfckDataSearch")

        # 保存库存查询接口地址。
        self.stock_query_url = account.get("库存接口", f"{self.base_url}/actionapi/Users/sphwphDataSearch")

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 100))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录响应。
        self.login_result = {}

        # 初始化平台用户 ID。
        self.user_id = account.get("用户ID")

        # 初始化登录状态。
        self.logged_in = False

    # 定义河南慧采医药登录方法。
    def login(self):

        """
        执行河南慧采医药协议登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("河南慧采医药登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 访问登录页以初始化站点 Cookie。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 构建登录请求体。
        payload = self.build_login_payload()

        # 打印登录账号日志。
        self.ph.info(self.ph.join(self.ph.kv("登录账号", self.clean_login_name(), "cyan"), self.ph.kv("密码长度", len(self.clean_login_password()), "cyan")))

        # 打印正式登录请求日志。
        self.ph.info("提交正式登录请求: /actionapi/Users/UserLogin")

        # 提交登录请求。
        result = self.post_json(self.login_api_url, payload, self.login_page_url)

        # 解析登录响应。
        login_result = self.decode_json_response(result)

        # 校验登录结果。
        if not self.is_success_response(login_result):

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 正式登录失败：{login_result.get('Message') or login_result}")

        # 保存登录响应。
        self.login_result = login_result

        # 保存平台用户 ID。
        self.user_id = self.resolve_user_id(login_result)

        # 保存登录状态。
        self.logged_in = True

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("正式登录成功", login_result.get("Message", "登录成功"), "green"), self.ph.kv("USID", self.user_id, "cyan")))

        # 返回登录响应。
        return login_result

    # 定义登录请求体构建方法。
    def build_login_payload(self):

        """
        构建登录请求体。

        :return: 登录请求体。
        :rtype: dict
        """

        # 返回登录字段。
        return {"UserName": self.clean_login_name(), "PassWord": self.clean_login_password()}



    # 定义平台用户 ID 解析方法。
    def resolve_user_id(self, login_result):

        """
        解析平台用户 ID。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 平台用户 ID。
        :rtype: int | str
        """

        # 判断配置中是否已经提供平台用户 ID。
        if self.user_id not in [None, ""]:

            # 返回配置中的平台用户 ID。
            return self.user_id

        # 读取登录响应业务数据。
        data = login_result.get("Data")

        # 从字典结构中读取平台用户 ID。
        if isinstance(data, dict):

            # 返回常见字段中的平台用户 ID。
            return data.get("USID") or data.get("usid") or data.get("UserID") or data.get("ID")

        # 从列表首行中读取平台用户 ID。
        if isinstance(data, list) and data and isinstance(data[0], dict):

            # 返回常见字段中的平台用户 ID。
            return data[0].get("USID") or data[0].get("usid") or data[0].get("UserID") or data[0].get("ID")

        # 抛出平台用户 ID 缺失异常。
        raise RuntimeError(f"{self.company_name} 登录响应未返回 USID，请在账号配置中补充 用户ID")


    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询河南慧采医药业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 查询指定类型完整数据。
        query_result = self.query_hnhcyy_data(data_type, start_date, end_date)

        # 返回查询结果数据行。
        return query_result["rows"]


    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询河南慧采医药业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定业务数据。
        query_result = self.query_hnhcyy_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义河南慧采医药数据查询方法。
    def query_hnhcyy_data(self, data_type, start_date, end_date):

        """
        查询河南慧采医药指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 校验流向类查询开始日期。
        if data_type in ["购进流向", "配送流向"] and not start_date:

            # 抛出查询开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验流向类查询结束日期。
        if data_type in ["购进流向", "配送流向"] and not end_date:

            # 抛出查询结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 初始化查询结果行。
        all_rows = []

        # 初始化响应表头。
        headers = []

        # 初始化当前页码。
        page_index = 1

        # 初始化总记录数。
        total_count = None

        # 循环读取所有分页。
        while True:

            # 构建当前分页请求体。
            payload = self.build_query_payload(data_type, start_date, end_date, page_index)

            # 打印当前分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"pageIndex={page_index}", "cyan"), self.ph.kv("pageSize", self.page_size, "cyan")))

            # 发送当前分页查询请求。
            result = self.post_json(self.build_query_api_url(data_type), payload, self.login_page_url)

            # 解析当前分页响应。
            response_data = self.decode_json_response(result)

            # 校验当前分页响应状态。
            if not self.is_success_response(response_data):

                # 抛出查询失败异常。
                raise RuntimeError(f"{self.company_name} {data_type}查询失败：{response_data.get('Message') or response_data}")

            # 读取响应表头。
            headers = headers or self.extract_headers(response_data, data_type)

            # 读取当前页数据行。
            current_rows = self.extract_rows(response_data)

            # 读取当前页总记录数。
            total_count = self.resolve_total_count(current_rows, total_count)

            # 合并当前页数据行。
            all_rows.extend(current_rows)

            # 打印当前页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(current_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count if total_count is not None else "未知", "cyan")))

            # 判断总数是否已经取完。
            if total_count is not None and len(all_rows) >= total_count:

                # 结束分页循环。
                break

            # 判断当前页是否已经不足一页。
            if len(current_rows) < self.page_size:

                # 结束分页循环。
                break

            # 递增页码。
            page_index += 1

        # 返回查询数据与表头。
        return {"rows": all_rows, "headers": headers or self.build_default_headers(data_type)}


    # 定义查询阶段标题打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印查询阶段标题。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断是否为库存查询。
        if data_type == "库存":

            # 打印库存查询标题。
            self.ph.section(f"{data_type}查询", "库存接口不传递起止日期")

            # 结束标题打印。
            return

        # 打印流向查询标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

    # 定义查询请求体构建方法。
    def build_query_payload(self, data_type, start_date, end_date, page_index):

        """
        构建业务查询请求体。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param page_index: 当前页码。
        :type page_index: int
        :return: 业务查询请求体。
        :rtype: dict
        """

        # 构建通用查询字段。
        payload = {"dwmch": "", "tongym": "", "shengccj": "", "USID": self.user_id, "pageIndex": page_index, "pageSize": self.page_size}

        # 判断是否为流向类查询。
        if data_type in ["购进流向", "配送流向"]:

            # 写入查询开始日期。
            payload["begintime"] = start_date

            # 写入查询结束日期。
            payload["endtime"] = end_date

        # 返回查询请求体。
        return payload

    # 定义查询接口地址构建方法。
    def build_query_api_url(self, data_type):

        """
        构建数据类型对应的查询接口地址。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 查询接口地址。
        :rtype: str
        """

        # 定义数据类型与接口地址映射。
        mapping = {"购进流向": self.purchase_query_url, "配送流向": self.distribution_query_url, "库存": self.stock_query_url}

        # 返回当前数据类型查询接口。
        return mapping[data_type]

    # 定义 JSON POST 请求方法。
    def post_json(self, url, payload, referer):

        """
        提交 JSON 请求。

        :param url: 请求地址。
        :type url: str
        :param payload: 请求体字段。
        :type payload: dict
        :param referer: 来源页面地址。
        :type referer: str
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 序列化 JSON 请求体。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建 JSON 请求头。
        headers = self.build_json_headers(referer)

        # 发送 POST 请求。
        return self.http.request("POST", url, headers=headers, data=body, timeout=90)

    # 定义 JSON 请求头构建方法。
    def build_json_headers(self, referer):

        """
        构建 JSON 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: JSON 请求头。
        :rtype: dict
        """

        # 返回 JSON 请求头。
        return {"Accept": "*/*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/json;charset=UTF-8", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应文本。
        response_text = result.body.decode("utf-8", errors="replace").strip()

        # 处理空响应。
        if not response_text:

            # 返回空响应结构。
            return {}

        # 返回 JSON 响应数据。
        return json.loads(response_text)

    # 定义响应成功判断方法。
    def is_success_response(self, response_data):

        """
        判断接口响应是否成功。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 是否成功。
        :rtype: bool
        """

        # 返回接口状态判断结果。
        return str(response_data.get("Code")) == "200"

    # 定义响应业务数据读取方法。
    def extract_response_data(self, response_data):

        """
        读取响应中的业务数据对象。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 响应业务数据对象。
        :rtype: dict
        """

        # 读取响应业务数据。
        data = response_data.get("Data") or {}

        # 返回字典结构业务数据。
        if isinstance(data, dict):

            # 返回业务数据字典。
            return data

        # 返回空业务数据。
        return {}

    # 定义表格行提取方法。
    def extract_rows(self, response_data):

        """
        提取响应数据行。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 读取响应业务数据。
        data = self.extract_response_data(response_data)

        # 读取表格数据行。
        rows = data.get("data1") or []

        # 返回字典数据行列表。
        return [row for row in rows if isinstance(row, dict)]

    # 定义响应表头提取方法。
    def extract_headers(self, response_data, data_type):

        """
        提取响应表头。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 读取响应业务数据。
        data = self.extract_response_data(response_data)

        # 读取平台表头定义。
        raw_headers = data.get("data2") or []

        # 初始化表头配置。
        headers = []

        # 遍历按序号排序后的平台表头。
        for header in sorted([item for item in raw_headers if isinstance(item, dict)], key=lambda item: item.get("id") or 0):

            # 读取字段名称。
            field_name = str(header.get("fieldname") or "").strip()

            # 跳过空字段名称。
            if not field_name:

                # 进入下一个表头。
                continue

            # 读取字段中文标题。
            header_name = str(header.get("annotations") or field_name).strip()

            # 保存表头配置。
            headers.append((header_name, field_name))

        # 返回响应表头或默认表头。
        return headers or self.build_default_headers(data_type)

    # 定义总记录数解析方法。
    def resolve_total_count(self, rows, current_total):

        """
        解析分页总记录数。

        :param rows: 当前页数据行。
        :type rows: list[dict]
        :param current_total: 已解析的总记录数。
        :type current_total: int | None
        :return: 总记录数。
        :rtype: int | None
        """

        # 返回已存在的总记录数。
        if current_total is not None:

            # 返回当前总记录数。
            return current_total

        # 判断当前页是否存在数据行。
        if not rows:

            # 返回零记录数。
            return 0

        # 读取首行总记录数字段。
        total_count = rows[0].get("TotalCount")

        # 返回数字型总记录数。
        if total_count not in [None, ""]:

            # 返回整型总记录数。
            return int(float(total_count))

        # 返回未知总记录数。
        return None

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验河南慧采医药数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已接入。
        if data_type not in ["购进流向", "配送流向", "库存"]:

            # 抛出数据类型未接入异常。
            raise ValueError(f"河南慧采医药暂未接入数据类型：{data_type}")

    # 定义默认表头构建方法。
    def build_default_headers(self, data_type):

        """
        构建默认输出表头配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 定义流向类默认表头。
        flow_headers = [("行数", "RowNum"), ("业务员", "ywy"), ("日期", "rq"), ("单据编号", "djbh"), ("商品编号", "spbh"), ("通用名", "tongym"), ("规格", "shpgg"), ("生产厂家", "shengccj"), ("单位", "dw"), ("剂型", "jixing"), ("批准文号", "pizwh"), ("数量", "shl"), ("批号", "pihao"), ("保质期至", "baozhiqi"), ("失效日期", "sxrq"), ("含税价", "hshj"), ("含税金额", "hsje"), ("制单人", "zhidr"), ("单位名称", "dwmch"), ("客户名称", "customerName")]

        # 判断当前类型是否为库存。
        if data_type == "库存":

            # 返回库存默认表头。
            return [("行数", "RowNum"), ("商品编号", "spbh"), ("通用名", "tongym"), ("批号", "pihao"), ("规格", "shpgg"), ("生产厂家", "shengccj"), ("单位", "dw"), ("数量", "shl"), ("含税价", "hshj"), ("库存金额", "kcje"), ("保质期至", "baozhiqi"), ("失效日期", "sxrq"), ("单位名称", "dwmch"), ("客户名称", "customerName")]

        # 返回流向类默认表头。
        return flow_headers




    # 定义通用 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(field_name, "") if isinstance(row, dict) else "" for _, field_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 读取当前表头名称。
            header_name = header_item[0]

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error


