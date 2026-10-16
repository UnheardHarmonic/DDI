# -*- coding: utf-8 -*-
"""
漳州市天健医药流向查询平台协议适配器。

本模块负责漳州市天健医药流向查询平台的账号登录、分页查询和本地 Excel 生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入 URL 查询参数编码函数。
from urllib.parse import urlencode

# 导入下载命名工具。
from ddi_framework import DownloadWatcher

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义漳州天健医药流向查询平台协议适配器。
class ZztjPlatform(BasePlatform):

    """
    漳州天健医药流向查询平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 漳州天健医药平台适配器实例。
    :rtype: ZztjPlatform
    """

    # 定义漳州天健医药适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化漳州天健医药平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "https://zz.wan2.x-tj.cn:18116/index.html?id=1")

        # 保存平台根地址。
        self.base_url = account.get("接口根地址", "https://zz.wan2.x-tj.cn:18116")

        # 保存核心接口地址。
        self.core_api_url = account.get("核心接口", f"{self.base_url}/CoreAPI.php")

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 500))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录状态。
        self.logged_in = False

    # 定义漳州天健医药登录方法。
    def login(self):

        """
        执行漳州天健医药协议登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("漳州天健医药登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 访问登录页以初始化站点上下文。
        self.http.get(self.login_page_url)

        # 构建登录接口地址。
        login_url = self.build_api_url({"ct": "login", "u": self.clean_login_name(), "p": self.clean_login_password()})

        # 打印登录账号日志。
        self.ph.info(self.ph.join(self.ph.kv("登录账号", self.clean_login_name(), "cyan"), self.ph.kv("密码长度", len(self.clean_login_password()), "cyan")))

        # 发送登录请求。
        result = self.http.get(login_url, referer=self.login_page_url)

        # 解码登录响应文本。
        response_text = self.decode_text(result)

        # 校验登录响应内容。
        if response_text.lower() != "ok":

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 正式登录失败：{response_text}")

        # 保存登录状态。
        self.logged_in = True

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("正式登录成功", "ok", "green"), self.ph.kv("Cookie数", len(list(self.http.cookie_jar)), "cyan")))

        # 返回登录结果。
        return {"code": 0, "msg": "ok"}


    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询漳州天健医药业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 查询指定类型完整数据。
        query_result = self.query_zztj_data(data_type, start_date, end_date)

        # 返回查询结果数据行。
        return query_result["rows"]


    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询漳州天健医药业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定类型完整数据。
        query_result = self.query_zztj_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义漳州天健医药数据查询方法。
    def query_zztj_data(self, data_type, start_date, end_date):

        """
        查询漳州天健医药指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 校验流向类查询开始日期。
        if data_type in ["购进流向", "配送流向"] and not start_date:

            # 抛出查询开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验流向类查询结束日期。
        if data_type in ["购进流向", "配送流向"] and not end_date:

            # 抛出查询结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 初始化查询结果行。
        all_rows = []

        # 初始化当前页码。
        page_index = 1

        # 初始化总记录数。
        total_count = None

        # 循环读取所有分页。
        while True:

            # 构建当前分页请求参数。
            query_params = self.build_query_params(data_type, start_date, end_date, page_index)

            # 构建当前分页接口地址。
            query_url = self.build_api_url(query_params)

            # 打印当前分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"page={page_index}", "cyan"), self.ph.kv("limit", query_params["limit"], "cyan")))

            # 发送当前分页查询请求。
            result = self.http.get(query_url, referer=self.build_referer(data_type))

            # 解析当前分页响应。
            response_data = self.decode_json_response(result)

            # 校验当前分页响应状态。
            if str(response_data.get("code")) != "0":

                # 抛出查询失败异常。
                raise RuntimeError(f"{self.company_name} {data_type}查询失败：{response_data.get('msg') or response_data}")

            # 提取当前页数据行。
            current_rows = self.extract_rows(response_data)

            # 解析当前页总记录数。
            total_count = self.resolve_total_count(response_data, total_count)

            # 合并当前页数据行。
            all_rows.extend(current_rows)

            # 打印当前页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(current_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count, "cyan")))

            # 判断总数是否已经取完。
            if len(all_rows) >= total_count:

                # 结束分页循环。
                break

            # 判断当前页是否已经不足一页。
            if len(current_rows) < self.page_size:

                # 结束分页循环。
                break

            # 递增页码。
            page_index += 1

        # 推断当前数据表头。
        headers = self.infer_headers(all_rows)

        # 返回查询数据与表头。
        return {"rows": all_rows, "headers": headers or self.build_default_headers(data_type)}


    # 定义查询阶段标题打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印查询阶段标题。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断是否为库存查询。
        if data_type == "库存":

            # 打印库存查询标题。
            self.ph.section(f"{data_type}查询", "库存接口不传递起止日期")

            # 结束标题打印。
            return

        # 打印流向查询标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

    # 定义查询参数构建方法。
    def build_query_params(self, data_type, start_date, end_date, page_index):

        """
        构建业务查询请求参数。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param page_index: 当前页码。
        :type page_index: int
        :return: 业务查询请求参数。
        :rtype: dict
        """

        # 构建通用查询字段。
        params = {"ct": self.build_query_ct(data_type), "th": "", "k": "", "page": page_index, "limit": self.page_size}

        # 判断是否为流向类查询。
        if data_type in ["购进流向", "配送流向"]:

            # 写入查询开始日期。
            params["d1"] = start_date

            # 写入查询结束日期。
            params["d2"] = end_date

        # 返回查询请求参数。
        return params

    # 定义查询类型编码构建方法。
    def build_query_ct(self, data_type):

        """
        构建数据类型对应的接口动作编码。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 接口动作编码。
        :rtype: str
        """

        # 定义内部数据类型与接口动作映射。
        mapping = {"购进流向": "show_cg_true", "配送流向": "show_xs_true", "库存": "show_kc_true"}

        # 返回当前数据类型接口动作。
        return mapping[data_type]

    # 定义来源页面构建方法。
    def build_referer(self, data_type):

        """
        构建数据类型对应的来源页面。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 来源页面地址。
        :rtype: str
        """

        # 定义内部数据类型与来源页面映射。
        mapping = {"购进流向": "cg_true.php", "配送流向": "xs_true.php", "库存": "kc_true.php"}

        # 返回当前数据类型来源页面。
        return f"{self.base_url}/{mapping[data_type]}"

    # 定义核心接口地址构建方法。
    def build_api_url(self, params):

        """
        构建核心接口完整地址。

        :param params: 查询参数。
        :type params: dict
        :return: 完整接口地址。
        :rtype: str
        """

        # 返回带查询参数的接口地址。
        return f"{self.core_api_url}?{urlencode(params)}"



    # 定义文本响应解析方法。
    def decode_text(self, result):

        """
        解析文本响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 响应文本。
        :rtype: str
        """

        # 返回清洗后的响应文本。
        return result.body.decode("utf-8", errors="replace").strip()

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应文本。
        response_text = self.decode_text(result)

        # 处理空响应。
        if not response_text:

            # 返回空响应结构。
            return {}

        # 返回 JSON 响应数据。
        return json.loads(response_text)

    # 定义响应数据行提取方法。
    def extract_rows(self, response_data):

        """
        提取响应数据行。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 读取响应数据行。
        rows = response_data.get("data") or []

        # 返回字典数据行列表。
        return [row for row in rows if isinstance(row, dict)]

    # 定义总记录数解析方法。
    def resolve_total_count(self, response_data, current_total):

        """
        解析分页总记录数。

        :param response_data: 当前页响应数据。
        :type response_data: dict
        :param current_total: 已解析的总记录数。
        :type current_total: int | None
        :return: 总记录数。
        :rtype: int
        """

        # 返回已存在的总记录数。
        if current_total is not None:

            # 返回当前总记录数。
            return current_total

        # 读取响应总记录数字段。
        count_value = response_data.get("count", 0)

        # 返回整型总记录数。
        return int(float(count_value or 0))

    # 定义表头推断方法。
    def infer_headers(self, rows):

        """
        根据数据行推断 Excel 表头。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :return: 表头列表。
        :rtype: list[str]
        """

        # 初始化表头结果列表。
        headers = []

        # 遍历查询结果数据行。
        for row in rows:

            # 跳过非字典数据行。
            if not isinstance(row, dict):

                # 继续处理下一行。
                continue

            # 遍历当前数据行字段名称。
            for header_name in row.keys():

                # 判断字段名称是否已经记录。
                if header_name in headers:

                    # 继续处理下一字段。
                    continue

                # 追加当前字段名称。
                headers.append(header_name)

        # 返回完整字段顺序。
        return headers

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验漳州天健医药数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已接入。
        if data_type not in ["购进流向", "配送流向", "库存"]:

            # 抛出数据类型未接入异常。
            raise ValueError(f"漳州天健医药暂未接入数据类型：{data_type}")

    # 定义默认表头构建方法。
    def build_default_headers(self, data_type):

        """
        构建默认输出表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 默认表头列表。
        :rtype: list[str]
        """

        # 定义购进流向默认表头。
        purchase_headers = ["单据编号", "日期", "单位名称", "商品编号", "商品名称", "通用名", "商品规格", "生产厂家", "单位", "批号", "有效期至", "出库数量", "库存数量", "进价", "总金额", "零售价", "应收金额", "单据类型", "在库天数", "含税进价", "失效天数", "最后供应商", "RowNum_x"]

        # 定义配送流向默认表头。
        distribution_headers = ["单据编号", "日期", "开单员", "单位名称", "商品编号", "商品名称", "通用名", "商品规格", "生产厂家", "单位", "批号", "有效期至", "出库数量", "库存数量", "单价", "小计金额", "单据类型", "在库天数", "含税进价", "失效天数", "最后供应商", "库区", "本次累计金额", "RowNum_x"]

        # 定义库存默认表头。
        stock_headers = ["商品名称", "通用名", "商品规格", "计量规格", "生产厂家", "单位", "批号", "有效期至", "库存数量", "含税单价", "含税金额", "包装规格", "批准文号", "在库天数", "含税进价", "失效天数", "最后供应商", "库区", "RowNum_x"]

        # 定义数据类型与默认表头映射。
        mapping = {"购进流向": purchase_headers, "配送流向": distribution_headers, "库存": stock_headers}

        # 返回当前数据类型默认表头。
        return mapping[data_type]


    # 定义通用 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头列表。
        :type headers: list[str]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(header_name, "") if isinstance(row, dict) else "" for header_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

        # 关闭导出工作簿。
        workbook.close()

