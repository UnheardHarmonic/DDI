# -*- coding: utf-8 -*-
"""
洪兴流向系统协议适配器。

本模块负责洪兴流向系统的会话初始化、WebForms 表单查询、HTML 表格解析和 Excel 文件生成。
"""

# 导入 HTML 实体解码函数。
from html import unescape

# 导入 HTML 标准解析器。
from html.parser import HTMLParser

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入正则表达式模块。
import re

# 导入日期时间解析对象。
from datetime import datetime

# 导入时区对象。
from datetime import timezone

# 导入 Cookie 对象。
from http.cookiejar import Cookie

# 导入 URL 编码函数。
from urllib.parse import quote

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入表单编码函数。
from urllib.parse import urlencode

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入下载命名工具。
from ddi_framework import DownloadWatcher

# 导入日志打印器。
from ddi_framework import PrintHandler

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient


# 定义洪兴流向系统表单解析器。
class HxlxFormParser(HTMLParser):

    """
    洪兴流向系统 HTML 表单解析器。

    :return: HTML 表单解析器实例。
    :rtype: HxlxFormParser
    """

    # 定义解析器初始化方法。
    def __init__(self):

        """
        初始化 HTML 表单解析器。

        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(convert_charrefs=True)

        # 初始化输入字段字典。
        self.fields = {}

        # 初始化按钮字段字典。
        self.buttons = {}

        # 初始化选择字段字典。
        self.select_values = {}

        # 初始化当前选择字段名称。
        self.current_select_name = ""

    # 定义开始标签处理方法。
    def handle_starttag(self, tag, attrs):

        """
        处理 HTML 开始标签。

        :param tag: 标签名称。
        :type tag: str
        :param attrs: 标签属性列表。
        :type attrs: list[tuple[str, str | None]]
        :return: None。
        :rtype: None
        """

        # 转换标签属性为字典。
        attr_map = {name: value or "" for name, value in attrs}

        # 处理输入控件标签。
        if tag.lower() == "input":

            # 读取输入控件名称。
            field_name = attr_map.get("name", "")

            # 读取输入控件类型。
            field_type = attr_map.get("type", "text").lower()

            # 读取输入控件值。
            field_value = attr_map.get("value", "")

            # 判断输入控件名称是否存在。
            if field_name:

                # 保存输入控件字段。
                self.fields[field_name] = field_value

            # 判断输入控件是否为按钮。
            if field_name and field_type in ("submit", "button"):

                # 保存按钮字段。
                self.buttons[field_name] = field_value

        # 处理选择控件标签。
        if tag.lower() == "select":

            # 保存当前选择控件名称。
            self.current_select_name = attr_map.get("name", "")

            # 判断选择控件名称是否存在。
            if self.current_select_name:

                # 初始化选择控件默认值。
                self.select_values.setdefault(self.current_select_name, "")

        # 处理选择控件选项标签。
        if tag.lower() == "option" and self.current_select_name:

            # 读取选项值。
            option_value = attr_map.get("value", "")

            # 读取选项选中状态。
            selected_flag = "selected" in attr_map

            # 判断选项是否应作为当前值。
            if selected_flag or not self.select_values.get(self.current_select_name):

                # 保存选择控件当前值。
                self.select_values[self.current_select_name] = option_value

    # 定义结束标签处理方法。
    def handle_endtag(self, tag):

        """
        处理 HTML 结束标签。

        :param tag: 标签名称。
        :type tag: str
        :return: None。
        :rtype: None
        """

        # 判断结束标签是否为选择控件。
        if tag.lower() == "select":

            # 清理当前选择控件名称。
            self.current_select_name = ""

    # 定义字段构建方法。
    def build_fields(self):

        """
        构建页面表单字段。

        :return: 页面表单字段。
        :rtype: dict
        """

        # 复制输入控件字段。
        result = dict(self.fields)

        # 遍历选择控件字段。
        for field_name, field_value in self.select_values.items():

            # 写入选择控件字段。
            result[field_name] = field_value

        # 返回完整表单字段。
        return result


# 定义洪兴流向系统表格解析器。
class HxlxTableParser(HTMLParser):

    """
    洪兴流向系统 HTML 表格解析器。

    :return: HTML 表格解析器实例。
    :rtype: HxlxTableParser
    """

    # 定义解析器初始化方法。
    def __init__(self):

        """
        初始化 HTML 表格解析器。

        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(convert_charrefs=True)

        # 初始化表格集合。
        self.tables = []

        # 初始化当前表格。
        self.current_table = None

        # 初始化当前行。
        self.current_row = None

        # 初始化当前单元格。
        self.current_cell = None

        # 初始化表格嵌套深度。
        self.table_depth = 0

    # 定义开始标签处理方法。
    def handle_starttag(self, tag, attrs):

        """
        处理 HTML 开始标签。

        :param tag: 标签名称。
        :type tag: str
        :param attrs: 标签属性列表。
        :type attrs: list[tuple[str, str | None]]
        :return: None。
        :rtype: None
        """

        # 转换标签名称。
        tag_name = tag.lower()

        # 处理表格开始标签。
        if tag_name == "table":

            # 增加表格深度。
            self.table_depth += 1

            # 判断当前是否为顶层表格。
            if self.table_depth == 1:

                # 初始化当前表格。
                self.current_table = []

        # 处理表格行开始标签。
        if tag_name == "tr" and self.current_table is not None:

            # 初始化当前行。
            self.current_row = []

        # 处理表格单元格开始标签。
        if tag_name in ("td", "th") and self.current_row is not None:

            # 初始化当前单元格。
            self.current_cell = []

    # 定义结束标签处理方法。
    def handle_endtag(self, tag):

        """
        处理 HTML 结束标签。

        :param tag: 标签名称。
        :type tag: str
        :return: None。
        :rtype: None
        """

        # 转换标签名称。
        tag_name = tag.lower()

        # 处理表格单元格结束标签。
        if tag_name in ("td", "th") and self.current_cell is not None and self.current_row is not None:

            # 拼接单元格文本。
            cell_text = "".join(self.current_cell)

            # 清理单元格文本。
            clean_text = self.clean_cell_text(cell_text)

            # 追加单元格文本。
            self.current_row.append(clean_text)

            # 清理当前单元格。
            self.current_cell = None

        # 处理表格行结束标签。
        if tag_name == "tr" and self.current_row is not None and self.current_table is not None:

            # 判断当前行是否存在有效单元格。
            if any(cell for cell in self.current_row):

                # 追加当前行。
                self.current_table.append(self.current_row)

            # 清理当前行。
            self.current_row = None

        # 处理表格结束标签。
        if tag_name == "table" and self.table_depth:

            # 判断当前是否离开顶层表格。
            if self.table_depth == 1 and self.current_table is not None:

                # 保存当前表格。
                self.tables.append(self.current_table)

                # 清理当前表格。
                self.current_table = None

            # 减少表格深度。
            self.table_depth -= 1

    # 定义文本处理方法。
    def handle_data(self, data):

        """
        处理 HTML 文本。

        :param data: 文本内容。
        :type data: str
        :return: None。
        :rtype: None
        """

        # 判断当前单元格是否正在采集。
        if self.current_cell is not None:

            # 追加文本内容。
            self.current_cell.append(data)

    # 定义单元格清理方法。
    def clean_cell_text(self, value):

        """
        清理表格单元格文本。

        :param value: 原始单元格文本。
        :type value: object
        :return: 清理后的单元格文本。
        :rtype: str
        """

        # 解码 HTML 实体。
        decoded_text = unescape(str(value))

        # 替换不间断空格。
        normalized_text = decoded_text.replace("\xa0", " ")

        # 压缩连续空白。
        compact_text = " ".join(normalized_text.split())

        # 返回清理后的文本。
        return compact_text.strip()


# 定义洪兴流向系统协议适配器。
class HxlxPlatform(BasePlatform):

    """
    洪兴流向系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 洪兴流向系统协议适配器实例。
    :rtype: HxlxPlatform
    """

    # 定义 DevExpress 脚本状态字段。
    DEVEXPRESS_SCRIPT_STATE = "1_230,1_134,1_223,1_168,1_131,1_166,1_174,1_138,1_179,1_165,1_164,1_215,1_213,1_156,1_153,1_193,1_195"

    # 定义 DevExpress 查询后脚本状态字段。
    DEVEXPRESS_GRID_SCRIPT_STATE = "1_230,1_134,1_223,1_168,1_131,1_166,1_174,1_138,1_179,1_165,1_164,1_215,1_213,1_156,1_153,1_193,1_195,1_211"

    # 定义 DevExpress 样式状态字段。
    DEVEXPRESS_STYLE_STATE = "1_29,1_33,1_35,1_3,1_17,1_18,1_16,1_22,1_21"

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化洪兴流向系统协议适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://117.40.251.4:8099/MainForm.html")

        # 保存站点根地址。
        self.base_url = account.get("接口根地址", self.build_origin(self.login_page_url))

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 5000))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录结果。
        self.login_result = {}

        # 初始化登录状态。
        self.logged_in = False

    # 定义平台登录方法。
    def login(self):

        """
        执行洪兴流向系统登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("洪兴流向登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 写入浏览器登录 Cookie。
        self.install_login_cookies()

        # 构建入口页面地址。
        entry_url = self.build_login_entry_url()

        # 打印入口访问日志。
        self.ph.info(self.ph.kv("访问入口页", entry_url, "cyan"))

        # 请求入口页面。
        result = self.http.get(entry_url, referer=self.base_url)

        # 解码入口页面文本。
        response_text = self.decode_text(result)

        # 判断入口页面是否有效。
        if "frameset" not in response_text.lower() and "rightframe" not in response_text:

            # 抛出入口页面异常。
            raise RuntimeError(f"{self.company_name} 入口页未返回有效系统页面")

        # 保存登录结果。
        self.login_result = {"登录状态": "成功", "入口页": entry_url, "Cookie数量": len(list(self.http.cookie_jar))}

        # 标记登录状态。
        self.logged_in = True

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("登录状态", "成功", "green"), self.ph.kv("Cookie", len(list(self.http.cookie_jar)), "cyan")))

        # 返回登录结果。
        return self.login_result

    # 定义平台导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        导出洪兴流向系统数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 返回平台导出文件。
        return self.download_flow_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义平台查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询洪兴流向系统数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果数据行。
        :rtype: list[dict]
        """

        # 下载平台 Excel 文件。
        excel_path = self.download_flow_excel(data_type, start_date=start_date, end_date=end_date, output_dir=None)

        # 返回 Excel 数据行。
        return self.read_rows_from_excel(excel_path)

    # 定义查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询洪兴流向系统数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 返回平台导出文件。
        return self.download_flow_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义平台 Excel 下载方法。
    def download_flow_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        下载洪兴流向系统平台 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 校验数据类型。
        self.validate_data_type(data_type)

        # 读取页面配置。
        page_config = self.resolve_page_config(data_type)

        # 构建页面地址。
        page_url = self.build_page_url(page_config)

        # 构建下载地址。
        download_url = self.build_download_url(page_url)

        # 打印数据类型查询分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 打印导出页加载日志。
        self.ph.info(self.ph.join(self.ph.kv("加载导出页", page_config["页面名称"], "cyan"), self.ph.kv("端点", page_config["路径"], "yellow")))

        # 请求业务页面。
        page_result = self.http.get(page_url, referer=self.login_page_url)

        # 解码业务页面。
        page_text = self.decode_text(page_result)

        # 解析初始表单字段。
        payload = self.parse_form_fields(page_text)

        # 写入查询按钮。
        query_payload = self.apply_submit_button(payload, page_text, "查看")

        # 写入日期字段。
        query_payload = self.apply_date_fields(query_payload, start_date, end_date)

        # 写入 DevExpress 运行时字段。
        query_payload = self.apply_devexpress_runtime_fields(query_payload)

        # 写入 DevExpress 提交状态字段。
        query_payload = self.apply_devexpress_submit_state_fields(query_payload, page_text)

        # 打印平台查询日志。
        self.ph.info(self.ph.join(self.ph.kv("提交平台查询", data_type, "cyan"), self.ph.kv("日期", f"{start_date} 至 {end_date}", "yellow")))

        # 提交平台查询表单。
        query_result = self.post_webforms(download_url, query_payload, download_url)

        # 解码平台查询页面。
        query_text = self.decode_text(query_result)

        # 解析查询后表单字段。
        export_payload = self.parse_form_fields(query_text)

        # 写入下载按钮。
        export_payload = self.apply_submit_button(export_payload, query_text, "下载")

        # 写入日期字段。
        export_payload = self.apply_date_fields(export_payload, start_date, end_date)

        # 写入 DevExpress 运行时字段。
        export_payload = self.apply_devexpress_runtime_fields(export_payload, grid_loaded=True)

        # 写入 DevExpress 提交状态字段。
        export_payload = self.apply_devexpress_submit_state_fields(export_payload, query_text)

        # 打印平台导出日志。
        self.ph.info(self.ph.join(self.ph.kv("提交平台导出", data_type, "cyan"), self.ph.kv("日期", f"{start_date} 至 {end_date}", "yellow")))

        # 初始化导出文件字节。
        export_file_bytes = b""

        # 初始化导出校验异常。
        export_validate_error = None

        # 遍历导出重试次数。
        for attempt_index in range(1, 4):

            # 提交平台导出表单。
            export_result = self.post_webforms(download_url, export_payload, download_url)

            # 读取导出响应字节。
            export_file_bytes = export_result.body

            # 捕获导出文件结构异常。
            try:

                # 校验导出文件结构。
                self.validate_xls_file_bytes(export_file_bytes)

                # 清理导出校验异常。
                export_validate_error = None

                # 结束导出重试循环。
                break

            # 记录导出文件结构异常。
            except Exception as error:

                # 保存导出校验异常。
                export_validate_error = error

                # 判断是否仍有重试机会。
                if attempt_index < 3:

                    # 打印导出重试日志。
                    self.ph.warning(self.ph.join(self.ph.kv("导出文件不完整", f"第 {attempt_index} 次", "yellow"), self.ph.kv("重新下载", data_type, "cyan")))

        # 判断导出文件是否仍然无效。
        if export_validate_error is not None:

            # 抛出导出文件异常。
            raise RuntimeError(f"{self.company_name} 平台导出 Excel 文件连续校验失败：{export_validate_error}") from export_validate_error

        # 保存平台 xls 文件。
        xls_path = self.save_platform_xls(export_file_bytes, data_type, start_date, end_date, output_dir)

        # 转换平台 xls 文件。
        xlsx_path = self.convert_xls_to_xlsx(xls_path)

        # 打印 Excel 文件生成日志。
        self.ph.success(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(xlsx_path), "green"))

        # 返回 xlsx 文件路径。
        return xlsx_path

    # 定义下载地址构建方法。
    def build_download_url(self, page_url):

        """
        构建平台导出地址。

        :param page_url: 页面地址。
        :type page_url: str
        :return: 平台导出地址。
        :rtype: str
        """

        # 判断页面地址是否以 .aspx 结尾。
        if page_url.lower().endswith(".aspx"):

            # 返回去除扩展名后的地址。
            return page_url[:-5]

        # 返回原始页面地址。
        return page_url

    # 定义 xls 字节结构校验方法。
    def validate_xls_file_bytes(self, file_bytes):

        """
        校验平台导出 xls 文件字节结构。

        :param file_bytes: 文件字节。
        :type file_bytes: bytes
        :return: None。
        :rtype: None
        """

        # 校验 Excel 复合文档头。
        if not file_bytes.startswith(b"\xd0\xcf\x11\xe0"):

            # 抛出导出异常。
            raise RuntimeError(f"{self.company_name} 平台未返回有效 Excel 文件：{file_bytes[:120]!r}")

        # 导入旧版 xls 读取模块。
        import xlrd

        # 打开内存 xls 工作簿。
        workbook = xlrd.open_workbook(file_contents=file_bytes, ignore_workbook_corruption=True)

        # 判断工作簿是否包含工作表。
        if workbook.nsheets < 1:

            # 抛出空工作簿异常。
            raise RuntimeError(f"{self.company_name} 平台导出 Excel 未包含工作表")

        # 读取首个工作表。
        worksheet = workbook.sheet_by_index(0)

        # 判断工作表是否包含有效单元格区域。
        if worksheet.nrows < 1 or worksheet.ncols < 1:

            # 抛出空工作表异常。
            raise RuntimeError(f"{self.company_name} 平台导出 Excel 首个工作表为空")

        # 读取首行数据触发 OLE 数据流完整性校验。
        worksheet.row_values(0)

        # 返回空结果。
        return None

    # 定义平台 xls 保存方法。
    def save_platform_xls(self, file_bytes, data_type, start_date, end_date, output_dir):

        """
        保存平台 xls 文件。

        :param file_bytes: 文件字节。
        :type file_bytes: bytes
        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: xls 文件路径。
        :rtype: str
        """

        # 校验 Excel 复合文档头。
        if not file_bytes.startswith(b"\xd0\xcf\x11\xe0"):

            # 抛出导出异常。
            raise RuntimeError(f"{self.company_name} 平台未返回有效 Excel 文件：{file_bytes[:120]!r}")

        # 解析输出目录。
        target_dir = output_dir or os.getcwd()

        # 创建输出目录。
        os.makedirs(target_dir, exist_ok=True)

        # 构建 xlsx 文件名。
        xlsx_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 xls 文件名。
        xls_name = os.path.splitext(xlsx_name)[0] + ".xls"

        # 构建 xls 文件路径。
        xls_path = os.path.join(target_dir, xls_name)

        # 判断旧 xls 文件是否存在。
        if os.path.exists(xls_path):

            # 删除旧 xls 文件。
            os.remove(xls_path)

        # 打开 xls 输出文件。
        with open(xls_path, "wb") as file_obj:

            # 写入 xls 文件字节。
            file_obj.write(file_bytes)

        # 返回 xls 文件路径。
        return xls_path

    # 定义 xls 转 xlsx 方法。
    def convert_xls_to_xlsx(self, xls_path):

        """
        将平台 xls 文件转换为 xlsx 文件。

        :param xls_path: xls 文件路径。
        :type xls_path: str
        :return: xlsx 文件路径。
        :rtype: str
        """

        # 构建 xlsx 文件路径。
        xlsx_path = os.path.splitext(xls_path)[0] + ".xlsx"

        # 构建解析器转换标记异常。
        source_error = RuntimeError("平台 xls 使用解析器转换")

        # 使用 xls 解析器转换文件。
        return self.convert_xls_to_xlsx_with_xlrd(xls_path, xlsx_path, source_error)

    # 定义 xls 解析兜底转换方法。
    def convert_xls_to_xlsx_with_xlrd(self, xls_path, xlsx_path, source_error):

        """
        使用 xlrd 解析 xls 并重写为 xlsx。

        :param xls_path: xls 文件路径。
        :type xls_path: str
        :param xlsx_path: xlsx 文件路径。
        :type xlsx_path: str
        :param source_error: 原始转换异常。
        :type source_error: Exception
        :return: xlsx 文件路径。
        :rtype: str
        """

        # 捕获解析转换异常。
        try:

            # 导入旧版 xls 读取模块。
            import xlrd

            # 导入 xlsx 工作簿对象。
            from openpyxl import Workbook

            # 打开旧版 xls 工作簿。
            source_workbook = xlrd.open_workbook(xls_path, ignore_workbook_corruption=True)

            # 创建目标 xlsx 工作簿。
            target_workbook = Workbook(write_only=False)

            # 读取默认工作表。
            default_sheet = target_workbook.active

            # 删除默认工作表。
            target_workbook.remove(default_sheet)

            # 遍历源工作表。
            for sheet_index in range(source_workbook.nsheets):

                # 读取源工作表。
                source_sheet = source_workbook.sheet_by_index(sheet_index)

                # 创建目标工作表。
                target_sheet = target_workbook.create_sheet(source_sheet.name[:31] or f"Sheet{sheet_index + 1}")

                # 遍历源工作表行。
                for row_index in range(source_sheet.nrows):

                    # 写入当前行数据。
                    target_sheet.append(source_sheet.row_values(row_index))

            # 判断旧 xlsx 文件是否存在。
            if os.path.exists(xlsx_path):

                # 删除旧 xlsx 文件。
                os.remove(xlsx_path)

            # 保存目标工作簿。
            target_workbook.save(xlsx_path)

            # 关闭目标工作簿。
            target_workbook.close()

        # 处理解析转换异常。
        except Exception as error:

            # 抛出转换异常。
            raise RuntimeError(f"{self.company_name} 平台 Excel 转换 xlsx 失败：{source_error}；兜底转换失败：{error}") from error

        # 返回 xlsx 文件路径。
        return xlsx_path

    # 定义 Excel 行读取方法。
    def read_rows_from_excel(self, excel_path):

        """
        读取 Excel 数据行。

        :param excel_path: Excel 文件路径。
        :type excel_path: str
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 导入工作簿加载函数。
        from openpyxl import load_workbook

        # 加载工作簿。
        workbook = load_workbook(excel_path, read_only=True, data_only=True)

        # 读取活动工作表。
        worksheet = workbook.active

        # 读取全部行。
        raw_rows = list(worksheet.iter_rows(values_only=True))

        # 关闭工作簿。
        workbook.close()

        # 判断原始行是否为空。
        if not raw_rows:

            # 返回空列表。
            return []

        # 构建表头。
        headers = [str(value or f"字段{index}").strip() for index, value in enumerate(raw_rows[0], start=1)]

        # 初始化结果行列表。
        rows = []

        # 遍历原始数据行。
        for raw_row in raw_rows[1:]:

            # 构建数据行。
            row_data = {headers[index]: raw_row[index] if index < len(raw_row) else "" for index in range(len(headers))}

            # 判断数据行是否存在有效值。
            if any(value not in ("", None) for value in row_data.values()):

                # 追加数据行。
                rows.append(row_data)

        # 返回数据行。
        return rows

    # 定义页面查询方法。
    def query_webforms_page(self, data_type, start_date, end_date):

        """
        查询 WebForms 业务页面。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果结构。
        :rtype: dict
        """

        # 校验数据类型。
        self.validate_data_type(data_type)

        # 读取页面配置。
        page_config = self.resolve_page_config(data_type)

        # 构建页面地址。
        page_url = self.build_page_url(page_config)

        # 打印页面加载日志。
        self.ph.info(self.ph.join(self.ph.kv("加载业务页", page_config["页面名称"], "cyan"), self.ph.kv("端点", page_config["路径"], "yellow")))

        # 请求业务页面。
        page_result = self.http.get(page_url, referer=self.login_page_url)

        # 解码业务页面。
        page_text = self.decode_text(page_result)

        # 解析表单字段。
        payload = self.parse_form_fields(page_text)

        # 写入查询按钮。
        payload = self.apply_submit_button(payload, page_text, "查看")

        # 写入日期字段。
        payload = self.apply_date_fields(payload, start_date, end_date)

        # 写入 DevExpress 运行时字段。
        payload = self.apply_devexpress_runtime_fields(payload)

        # 写入 DevExpress 提交状态字段。
        payload = self.apply_devexpress_submit_state_fields(payload, page_text)

        # 打印查询提交日志。
        self.ph.info(self.ph.join(self.ph.kv("提交查询", data_type, "cyan"), self.ph.kv("日期", f"{start_date} 至 {end_date}", "yellow")))

        # 提交 WebForms 表单。
        query_result = self.post_webforms(page_url, payload, page_url)

        # 解码查询响应。
        query_text = self.decode_text(query_result)

        # 解析查询表格。
        rows, headers = self.parse_result_rows(query_text)

        # 打印解析结果日志。
        self.ph.success(self.ph.join(self.ph.kv("解析完成", f"{len(rows)} 条", "green"), self.ph.kv("表头", len(headers), "cyan")))

        # 返回查询结果。
        return {"rows": rows, "headers": headers}

    # 定义登录 Cookie 写入方法。
    def install_login_cookies(self):

        """
        写入洪兴流向系统登录 Cookie。

        :return: None。
        :rtype: None
        """

        # 解析站点主机。
        host_name = urlparse(self.base_url).hostname or "117.40.251.4"

        # 构建登录 Cookie 字段。
        cookie_values = {"bxnwebUserCode": self.clean_login_name().upper(), "bxnwebPassword": self.clean_login_password()}

        # 遍历登录 Cookie 字段。
        for cookie_name, cookie_value in cookie_values.items():

            # 构建 Cookie 对象。
            cookie_obj = Cookie(version=0, name=cookie_name, value=cookie_value, port=None, port_specified=False, domain=host_name, domain_specified=False, domain_initial_dot=False, path="/", path_specified=True, secure=False, expires=None, discard=True, comment=None, comment_url=None, rest={}, rfc2109=False)

            # 写入 Cookie 容器。
            self.http.cookie_jar.set_cookie(cookie_obj)

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie。

        :param output_dir: Cookie 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建 Cookie 输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件名。
        file_name = f"{self.company_name}_{self.account.get('平台')}_cookie.json"

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, file_name)

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义表单提交方法。
    def post_webforms(self, url, payload, referer):

        """
        提交 WebForms 表单。

        :param url: 请求地址。
        :type url: str
        :param payload: 表单字段。
        :type payload: dict
        :param referer: 来源页面。
        :type referer: str
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 构建请求头。
        headers = self.build_webforms_headers(referer)

        # 编码请求体。
        data = urlencode(payload).encode("utf-8")

        # 发送表单请求。
        return self.http.request("POST", url, headers=headers, data=data, timeout=90)

    # 定义请求头构建方法。
    def build_webforms_headers(self, referer):

        """
        构建 WebForms 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回浏览器表单请求头。
        return {"Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3;q=0.7", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Cache-Control": "max-age=0", "Connection": "keep-alive", "Content-Type": "application/x-www-form-urlencoded", "Origin": self.base_url, "Referer": referer, "Upgrade-Insecure-Requests": "1", "User-Agent": self.http.user_agent()}

    # 定义表单字段解析方法。
    def parse_form_fields(self, html_text):

        """
        解析 HTML 表单字段。

        :param html_text: HTML 页面文本。
        :type html_text: str
        :return: 表单字段字典。
        :rtype: dict
        """

        # 初始化表单解析器。
        parser = HxlxFormParser()

        # 输入 HTML 页面文本。
        parser.feed(html_text)

        # 构建表单字段。
        fields = parser.build_fields()

        # 校验页面状态字段。
        if "__VIEWSTATE" not in fields:

            # 抛出页面状态异常。
            raise RuntimeError(f"{self.company_name} 页面未返回 __VIEWSTATE")

        # 返回表单字段。
        return fields

    # 定义查询按钮写入方法。
    def apply_submit_button(self, payload, html_text, button_text="下载"):

        """
        写入提交按钮字段。

        :param payload: 表单字段。
        :type payload: dict
        :param html_text: HTML 页面文本。
        :type html_text: str
        :param button_text: 按钮文本。
        :type button_text: str
        :return: 更新后的表单字段。
        :rtype: dict
        """

        # 初始化表单解析器。
        parser = HxlxFormParser()

        # 输入 HTML 页面文本。
        parser.feed(html_text)

        # 复制表单字段。
        result = dict(payload)

        # 移除浏览器不会自动提交的按钮字段。
        for parser_button_name in parser.buttons:

            # 删除非目标按钮残留值。
            result.pop(parser_button_name, None)

        # 查找目标按钮名称。
        button_name = self.find_button_name(parser.buttons, button_text)

        # 判断目标按钮名称是否存在。
        if button_name:

            # 写入目标按钮值。
            result[button_name] = parser.buttons.get(button_name, button_text) or button_text

        # 返回表单字段。
        return result

    # 定义按钮查找方法。
    def find_button_name(self, buttons, button_text):

        """
        查找指定文本按钮名称。

        :param buttons: 按钮字段字典。
        :type buttons: dict
        :param button_text: 按钮文本。
        :type button_text: str
        :return: 按钮字段名称。
        :rtype: str
        """

        # 遍历按钮字段。
        for button_name, button_value in buttons.items():

            # 判断按钮文本是否匹配。
            if str(button_value).strip() == button_text:

                # 返回按钮字段名称。
                return button_name

        # 判断目标按钮是否为查看按钮。
        if button_text == "查看":

            # 返回默认查看按钮名称。
            return "ASPxRoundPanel1$Button1"

        # 返回默认下载按钮名称。
        return "ASPxRoundPanel1$Button2"

    # 定义日期字段写入方法。
    def apply_date_fields(self, payload, start_date, end_date):

        """
        写入页面查询日期字段。

        :param payload: 表单字段。
        :type payload: dict
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 更新后的表单字段。
        :rtype: dict
        """

        # 复制表单字段。
        result = dict(payload)

        # 解析日期字段名称。
        date_field_names = self.resolve_date_field_names(result)

        # 判断页面是否没有日期字段。
        if not date_field_names:

            # 返回原始表单字段。
            return result

        # 判断日期字段数量是否不足。
        if len(date_field_names) < 2:

            # 抛出日期字段异常。
            raise RuntimeError(f"{self.company_name} 未识别到两个日期输入字段：{date_field_names}")

        # 写入开始日期。
        result[date_field_names[0]] = self.format_page_date(start_date)

        # 写入结束日期。
        result[date_field_names[1]] = self.format_page_date(end_date)

        # 返回表单字段。
        return result

    # 定义日期字段名称解析方法。
    def resolve_date_field_names(self, payload):

        """
        解析页面日期输入字段名称。

        :param payload: 表单字段。
        :type payload: dict
        :return: 日期输入字段名称列表。
        :rtype: list[str]
        """

        # 初始化日期字段列表。
        date_fields = []

        # 遍历表单字段。
        for field_name, field_value in payload.items():

            # 跳过 DevExpress 日期内部字段。
            if "$DDD" in field_name or field_name.endswith("$State") or field_name.endswith("$DDDSC"):

                # 继续处理下一个字段。
                continue

            # 判断字段是否为日期输入。
            if "DateEdit" in field_name:

                # 追加日期字段名称。
                date_fields.append(field_name)

        # 返回排序后的日期字段。
        return sorted(date_fields)

    # 定义 DevExpress 运行时字段写入方法。
    def apply_devexpress_runtime_fields(self, payload, grid_loaded=False):

        """
        写入 DevExpress 运行时表单字段。

        :param payload: 表单字段。
        :type payload: dict
        :param grid_loaded: GridView 是否已经完成查询加载。
        :type grid_loaded: bool
        :return: 更新后的表单字段。
        :rtype: dict
        """

        # 复制表单字段。
        result = dict(payload)

        # 选择脚本资源状态字段。
        script_state = self.DEVEXPRESS_GRID_SCRIPT_STATE if grid_loaded else self.DEVEXPRESS_SCRIPT_STATE

        # 写入脚本资源状态字段。
        result.setdefault("DXScript", script_state)

        # 写入样式资源状态字段。
        result.setdefault("DXCss", self.DEVEXPRESS_STYLE_STATE)

        # 返回表单字段。
        return result

    # 定义 DevExpress 提交状态字段写入方法。
    def apply_devexpress_submit_state_fields(self, payload, html_text):

        """
        写入 DevExpress 控件提交状态字段。

        :param payload: 表单字段。
        :type payload: dict
        :param html_text: HTML 页面文本。
        :type html_text: str
        :return: 更新后的表单字段。
        :rtype: dict
        """

        # 复制表单字段。
        result = dict(payload)

        # 遍历日期字段。
        for date_field_name in self.resolve_date_field_names(result):

            # 写入日期控件状态字段。
            result = self.apply_devexpress_date_submit_state(result, date_field_name)

        # 遍历筛选日期字段。
        for filter_field_name in self.resolve_filter_date_field_names(html_text):

            # 写入筛选日期控件状态字段。
            result = self.apply_devexpress_filter_date_state(result, filter_field_name)

        # 解析 GridView 控件状态。
        grid_state = self.extract_grid_state_object(html_text)

        # 判断 GridView 控件状态是否存在。
        if grid_state:

            # 写入 GridView 控件状态字段。
            result[grid_state["字段名"]] = grid_state["字段值"]

        # 返回表单字段。
        return result

    # 定义日期提交状态写入方法。
    def apply_devexpress_date_submit_state(self, payload, field_name):

        """
        写入 DevExpress 日期控件提交状态。

        :param payload: 表单字段。
        :type payload: dict
        :param field_name: 日期字段名称。
        :type field_name: str
        :return: 更新后的表单字段。
        :rtype: dict
        """

        # 复制表单字段。
        result = dict(payload)

        # 读取日期字段值。
        date_value = result.get(field_name)

        # 写入日期原始值状态。
        result[f"{field_name}$State"] = self.build_entity_json({"rawValue": self.build_devexpress_raw_date(date_value)})

        # 写入日期下拉可见日期状态。
        result[f"{field_name}$DDD$C"] = self.build_entity_json({"visibleDate": self.format_devexpress_selected_date(date_value), "selectedDates": [self.format_devexpress_selected_date(date_value)]})

        # 构建弹窗窗口状态。
        popup_state = self.build_entity_json({"windowsState": "0:0:-1:0:0:0:-10000:-10000:1:0:0:0"})

        # 写入日期下拉窗口状态。
        result[f"{field_name}$DDDState"] = popup_state

        # 写入日期下拉浮层状态。
        result[f"{field_name}$DDD$C$FNPState"] = popup_state

        # 返回表单字段。
        return result

    # 定义筛选日期字段名称解析方法。
    def resolve_filter_date_field_names(self, html_text):

        """
        解析 GridView 筛选日期字段名称。

        :param html_text: HTML 页面文本。
        :type html_text: str
        :return: 筛选日期字段名称列表。
        :rtype: list[str]
        """

        # 初始化字段名称列表。
        field_names = []

        # 遍历页面日期控件初始化脚本。
        for match_obj in re.finditer(r"new ASPxClientDateEdit\('([^']+)'\)", html_text):

            # 读取客户端控件名称。
            client_name = match_obj.group(1)

            # 判断控件是否为 GridView 筛选日期控件。
            if "DXFREditorcol" in client_name:

                # 转换为服务端表单字段名称。
                field_name = client_name.replace("_", "$")

                # 追加字段名称。
                field_names.append(field_name)

        # 返回日期筛选字段名称列表。
        return field_names

    # 定义筛选日期状态写入方法。
    def apply_devexpress_filter_date_state(self, payload, field_name):

        """
        写入 GridView 筛选日期控件状态。

        :param payload: 表单字段。
        :type payload: dict
        :param field_name: 筛选字段名称。
        :type field_name: str
        :return: 更新后的表单字段。
        :rtype: dict
        """

        # 复制表单字段。
        result = dict(payload)

        # 构建空日期原始值状态。
        result[f"{field_name}$State"] = self.build_entity_json({"rawValue": "N"})

        # 写入日期下拉可见日期状态。
        result[f"{field_name}$DDD$C"] = self.build_entity_json({"visibleDate": self.format_devexpress_visible_date()})

        # 构建弹窗窗口状态。
        popup_state = self.build_entity_json({"windowsState": "0:0:-1:0:0:0:-10000:-10000:1:0:0:0"})

        # 写入日期下拉窗口状态。
        result[f"{field_name}$DDDState"] = popup_state

        # 写入日期下拉浮层状态。
        result[f"{field_name}$DDD$C$FNPState"] = popup_state

        # 返回表单字段。
        return result

    # 定义 GridView 状态解析方法。
    def extract_grid_state_object(self, html_text):

        """
        解析 DevExpress GridView 控件状态。

        :param html_text: HTML 页面文本。
        :type html_text: str
        :return: GridView 控件状态。
        :rtype: dict
        """

        # 匹配 GridView 初始化脚本块。
        grid_match = re.search(r"new ASPxClientGridView\('([^']+)'\);(?P<body>.*?)dxo\.AfterCreate\(\);", html_text, re.S)

        # 判断 GridView 初始化脚本块是否存在。
        if not grid_match:

            # 返回空状态。
            return {}

        # 读取 GridView 初始化脚本内容。
        grid_body = grid_match.group("body")

        # 匹配 GridView 唯一字段名称。
        unique_match = re.search(r"dxo\.uniqueID\s*=\s*'([^']+)';", grid_body)

        # 匹配 GridView 回调状态。
        callback_match = re.search(r"dxo\.stateObject\s*=\s*\(\{'keys':\[\],'callbackState':'([^']*)','selection':'([^']*)'\}\);", grid_body)

        # 判断必要状态是否存在。
        if not unique_match or not callback_match:

            # 返回空状态。
            return {}

        # 读取 GridView 唯一字段名称。
        field_name = unique_match.group(1)

        # 读取 GridView 回调状态。
        callback_state = callback_match.group(1)

        # 读取 GridView 选择状态。
        selection_state = callback_match.group(2)

        # 构建 GridView 状态值。
        field_value = self.build_entity_json({"keys": [], "callbackState": callback_state, "selection": selection_state})

        # 返回 GridView 控件状态。
        return {"字段名": field_name, "字段值": field_value}

    # 定义实体 JSON 构建方法。
    def build_entity_json(self, data):

        """
        构建 DevExpress 表单实体 JSON 文本。

        :param data: 状态数据。
        :type data: dict
        :return: 实体 JSON 文本。
        :rtype: str
        """

        # 序列化紧凑 JSON 文本。
        json_text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))

        # 返回双引号实体化文本。
        return json_text.replace('"', "&quot;")

    # 定义 DevExpress 可见日期格式化方法。
    def format_devexpress_visible_date(self):

        """
        格式化 DevExpress 日历可见日期。

        :return: 日历可见日期。
        :rtype: str
        """

        # 读取当前日期对象。
        date_obj = datetime.now()

        # 返回月份日期年份格式。
        return f"{date_obj.month:02d}/{date_obj.day:02d}/{date_obj.year}"

    # 定义日期状态同步方法。
    def apply_date_state_fields(self, payload, field_name, date_value):

        """
        同步 DevExpress 日期状态字段。

        :param payload: 表单字段。
        :type payload: dict
        :param field_name: 日期字段名称。
        :type field_name: str
        :param date_value: 日期文本。
        :type date_value: str | None
        :return: 更新后的表单字段。
        :rtype: dict
        """

        # 复制表单字段。
        result = dict(payload)

        # 构建状态字段名称。
        state_name = f"{field_name}$State"

        # 写入日期状态字段。
        result[state_name] = json.dumps({"rawValue": self.build_devexpress_raw_date(date_value)}, separators=(",", ":"))

        # 构建日期选择字段名称。
        selection_name = f"{field_name}$DDDSC"

        # 写入日期选择字段。
        result[selection_name] = json.dumps({"visibleDate": self.format_devexpress_selected_date(date_value), "selectedDates": [self.format_devexpress_selected_date(date_value)]}, separators=(",", ":"))

        # 返回表单字段。
        return result

    # 定义查询结果解析方法。
    def parse_result_rows(self, html_text):

        """
        解析查询响应 HTML 表格。

        :param html_text: 查询响应 HTML。
        :type html_text: str
        :return: 查询结果行和表头。
        :rtype: tuple[list[dict], list[str]]
        """

        # 初始化表格解析器。
        parser = HxlxTableParser()

        # 输入 HTML 页面文本。
        parser.feed(html_text)

        # 选择业务数据表格。
        table_rows = self.select_data_table(parser.tables)

        # 判断业务表格是否为空。
        if not table_rows:

            # 返回空结果。
            return [], []

        # 解析表头。
        headers = self.resolve_table_headers(table_rows)

        # 解析数据起始下标。
        data_start_index = self.resolve_data_start_index(table_rows, headers)

        # 初始化结果行列表。
        rows = []

        # 遍历表格数据行。
        for raw_row in table_rows[data_start_index:]:

            # 判断当前行是否为数据行。
            if not self.is_data_row(raw_row, headers):

                # 继续处理下一行。
                continue

            # 构建行数据。
            row_data = self.build_row_data(headers, raw_row)

            # 判断行数据是否存在有效值。
            if any(value for value in row_data.values()):

                # 追加结果行。
                rows.append(row_data)

        # 返回查询结果。
        return rows, headers

    # 定义业务表格选择方法。
    def select_data_table(self, tables):

        """
        选择业务数据表格。

        :param tables: 页面表格集合。
        :type tables: list[list[list[str]]]
        :return: 业务数据表格。
        :rtype: list[list[str]]
        """

        # 初始化候选表格列表。
        candidates = []

        # 遍历页面表格。
        for table_rows in tables:

            # 计算表格最大列数。
            max_column_count = max([len(row) for row in table_rows] or [0])

            # 判断表格是否具备业务列数。
            if max_column_count >= 5:

                # 追加候选表格。
                candidates.append(table_rows)

        # 判断候选表格是否存在。
        if candidates:

            # 返回规模最大的候选表格。
            return max(candidates, key=lambda table_rows: len(table_rows) * max([len(row) for row in table_rows] or [0]))

        # 返回空表格。
        return []

    # 定义表头解析方法。
    def resolve_table_headers(self, table_rows):

        """
        解析业务表格表头。

        :param table_rows: 表格行。
        :type table_rows: list[list[str]]
        :return: 表头列表。
        :rtype: list[str]
        """

        # 遍历全部表格行。
        for row in table_rows:

            # 判断当前行是否为表头。
            if self.is_header_row(row):

                # 返回去重后的表头。
                return self.deduplicate_headers(row)

        # 读取兜底表头行。
        fallback_row = table_rows[0] if table_rows else []

        # 返回兜底表头。
        return self.deduplicate_headers([value or f"字段{index}" for index, value in enumerate(fallback_row, start=1)])

    # 定义数据起始下标解析方法。
    def resolve_data_start_index(self, table_rows, headers):

        """
        解析业务数据起始下标。

        :param table_rows: 表格行。
        :type table_rows: list[list[str]]
        :param headers: 表头列表。
        :type headers: list[str]
        :return: 数据起始下标。
        :rtype: int
        """

        # 遍历表格行。
        for row_index, row in enumerate(table_rows):

            # 判断当前行是否为表头行。
            if self.is_header_row(row) or self.rows_match_headers(row, headers):

                # 返回表头下一行下标。
                return row_index + 1

        # 返回默认起始下标。
        return 1

    # 定义表头行判断方法。
    def is_header_row(self, row):

        """
        判断表格行是否为表头行。

        :param row: 表格行。
        :type row: list[str]
        :return: 是否为表头行。
        :rtype: bool
        """

        # 定义表头关键词。
        header_keywords = ["单号", "日期", "商品", "品名", "批号", "数量", "库存", "门店", "供应商", "结算"]

        # 统计关键词命中数量。
        match_count = sum(1 for cell in row if any(keyword in cell for keyword in header_keywords))

        # 返回表头判断结果。
        return match_count >= 2

    # 定义表头匹配方法。
    def rows_match_headers(self, row, headers):

        """
        判断表格行是否匹配表头。

        :param row: 表格行。
        :type row: list[str]
        :param headers: 表头列表。
        :type headers: list[str]
        :return: 是否匹配表头。
        :rtype: bool
        """

        # 返回表头匹配结果。
        return [str(item).strip() for item in row[:len(headers)]] == [str(item).strip() for item in headers]

    # 定义数据行判断方法。
    def is_data_row(self, row, headers):

        """
        判断表格行是否为业务数据行。

        :param row: 表格行。
        :type row: list[str]
        :param headers: 表头列表。
        :type headers: list[str]
        :return: 是否为业务数据行。
        :rtype: bool
        """

        # 判断当前行列数是否不足。
        if len(row) < max(2, min(len(headers), 3)):

            # 返回非数据行状态。
            return False

        # 判断当前行是否匹配表头。
        if self.rows_match_headers(row, headers):

            # 返回非数据行状态。
            return False

        # 返回数据行状态。
        return True

    # 定义行数据构建方法。
    def build_row_data(self, headers, row):

        """
        构建标准行数据。

        :param headers: 表头列表。
        :type headers: list[str]
        :param row: 表格行。
        :type row: list[str]
        :return: 标准行数据。
        :rtype: dict
        """

        # 初始化行数据字典。
        row_data = {}

        # 遍历表头字段。
        for column_index, header_name in enumerate(headers):

            # 读取当前单元格值。
            cell_value = row[column_index] if column_index < len(row) else ""

            # 写入当前字段值。
            row_data[header_name] = cell_value

        # 返回行数据。
        return row_data

    # 定义表头去重方法。
    def deduplicate_headers(self, headers):

        """
        去重并补齐表头。

        :param headers: 原始表头列表。
        :type headers: list[str]
        :return: 标准表头列表。
        :rtype: list[str]
        """

        # 初始化表头计数字典。
        header_counts = {}

        # 初始化结果表头列表。
        result = []

        # 遍历原始表头。
        for column_index, header_name in enumerate(headers, start=1):

            # 清理表头文本。
            clean_name = str(header_name).strip() or f"字段{column_index}"

            # 累加表头出现次数。
            header_counts[clean_name] = header_counts.get(clean_name, 0) + 1

            # 判断表头是否重复。
            if header_counts[clean_name] > 1:

                # 追加带序号表头。
                result.append(f"{clean_name}_{header_counts[clean_name]}")

            # 处理首次出现表头。
            else:

                # 追加原始表头。
                result.append(clean_name)

        # 返回标准表头。
        return result

    # 定义页面配置解析方法。
    def resolve_page_config(self, data_type):

        """
        解析数据类型页面配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 页面配置。
        :rtype: dict
        """

        # 读取页面映射。
        page_map = self.account.get("页面映射", {})

        # 判断页面映射是否存在。
        if data_type not in page_map:

            # 抛出页面配置异常。
            raise ValueError(f"洪兴流向系统暂未配置数据类型：{data_type}")

        # 返回页面配置。
        return page_map[data_type]

    # 定义页面地址构建方法。
    def build_page_url(self, page_config):

        """
        构建业务页面地址。

        :param page_config: 页面配置。
        :type page_config: dict
        :return: 业务页面地址。
        :rtype: str
        """

        # 读取页面路径。
        page_path = page_config.get("路径", "")

        # 判断页面路径是否为完整地址。
        if page_path.startswith("http://") or page_path.startswith("https://"):

            # 返回完整页面地址。
            return page_path

        # 返回拼接后的页面地址。
        return f"{self.base_url}{page_path}"

    # 定义入口地址构建方法。
    def build_login_entry_url(self):

        """
        构建登录入口地址。

        :return: 登录入口地址。
        :rtype: str
        """

        # 判断登录页是否已包含用户编码。
        if "usercode=" in self.login_page_url.lower():

            # 返回登录页地址。
            return self.login_page_url

        # 返回带用户编码的入口地址。
        return f"{self.login_page_url}?usercode={quote(self.clean_login_name().upper())}"



    # 定义密码读取方法。
    def clean_login_password(self):

        """
        读取登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回登录密码文本。
        return str(self.account.get("密码", ""))

    # 定义响应文本解码方法。
    def decode_text(self, result):

        """
        解码 HTTP 响应文本。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 响应文本。
        :rtype: str
        """

        # 返回 UTF-8 文本。
        return result.body.decode("utf-8", errors="replace")

    # 定义页面日期格式化方法。
    def format_page_date(self, date_value):

        """
        格式化页面日期。

        :param date_value: 日期文本。
        :type date_value: str | None
        :return: 页面日期文本。
        :rtype: str
        """

        # 解析日期对象。
        date_obj = self.parse_date(date_value)

        # 返回页面日期格式。
        return f"{date_obj.year}/{date_obj.month}/{date_obj.day}"

    # 定义日期解析方法。
    def parse_date(self, date_value):

        """
        解析日期文本。

        :param date_value: 日期文本。
        :type date_value: str | None
        :return: 日期对象。
        :rtype: datetime
        """

        # 判断日期文本是否为空。
        if not date_value:

            # 返回当前日期。
            return datetime.now()

        # 遍历支持的日期格式。
        for date_format in ("%Y-%m-%d", "%Y/%m/%d"):

            # 捕获日期解析异常。
            try:

                # 返回解析后的日期。
                return datetime.strptime(str(date_value), date_format)

            # 处理日期解析异常。
            except ValueError:

                # 继续尝试下一种格式。
                continue

        # 抛出日期格式异常。
        raise ValueError(f"无法解析日期：{date_value}")

    # 定义日期特征判断方法。
    def looks_like_date(self, value):

        """
        判断文本是否接近日期。

        :param value: 原始文本。
        :type value: object
        :return: 是否接近日期。
        :rtype: bool
        """

        # 转换文本。
        text = str(value or "").strip()

        # 返回日期特征判断结果。
        return "/" in text and any(char.isdigit() for char in text)

    # 定义 DevExpress 日期值构建方法。
    def build_devexpress_raw_date(self, date_value):

        """
        构建 DevExpress 日期毫秒值。

        :param date_value: 日期文本。
        :type date_value: str | None
        :return: 日期毫秒值。
        :rtype: str
        """

        # 解析日期对象。
        date_obj = self.parse_date(date_value)

        # 构建 UTC 日期对象。
        utc_date_obj = datetime(date_obj.year, date_obj.month, date_obj.day, tzinfo=timezone.utc)

        # 转换为 UTC 毫秒时间戳。
        timestamp_value = int(utc_date_obj.timestamp() * 1000)

        # 返回毫秒值文本。
        return str(timestamp_value)

    # 定义 DevExpress 选择日期格式化方法。
    def format_devexpress_selected_date(self, date_value):

        """
        格式化 DevExpress 选择日期。

        :param date_value: 日期文本。
        :type date_value: str | None
        :return: DevExpress 选择日期。
        :rtype: str
        """

        # 解析日期对象。
        date_obj = self.parse_date(date_value)

        # 返回月份日期年份格式。
        return f"{date_obj.month:02d}/{date_obj.day:02d}/{date_obj.year}"

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验洪兴流向系统数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 读取支持的数据类型列表。
        supported_types = self.account.get("数据类型列表", [])

        # 判断当前数据类型是否支持。
        if data_type not in supported_types:

            # 抛出数据类型异常。
            raise ValueError(f"洪兴流向系统暂未接入数据类型：{data_type}")


    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头列表。
        :type headers: list[str]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列字母工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 判断表头是否为空。
        if not headers:

            # 写入空结果表头。
            headers = ["查询结果"]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果。
        for row in rows:

            # 写入当前数据行。
            worksheet.append([row.get(header_name, "") for header_name in headers])

        # 遍历表头设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取列字母。
            column_letter = get_column_letter(column_index)

            # 设置列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 8))

        # 冻结表头。
        worksheet.freeze_panes = "A2"

        # 设置自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 判断同名文件是否存在。
        if os.path.exists(output_path):

            # 删除同名文件。
            os.remove(output_path)

        # 保存 Excel 文件。
        workbook.save(output_path)

        # 关闭工作簿。
        workbook.close()
