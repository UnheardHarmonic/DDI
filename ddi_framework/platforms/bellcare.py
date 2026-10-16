# -*- coding: utf-8 -*-
"""
贝尔康流向采集服务平台协议适配器。

本模块通过密码加密端点、登录接口和带 accesskey 的分页查询接口完成采购、销售和库存数据采集，并将标准字段与全部原始字段写入本地 Excel。

:return: 模块提供贝尔康流向采集服务平台协议适配器。
:rtype: None
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期时间处理对象。
from datetime import datetime

# 导入 URL 查询字符串编码函数。
from urllib.parse import urlencode

# 导入 Excel 工作簿对象。
from openpyxl import Workbook

# 导入 Excel 列标识转换函数。
from openpyxl.utils import get_column_letter

# 导入平台适配器基类。
from ddi_framework import BasePlatform

# 导入 Excel 统一样式函数。
from ddi_framework import apply_standard_excel_style

# 导入公共日志打印器。
from ddi_framework import PrintHandler

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient


# 定义数据类型与查询页面、接口及筛选字段的协议映射。
DATA_TYPE_CONFIG = {
    "购进流向": {
        "page_path": "/views/supply/flow/buy/buy",
        "api_path": "/views/supply/flow/buy/buy_api",
        "filters": {"province": "", "cust": "", "custmc": "", "goods": "", "batch": ""},
    },
    "配送流向": {
        "page_path": "/views/supply/flow/sale/sale",
        "api_path": "/views/supply/flow/sale/sale_api",
        "filters": {"state": "", "province": "", "cust": "", "custmc": "", "agent": "", "goods": "", "batch": "", "sh_state": ""},
    },
    "库存": {
        "page_path": "/views/supply/flow/stock/stock",
        "api_path": "/views/supply/flow/stock/stock_api",
        "filters": {"province": "", "cust": "", "custmc": "", "goods": "", "batch": ""},
    },
}

# 定义各数据类型的标准字段与服务端原始字段映射。
STANDARD_FIELD_MAP = {
    "购进流向": [("省份", "province"), ("购进日期", "dayx"), ("客户名称", "cust"), ("供应商", "agent"), ("商品名称", "goods"), ("规格", "spec"), ("批号", "batch"), ("数量", "num"), ("单价", "price"), ("金额", "money"), ("匹配状态", "pp_state")],
    "配送流向": [("省份", "province"), ("流向日期", "dayx"), ("客户名称", "cust"), ("配送单位", "agent"), ("终端名称", "agent1"), ("终端等级", "zdgrade"), ("商品名称", "goods"), ("规格", "spec"), ("批号", "batch"), ("数量", "num"), ("单价", "price"), ("金额", "money"), ("匹配状态", "pp_state")],
    "库存": [("省份", "province"), ("库存日期", "dayx"), ("客户名称", "cust"), ("终端名称", "agent1"), ("商品名称", "goods"), ("规格", "spec"), ("批号", "batch"), ("库存数量", "num"), ("单价", "price"), ("金额", "money")],
}

# 定义常见原始字段的可读标题映射。
RAW_FIELD_TITLES = {
    "province": "省份",
    "dayx": "业务日期",
    "cust": "客户名称",
    "agent": "供应商/配送单位",
    "agent1": "终端名称",
    "zdgrade": "终端等级",
    "goods": "商品名称",
    "spec": "规格",
    "batch": "批号",
    "num": "数量",
    "price": "单价",
    "money": "金额",
    "pp_state": "匹配状态",
}


# 定义贝尔康流向采集服务平台协议适配器。
class BellcarePlatform(BasePlatform):

    """
    贝尔康流向采集服务平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 贝尔康流向采集服务平台协议适配器实例。
    :rtype: BellcarePlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化带 Cookie 的 HTTP 会话、接口地址和查询状态。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化公共日志打印器。
        self.ph = printer or PrintHandler()

        # 读取平台 HTTP 服务根地址。
        self.base_url = str(account.get("接口根地址", "")).rstrip("/")

        # 读取登录页面地址。
        self.login_page_url = str(account.get("登录页", f"{self.base_url}/views/user/user/1005.html"))

        # 读取密码加密接口地址。
        self.password_encrypt_url = str(account.get("密码加密接口", f"{self.base_url}/management/api/login/md.php"))

        # 读取登录接口地址。
        self.login_api_url = str(account.get("登录接口", f"{self.base_url}/management/api/login/login"))

        # 读取平台单页最大查询数量。
        self.page_size = int(account.get("分页大小", 1000))

        # 初始化带 Cookie 管理的协议客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录状态标识。
        self.logged_in = False

        # 初始化查询 accesskey。
        self.access_key = ""

    # 定义平台登录方法。
    def login(self):

        """
        通过密码加密端点和登录接口建立协议会话。

        :return: 不包含凭据和完整 accesskey 的登录结果摘要。
        :rtype: dict
        """

        # 打印登录阶段分隔线。
        self.ph.section("贝尔康流向采集服务平台登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 初始化登录页面以建立 PHP 会话 Cookie。
        self.http.get(self.login_page_url, referer=self.base_url)

        # 读取配置中的登录密码。
        password = str(self.account.get("密码", ""))

        # 校验登录密码已配置。
        if not password:

            # 抛出缺少密码异常。
            raise RuntimeError("登录失败：未配置登录密码。")

        # 调用站点密码加密接口生成登录所需摘要。
        encrypt_result = self.http.post_form(self.password_encrypt_url, {"d": password}, self.login_page_url)

        # 解析密码加密接口响应文本。
        encrypted_password = encrypt_result.body.decode("utf-8", errors="replace").strip()

        # 校验加密接口返回了可用摘要。
        if len(encrypted_password) != 32:

            # 抛出密码加密失败异常。
            raise RuntimeError("登录失败：密码加密接口未返回有效摘要。")

        # 构建登录接口查询参数。
        login_params = {
            "username": str(self.account.get("账号", "")),
            "password": "",
            "Repassword": encrypted_password,
            "userType": "0",
            "vercode": "DDI",
            "accesskey": "0",
        }

        # 拼接带编码查询参数的登录接口地址。
        login_url = self.build_query_url(self.login_api_url, login_params)

        # 构建 AJAX 登录请求头。
        login_headers = self.http.ajax_form_headers(self.login_page_url)

        # 提交登录 GET 请求。
        login_result = self.http.request("GET", login_url, headers=login_headers, timeout=60)

        # 解析登录接口 JSON 响应。
        login_data = self.decode_json_response("登录", login_result.body)

        # 校验登录接口返回成功状态。
        if str(login_data.get("code", "")) != "0":

            # 读取服务端失败提示。
            message = str(login_data.get("msg", "未知错误"))

            # 抛出不包含敏感值的登录异常。
            raise RuntimeError(f"登录失败：{message}")

        # 读取登录数据对象。
        login_payload = login_data.get("data") or {}

        # 读取用于后续查询的 accesskey。
        self.access_key = str(login_payload.get("accesskey", "")).strip()

        # 校验查询 accesskey 已返回。
        if not self.access_key:

            # 抛出缺少查询会话参数异常。
            raise RuntimeError("登录失败：服务端未返回查询 accesskey。")

        # 标记当前协议会话已登录。
        self.logged_in = True

        # 读取登录身份名称用于账号归属核对。
        account_user_name = str(login_payload.get("username", ""))

        # 读取服务端返回的机构名称用于账号归属核对。
        account_company_name = str(login_payload.get("custname", ""))

        # 打印不含敏感值的登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("登录状态", "成功", "green"), self.ph.kv("Cookie", len(list(self.http.cookie_jar)), "cyan"), self.ph.kv("查询密钥长度", len(self.access_key), "cyan")))

        # 返回不含密码和完整 accesskey 的登录摘要。
        return {"登录状态": "成功", "账号姓名": account_user_name, "服务端机构": account_company_name, "Cookie数量": len(list(self.http.cookie_jar)), "查询密钥长度": len(self.access_key)}

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前协议会话 Cookie 供公共任务服务记录。

        :param output_dir: Cookie 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建 Cookie 输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 输出文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存当前协议会话 Cookie。
        self.http.save_cookie_file(cookie_path)

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义原始数据查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定数据类型的全部原始业务记录。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 原始记录列表。
        :rtype: list[dict]
        """

        # 查询当前数据类型的全部分页数据。
        query_result = self.query_platform_data(data_type, start_date, end_date)

        # 返回未裁剪的原始记录列表。
        return query_result["rows"]

    # 定义查询并生成 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定数据类型并生成标准字段加原始字段 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: Excel 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询当前数据类型的全部分页数据。
        query_result = self.query_platform_data(data_type, start_date, end_date)

        # 解析默认 Excel 输出目录。
        target_directory = output_dir or os.getcwd()

        # 创建 Excel 输出目录。
        os.makedirs(target_directory, exist_ok=True)

        # 构建公共标准 Excel 文件名。
        file_name = self.build_output_file_name(data_type, query_result["start_date"], query_result["end_date"])

        # 拼接 Excel 完整输出路径。
        output_path = os.path.join(target_directory, file_name)

        # 写入标准字段和服务端全部原始字段。
        self.write_rows_to_excel(query_result["rows"], query_result["raw_headers"], data_type, output_path)

        # 打印 Excel 落地结果摘要。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan"), self.ph.kv("原始字段数", len(query_result["raw_headers"]), "cyan")))

        # 返回 Excel 文件路径。
        return output_path


    # 定义平台主查询方法。
    def query_platform_data(self, data_type, start_date, end_date):

        """
        建立登录态、提交带 accesskey 的筛选条件并读取全部分页记录。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询记录、原始表头与实际日期范围。
        :rtype: dict
        """

        # 校验当前数据类型存在协议映射。
        if data_type not in DATA_TYPE_CONFIG:

            # 抛出未支持数据类型异常。
            raise ValueError(f"不支持的数据类型：{data_type}")

        # 确保当前协议会话已登录。
        if not self.logged_in:

            # 执行平台登录。
            self.login()

        # 规范化查询开始日期。
        normalized_start_date = self.normalize_date(start_date, "开始日期")

        # 规范化查询结束日期。
        normalized_end_date = self.normalize_date(end_date, "结束日期")

        # 校验查询日期范围顺序。
        if normalized_start_date > normalized_end_date:

            # 抛出日期范围异常。
            raise ValueError("查询开始日期不能晚于结束日期。")

        # 读取当前数据类型协议配置。
        type_config = DATA_TYPE_CONFIG[data_type]

        # 拼接当前数据类型查询页面地址。
        page_url = f"{self.base_url}{type_config['page_path']}"

        # 拼接当前数据类型查询接口地址。
        api_url = f"{self.base_url}{type_config['api_path']}"

        # 打印当前数据类型查询分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", normalized_start_date, "cyan"), self.ph.kv("结束日期", normalized_end_date, "cyan")))

        # 初始化累计原始记录集合。
        rows = []

        # 初始化原始字段顺序集合。
        raw_headers = []

        # 初始化请求页码。
        page_index = 1

        # 初始化服务端总记录数。
        total_count = None

        # 循环读取全部分页数据。
        while True:

            # 构建当前页公共筛选参数。
            params = {"ac": self.access_key, "page": page_index, "limit": self.page_size, "st": normalized_start_date, "et": normalized_end_date}

            # 追加当前数据类型的空筛选字段。
            params.update(type_config["filters"])

            # 拼接当前页查询地址。
            request_url = self.build_query_url(api_url, params)

            # 构建 AJAX 查询请求头。
            query_headers = self.http.ajax_form_headers(page_url)

            # 提交当前页查询请求。
            query_result = self.http.request("GET", request_url, headers=query_headers, timeout=90)

            # 解析当前页 JSON 响应。
            response_data = self.decode_json_response(f"{data_type}查询", query_result.body)

            # 校验当前页业务响应状态。
            if str(response_data.get("code", "")) != "0":

                # 读取服务端失败提示。
                message = str(response_data.get("msg", "未知错误"))

                # 抛出查询失败异常。
                raise RuntimeError(f"{data_type}查询失败：{message}")

            # 读取当前页数据列表。
            page_rows = response_data.get("data") or []

            # 校验当前页数据结构为列表。
            if not isinstance(page_rows, list):

                # 抛出非预期数据结构异常。
                raise RuntimeError(f"{data_type}查询失败：data 字段不是列表。")

            # 读取服务端总记录数。
            current_total = self.parse_total_count(response_data.get("count"))

            # 保存首次有效总记录数。
            if total_count is None and current_total is not None:

                # 写入服务端总记录数。
                total_count = current_total

            # 遍历当前页记录以维护原始字段顺序。
            for row in page_rows:

                # 跳过非字典型异常记录。
                if not isinstance(row, dict):

                    # 进入下一条记录。
                    continue

                # 遍历当前记录原始字段。
                for field_name in row.keys():

                    # 追加首次出现的原始字段。
                    if field_name not in raw_headers:

                        # 保存原始字段顺序。
                        raw_headers.append(field_name)

            # 追加当前页有效记录。
            rows.extend([row for row in page_rows if isinstance(row, dict)])

            # 打印当前页查询进度。
            self.ph.info(self.ph.join(self.ph.kv("页码", page_index, "cyan"), self.ph.kv("本页记录", len(page_rows), "cyan"), self.ph.kv("累计记录", len(rows), "cyan"), self.ph.kv("服务端总数", total_count if total_count is not None else "未返回", "cyan")))

            # 判断服务端总数已全部读取。
            if total_count is not None and len(rows) >= total_count:

                # 结束分页读取。
                break

            # 判断当前页数据不足一页。
            if len(page_rows) < self.page_size:

                # 结束分页读取。
                break

            # 进入下一页。
            page_index += 1

        # 校验返回记录的业务日期范围。
        self.validate_business_dates(data_type, rows, normalized_start_date, normalized_end_date)

        # 返回查询结果对象。
        return {"rows": rows, "raw_headers": raw_headers, "start_date": normalized_start_date, "end_date": normalized_end_date}

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, raw_headers, data_type, output_path):

        """
        写入标准字段、全部原始字段和统一 Excel 样式。

        :param rows: 原始记录列表。
        :type rows: list[dict]
        :param raw_headers: 服务端原始字段顺序。
        :type raw_headers: list[str]
        :param data_type: 数据类型名称。
        :type data_type: str
        :param output_path: Excel 输出路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 创建新的 Excel 工作簿。
        workbook = Workbook()

        # 获取默认活动工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = data_type[:31]

        # 读取当前数据类型标准字段映射。
        standard_fields = STANDARD_FIELD_MAP[data_type]

        # 构建标准字段标题列表。
        standard_headers = [field_title for field_title, _ in standard_fields]

        # 构建保留服务端字段名的原始字段标题列表。
        raw_header_titles = [f"原始_{field_name}" for field_name in raw_headers]

        # 写入标准字段和全部原始字段表头。
        worksheet.append(standard_headers + raw_header_titles)

        # 遍历每条原始业务记录。
        for row in rows:

            # 构建当前记录的标准字段值。
            standard_values = [self.normalize_excel_value(row.get(field_name, "")) for _, field_name in standard_fields]

            # 构建当前记录的全部原始字段值。
            raw_values = [self.normalize_excel_value(row.get(field_name, "")) for field_name in raw_headers]

            # 写入当前记录完整行。
            worksheet.append(standard_values + raw_values)

        # 遍历所有 Excel 列以调整显示宽度。
        for column_index in range(1, worksheet.max_column + 1):

            # 初始化当前列最大文本长度。
            max_length = 0

            # 遍历当前列写入单元格。
            for row_index in range(1, worksheet.max_row + 1):

                # 读取当前单元格文本值。
                cell_text = str(worksheet.cell(row=row_index, column=column_index).value or "")

                # 更新当前列最大文本长度。
                max_length = max(max_length, len(cell_text))

            # 设置限制后的当前列宽度。
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 36)

        # 保存 Excel 工作簿。
        workbook.save(output_path)

        # 关闭 Excel 工作簿。
        workbook.close()

        # 应用项目统一 Excel 样式。
        apply_standard_excel_style(output_path)

    # 定义查询 URL 构建方法。
    def build_query_url(self, endpoint_url, params):

        """
        对查询参数进行 URL 编码并构建请求地址。

        :param endpoint_url: 不含查询字符串的接口地址。
        :type endpoint_url: str
        :param params: 查询参数字典。
        :type params: dict
        :return: 带查询字符串的接口地址。
        :rtype: str
        """

        # 编码接口查询参数。
        encoded_params = urlencode(params)

        # 返回完整查询地址。
        return f"{endpoint_url}?{encoded_params}"

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, stage_name, body):

        """
        将接口响应字节解析为 JSON 字典。

        :param stage_name: 当前协议阶段名称。
        :type stage_name: str
        :param body: HTTP 响应字节。
        :type body: bytes
        :return: JSON 响应字典。
        :rtype: dict
        """

        # 解码 UTF-8 响应文本。
        response_text = body.decode("utf-8-sig", errors="replace").strip()

        # 捕获 JSON 解析异常。
        try:

            # 解析响应 JSON。
            response_data = json.loads(response_text)

        # 处理 JSON 解析异常。
        except json.JSONDecodeError as error:

            # 抛出包含有限响应摘要的协议异常。
            raise RuntimeError(f"{stage_name}失败：响应不是有效 JSON：{response_text[:200]}") from error

        # 校验响应结构为字典。
        if not isinstance(response_data, dict):

            # 抛出非字典响应异常。
            raise RuntimeError(f"{stage_name}失败：响应 JSON 不是对象。")

        # 返回解析后的响应字典。
        return response_data

    # 定义总记录数解析方法。
    def parse_total_count(self, value):

        """
        将服务端返回的总记录数字段转换为整数。

        :param value: 服务端 count 字段值。
        :type value: object
        :return: 总记录数或空值。
        :rtype: int | None
        """

        # 判断服务端未返回总记录数。
        if value is None or str(value).strip() == "":

            # 返回空值。
            return None

        # 捕获整数转换异常。
        try:

            # 返回转换后的总记录数。
            return int(value)

        # 处理非数字总记录数。
        except (TypeError, ValueError):

            # 返回空值。
            return None

    # 定义日期规范化方法。
    def normalize_date(self, value, field_name):

        """
        将输入日期规范化为 YYYY-MM-DD 格式。

        :param value: 输入日期文本。
        :type value: str | None
        :param field_name: 日期字段名称。
        :type field_name: str
        :return: 规范化日期文本。
        :rtype: str
        """

        # 读取待处理日期文本。
        date_text = str(value or "").strip()

        # 判断日期未传入。
        if not date_text:

            # 返回当天日期。
            return datetime.now().strftime("%Y-%m-%d")

        # 捕获日期格式解析异常。
        try:

            # 按平台要求解析日期。
            parsed_date = datetime.strptime(date_text, "%Y-%m-%d")

        # 处理不符合格式的日期文本。
        except ValueError as error:

            # 抛出日期格式异常。
            raise ValueError(f"{field_name}格式必须为 YYYY-MM-DD：{date_text}") from error

        # 返回规范化日期文本。
        return parsed_date.strftime("%Y-%m-%d")

    # 定义业务日期范围验证方法。
    def validate_business_dates(self, data_type, rows, start_date, end_date):

        """
        校验返回记录中的业务日期处于请求范围内并记录实际范围。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 原始记录列表。
        :type rows: list[dict]
        :param start_date: 请求开始日期。
        :type start_date: str
        :param end_date: 请求结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 提取可解析的业务日期文本。
        date_values = [str(row.get("dayx", "")).strip() for row in rows if str(row.get("dayx", "")).strip()]

        # 判断当前查询未返回可验证日期。
        if not date_values:

            # 打印空结果日期校验说明。
            self.ph.info(self.ph.kv("业务日期校验", "无返回记录", "yellow"))

            # 结束日期校验。
            return

        # 读取最小业务日期。
        minimum_date = min(date_values)

        # 读取最大业务日期。
        maximum_date = max(date_values)

        # 判断实际日期是否超出请求范围。
        if minimum_date < start_date or maximum_date > end_date:

            # 抛出业务日期范围异常。
            raise RuntimeError(f"{data_type}查询日期越界：实际 {minimum_date} 至 {maximum_date}，请求 {start_date} 至 {end_date}。")

        # 打印实际业务日期范围。
        self.ph.info(self.ph.join(self.ph.kv("业务日期最小值", minimum_date, "cyan"), self.ph.kv("业务日期最大值", maximum_date, "cyan")))

    # 定义 Excel 单元格值规范化方法。
    def normalize_excel_value(self, value):

        """
        将复杂原始字段转换为可写入 Excel 的稳定值。

        :param value: 原始字段值。
        :type value: object
        :return: Excel 单元格值。
        :rtype: str | int | float | bool | None
        """

        # 判断原始值为字典或列表。
        if isinstance(value, (dict, list)):

            # 返回不丢失结构的 JSON 文本。
            return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

        # 返回普通 Excel 可写值。
        return value
