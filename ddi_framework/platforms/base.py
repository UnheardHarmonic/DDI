# -*- coding: utf-8 -*-
"""
平台协议适配器基类。

本模块定义各商业平台适配器需要实现的统一方法，入口程序通过这些方法执行登录、查询、导出和 Excel 落地。
"""

# 导入操作系统路径模块。
import os

# 导入 URL 解析模块。
from urllib.parse import urlparse

# 导入日期时间模块。
from datetime import datetime

# 导入公共下载命名工具。
from ddi_framework.download_watcher import DownloadWatcher


# 定义平台协议适配器基类。
class BasePlatform:

    """
    平台协议适配器基类。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 平台协议适配器实例。
    :rtype: BasePlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化平台协议适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 保存商业公司名称。
        self.company_name = company_name

        # 保存商业账号配置。
        self.account = account

        # 保存日志打印器。
        self.printer = printer

    # 定义平台登录抽象方法。
    def login(self):

        """
        执行平台登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 要求子类实现登录流程。
        raise NotImplementedError

    # 定义平台能力声明读取方法。
    def capabilities(self):

        """
        返回平台适配器的声明能力。

        :return: 平台能力字典。
        :rtype: dict
        """

        # 读取账号配置声明的数据类型。
        declared_types = tuple(self.account.get("数据类型列表", []) or [])

        # 返回公共层使用的平台能力。
        return {
            "data_types": declared_types,
            "supports_date_filter": not bool(self.account.get("库存快照", False)),
            "inventory_snapshot": bool(self.account.get("库存快照", False)),
            "supports_paging": bool(self.account.get("支持分页", True)),
            "supports_native_export": bool(self.account.get("支持原生导出", False)),
            "requires_captcha": bool(self.account.get("需要验证码", False)),
            "dynamic_fields": bool(self.account.get("动态字段", False)),
            "max_query_days": self.account.get("最大查询天数"),
            "allow_company_code_split": bool(self.account.get("允许多公司代码拆分", False)),
            "temporary_browser_fallback": bool(self.account.get("临时浏览器兜底", False)),
        }

    # 定义统一查询兼容入口。
    def query(self, data_type, date_range):

        """
        通过统一日期范围对象执行旧版查询适配器。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param date_range: 公共日期范围对象。
        :type date_range: DateRangeSpec
        :return: 平台原始查询结果。
        :rtype: object
        """

        # 调用既有协议查询实现。
        return self.query_flow(data_type, start_date=date_range.start_date, end_date=date_range.end_date)

    # 定义统一查询落地入口。
    def query_to_excel(self, data_type, date_range, output_dir=None):

        """
        通过统一日期范围对象执行查询并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param date_range: 公共日期范围对象。
        :type date_range: DateRangeSpec
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径或路径集合。
        :rtype: str | list[str]
        """

        # 调用既有查询落地实现保持商业兼容。
        return self.query_flow_to_excel(data_type, start_date=date_range.start_date, end_date=date_range.end_date, output_dir=output_dir)

    # 定义统一导出兼容入口。
    def export(self, data_type, date_range, output_dir=None):

        """
        通过统一日期范围对象执行旧版导出适配器。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param date_range: 公共日期范围对象。
        :type date_range: DateRangeSpec
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: 平台导出结果。
        :rtype: object
        """

        # 调用既有协议导出实现。
        return self.export_flow(data_type, start_date=date_range.start_date, end_date=date_range.end_date, output_dir=output_dir)

    # 定义平台查询抽象方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定类型流向数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 要求子类实现查询流程。
        raise NotImplementedError

    # 定义平台查询落地抽象方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定类型流向数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 要求子类实现查询落地流程。
        raise NotImplementedError

    # 定义平台导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        使用协议查询结果生成本地 Excel 文件。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: Excel 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 调用查询并落地 Excel 的统一流程。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义登录态保障方法。
    def ensure_login(self):

        """
        确保当前协议客户端已经建立有效登录态。

        :return: None。
        :rtype: None
        """

        # 判断当前会话已经完成登录。
        if self.logged_in:

            # 结束重复登录处理。
            return

        # 执行平台登录流程。
        self.login()

    # 定义登录态检查方法。
    def ensure_logged_in(self):

        """
        确保当前协议会话已经登录。

        :return: None。
        :rtype: None
        """

        # 判断当前是否未登录。
        if not self.logged_in:

            # 执行登录流程。
            self.login()

        # 返回空结果。
        return None

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前协议会话 Cookie 供公共任务服务记录。

        :param output_dir: Cookie 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建 Cookie 输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 输出文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存当前协议会话 Cookie。
        self.http.save_cookie_file(cookie_path)

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义查询日期范围校验方法。
    def validate_query_range(self, start_date, end_date):

        """
        校验查询日期范围。

        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 解析查询开始日期。
        start_value = datetime.strptime(start_date, "%Y-%m-%d")

        # 解析查询结束日期。
        end_value = datetime.strptime(end_date, "%Y-%m-%d")

        # 判断开始日期是否不晚于结束日期。
        if start_value <= end_value:

            # 日期范围合法时直接返回。
            return

        # 抛出日期范围非法异常。
        raise ValueError(f"{self.company_name} 查询开始日期不能晚于结束日期：{start_date} > {end_date}")

    # 定义查询分隔线打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印数据类型查询分隔线。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 打印标准阶段分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 返回空结果。
        return None

    # 定义已有文件删除方法。
    def remove_existing_file(self, output_path):

        """
        删除已有同名输出文件。

        :param output_path: 输出文件路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 判断目标文件是否存在。
        if not os.path.exists(output_path):

            # 无需删除不存在的文件。
            return

        # 捕获文件删除异常。
        try:

            # 删除已有目标文件。
            os.remove(output_path)

        # 处理文件被占用导致的删除异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

    # 定义通用 Excel 写表方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(field_name, "") if isinstance(row, dict) else "" for _, field_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 读取当前表头名称。
            header_name = header_item[0]

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

    # 定义登录账号读取方法。
    def clean_login_name(self):

        """
        读取登录账号。

        :return: 登录账号。
        :rtype: str
        """

        # 返回清洗后的账号。
        return str(self.account.get("账号", "")).strip()

    # 定义登录密码读取方法。
    def clean_login_password(self):

        """
        读取登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回清洗后的密码。
        return str(self.account.get("密码", "")).strip()

    # 定义站点根地址构建方法。
    def build_origin(self, url):

        """
        构建站点源地址。

        :param url: 任意站点地址。
        :type url: str
        :return: 源地址。
        :rtype: str
        """

        # 解析 URL 地址。
        parsed_url = urlparse(url)

        # 返回协议与主机地址。
        return f"{parsed_url.scheme}://{parsed_url.netloc}"

    # 定义标准输出文件名构建方法。
    def build_output_file_name(self, data_type, start_date=None, end_date=None):

        """
        构建标准输出文件名。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 标准输出文件名。
        :rtype: str
        """

        # 读取公共任务层传入的日期范围上下文。
        date_range = getattr(self, "_ddi_date_range", None)

        # 读取当前任务日期模式。
        date_mode = getattr(date_range, "mode", "custom")

        # 读取当前任务开始日期。
        effective_start_date = getattr(date_range, "start_date", "") or start_date

        # 读取当前任务结束日期。
        effective_end_date = getattr(date_range, "end_date", "") or end_date

        # 判断当前数据类型是否为库存类。
        is_inventory_type = "库存" in data_type

        # 初始化公共下载命名工具。
        watcher = DownloadWatcher("")

        # 返回公共命名工具生成的标准文件名（库存类统一按导出当天快照日期命名）。
        return watcher.build_target_name(self.company_name, data_type, self.resolve_output_query_date(data_type, effective_start_date, effective_end_date), ".xlsx", start_date=effective_start_date, end_date=effective_end_date, date_mode=date_mode, snapshot_date=self.resolve_output_query_date(data_type, effective_start_date, effective_end_date) if is_inventory_type else None)

    # 定义标准输出命名日期解析方法。
    def resolve_output_query_date(self, data_type, start_date=None, end_date=None):

        """
        解析标准输出文件名使用的业务日期。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 标准输出文件名使用的业务日期。
        :rtype: str | None
        """

        # 判断当前数据类型是否为库存类。
        if "库存" in data_type:

            # 返回库存类命名使用的导出当天日期。
            return datetime.now().strftime("%Y-%m-%d")

        # 返回流向类优先使用的查询开始日期。
        return start_date or end_date

    # 定义导出文件商业名前缀构建方法。
    def build_export_company_name(self):

        """
        构建导出文件使用的商业名称前缀。

        :return: 带 DDI 标识的商业名称。
        :rtype: str
        """

        # 返回带 DDI 标识的商业名称。
        return f"「DDI」{self.company_name}"
