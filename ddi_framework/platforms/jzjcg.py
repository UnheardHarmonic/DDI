# -*- coding: utf-8 -*-
"""
健之佳采购云销售明细查询协议适配器。

本模块负责健之佳采购云的账号登录、销售明细分页查询、周期交集过滤、
原始字段保留、标准字段前置和 Excel 文件生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统接口模块。
import os

# 导入日期时间模块。
from datetime import date
from datetime import datetime


# 导入 URL 编码工具。
from urllib.parse import urlencode

# 导入 URL 错误类型。
from urllib.error import HTTPError

from urllib.request import Request

# 导入 Excel 统一样式函数。
from ddi_framework import apply_standard_excel_style
# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义健之佳采购云支持的数据类型。
JZJCG_SUPPORTED_DATA_TYPES = {"配送流向"}


# 定义健之佳采购云销售明细标准字段。
JZJCG_STANDARD_FIELDS = [
    ("周期开始日期", "beginTime"),
    ("周期结束日期", "enddaTime"),
    ("供应商编码", "supplierNumber"),
    ("供应商名称", "supplierName"),
    ("门店编码", "storeId"),
    ("门店名称", "storeName"),
    ("省份/分部", "province"),
    ("市级分部", "cityBranch"),
    ("商品编码", "goodsId"),
    ("商品名称", "goodsName"),
    ("规格", "goodsNorms"),
    ("单位", "goodsUnit"),
    ("数量", "volume"),
    ("明细数量", "mxVolume"),
    ("OTO数量", "otoVolume"),
    ("调拨数量", "allocation"),
    ("请求开始日期", "__request_start_date"),
    ("请求结束日期", "__request_end_date"),
    ("周期超出请求范围", "__period_out_of_request"),
]


# 定义健之佳采购云平台适配器。
class JzjcgPlatform(BasePlatform):

    """
    健之佳采购云销售明细查询协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 平台适配器实例。
    :rtype: JzjcgPlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化健之佳采购云协议适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用平台基类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存站点根地址。
        self.site_root_url = account.get("接口根地址", "http://220.163.118.50:9994").rstrip("/")

        # 保存接口根地址。
        self.api_root_url = account.get("接口根地址", self.site_root_url).rstrip("/")

        # 保存登录接口地址。
        self.login_url = account.get("登录接口", f"{self.api_root_url}/srm/base/api/login/userName")

        # 保存区间接口地址。
        self.section_url = account.get("区间接口", f"{self.api_root_url}/srm/customer/api/cellInfo/getSection")

        # 保存查询接口地址。
        self.query_url = account.get("查询接口", f"{self.api_root_url}/srm/customer/api/cellSupplierInfo/getListForPageSellSupplier")

        # 保存登录账号。
        self.username = account.get("账号", "")

        # 保存登录密码。
        self.password = account.get("密码", "")

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 500))

        # 初始化协议请求客户端并托管平台会话。
        self.http = ProtocolHttpClient(self.site_root_url)

        # 引用协议客户端管理的 Cookie 容器。
        self.cookie_jar = self.http.cookie_jar

        # 引用协议客户端管理的请求打开器。
        self.opener = self.http.opener

        # 初始化登录令牌。
        self.token = ""

        # 初始化登录返回用户信息。
        self.login_data = {}

        # 初始化登录状态。
        self.logged_in = False

    # 定义平台登录方法。
    def login(self):

        """
        登录健之佳采购云系统。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段分隔线。
        self.ph.section("健之佳采购云登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 构建登录表单数据。
        payload = {"userName": self.username, "passWord": self.password}

        # 请求登录接口。
        response_data = self.post_form(self.login_url, payload, token_required=False, referer=f"{self.site_root_url}/index/login.html")

        # 读取登录返回数据。
        login_data = response_data.get("data") if isinstance(response_data.get("data"), dict) else {}

        # 读取登录令牌。
        self.token = login_data.get("token", "")

        # 校验登录令牌是否存在。
        if not self.token:

            # 抛出登录失败异常。
            raise RuntimeError(f"健之佳采购云登录失败：{response_data.get('msg') or response_data}")

        # 保存登录返回用户信息。
        self.login_data = login_data

        # 标记当前会话已登录。
        self.logged_in = True

        # 读取登录返回公司名称。
        real_name = login_data.get("realName", "")

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("健之佳采购云登录成功", real_name or self.username, "green"), self.ph.kv("账号", self.username, "cyan")))

        # 记录账号归属差异。
        if real_name and real_name != self.company_name:

            # 打印账号归属差异提醒。
            self.ph.warning(self.ph.join(self.ph.kv("账号返回名称", real_name, "yellow"), self.ph.kv("配置商业", self.company_name, "cyan")))

        # 返回登录摘要信息。
        return {"success": True, "data": login_data}

    # 定义会话信息保存方法。
    def save_cookie(self, output_dir):

        """
        保存健之佳采购云当前登录会话信息。

        :param output_dir: 会话信息输出目录。
        :type output_dir: str
        :return: 会话信息文件路径。
        :rtype: str
        """

        # 创建会话信息输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建会话信息文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 组装当前会话信息。
        cookie_data = {"token": self.token, "user": self.login_data, "api_root": self.api_root_url, "site_root": self.site_root_url}

        # 写入会话信息文件。
        with open(cookie_path, "w", encoding="utf-8") as cookie_file:

            # 序列化会话信息。
            json.dump(cookie_data, cookie_file, ensure_ascii=False, indent=2)

        # 打印会话信息保存日志。
        self.ph.success(self.ph.kv("会话信息已保存", self.ph.path(cookie_path), "green"))

        # 返回会话信息文件路径。
        return cookie_path

    # 定义数据查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行列表。
        :rtype: list[dict]
        """

        # 打印数据类型查询分隔线。
        self.print_query_section(data_type, start_date, end_date)

        # 返回协议分页查询结果。
        return self.fetch_protocol_rows(data_type, start_date, end_date)


    # 定义查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定类型业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 打印数据类型查询分隔线。
        self.print_query_section(data_type, start_date, end_date)

        # 查询协议数据行。
        rows = self.fetch_protocol_rows(data_type, start_date, end_date)

        # 读取输出目录。
        save_dir = output_dir or "downloads"

        # 创建输出目录。
        os.makedirs(save_dir, exist_ok=True)

        # 构建目标文件名。
        target_name = self.build_output_file_name(data_type, start_date, end_date)

        # 拼接目标文件路径。
        target_path = os.path.join(save_dir, target_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(data_type, rows, target_path)

        # 应用统一 Excel 样式。
        apply_standard_excel_style(target_path)

        # 打印 Excel 生成日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(target_path), "green"), self.ph.kv("记录数", f"{len(rows)} 条", "cyan")))

        # 返回 Excel 文件路径。
        return target_path


    # 定义协议分页查询方法。
    def fetch_protocol_rows(self, data_type, start_date, end_date):

        """
        通过健之佳采购云接口分页读取销售明细。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 标准字段与原始字段合并后的数据行。
        :rtype: list[dict]
        """

        # 校验数据类型是否受支持。
        self.validate_data_type(data_type)

        # 确保当前会话已登录。
        self.ensure_logged_in()

        # 初始化页码。
        page_number = 1

        # 初始化原始数据行容器。
        raw_rows = []

        # 初始化接口总数。
        total_count = None

        # 持续分页直到接口没有更多数据。
        while True:

            # 构建当前页查询参数。
            query_params = self.build_query_params(page_number)

            # 请求当前页数据。
            response_data = self.get_json(self.query_url, query_params, referer=f"{self.site_root_url}/web/page/oneCard/cells/saleManageOffer.html")

            # 解析当前页数据和总数。
            page_rows, parsed_total = self.extract_rows_and_total(response_data)

            # 更新接口返回总数。
            total_count = parsed_total if parsed_total is not None else total_count

            # 追加当前页原始数据。
            raw_rows.extend(page_rows)

            # 打印当前页查询日志。
            self.ph.info(self.ph.join(self.ph.kv("销售明细分页", page_number, "cyan"), self.ph.kv("本页", len(page_rows), "cyan"), self.ph.kv("累计", len(raw_rows), "cyan"), self.ph.kv("接口总数", "" if total_count is None else total_count, "cyan")))

            # 判断当前页是否已经为空。
            if not page_rows:

                # 结束分页循环。
                break

            # 判断是否已经达到接口返回总数。
            if total_count is not None and len(raw_rows) >= total_count:

                # 结束分页循环。
                break

            # 判断当前页数量是否小于分页大小。
            if len(page_rows) < self.page_size:

                # 结束分页循环。
                break

            # 增加页码。
            page_number += 1

        # 按请求日期过滤周期交集。
        filtered_rows = self.filter_rows_by_period(raw_rows, start_date, end_date)

        # 转换查询结果行。
        result_rows = [self.transform_row(row_data, start_date, end_date) for row_data in filtered_rows]

        # 打印自然月处理提醒。
        self.print_period_scope_warning(filtered_rows, start_date, end_date)

        # 打印查询完成日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type}查询完成", f"{len(result_rows)} 条", "green"), self.ph.kv("接口总数", "" if total_count is None else total_count, "cyan"), self.ph.kv("过滤前", len(raw_rows), "cyan")))

        # 返回结果行列表。
        return result_rows


    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验数据类型是否已经接入。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否受支持。
        if data_type not in JZJCG_SUPPORTED_DATA_TYPES:

            # 抛出未支持数据类型异常。
            raise NotImplementedError(f"健之佳采购云暂未接入{data_type}")

        # 返回空结果。
        return None

    # 定义查询参数构建方法。
    def build_query_params(self, page_number):

        """
        构建销售明细分页查询参数。

        :param page_number: 页码。
        :type page_number: int
        :return: 查询参数。
        :rtype: dict
        """

        # 返回分页查询参数。
        return {"goodsId": "", "modalType": "apply", "multiSort": "", "page": page_number, "rows": self.page_size}

    # 定义接口数据解析方法。
    def extract_rows_and_total(self, response_data):

        """
        从接口响应中解析数据行与总数。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 数据行列表与总数。
        :rtype: tuple[list[dict], int | None]
        """

        # 读取响应数据节点。
        data = response_data.get("data") if isinstance(response_data, dict) else {}

        # 判断数据节点是否为分页字典。
        if isinstance(data, dict):

            # 读取分页行列表。
            rows = data.get("rows") or data.get("list") or data.get("data") or []

            # 读取分页总数。
            total = data.get("total") or data.get("count") or data.get("totalCount")

            # 返回分页行和总数。
            return rows if isinstance(rows, list) else [], total

        # 判断数据节点是否为列表。
        if isinstance(data, list):

            # 返回列表数据和列表长度。
            return data, len(data)

        # 返回空列表和空总数。
        return [], None

    # 定义周期交集过滤方法。
    def filter_rows_by_period(self, rows, start_date, end_date):

        """
        按接口周期与请求日期的交集过滤销售明细。

        :param rows: 接口原始数据行。
        :type rows: list[dict]
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 命中请求日期范围的数据行。
        :rtype: list[dict]
        """

        # 解析请求开始日期。
        request_start = self.parse_date_text(start_date)

        # 解析请求结束日期。
        request_end = self.parse_date_text(end_date or start_date)

        # 无日期条件时保留全部数据。
        if request_start is None and request_end is None:

            # 返回全部原始数据。
            return rows

        # 单边开始日期使用结束日期补齐。
        request_start = request_start or request_end

        # 单边结束日期使用开始日期补齐。
        request_end = request_end or request_start

        # 初始化过滤结果。
        filtered_rows = []

        # 遍历接口原始数据行。
        for row_data in rows:

            # 解析当前行周期开始日期。
            period_start = self.parse_date_text(row_data.get("beginTime"))

            # 解析当前行周期结束日期。
            period_end = self.parse_date_text(row_data.get("enddaTime") or row_data.get("endTime"))

            # 无周期日期的记录保留以避免误删接口字段异常数据。
            if period_start is None or period_end is None:

                # 追加当前行。
                filtered_rows.append(row_data)

                # 进入下一行。
                continue

            # 判断当前周期是否与请求范围存在交集。
            if period_start <= request_end and period_end >= request_start:

                # 追加命中周期的数据行。
                filtered_rows.append(row_data)

        # 返回过滤结果。
        return filtered_rows

    # 定义周期范围提醒方法。
    def print_period_scope_warning(self, rows, start_date, end_date):

        """
        打印周期口径与请求日期范围的差异提醒。

        :param rows: 已过滤的数据行。
        :type rows: list[dict]
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 解析请求开始日期。
        request_start = self.parse_date_text(start_date)

        # 解析请求结束日期。
        request_end = self.parse_date_text(end_date or start_date)

        # 无完整日期时跳过范围提醒。
        if request_start is None or request_end is None:

            # 返回空结果。
            return None

        # 统计超出请求日期范围的周期行数。
        out_of_scope_count = sum(1 for row_data in rows if self.is_period_out_of_request(row_data, request_start, request_end))

        # 判断是否存在超出请求日期范围的周期。
        if out_of_scope_count:

            # 打印周期口径提醒。
            self.ph.warning(self.ph.join(self.ph.kv("平台周期口径", "26日至次月25日", "yellow"), self.ph.kv("超出请求范围行数", out_of_scope_count, "yellow"), self.ph.kv("处理方式", "按周期交集保留", "cyan")))

        # 返回空结果。
        return None

    # 定义周期超范围判断方法。
    def is_period_out_of_request(self, row_data, request_start, request_end):

        """
        判断当前行周期是否超出请求日期范围。

        :param row_data: 接口原始数据行。
        :type row_data: dict
        :param request_start: 请求开始日期。
        :type request_start: date
        :param request_end: 请求结束日期。
        :type request_end: date
        :return: 周期是否超出请求范围。
        :rtype: bool
        """

        # 解析当前行周期开始日期。
        period_start = self.parse_date_text(row_data.get("beginTime"))

        # 解析当前行周期结束日期。
        period_end = self.parse_date_text(row_data.get("enddaTime") or row_data.get("endTime"))

        # 无周期日期时按未超范围处理。
        if period_start is None or period_end is None:

            # 返回未超范围。
            return False

        # 返回周期是否超出请求范围。
        return period_start < request_start or period_end > request_end

    # 定义数据行转换方法。
    def transform_row(self, row_data, start_date, end_date):

        """
        将接口原始行转换为标准字段前置的输出行。

        :param row_data: 接口原始行。
        :type row_data: dict
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 输出行。
        :rtype: dict
        """

        # 展平接口原始行。
        flat_row = self.flatten_row(row_data)

        # 写入请求开始日期。
        flat_row["__request_start_date"] = start_date or ""

        # 写入请求结束日期。
        flat_row["__request_end_date"] = end_date or start_date or ""

        # 解析请求开始日期。
        request_start = self.parse_date_text(start_date)

        # 解析请求结束日期。
        request_end = self.parse_date_text(end_date or start_date)

        # 写入周期超范围标记。
        flat_row["__period_out_of_request"] = "是" if request_start and request_end and self.is_period_out_of_request(row_data, request_start, request_end) else "否"

        # 初始化输出行。
        output_row = {}

        # 遍历标准字段定义。
        for header_name, source_key in JZJCG_STANDARD_FIELDS:

            # 写入标准字段值。
            output_row[header_name] = flat_row.get(source_key, "")

        # 遍历原始展平字段。
        for field_name, field_value in flat_row.items():

            # 判断当前字段是否为内部补充字段。
            if field_name.startswith("__"):

                # 跳过内部补充字段。
                continue

            # 判断当前字段是否已经作为标准字段写入。
            if field_name not in output_row:

                # 写入原始字段。
                output_row[field_name] = field_value

        # 返回输出行。
        return output_row

    # 定义字典展平方法。
    def flatten_row(self, value, prefix=""):

        """
        展平嵌套字典和列表字段。

        :param value: 原始值。
        :type value: dict | list | str | int | float | None
        :param prefix: 字段前缀。
        :type prefix: str
        :return: 展平后的字段字典。
        :rtype: dict
        """

        # 初始化展平结果。
        result = {}

        # 判断当前值是否为字典。
        if isinstance(value, dict):

            # 遍历字典字段。
            for key, child_value in value.items():

                # 构建子字段名称。
                child_key = f"{prefix}.{key}" if prefix else str(key)

                # 合并子字段展平结果。
                result.update(self.flatten_row(child_value, child_key))

            # 返回字典展平结果。
            return result

        # 判断当前值是否为列表。
        if isinstance(value, list):

            # 将列表按 JSON 文本保留。
            result[prefix] = json.dumps(value, ensure_ascii=False)

            # 返回列表字段结果。
            return result

        # 写入普通字段值。
        result[prefix] = "" if value is None else value

        # 返回普通字段结果。
        return result

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, data_type, rows, target_path):

        """
        将查询结果写入 Excel 文件。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果行。
        :type rows: list[dict]
        :param target_path: 目标文件路径。
        :type target_path: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列号转换函数。
        from openpyxl.utils import get_column_letter

        # 构建输出表头。
        headers = self.build_excel_headers(rows)

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = data_type[:31]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果行。
        for row_data in rows:

            # 按表头顺序写入当前行。
            worksheet.append([row_data.get(header, "") for header in headers])

        # 遍历表头列设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 计算当前列最大文本长度。
            width = max([len(str(header_name))] + [len(str(row_data.get(header_name, ""))) for row_data in rows[:200]])

            # 设置当前列宽。
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(width + 2, 12), 45)

        # 删除已存在的目标文件。
        if os.path.exists(target_path):

            # 移除旧文件。
            os.remove(target_path)

        # 保存工作簿。
        workbook.save(target_path)

        # 关闭工作簿。
        workbook.close()

        # 返回空结果。
        return None

    # 定义 Excel 表头构建方法。
    def build_excel_headers(self, rows):

        """
        构建标准字段前置的 Excel 表头。

        :param rows: 查询结果行。
        :type rows: list[dict]
        :return: 表头列表。
        :rtype: list[str]
        """

        # 初始化表头列表。
        headers = [header_name for header_name, _ in JZJCG_STANDARD_FIELDS]

        # 遍历全部结果行。
        for row_data in rows:

            # 遍历当前行字段。
            for field_name in row_data.keys():

                # 判断字段是否已经存在。
                if field_name not in headers:

                    # 追加原始字段。
                    headers.append(field_name)

        # 返回表头列表。
        return headers

    # 定义表单 POST 请求方法。
    def post_form(self, request_url, payload, token_required=True, referer=None):

        """
        发起表单 POST 请求并解析响应。

        :param request_url: 接口地址。
        :type request_url: str
        :param payload: 请求体。
        :type payload: dict
        :param token_required: 是否携带登录令牌。
        :type token_required: bool
        :param referer: 来源页面地址。
        :type referer: str | None
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 序列化表单请求体。
        request_body = urlencode(payload).encode("utf-8")

        # 构建请求头。
        headers = self.build_headers(token_required=token_required, referer=referer, content_type="application/x-www-form-urlencoded; charset=UTF-8")

        # 构建请求对象。
        request = Request(request_url, data=request_body, headers=headers, method="POST")

        # 返回 JSON 响应。
        return self.open_json(request)

    # 定义 JSON GET 请求方法。
    def get_json(self, request_url, query_params, referer=None):

        """
        发起 GET 请求并解析 JSON 响应。

        :param request_url: 接口地址。
        :type request_url: str
        :param query_params: 查询参数。
        :type query_params: dict
        :param referer: 来源页面地址。
        :type referer: str | None
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 拼接查询字符串。
        query_string = urlencode(query_params)

        # 拼接完整请求地址。
        full_url = f"{request_url}?{query_string}" if query_string else request_url

        # 构建请求头。
        headers = self.build_headers(token_required=True, referer=referer, content_type=None)

        # 构建请求对象。
        request = Request(full_url, headers=headers, method="GET")

        # 返回 JSON 响应。
        return self.open_json(request)

    # 定义请求执行方法。
    def open_json(self, request):

        """
        执行 HTTP 请求并解析 JSON 响应。

        :param request: 请求对象。
        :type request: Request
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 捕获接口错误响应。
        try:

            # 发送 HTTP 请求。
            with self.opener.open(request, timeout=90) as response:

                # 读取响应文本。
                response_text = response.read().decode("utf-8", errors="replace")

        # 处理 HTTP 错误响应。
        except HTTPError as error:

            # 读取错误响应文本。
            response_text = error.read().decode("utf-8", errors="replace")

        # 解析 JSON 响应。
        response_data = json.loads(response_text)

        # 校验业务状态码。
        if response_data.get("code") not in (0, None):

            # 抛出业务异常。
            raise RuntimeError(f"健之佳采购云接口返回异常：{response_data.get('msg') or response_data}")

        # 返回接口响应。
        return response_data

    # 定义请求头构建方法。
    def build_headers(self, token_required=True, referer=None, content_type=None):

        """
        构建健之佳采购云请求头。

        :param token_required: 是否携带登录令牌。
        :type token_required: bool
        :param referer: 来源页面地址。
        :type referer: str | None
        :param content_type: 请求体内容类型。
        :type content_type: str | None
        :return: 请求头字典。
        :rtype: dict
        """

        # 初始化基础请求头。
        headers = {"Accept": "application/json, text/javascript, */*; q=0.01", "Origin": self.site_root_url, "Referer": referer or f"{self.site_root_url}/", "User-Agent": "Mozilla/5.0", "X-Requested-With": "XMLHttpRequest"}

        # 判断是否存在请求体内容类型。
        if content_type:

            # 写入请求体内容类型。
            headers["Content-Type"] = content_type

        # 判断是否需要登录令牌。
        if token_required:

            # 写入平台要求的原始令牌。
            headers["Authorization"] = self.token

        # 返回请求头。
        return headers

    # 定义日期文本解析方法。
    def parse_date_text(self, date_text):

        """
        解析平台日期文本。

        :param date_text: 日期文本。
        :type date_text: str | None
        :return: 日期对象。
        :rtype: date | None
        """

        # 判断日期文本是否为空。
        if not date_text:

            # 返回空日期。
            return None

        # 转换日期文本。
        text = str(date_text).strip()

        # 处理标准日期格式。
        for date_format in ("%Y-%m-%d", "%Y%m%d"):

            # 捕获日期解析异常。
            try:

                # 返回解析后的日期。
                return datetime.strptime(text[:10] if date_format == "%Y-%m-%d" else text[:8], date_format).date()

            # 忽略当前格式解析失败。
            except ValueError:

                # 继续尝试下一种格式。
                continue

        # 返回空日期。
        return None

