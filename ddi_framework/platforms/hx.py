# -*- coding: utf-8 -*-
"""
海翔供应链平台协议适配器。

本模块负责海翔供应链平台的账号登录、分页查询和本地 Excel 生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入 UUID 生成模块。
import uuid

# 导入下载命名工具。
from ddi_framework import DownloadWatcher

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义海翔供应链平台协议适配器。
class HxPlatform(BasePlatform):

    """
    海翔供应链平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 海翔供应链平台适配器实例。
    :rtype: HxPlatform
    """

    # 定义海翔供应链平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化海翔供应链平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://fn.seaflysoft.com/")

        # 保存平台根地址。
        self.base_url = account.get("接口根地址", "http://fn.seaflysoft.com")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.base_url}/login.do")

        # 保存报表查询接口地址。
        self.report_api_url = account.get("报表接口", f"{self.base_url}/getreportdata.do")

        # 保存报表来源页地址。
        self.report_page_url = account.get("报表页", f"{self.base_url}/report?reportName=rp_supplierflow_dz")

        # 保存报表名称。
        self.report_name = account.get("报表名称", "rp_supplierflow_dz")

        # 保存账套版本。
        self.version_flag = str(account.get("账套版本编码", "17")).strip()

        # 保存账套名称。
        self.account_name = str(account.get("账套", "蜂鸟医药")).strip()

        # 保存账套编号。
        self.account_id = str(account.get("账套编号", "2")).strip()

        # 保存登录机构名称。
        self.organization_name = str(account.get("登录机构", "总公司")).strip()

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 200))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录状态。
        self.logged_in = False

    # 定义海翔供应链平台登录方法。
    def login(self):

        """
        执行海翔供应链平台协议登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("海翔供应链平台登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 访问登录页以初始化站点 Cookie。
        self.http.get(self.login_page_url)

        # 构建登录请求体。
        payload = self.build_login_payload()

        # 打印登录账号日志。
        self.ph.info(self.ph.join(self.ph.kv("登录账号", self.clean_login_name(), "cyan"), self.ph.kv("账套", self.account_name, "cyan"), self.ph.kv("登录机构", self.organization_name, "cyan")))

        # 提交登录表单请求。
        result = self.http.post_form(self.login_api_url, payload, self.login_page_url)

        # 解析登录响应。
        login_result = self.decode_json_response(result)

        # 校验登录结果。
        if str(login_result.get("status")) != "200":

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 正式登录失败：{login_result.get('msg') or login_result}")

        # 保存登录状态。
        self.logged_in = True

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("正式登录成功", login_result.get("msg", "OK"), "green"), self.ph.kv("Cookie数", len(list(self.http.cookie_jar)), "cyan")))

        # 返回登录响应。
        return login_result

    # 定义登录请求体构建方法。
    def build_login_payload(self):

        """
        构建登录请求体。

        :return: 登录请求体。
        :rtype: dict
        """

        # 生成本次登录机器标识。
        machine_uuid = str(self.account.get("机器标识") or uuid.uuid4())

        # 返回登录表单字段。
        return {"verflag": self.version_flag, "accountname": self.account_name, "yid": self.account_id, "yidim": "", "code": self.clean_login_name(), "ecode": self.clean_login_name(), "password": self.clean_login_password(), "remebers": "1", "phone": "", "yzm": "", "macuuid": machine_uuid}


    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询海翔供应链平台业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 查询指定类型完整数据。
        query_result = self.query_hx_data(data_type, start_date, end_date)

        # 返回查询结果数据行。
        return query_result["rows"]


    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询海翔供应链平台业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定类型完整数据。
        query_result = self.query_hx_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义海翔供应链平台数据查询方法。
    def query_hx_data(self, data_type, start_date, end_date):

        """
        查询海翔供应链平台指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 校验查询开始日期。
        if not start_date:

            # 抛出查询开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验查询结束日期。
        if not end_date:

            # 抛出查询结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 初始化查询结果行。
        all_rows = []

        # 初始化当前页码。
        page_index = 1

        # 初始化总记录数。
        total_count = None

        # 循环读取所有分页。
        while True:

            # 构建当前分页请求体。
            payload = self.build_report_payload(data_type, start_date, end_date, page_index)

            # 打印当前分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"page={page_index}", "cyan"), self.ph.kv("rows", self.page_size, "cyan")))

            # 发送当前分页查询请求。
            result = self.http.post_form(self.report_api_url, payload, self.report_page_url)

            # 解析当前分页响应。
            response_data = self.decode_json_response(result)

            # 提取当前页数据行。
            current_rows = self.extract_rows(response_data)

            # 解析当前页总记录数。
            total_count = self.resolve_total_count(response_data, total_count)

            # 合并当前页数据行。
            all_rows.extend(current_rows)

            # 打印当前页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(current_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count, "cyan")))

            # 判断总数是否已经取完。
            if len(all_rows) >= total_count:

                # 结束分页循环。
                break

            # 判断当前页是否已经不足一页。
            if len(current_rows) < self.page_size:

                # 结束分页循环。
                break

            # 递增页码。
            page_index += 1

        # 推断当前数据表头。
        headers = self.infer_headers(all_rows)

        # 返回查询数据与表头。
        return {"rows": all_rows, "headers": headers or self.build_default_headers(data_type)}

    # 定义报表请求体构建方法。
    def build_report_payload(self, data_type, start_date, end_date, page_index):

        """
        构建报表查询请求体。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :param page_index: 当前页码。
        :type page_index: int
        :return: 报表查询请求体。
        :rtype: dict
        """

        # 构建报表业务参数。
        report_params = self.build_report_params(start_date, end_date)

        # 返回报表表单字段。
        return {"tabID": self.build_tab_id(data_type), "reportName": self.report_name, "gridIndex": "1", "parJson": json.dumps(report_params, ensure_ascii=False, separators=(",", ":")), "page": page_index, "rows": self.page_size}

    # 定义报表业务参数构建方法。
    def build_report_params(self, start_date, end_date):

        """
        构建报表业务参数。

        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 报表业务参数。
        :rtype: dict
        """

        # 返回报表业务参数。
        params = {"begindate": start_date, "enddate": end_date, "pid": "0", "pdim": "", "cid": 0, "y_id": int(self.account_id), "supplytype": "0", "supplyid": 0, "classid": "", "customclassid": ""}

        # 门店版部署对 supplyid 字段返回空响应，按账号配置剔除。
        if not bool(self.account.get("携带supplyid", True)):

            # 移除 supplyid 字段。
            params.pop("supplyid", None)

        # 返回报表业务参数。
        return params

    # 定义数据类型页签编号构建方法。
    def build_tab_id(self, data_type):

        """
        构建数据类型对应的报表页签编号。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 报表页签编号。
        :rtype: str
        """

        # 定义数据类型与页签编号默认映射（蜂鸟系采购/销售/库存顺序）。
        mapping = {"购进流向": "1", "配送流向": "2", "库存": "3"}

        # 门店版等部署的页签顺序可能不同，按账号配置覆盖。
        mapping.update({name: str(tab_id) for name, tab_id in (self.account.get("页签映射") or {}).items()})

        # 返回当前数据类型页签编号。
        return mapping[data_type]


    # 定义查询阶段标题打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印查询阶段标题。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 打印流向查询标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date, "cyan"), self.ph.kv("结束日期", end_date, "cyan")))



    # 定义文本响应解析方法。
    def decode_text(self, result):

        """
        解析文本响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 响应文本。
        :rtype: str
        """

        # 解码 UTF-8 响应文本。
        return result.body.decode("utf-8", errors="replace").strip()

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应文本。
        response_text = self.decode_text(result)

        # 处理空响应。
        if not response_text:

            # 返回空响应结构。
            return {}

        # 返回 JSON 响应数据。
        return json.loads(response_text)

    # 定义响应数据行提取方法。
    def extract_rows(self, response_data):

        """
        提取响应数据行。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 读取响应数据行。
        rows = response_data.get("rows") or []

        # 返回字典数据行列表。
        return [row for row in rows if isinstance(row, dict)]

    # 定义总记录数解析方法。
    def resolve_total_count(self, response_data, current_total):

        """
        解析分页总记录数。

        :param response_data: 当前页响应数据。
        :type response_data: dict
        :param current_total: 已解析的总记录数。
        :type current_total: int | None
        :return: 总记录数。
        :rtype: int
        """

        # 返回已存在的总记录数。
        if current_total is not None:

            # 返回当前总记录数。
            return current_total

        # 读取响应总记录数字段。
        count_value = response_data.get("total", 0)

        # 返回整型总记录数。
        return int(float(count_value or 0))

    # 定义表头推断方法。
    def infer_headers(self, rows):

        """
        根据数据行推断 Excel 表头。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :return: 表头列表。
        :rtype: list[str]
        """

        # 初始化表头结果列表。
        headers = []

        # 遍历查询结果数据行。
        for row in rows:

            # 跳过非字典数据行。
            if not isinstance(row, dict):

                # 继续处理下一行。
                continue

            # 遍历当前数据行字段名称。
            for header_name in row.keys():

                # 判断字段名称是否已经记录。
                if header_name in headers:

                    # 继续处理下一字段。
                    continue

                # 追加当前字段名称。
                headers.append(header_name)

        # 返回完整字段顺序。
        return headers

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验海翔供应链平台数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已接入。
        if data_type not in ["购进流向", "配送流向", "库存"]:

            # 抛出数据类型未接入异常。
            raise ValueError(f"海翔供应链平台暂未接入数据类型：{data_type}")

    # 定义默认表头构建方法。
    def build_default_headers(self, data_type):

        """
        构建默认输出表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 默认表头列表。
        :rtype: list[str]
        """

        # 定义流向类默认表头。
        flow_headers = ["billdate", "billnumber", "billname", "clientcode", "clientname", "code", "name", "standard", "factory", "unit", "qty", "batchcode", "productdate", "validdate", "price", "total", "taxprice", "taxtotal", "sname", "uname", "empname", "yname"]

        # 定义库存默认表头。
        stock_headers = ["code", "name", "standard", "factory", "unit", "qty", "batchcode", "productdate", "validdate", "price", "total", "taxprice", "taxtotal", "sname", "yname"]

        # 判断当前数据类型是否为库存。
        if data_type == "库存":

            # 返回库存默认表头。
            return stock_headers

        # 返回流向类默认表头。
        return flow_headers


    # 定义通用 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头列表。
        :type headers: list[str]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(header_name, "") if isinstance(row, dict) else "" for header_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

        # 关闭导出工作簿。
        workbook.close()

