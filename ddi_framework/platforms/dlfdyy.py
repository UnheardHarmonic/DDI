# -*- coding: utf-8 -*-
"""
大连方大药业有限公司流向平台协议适配器。

本模块负责大连方大药业有限公司流向平台的 MD5 登录、业务数据查询和本地 Excel 生成。
"""

# 导入哈希摘要模块。
import hashlib

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期时间对象。
from datetime import datetime

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义大连方大药业流向平台协议适配器。
class DlfdyyPlatform(BasePlatform):

    """
    大连方大药业流向平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 大连方大药业平台适配器实例。
    :rtype: DlfdyyPlatform
    """

    # 定义大连方大药业适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化大连方大药业平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://dlfdyy.vicp.net:8088/fdweb/fdbus/index.w")

        # 保存平台根地址。
        self.base_url = self.build_origin(self.login_page_url)

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.base_url}/baas/fdbus/username/userdl")

        # 保存登录日志接口地址。
        self.login_log_api_url = account.get("登录日志接口", f"{self.base_url}/baas/fdbus/username/dlrz")

        # 保存购进流向查询接口地址。
        self.purchase_query_url = account.get("购进流向接口", f"{self.base_url}/baas/fdbus/splx/splxCxjgKhCg")

        # 保存配送流向查询接口地址。
        self.dist_query_url = account.get("配送流向接口", f"{self.base_url}/baas/fdbus/splx/splxCxjgKhXs")

        # 保存库存查询接口地址。
        self.stock_query_url = account.get("库存接口", f"{self.base_url}/baas/fdbus/splx/splxCxjgKhKc")

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录响应。
        self.login_result = {}

        # 初始化登录用户行。
        self.login_row = {}

        # 初始化登录状态。
        self.logged_in = False

        # 初始化查询商品编码。
        self.query_item_code = str(account.get("商品编码", "")).strip()

        # 初始化查询供应商过滤值。
        self.query_supplier = str(account.get("供应商过滤", "")).strip()

    # 定义大连方大药业登录方法。
    def login(self):

        """
        执行大连方大药业协议登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("大连方大药业登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 访问登录页以初始化站点 Cookie。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 构建登录请求体。
        payload = self.build_login_payload()

        # 打印登录账号日志。
        self.ph.info(self.ph.join(self.ph.kv("登录账号", self.clean_login_name(), "cyan"), self.ph.kv("密码加密", "MD5", "yellow")))

        # 打印正式登录请求日志。
        self.ph.info("提交正式登录请求: /baas/fdbus/username/userdl")

        # 提交登录请求。
        result = self.post_json(self.login_api_url, payload, self.login_page_url)

        # 解析登录响应。
        login_result = self.decode_table_response(result)

        # 校验登录结果。
        if not self.is_login_success(login_result):

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 正式登录失败：用户或密码有误")

        # 保存登录响应。
        self.login_result = login_result

        # 保存登录用户行。
        self.login_row = self.extract_first_row(login_result)

        # 保存登录状态。
        self.logged_in = True

        # 补充查询商品编码。
        self.query_item_code = self.query_item_code or str(self.login_row.get("f_desc") or "").strip()

        # 写入登录日志。
        self.write_login_log()

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("正式登录成功", "用户校验通过", "green"), self.ph.kv("查询商品编码", self.query_item_code, "cyan")))

        # 返回登录响应。
        return login_result

    # 定义登录请求体构建方法。
    def build_login_payload(self):

        """
        构建大连方大药业登录请求体。

        :return: 登录请求体。
        :rtype: dict
        """

        # 返回登录请求字段。
        return {"username": self.clean_login_name(), "password": self.build_password_hash()}

    # 定义登录密码摘要构建方法。
    def build_password_hash(self):

        """
        构建登录密码 MD5 摘要。

        :return: 登录密码 MD5 十六进制摘要。
        :rtype: str
        """

        # 读取清洗后的登录密码。
        password_text = self.clean_login_password()

        # 返回前端一致的 MD5 摘要。
        return hashlib.md5(password_text.encode("utf-8")).hexdigest()



    # 定义登录日志写入方法。
    def write_login_log(self):

        """
        写入平台登录日志。

        :return: None。
        :rtype: None
        """

        # 判断登录日志是否启用。
        if not self.account.get("写入登录日志", True):

            # 跳过登录日志写入。
            return

        # 构建登录日志请求体。
        payload = {"npname": self.clean_login_name()}

        # 打印登录日志请求。
        self.ph.info("写入登录日志: /baas/fdbus/username/dlrz")

        # 捕获登录日志写入异常。
        try:

            # 提交登录日志请求。
            self.post_json(self.login_log_api_url, payload, self.login_page_url)

        # 处理登录日志写入失败。
        except RuntimeError as error:

            # 打印登录日志降级提示。
            self.ph.warning(f"登录日志写入失败，继续执行查询: {error}")


    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询大连方大药业业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 查询指定类型完整数据。
        query_result = self.query_dlfdyy_data(data_type, start_date, end_date)

        # 返回查询结果数据行。
        return query_result["rows"]


    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询大连方大药业业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定业务数据。
        query_result = self.query_dlfdyy_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义大连方大药业数据查询方法。
    def query_dlfdyy_data(self, data_type, start_date, end_date):

        """
        查询大连方大药业指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 校验流向类查询开始日期。
        if data_type in ["购进流向", "配送流向"] and not start_date:

            # 抛出查询开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验流向类查询结束日期。
        if data_type in ["购进流向", "配送流向"] and not end_date:

            # 抛出查询结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 构建查询请求体。
        payload = self.build_query_payload(data_type, start_date, end_date)

        # 打印查询请求日志。
        self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}", self.build_query_path(data_type), "cyan"), self.ph.kv("商品编码", payload.get("spbh", ""), "yellow")))

        # 发送查询请求。
        result = self.post_json(self.build_query_api_url(data_type), payload, self.login_page_url)

        # 解析查询响应。
        response_data = self.decode_table_response(result)

        # 规范化查询数据行。
        rows = self.extract_table_rows(response_data)

        # 打印查询解析日志。
        self.ph.success(self.ph.join(self.ph.kv("查询解析完成", f"{len(rows)} 条", "green"), self.ph.kv("表头数", len(self.build_headers(data_type)), "cyan")))

        # 返回查询数据与表头。
        return {"rows": rows, "headers": self.build_headers(data_type)}


    # 定义查询阶段标题打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印查询阶段标题。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断是否为库存查询。
        if data_type == "库存":

            # 打印库存查询标题。
            self.ph.section(f"{data_type}查询", "库存接口不传递起止日期")

            # 结束标题打印。
            return

        # 打印流向查询标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

    # 定义查询请求体构建方法。
    def build_query_payload(self, data_type, start_date, end_date):

        """
        构建业务查询请求体。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 业务查询请求体。
        :rtype: dict
        """

        # 构建通用查询字段。
        payload = {"user": self.clean_login_name(), "spbh": self.get_query_item_code(), "supplier": self.query_supplier}

        # 判断是否为流向类查询。
        if data_type in ["购进流向", "配送流向"]:

            # 写入查询开始日期。
            payload["qdate"] = start_date

            # 写入查询结束日期。
            payload["zdate"] = end_date

        # 返回查询请求体。
        return payload

    # 定义查询商品编码读取方法。
    def get_query_item_code(self):

        """
        读取查询商品编码。

        :return: 查询商品编码。
        :rtype: str
        """

        # 判断查询商品编码是否已经存在。
        if self.query_item_code:

            # 返回已保存的查询商品编码。
            return self.query_item_code

        # 抛出查询商品编码缺失异常。
        raise RuntimeError(f"{self.company_name} 登录响应未返回查询商品编码")

    # 定义查询接口路径构建方法。
    def build_query_path(self, data_type):

        """
        构建查询接口路径。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 查询接口路径。
        :rtype: str
        """

        # 定义数据类型与接口路径映射。
        mapping = {"购进流向": "/baas/fdbus/splx/splxCxjgKhCg", "配送流向": "/baas/fdbus/splx/splxCxjgKhXs", "库存": "/baas/fdbus/splx/splxCxjgKhKc"}

        # 返回当前数据类型查询接口路径。
        return mapping[data_type]

    # 定义查询接口地址构建方法。
    def build_query_api_url(self, data_type):

        """
        构建数据类型对应的查询接口地址。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 查询接口地址。
        :rtype: str
        """

        # 定义数据类型与接口地址映射。
        mapping = {"购进流向": self.purchase_query_url, "配送流向": self.dist_query_url, "库存": self.stock_query_url}

        # 返回当前数据类型查询接口。
        return mapping[data_type]

    # 定义 JSON POST 请求方法。
    def post_json(self, url, payload, referer):

        """
        提交 JSON 请求。

        :param url: 请求地址。
        :type url: str
        :param payload: 请求体字段。
        :type payload: dict
        :param referer: 来源页面地址。
        :type referer: str
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 序列化 JSON 请求体。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建 JSON 请求头。
        headers = self.build_json_headers(referer)

        # 发送 POST 请求。
        return self.http.request("POST", url, headers=headers, data=body, timeout=90)

    # 定义 JSON 请求头构建方法。
    def build_json_headers(self, referer):

        """
        构建 JSON 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: JSON 请求头。
        :rtype: dict
        """

        # 返回 JSON 请求头。
        return {"Accept": "application/json, text/javascript, */*; q=0.01", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/json;charset=UTF-8", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    # 定义表格响应解析方法。
    def decode_table_response(self, result):

        """
        解析平台表格 JSON 响应。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 表格响应数据。
        :rtype: dict
        """

        # 解码响应文本。
        response_text = result.body.decode("utf-8", errors="replace").strip()

        # 处理空响应。
        if not response_text:

            # 返回空表格结构。
            return {"@type": "table", "rows": []}

        # 返回 JSON 响应数据。
        return json.loads(response_text)

    # 定义登录成功判断方法。
    def is_login_success(self, login_result):

        """
        判断登录响应是否成功。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 是否登录成功。
        :rtype: bool
        """

        # 读取登录响应行。
        rows = login_result.get("rows") or []

        # 返回登录成功判断结果。
        return bool(rows)

    # 定义首行数据提取方法。
    def extract_first_row(self, response_data):

        """
        提取响应首行数据。

        :param response_data: 表格响应数据。
        :type response_data: dict
        :return: 标准化首行数据。
        :rtype: dict
        """

        # 提取标准化表格数据行。
        rows = self.extract_table_rows(response_data)

        # 返回首行或空字典。
        return rows[0] if rows else {}

    # 定义表格行提取方法。
    def extract_table_rows(self, response_data):

        """
        提取并标准化表格数据行。

        :param response_data: 表格响应数据。
        :type response_data: dict
        :return: 标准化数据行列表。
        :rtype: list[dict]
        """

        # 读取原始数据行。
        raw_rows = response_data.get("rows") or []

        # 返回标准化数据行列表。
        return [self.normalize_table_row(row) for row in raw_rows]

    # 定义表格行标准化方法。
    def normalize_table_row(self, row):

        """
        标准化平台表格数据行。

        :param row: 原始表格数据行。
        :type row: dict
        :return: 标准化数据行。
        :rtype: dict
        """

        # 初始化标准化数据行。
        normalized_row = {}

        # 遍历原始字段。
        for field_name, field_value in row.items():

            # 跳过 Justep 行状态字段。
            if field_name == "userdata":

                # 进入下一个字段。
                continue

            # 读取单元格包装值。
            if isinstance(field_value, dict) and "value" in field_value:

                # 写入单元格值。
                normalized_row[field_name] = field_value.get("value")

            # 处理普通字段值。
            else:

                # 写入普通字段值。
                normalized_row[field_name] = field_value

        # 返回标准化数据行。
        return normalized_row

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验大连方大药业数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已接入。
        if data_type not in ["购进流向", "配送流向", "库存"]:

            # 抛出数据类型未接入异常。
            raise ValueError(f"大连方大药业暂未接入数据类型：{data_type}")

    # 定义表头构建方法。
    def build_headers(self, data_type):

        """
        构建输出表头配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 构建平台通用表头。
        common_headers = [("记录日期", "addDate"), ("业务类型", "zt"), ("业务日期", "f_date"), ("商品编码", "f_spbh"), ("商品名称", "f_spmc"), ("商品规格", "f_spgg"), ("单位", "f_dw"), ("生产厂家", "f_cd"), ("供应商编码", "f_supplier"), ("商业名称", "f_desc"), ("往来单位", "f_wldw"), ("批号", "f_ph"), ("数量", "f_sl"), ("单价", "f_dj"), ("金额", "f_je"), ("生产日期", "f_scdate"), ("有效期", "f_dqdate"), ("登录用户", "f_user"), ("单据号", "f_djh"), ("平台用户ID", "f_userId"), ("客户状态", "f_status_client"), ("批次定价", "f_pcdjs")]

        # 判断是否为库存查询。
        if data_type == "库存":

            # 返回库存表头。
            return common_headers

        # 返回流向表头。
        return common_headers




    # 定义通用 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(field_name, "") if isinstance(row, dict) else "" for _, field_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 读取当前表头名称。
            header_name = header_item[0]

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error


