# -*- coding: utf-8 -*-
"""
仁和堂内蒙古供应商流向系统协议适配器。

本模块负责仁和堂内蒙古供应商流向系统的验证码登录、分页查询、响应解析和 Excel 文件生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入随机数模块。
import random

# 导入时间模块。
import time

# 导入日期时间解析对象。
from datetime import datetime

# 导入 URL 参数编码函数。
from urllib.parse import urlencode

# 导入 URL 拼接函数。
from urllib.parse import urljoin

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入验证码处理器。
from ddi_framework import CaptchaSolver

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义仁和堂内蒙古平台适配器。
class RhtnmPlatform(BasePlatform):

    """
    仁和堂内蒙古供应商流向系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 仁和堂内蒙古平台适配器实例。
    :rtype: RhtnmPlatform
    """

    # 定义业务类型与接口路径映射。
    DATA_TYPE_ENDPOINTS = {
        "购进流向": "/prod-api/chain/purchase/getDetails",
        "配送流向": "/prod-api/chain/distribution/getDetails",
        "库存": "/prod-api/chain/stock/getDetails",
        "纯销流向": "/prod-api/chain/sale/getDetails",
    }

    # 定义业务类型与平台页面路径映射。
    DATA_TYPE_REFERERS = {
        "购进流向": "/chain/chainpurchase",
        "配送流向": "/chain/chaindistribution",
        "库存": "/chain/chainstock",
        "纯销流向": "/chain/chainsale",
    }

    # 定义购进流向字段映射。
    PURCHASE_HEADERS = [
        ("业务日期", "credate"),
        ("保管账", "storagename"),
        ("货品ID", "goodsid"),
        ("商品名", "goodsname"),
        ("通用名", "currencyname"),
        ("规格", "goodstype"),
        ("厂家名称", "factoryname"),
        ("批号", "lotno"),
        ("采购数量", "goodsqty"),
        ("基本单位", "goodsunit"),
        ("单价", "unitprice"),
        ("金额", "total_line"),
        ("供应商", "supplyname"),
        ("生产日期", "proddate"),
        ("有效期", "invaliddate"),
    ]

    # 定义配送流向字段映射。
    DISTRIBUTION_HEADERS = [
        ("出入库总单ID", "inoutid"),
        ("业务日期", "credate"),
        ("单位ID", "companyid"),
        ("单位名称", "companyname"),
        ("类型", "comefrom"),
        ("出入库状态", "inout"),
        ("数量", "__quantity"),
        ("货品ID", "goodsid"),
        ("通用名", "goodsname"),
        ("产地", "prodarea"),
        ("规格", "goodstype"),
        ("生产厂商", "factoryname"),
        ("批号", "lotno"),
        ("基本单位", "goodsunit"),
        ("生产日期", "proddate"),
        ("有效期", "invaliddate"),
    ]

    # 定义库存字段映射。
    STOCK_HEADERS = [
        ("保管账ID", "storageid"),
        ("保管账", "storagename"),
        ("货品ID", "goodsid"),
        ("货品通用名", "goodsname"),
        ("规格", "goodstype"),
        ("生产厂商", "factoryname"),
        ("批号", "lotno"),
        ("库存数量", "goodsqty"),
        ("基本单位", "goodsunit"),
        ("生产日期", "proddate"),
        ("有效期", "invaliddate"),
    ]

    # 定义纯销流向字段映射。
    SALE_HEADERS = [
        ("区域", "areaname"),
        ("门店ID", "placepointid"),
        ("门店", "placepointname"),
        ("业务日期", "credate"),
        ("销售日期（时分秒）", "credate"),
        ("货品ID", "goodsid"),
        ("商品名", "goodsname"),
        ("规格", "goodstype"),
        ("生产厂商", "factoryname"),
        ("批号", "lotno"),
        ("销售数量", "goodsqty"),
        ("基本单位", "goodsunit"),
        ("销售单价", "unitprice"),
        ("销售金额", "realmoney"),
        ("有效期", "invaliddate"),
    ]

    # 定义业务类型与表头映射。
    DATA_TYPE_HEADERS = {
        "购进流向": PURCHASE_HEADERS,
        "配送流向": DISTRIBUTION_HEADERS,
        "库存": STOCK_HEADERS,
        "纯销流向": SALE_HEADERS,
    }

    # 定义接口原始字段标题映射。
    RAW_FIELD_TITLES = {
        "areaname": "区域名称",
        "area": "区域",
        "placepointid": "门店ID",
        "placepointname": "门店名称",
        "placepoint": "门店",
        "credate": "业务日期",
        "billid": "销售总单号",
        "billdtlid": "销售细单号",
        "goodsid": "货品ID",
        "goodsname": "货品名称",
        "currencyname": "通用名",
        "goodsopercode": "货品操作码",
        "opcode": "货品操作码",
        "goodstype": "规格",
        "factoryname": "生产厂家",
        "prodarea": "产地",
        "lotno": "批号",
        "goodsqty": "库存数量",
        "outqty": "出库数量",
        "inqty": "入库数量",
        "goodsunit": "基本单位",
        "unitprice": "销售单价",
        "realmoney": "销售金额",
        "total_line": "金额",
        "supplyname": "供应商",
        "proddate": "生产日期",
        "invaliddate": "有效期",
        "invaliddays": "距离失效天数",
        "days": "距离失效天数",
        "inoutid": "出入库总单ID",
        "companyid": "单位ID",
        "companyname": "单位名称",
        "comefrom": "类型",
        "inout": "出入库状态",
        "distributionorderid": "调拨单下单ID",
        "placereturndtlid": "调拨单下单ID",
        "iodtlid": "出入库细单ID",
        "sourceid": "销售总单号",
        "rsaid": "销售细单号",
    }

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化仁和堂内蒙古平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://flow.renhetang.cn/login?redirect=/index")

        # 保存平台根地址。
        self.base_url = self.build_origin(self.login_page_url)

        # 保存验证码接口地址。
        self.captcha_api_url = account.get("验证码接口", urljoin(self.base_url, "/prod-api/captchaImage"))

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", urljoin(self.base_url, "/prod-api/login"))

        # 保存默认分页大小。
        self.page_size = int(account.get("分页大小", 200))

        # 初始化登录令牌。
        self.login_token = ""

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化验证码处理器。
        self.captcha_solver = CaptchaSolver()

    # 定义平台登录方法。
    def login(self):

        """
        执行仁和堂内蒙古平台登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("仁和堂内蒙古登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 请求登录页初始化会话 Cookie。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 执行带验证码重试的登录流程。
        return self.login_with_captcha_retry()

    # 定义带验证码重试的登录方法。
    def login_with_captcha_retry(self, max_retries=8):

        """
        使用验证码重试机制执行登录。

        :param max_retries: 最大验证码刷新次数。
        :type max_retries: int
        :return: 登录响应数据。
        :rtype: dict
        """

        # 初始化最后一次失败原因。
        last_error = ""

        # 循环获取验证码并提交登录。
        for retry_index in range(max_retries):

            # 打印验证码尝试日志。
            self.ph.info(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{max_retries} 次", "cyan"))

            # 捕获验证码登录异常。
            try:

                # 获取验证码上下文。
                captcha_context = self.get_captcha_context()

                # 生成验证码候选列表。
                captcha_candidates = self.build_captcha_candidates(captcha_context["image_bytes"])

                # 初始化本轮候选通过状态。
                captcha_accepted = False

                # 遍历当前验证码候选。
                for captcha_code in captcha_candidates:

                    # 打印验证码识别结果。
                    self.ph.info(self.ph.kv("图片验证码候选", captcha_code, "cyan"))

                    # 构建登录请求体。
                    payload = self.build_login_payload(captcha_code, captcha_context["uuid"])

                    # 提交登录请求。
                    login_data = self.post_login(payload)

                    # 判断登录是否成功。
                    if self.is_login_success(login_data):

                        # 保存登录令牌。
                        self.login_token = str(login_data.get("token") or "")

                        # 标记当前验证码已经通过。
                        captcha_accepted = True

                        # 打印登录成功日志。
                        self.ph.success(self.ph.kv("正式登录成功", self.extract_message(login_data), "green"))

                        # 返回登录响应数据。
                        return login_data

                    # 记录登录失败原因。
                    last_error = self.extract_message(login_data)

                    # 打印候选未通过日志。
                    self.ph.warning(self.ph.kv("验证码候选未通过", captcha_code, "yellow"))

                # 判断当前验证码是否未通过。
                if not captcha_accepted:

                    # 记录候选全部未通过状态。
                    last_error = last_error or "验证码候选全部未通过"

            # 处理验证码登录异常。
            except Exception as error:

                # 记录当前异常原因。
                last_error = str(error)

                # 打印验证码处理失败日志。
                self.ph.warning(self.ph.kv("验证码处理未通过", last_error[:120], "yellow"))

            # 最后一轮失败后跳过等待。
            if retry_index >= max_retries - 1:

                # 进入下一轮循环结束判断。
                continue

            # 计算失败后的等待秒数。
            delay_seconds = random.uniform(2.0, 4.0)

            # 打印等待日志。
            self.ph.wait(self.ph.kv("验证码失败后等待，预计", f"{delay_seconds:.1f} 秒", "cyan"))

            # 等待后刷新验证码。
            time.sleep(delay_seconds)

        # 抛出登录失败异常。
        raise RuntimeError(f"{self.company_name} 登录失败：{last_error}")

    # 定义验证码上下文获取方法。
    def get_captcha_context(self):

        """
        获取登录验证码图片与验证码会话编号。

        :return: 验证码上下文字典。
        :rtype: dict
        """

        # 打印验证码请求日志。
        self.ph.info("请求图片验证码: /prod-api/captchaImage")

        # 请求验证码接口。
        result = self.get_json(self.captcha_api_url, self.login_page_url)

        # 解析验证码接口响应。
        data = self.http.decode_json(result)

        # 校验验证码接口响应状态。
        if str(data.get("code", "")) != "200":

            # 抛出验证码获取异常。
            raise RuntimeError(f"{self.company_name} 获取图片验证码失败：{data.get('msg') or data}")

        # 读取验证码图片 Base64 文本。
        image_base64 = str(data.get("img") or "")

        # 解码验证码图片。
        image_bytes = self.captcha_solver.decode_base64_image(image_base64)

        # 返回验证码上下文。
        return {"uuid": str(data.get("uuid") or ""), "image_bytes": image_bytes}

    # 定义验证码候选构建方法。
    def build_captcha_candidates(self, image_bytes):

        """
        构建验证码候选文本列表。

        :param image_bytes: 验证码图片字节流。
        :type image_bytes: bytes
        :return: 验证码候选文本列表。
        :rtype: list[str]
        """

        # 构建单数字操作数算式验证码答案候选。
        math_codes = self.captcha_solver.build_math_code_candidates(image_bytes, right_operand_max_digits=1)

        # 返回非空候选答案列表。
        if math_codes:

            # 返回候选答案列表。
            return math_codes[:3]

        # 返回空候选占位。
        return [""]

    # 定义登录请求体构建方法。
    def build_login_payload(self, captcha_code, captcha_uuid):

        """
        构建登录 JSON 请求体。

        :param captcha_code: 图片验证码。
        :type captcha_code: str
        :param captcha_uuid: 验证码会话编号。
        :type captcha_uuid: str
        :return: 登录请求体。
        :rtype: dict
        """

        # 返回登录请求字段。
        return {"username": self.clean_login_name(), "password": self.clean_login_password(), "code": captcha_code, "uuid": captcha_uuid}

    # 定义登录请求提交方法。
    def post_login(self, payload):

        """
        提交登录请求。

        :param payload: 登录请求体。
        :type payload: dict
        :return: 登录响应数据。
        :rtype: dict
        """

        # 序列化登录请求体。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建登录请求头。
        headers = self.build_json_headers(self.login_page_url, with_token=False)

        # 打印正式登录请求日志。
        self.ph.info("提交正式登录请求: /prod-api/login")

        # 发送登录请求。
        result = self.http.request("POST", self.login_api_url, headers=headers, data=body)

        # 返回登录 JSON 响应。
        return self.http.decode_json(result)


    # 定义查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果数据行。
        :rtype: list[dict]
        """

        # 查询指定业务数据。
        query_result = self.query_rhtnm_data(data_type, start_date, end_date)

        # 返回查询结果行。
        return query_result["rows"]


    # 定义查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定业务数据。
        query_result = self.query_rhtnm_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建输出文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建输出文件路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回输出文件路径。
        return output_path

    # 定义业务数据查询方法。
    def query_rhtnm_data(self, data_type, start_date, end_date):

        """
        查询仁和堂内蒙古指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 查询全部分页数据。
        rows = self.query_all_pages(data_type, start_date, end_date)

        # 构建业务输出表头。
        headers = self.build_output_headers(data_type, rows)

        # 打印查询解析日志。
        self.ph.success(self.ph.join(self.ph.kv("查询解析完成", f"{len(rows)} 条", "green"), self.ph.kv("表头数", len(headers), "cyan")))

        # 返回查询数据与表头。
        return {"rows": rows, "headers": headers}

    # 定义业务输出表头构建方法。
    def build_output_headers(self, data_type, rows):

        """
        构建固定业务字段与接口原始字段合并后的输出表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :return: 输出表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 复制当前数据类型固定表头。
        headers = list(self.DATA_TYPE_HEADERS[data_type])

        # 读取已存在字段集合。
        existing_fields = {field_name for _, field_name in headers}

        # 遍历查询结果数据行。
        for row in rows:

            # 跳过非字典数据行。
            if not isinstance(row, dict):

                # 进入下一条数据行。
                continue

            # 遍历当前数据行字段。
            for field_name in row.keys():

                # 读取字段名称文本。
                field_key = str(field_name or "").strip()

                # 跳过空字段名称。
                if not field_key:

                    # 进入下一项字段。
                    continue

                # 跳过已经输出的字段。
                if field_key in existing_fields:

                    # 进入下一项字段。
                    continue

                # 添加字段到去重集合。
                existing_fields.add(field_key)

                # 追加原始字段输出表头。
                headers.append((self.RAW_FIELD_TITLES.get(field_key, field_key), field_key))

        # 返回合并后的输出表头。
        return headers

    # 定义全部分页查询方法。
    def query_all_pages(self, data_type, start_date, end_date):

        """
        查询指定业务类型的全部分页数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果数据行。
        :rtype: list[dict]
        """

        # 初始化分页页码。
        page_number = 1

        # 初始化全部数据行。
        all_rows = []

        # 初始化平台总数。
        total_count = None

        # 循环读取分页数据。
        while True:

            # 请求当前分页。
            page_result = self.query_page(data_type, start_date, end_date, page_number)

            # 读取当前页数据行。
            page_rows = page_result["rows"]

            # 读取平台总数。
            total_count = page_result["total"] if total_count is None else total_count

            # 累加当前页数据行。
            all_rows.extend(page_rows)

            # 打印当前分页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(page_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count if total_count is not None else "未知", "cyan")))

            # 判断当前页是否已经到达末页。
            if not page_rows or len(page_rows) < self.page_size:

                # 返回全部数据行。
                return all_rows

            # 判断累计数据是否达到平台总数。
            if total_count is not None and len(all_rows) >= total_count:

                # 返回全部数据行。
                return all_rows

            # 递增分页页码。
            page_number += 1

    # 定义分页查询方法。
    def query_page(self, data_type, start_date, end_date, page_number):

        """
        查询指定业务类型的单页数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param page_number: 页码。
        :type page_number: int
        :return: 分页数据与总数。
        :rtype: dict
        """

        # 构建查询参数。
        params = {"pageNo": page_number, "pageSize": self.page_size, "startDate": start_date or "", "endDate": end_date or ""}

        # 构建查询地址。
        query_url = self.build_query_url(data_type, params)

        # 构建来源页地址。
        referer = urljoin(self.base_url, self.DATA_TYPE_REFERERS[data_type])

        # 打印分页请求日志。
        self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"pageNo={page_number}", "cyan"), self.ph.kv("pageSize", self.page_size, "cyan")))

        # 发送授权查询请求。
        result = self.get_json(query_url, referer, with_token=True)

        # 解析查询响应。
        response_data = self.http.decode_json(result)

        # 校验查询响应。
        if str(response_data.get("code", "")) != "200":

            # 抛出查询失败异常。
            raise RuntimeError(f"{self.company_name} {data_type}查询失败：{self.extract_message(response_data)}")

        # 返回分页解析结果。
        return self.extract_page_result(response_data)

    # 定义查询地址构建方法。
    def build_query_url(self, data_type, params):

        """
        构建分页查询地址。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param params: 查询参数。
        :type params: dict
        :return: 分页查询地址。
        :rtype: str
        """

        # 读取接口路径。
        endpoint = self.DATA_TYPE_ENDPOINTS[data_type]

        # 返回完整查询地址。
        return f"{urljoin(self.base_url, endpoint)}?{urlencode(params)}"

    # 定义分页响应解析方法。
    def extract_page_result(self, response_data):

        """
        解析分页查询响应。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :return: 分页数据与总数。
        :rtype: dict
        """

        # 读取分页根对象。
        page_data = response_data.get("data") if isinstance(response_data.get("data"), dict) else {}

        # 读取明细数据对象。
        detail_data = page_data.get("data") if isinstance(page_data.get("data"), dict) else {}

        # 读取明细列表。
        rows = detail_data.get("list") if isinstance(detail_data.get("list"), list) else []

        # 读取平台总数。
        total_value = page_data.get("total", detail_data.get("totalNum"))

        # 返回分页解析结果。
        return {"rows": rows, "total": self.safe_int(total_value)}

    # 定义授权 GET 请求方法。
    def get_json(self, url, referer, with_token=False):

        """
        发送 JSON GET 请求。

        :param url: 请求地址。
        :type url: str
        :param referer: 来源页地址。
        :type referer: str
        :param with_token: 是否携带登录令牌。
        :type with_token: bool
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 构建 JSON 请求头。
        headers = self.build_json_headers(referer, with_token=with_token)

        # 发送 GET 请求。
        return self.http.request("GET", url, headers=headers, timeout=120)

    # 定义 JSON 请求头构建方法。
    def build_json_headers(self, referer, with_token=False):

        """
        构建 JSON 请求头。

        :param referer: 来源页地址。
        :type referer: str
        :param with_token: 是否携带登录令牌。
        :type with_token: bool
        :return: 请求头。
        :rtype: dict
        """

        # 构建基础 JSON 请求头。
        headers = {"Accept": "application/json, text/plain, */*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/json;charset=UTF-8", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent()}

        # 判断是否需要登录令牌。
        if with_token and self.login_token:

            # 写入授权令牌。
            headers["Authorization"] = f"Bearer {self.login_token}"

        # 返回请求头。
        return headers

    # 定义登录成功判断方法。
    def is_login_success(self, response_data):

        """
        判断登录响应是否成功。

        :param response_data: 登录响应数据。
        :type response_data: dict
        :return: 是否登录成功。
        :rtype: bool
        """

        # 返回登录成功判断结果。
        return str(response_data.get("code", "")) == "200" and bool(response_data.get("token"))

    # 定义响应消息提取方法。
    def extract_message(self, response_data):

        """
        提取响应消息文本。

        :param response_data: 响应数据。
        :type response_data: dict
        :return: 响应消息文本。
        :rtype: str
        """

        # 返回常见消息字段或响应摘要。
        return str(response_data.get("msg") or response_data.get("message") or response_data)

    # 定义登录状态校验方法。
    def ensure_login(self):

        """
        确保当前适配器已经完成登录。

        :return: None。
        :rtype: None
        """

        # 判断当前会话是否已经登录。
        if self.login_token:

            # 直接返回已登录状态。
            return

        # 执行登录流程。
        self.login()

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验数据类型是否已接入。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已经接入。
        if data_type not in self.DATA_TYPE_ENDPOINTS:

            # 抛出数据类型未接入异常。
            raise ValueError(f"仁和堂内蒙古暂未接入数据类型：{data_type}")



    # 定义安全整数转换方法。
    def safe_int(self, value):

        """
        安全转换整数值。

        :param value: 原始值。
        :type value: object
        :return: 整数值或空值。
        :rtype: int | None
        """

        # 捕获整数转换异常。
        try:

            # 返回整数值。
            return int(value)

        # 处理无法转换的值。
        except (TypeError, ValueError):

            # 返回空值。
            return None




    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([self.resolve_export_value(row, field_name) for _, field_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 读取当前表头名称。
            header_name = header_item[0]

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        workbook.save(output_path)


    # 定义导出字段取值方法。
    def resolve_export_value(self, row, field_name):

        """
        解析导出字段值。

        :param row: 查询结果数据行。
        :type row: dict
        :param field_name: 字段名称。
        :type field_name: str
        :return: 导出字段值。
        :rtype: object
        """

        # 校验数据行类型。
        if not isinstance(row, dict):

            # 返回空字段值。
            return ""

        # 判断是否为合并数量字段。
        if field_name == "__quantity":

            # 返回出库数量或入库数量。
            return row.get("outqty") or row.get("inqty") or row.get("goodsqty") or ""

        # 返回普通字段值。
        return row.get(field_name, "")

