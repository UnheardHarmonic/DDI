# -*- coding: utf-8 -*-
"""
黑龙江省大众平安医药连锁有限公司流向平台协议适配器。

本模块负责黑龙江省大众平安医药连锁有限公司流向平台的验证码登录、令牌保存、分页查询和 Excel 落地。
"""

# 导入 Base64 编解码模块。
import base64

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入随机数模块。
import random

# 导入时间模块。
import time

# 导入日期时间对象。
from datetime import datetime

# 导入 URL 查询参数编码函数。
from urllib.parse import urlencode

# 导入验证码处理器。
from ddi_framework import CaptchaSolver

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义黑龙江大众平安流向平台协议适配器。
class HljdzpaPlatform(BasePlatform):

    """
    黑龙江大众平安流向平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 黑龙江大众平安平台适配器实例。
    :rtype: HljdzpaPlatform
    """

    # 定义黑龙江大众平安适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化黑龙江大众平安平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://123.57.178.31:58066/#/login")

        # 保存接口根地址。
        self.api_base_url = account.get("接口根地址", "http://123.57.178.31:8084")

        # 保存验证码接口地址。
        self.captcha_api_url = account.get("验证码接口", f"{self.api_base_url}/sys/getCode")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.api_base_url}/sys/login")

        # 保存查询接口地址。
        self.query_api_url = account.get("查询接口", f"{self.api_base_url}/form/generate/list")

        # 保存平台页面根地址。
        self.base_url = self.build_origin(self.login_page_url)

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.api_base_url)

        # 初始化验证码处理器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录响应。
        self.login_result = {}

        # 初始化登录令牌。
        self.login_token = ""

        # 保存默认分页大小。
        self.page_size = int(account.get("分页大小", 1000))

    # 定义黑龙江大众平安登录方法。
    def login(self):

        """
        执行黑龙江大众平安协议登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("黑龙江大众平安登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 请求登录页并初始化站点 Cookie。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 执行带验证码重试的登录流程。
        return self.login_with_captcha_retry()

    # 定义带验证码重试的登录方法。
    def login_with_captcha_retry(self, max_retries=8):

        """
        使用验证码重试机制执行登录。

        :param max_retries: 最大验证码刷新次数。
        :type max_retries: int
        :return: 登录结果。
        :rtype: dict
        """

        # 初始化最后一次失败原因。
        last_error = ""

        # 循环获取验证码并提交登录。
        for retry_index in range(max_retries):

            # 打印验证码尝试日志。
            self.ph.info(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{max_retries} 次", "cyan"))

            # 捕获当前验证码登录异常。
            try:

                # 获取验证码上下文。
                captcha_context = self.get_captcha_context()

                # 识别验证码答案。
                captcha_code = self.solve_captcha_code(captcha_context["image_bytes"])

                # 打印验证码识别完成日志。
                self.ph.info(self.ph.kv("图片验证码识别完成", f"长度={len(captcha_code)}", "cyan"))

                # 构建登录请求体。
                payload = self.build_login_payload(captcha_context["uuid"], captcha_code)

                # 打印正式登录请求日志。
                self.ph.info("提交正式登录请求: /sys/login")

                # 提交登录请求。
                result = self.post_login(payload)

                # 解析登录响应。
                login_result = self.decode_json_response(result)

                # 判断登录响应是否成功。
                if self.is_login_success(login_result):

                    # 保存登录响应。
                    self.login_result = login_result

                    # 保存登录令牌。
                    self.login_token = self.extract_login_token(login_result)

                    # 打印登录成功日志。
                    self.ph.success(self.ph.kv("正式登录成功", "token", "green"))

                    # 返回登录响应。
                    return login_result

                # 记录登录失败原因。
                last_error = self.extract_response_message(login_result)

                # 打印验证码或登录失败日志。
                self.ph.warning("登录校验未通过，准备刷新验证码")

            # 处理验证码识别或登录提交异常。
            except Exception as error:

                # 记录当前异常原因。
                last_error = str(error)

                # 打印验证码处理失败日志。
                self.ph.warning("图片验证码识别或登录校验未通过，准备刷新验证码")

            # 最后一轮失败后跳过等待。
            if retry_index >= max_retries - 1:

                # 进入失败处理。
                continue

            # 计算失败后的等待秒数。
            delay_seconds = random.uniform(2.2, 4.4)

            # 打印等待日志。
            self.ph.wait(self.ph.kv("验证码失败后等待，预计", f"{delay_seconds:.1f} 秒", "cyan"))

            # 等待后再刷新验证码。
            time.sleep(delay_seconds)

        # 抛出登录失败异常。
        raise RuntimeError(f"{self.company_name} 登录失败：{last_error}")

    # 定义验证码上下文获取方法。
    def get_captcha_context(self):

        """
        获取登录验证码图片与验证码会话编号。

        :return: 验证码上下文字典。
        :rtype: dict
        """

        # 打印验证码请求日志。
        self.ph.info("请求图片验证码: /sys/getCode")

        # 请求验证码接口。
        result = self.http.request("GET", self.captcha_api_url, headers=self.build_json_headers(self.login_page_url))

        # 解析验证码响应。
        response_data = self.decode_json_response(result)

        # 读取验证码图片文本。
        image_text = str(response_data.get("codeImg") or "")

        # 校验验证码图片文本。
        if not image_text:

            # 抛出验证码图片缺失异常。
            raise RuntimeError(f"验证码接口未返回 codeImg：{response_data}")

        # 解码验证码图片字节。
        image_bytes = base64.b64decode(image_text)

        # 读取验证码会话编号。
        uuid_value = str(response_data.get("uuid") or "")

        # 校验验证码会话编号。
        if not uuid_value:

            # 抛出验证码会话编号缺失异常。
            raise RuntimeError(f"验证码接口未返回 uuid：{response_data}")

        # 返回验证码上下文。
        return {"image_bytes": image_bytes, "uuid": uuid_value}

    # 定义验证码识别方法。
    def solve_captcha_code(self, image_bytes):

        """
        识别图片验证码文本。

        :param image_bytes: 验证码图片字节。
        :type image_bytes: bytes
        :return: 验证码文本。
        :rtype: str
        """

        # 使用 OCR 识别验证码。
        captcha_code = self.captcha_solver.solve_text_image(image_bytes, min_length=4)

        # 校验验证码识别结果。
        if not captcha_code:

            # 抛出验证码识别失败异常。
            raise RuntimeError("图片验证码识别为空")

        # 返回验证码识别结果。
        return captcha_code

    # 定义登录请求体构建方法。
    def build_login_payload(self, uuid_value, captcha_code):

        """
        构建黑龙江大众平安登录请求体。

        :param uuid_value: 验证码会话编号。
        :type uuid_value: str
        :param captcha_code: 验证码文本。
        :type captcha_code: str
        :return: 登录请求体。
        :rtype: dict
        """

        # 返回登录 JSON 字段。
        return {"username": self.clean_login_name(), "password": self.clean_login_password(), "uuid": uuid_value, "code": captcha_code}

    # 定义登录提交方法。
    def post_login(self, payload):

        """
        提交黑龙江大众平安登录请求。

        :param payload: 登录请求体。
        :type payload: dict
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 序列化登录请求体。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建登录请求头。
        headers = self.build_json_headers(self.login_page_url)

        # 发送登录请求。
        return self.http.request("POST", self.login_api_url, headers=headers, data=body)

    # 定义登录成功判断方法。
    def is_login_success(self, login_result):

        """
        判断登录响应是否成功。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 是否登录成功。
        :rtype: bool
        """

        # 返回令牌字段存在状态。
        return bool(login_result.get("token"))

    # 定义登录令牌提取方法。
    def extract_login_token(self, login_result):

        """
        提取登录令牌。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 登录令牌。
        :rtype: str
        """

        # 返回登录响应中的令牌字段。
        return str(login_result.get("token") or "")

    # 定义响应消息提取方法。
    def extract_response_message(self, response_data):

        """
        提取接口响应消息。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 接口响应消息。
        :rtype: str
        """

        # 返回接口消息或原始响应摘要。
        return str(response_data.get("msg") or response_data.get("message") or response_data)



    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie 与认证令牌。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 保存认证令牌文件。
        self.save_token_file(output_dir)

        # 打印 Cookie 保存日志。
        self.ph.success(self.ph.kv("登录 Cookie 已保存", self.ph.path(cookie_path), "green"))

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义令牌保存方法。
    def save_token_file(self, output_dir):

        """
        保存登录令牌到本地文件。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: 令牌文件路径。
        :rtype: str
        """

        # 构建令牌文件路径。
        token_path = os.path.join(output_dir, f"{self.company_name}_token.json")

        # 构建令牌文件内容。
        token_data = {"token": self.login_token, "login_result": self.login_result}

        # 写入令牌文件。
        with open(token_path, "w", encoding="utf-8") as file_obj:

            # 序列化令牌文件内容。
            json.dump(token_data, file_obj, ensure_ascii=False, indent=2)

        # 打印令牌保存日志。
        self.ph.success(self.ph.kv("登录令牌已保存", self.ph.path(token_path), "green"))

        # 返回令牌文件路径。
        return token_path

    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询黑龙江大众平安业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行。
        :rtype: list[dict]
        """

        # 查询指定类型完整数据。
        query_result = self.query_platform_data(data_type, start_date, end_date)

        # 返回查询结果行。
        return query_result["rows"]


    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询黑龙江大众平安业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定业务数据。
        query_result = self.query_platform_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义平台数据查询方法。
    def query_platform_data(self, data_type, start_date, end_date):

        """
        查询黑龙江大众平安指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 确保当前会话已经登录。
        self.ensure_login()

        # 解析查询月份。
        month_value = self.resolve_query_month(start_date, end_date)

        # 打印查询阶段标题。
        self.ph.section(f"{data_type}查询", self.ph.kv("月份", month_value, "cyan"))

        # 初始化当前页码。
        current_page = 1

        # 初始化累计数据行。
        all_rows = []

        # 初始化总记录数。
        total_count = None

        # 按页循环请求业务数据。
        while True:

            # 构建当前页查询参数。
            params = self.build_query_params(data_type, month_value, current_page)

            # 打印分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"current={current_page}", "cyan"), self.ph.kv("size", params["size"], "yellow")))

            # 发送当前页查询请求。
            result = self.get_query_page_with_retry(data_type, params, current_page)

            # 解析当前页查询响应。
            page_result = self.decode_query_page(result)

            # 读取当前页数据行。
            page_rows = page_result["rows"]

            # 保存总记录数。
            total_count = page_result["total"]

            # 追加当前页数据行。
            all_rows.extend(page_rows)

            # 打印当前页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(page_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count, "yellow")))

            # 判断是否已经完成全部分页。
            if len(all_rows) >= int(total_count or 0) or not page_rows:

                # 结束分页循环。
                break

            # 递增页码。
            current_page += 1

        # 返回查询数据与表头。
        return {"rows": all_rows, "headers": self.build_headers(data_type)}

    # 定义登录状态校验方法。
    def ensure_login(self):

        """
        确保当前适配器已经完成登录。

        :return: None。
        :rtype: None
        """

        # 判断登录令牌是否已经存在。
        if self.login_token:

            # 已登录时直接返回。
            return

        # 执行登录流程。
        self.login()

    # 定义查询参数构建方法。
    def build_query_params(self, data_type, month_value, current_page):

        """
        构建分页查询参数。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param month_value: 查询月份。
        :type month_value: str
        :param current_page: 当前页码。
        :type current_page: int
        :return: 查询参数。
        :rtype: dict
        """

        # 构建基础查询参数。
        params = {"formId": self.get_form_id(data_type), "params": self.build_filter_payload(data_type, month_value), "current": current_page, "size": self.page_size}

        # 返回查询参数。
        return params

    # 定义过滤参数构建方法。
    def build_filter_payload(self, data_type, month_value):

        """
        构建平台查询过滤 JSON 字符串。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param month_value: 查询月份。
        :type month_value: str
        :return: 平台过滤 JSON 字符串。
        :rtype: str
        """

        # 判断当前数据类型是否为库存。
        if data_type == "库存":

            # 返回库存无月份过滤参数。
            return "{}"

        # 返回月份过滤参数。
        return json.dumps({"period": month_value}, ensure_ascii=False, separators=(",", ":"))

    # 定义查询月份解析方法。
    def resolve_query_month(self, start_date, end_date):

        """
        解析平台查询月份。

        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 月份数字文本。
        :rtype: str
        """

        # 读取优先使用的日期文本。
        date_text = start_date or end_date or datetime.now().strftime("%Y-%m-%d")

        # 解析标准日期中的月份。
        if "-" in str(date_text):

            # 返回标准日期中的月份数字。
            return str(datetime.strptime(str(date_text), "%Y-%m-%d").month)

        # 返回直接传入的月份数字文本。
        return str(date_text).strip()

    # 定义查询页请求方法。
    def get_query_page(self, params):

        """
        请求平台分页查询接口。

        :param params: 查询参数。
        :type params: dict
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 编码查询字符串。
        query_string = urlencode(params)

        # 拼接分页查询地址。
        request_url = f"{self.query_api_url}?{query_string}"

        # 构建查询请求头。
        headers = self.build_query_headers()

        # 发送分页查询请求。
        return self.http.request("GET", request_url, headers=headers, timeout=120)

    # 定义查询页重试请求方法。
    def get_query_page_with_retry(self, data_type, params, current_page, max_retries=3):

        """
        使用重试机制请求平台分页查询接口。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param params: 查询参数。
        :type params: dict
        :param current_page: 当前页码。
        :type current_page: int
        :param max_retries: 最大请求次数。
        :type max_retries: int
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 初始化最后一次异常信息。
        last_error = None

        # 按最大次数请求当前分页。
        for retry_index in range(max_retries):

            # 捕获当前分页请求异常。
            try:

                # 返回当前分页请求结果。
                return self.get_query_page(params)

            # 处理当前分页请求异常。
            except RuntimeError as error:

                # 保存当前异常信息。
                last_error = error

                # 判断当前请求是否已经达到最大次数。
                if retry_index >= max_retries - 1:

                    # 结束重试循环。
                    break

                # 计算分页失败后的等待秒数。
                delay_seconds = random.uniform(2.2, 4.4)

                # 打印分页重试等待日志。
                self.ph.wait(self.ph.join(self.ph.kv(f"{data_type}第{current_page}页请求失败，准备重试", f"第 {retry_index + 2}/{max_retries} 次", "yellow"), self.ph.kv("等待", f"{delay_seconds:.1f} 秒", "cyan")))

                # 等待后重新请求当前分页。
                time.sleep(delay_seconds)

        # 抛出最后一次分页请求异常。
        raise last_error

    # 定义查询响应解析方法。
    def decode_query_page(self, result):

        """
        解析分页查询响应。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 标准分页结果。
        :rtype: dict
        """

        # 解析响应 JSON。
        response_data = self.decode_json_response(result)

        # 读取当前页记录。
        rows = response_data.get("records") or []

        # 读取总记录数。
        total = int(response_data.get("total") or 0)

        # 返回标准分页结果。
        return {"rows": rows, "total": total}

    # 定义数据类型对应表单编号读取方法。
    def get_form_id(self, data_type):

        """
        读取数据类型对应的平台表单编号。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 平台表单编号。
        :rtype: str
        """

        # 读取表单编号配置。
        form_ids = self.account.get("表单编号", {})

        # 返回当前数据类型表单编号。
        return str(form_ids[data_type])

    # 定义查询请求头构建方法。
    def build_query_headers(self):

        """
        构建分页查询请求头。

        :return: 分页查询请求头。
        :rtype: dict
        """

        # 构建 JSON 请求头。
        headers = self.build_json_headers(f"{self.base_url}/#/form/GenerateList")

        # 写入平台令牌请求头。
        headers["token"] = self.login_token

        # 返回分页查询请求头。
        return headers

    # 定义 JSON 请求头构建方法。
    def build_json_headers(self, referer):

        """
        构建平台 JSON 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: JSON 请求头。
        :rtype: dict
        """

        # 返回平台 JSON 请求头。
        return {"Accept": "application/json, text/plain, */*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/json;charset=UTF-8", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent()}

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应文本。
        response_text = result.body.decode("utf-8", errors="replace").strip()

        # 处理空响应。
        if not response_text:

            # 返回空响应错误结构。
            return {"message": "接口返回空响应"}

        # 返回 JSON 响应数据。
        return json.loads(response_text)

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验黑龙江大众平安数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已接入。
        if data_type not in ["购进流向", "配送流向", "库存"]:

            # 抛出数据类型未接入异常。
            raise NotImplementedError(f"当前黑龙江大众平安暂未接入: {data_type}")

    # 定义表头构建方法。
    def build_headers(self, data_type):

        """
        构建输出表头配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 判断是否为购进流向。
        if data_type == "购进流向":

            # 返回购进流向表头配置。
            return [("单据类型", "billcode"), ("单据编号", "billno"), ("业务机构", "busno"), ("执行日期", "execdate"), ("供应商编号", "vencusno"), ("供应商编码", "vencuscode"), ("供应商名称", "vencusname"), ("商品ID", "wareid"), ("商品编码", "warecode"), ("商品名称", "warename"), ("规格", "warespec"), ("单位", "wareunit"), ("生产厂家ID", "factoryid"), ("生产厂家", "factoryname"), ("批号ID", "batid"), ("批号", "makeno"), ("生产日期", "makedate"), ("有效期", "invalidate"), ("进价", "purprice"), ("税率", "purtax"), ("数量", "wareqty"), ("金额", "amt"), ("期间", "period"), ("机构名称", "orgname"), ("月份", "month_")]

        # 判断是否为配送流向。
        if data_type == "配送流向":

            # 返回配送流向表头配置。
            return [("单据类型", "billcode"), ("单据编号", "billno"), ("执行日期", "execdate"), ("来源机构", "srcbusno"), ("业务机构", "busno"), ("机构名称", "orgname"), ("供应商编号", "vencusno"), ("供应商编码", "vencuscode"), ("供应商名称", "vencusname"), ("商品ID", "wareid"), ("商品编码", "warecode"), ("商品名称", "warename"), ("规格", "warespec"), ("单位", "wareunit"), ("生产厂家ID", "factoryid"), ("生产厂家", "factoryname"), ("批号ID", "batid"), ("批号", "makeno"), ("生产日期", "makedate"), ("有效期", "invalidate"), ("数量", "wareqty"), ("配送价", "distprice"), ("金额", "amt"), ("期间", "period"), ("月份", "month_")]

        # 返回库存表头配置。
        return [("创建时间", "createtime"), ("业务机构", "busno"), ("机构名称", "orgname"), ("供应商编号", "vencusno"), ("供应商编码", "vencuscode"), ("供应商名称", "vencusname"), ("商品ID", "wareid"), ("商品编码", "warecode"), ("商品名称", "warename"), ("规格", "warespec"), ("单位", "wareunit"), ("生产厂家ID", "factoryid"), ("生产厂家", "factoryname"), ("批号ID", "batid"), ("批号", "makeno"), ("数量", "wareqty"), ("进价", "purprice"), ("金额", "amt"), ("生产日期", "makedate"), ("有效期", "invalidate")]




    # 定义通用 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(field_name, "") if isinstance(row, dict) else "" for _, field_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 读取当前表头名称。
            header_name = header_item[0]

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

    # 定义同名文件删除方法。
    def remove_existing_file(self, output_path):

        """
        删除已有同名输出文件。

        :param output_path: 输出文件路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 判断目标文件是否存在。
        if not os.path.exists(output_path):

            # 无需删除不存在的文件。
            return

        # 删除已存在的输出文件。
        os.remove(output_path)

    # 定义来源地址构建方法。
    def build_origin(self, url):

        """
        从页面地址构建协议来源地址。

        :param url: 页面地址。
        :type url: str
        :return: 协议来源地址。
        :rtype: str
        """

        # 导入 URL 解析函数。
        from urllib.parse import urlparse

        # 解析页面地址。
        parsed = urlparse(url)

        # 返回协议来源地址。
        return f"{parsed.scheme}://{parsed.netloc}"
