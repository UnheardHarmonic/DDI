# -*- coding: utf-8 -*-
"""
山西天泽诚信协同办公流向查询系统协议适配器。

本模块负责山西天泽诚信协同办公平台的登录、业务查询、响应解析和 Excel 文件生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入随机数模块。
import random

# 导入字符串常量模块。
import string

# 导入日期时间解析对象。
from datetime import datetime

# 导入 URL 编码函数。
from urllib.parse import urlencode

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义山西天泽诚信平台适配器。
class SxtzcxPlatform(BasePlatform):

    """
    山西天泽诚信协同办公平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 山西天泽诚信平台适配器实例。
    :rtype: SxtzcxPlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化山西天泽诚信平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://sxtzcx.xin:8089/login")

        # 保存站点根地址。
        self.base_url = self.build_origin(self.login_page_url)

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.base_url}/loginCheck")

        # 保存查询页面地址。
        self.query_page_url = account.get("查询页", f"{self.base_url}/Design_index?gnbh=1093")

        # 保存查询接口地址。
        self.query_api_url = account.get("查询接口", f"{self.base_url}/designquerydata")

        # 保存默认分页大小。
        self.page_size = int(account.get("分页大小", 5000))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录响应。
        self.login_result = {}

        # 初始化用户上下文。
        self.user_info = {}

        # 初始化功能设计元素。
        self.design_elements = []

        # 初始化登录状态。
        self.logged_in = False

    # 定义平台登录方法。
    def login(self):

        """
        执行山西天泽诚信平台登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("山西天泽诚信登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 请求登录页初始化会话 Cookie。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 构建登录请求体。
        payload = self.build_login_payload()

        # 打印登录账号日志。
        self.ph.info(self.ph.kv("登录账号", self.clean_login_name(), "cyan"))

        # 提交登录请求。
        result = self.post_form(self.login_api_url, payload, self.login_page_url)

        # 解析登录响应。
        login_data = self.decode_json_response(result)

        # 校验登录结果。
        if not self.is_login_success(login_data):

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 正式登录失败：{self.extract_message(login_data)}")

        # 保存登录响应数据。
        self.login_result = login_data

        # 标记当前会话已登录。
        self.logged_in = True

        # 初始化登录后的用户组织上下文。
        self.initialize_user_context()

        # 初始化流向查询功能设计配置。
        self.initialize_design_config()

        # 打印登录成功日志。
        self.ph.success(self.ph.kv("正式登录成功", self.extract_message(login_data) or "登录成功", "green"))

        # 返回登录响应数据。
        return login_data

    # 定义登录请求体构建方法。
    def build_login_payload(self):

        """
        构建登录表单请求体。

        :return: 登录表单字段。
        :rtype: dict
        """

        # 返回登录接口字段。
        return {"username": self.clean_login_name(), "password": self.clean_login_password(), "isValidatecode": "0", "validatecode": "", "isphone": "0", "uniqueIdentifier": self.build_unique_identifier(), "openId": ""}

    # 定义唯一标识构建方法。
    def build_unique_identifier(self):

        """
        构建登录请求唯一标识。

        :return: 前端登录请求唯一标识。
        :rtype: str
        """

        # 读取十六进制字符集合。
        alphabet = string.hexdigits.upper()[:16]

        # 返回十六位随机十六进制文本。
        return "".join(random.choice(alphabet) for _ in range(16))

    # 定义登录后查询页加载方法。
    def load_query_page(self):

        """
        加载流向查询页面。

        :return: 查询页面响应。
        :rtype: HttpResult
        """

        # 请求查询页面初始化功能上下文。
        return self.http.get(self.query_page_url, referer=self.login_page_url)

    # 定义用户上下文初始化方法。
    def initialize_user_context(self):

        """
        初始化登录后的用户组织上下文。

        :return: None。
        :rtype: None
        """

        # 构建登录后主页地址。
        home_url = f"{self.base_url}/?id={self.login_result.get('session', '')}&pcStatus=Y"

        # 请求登录后主页。
        home_result = self.http.get(home_url, referer=self.login_page_url)

        # 解码登录后主页文本。
        home_text = home_result.body.decode("utf-8", errors="replace")

        # 从主页脚本中提取基础用户上下文。
        self.user_info = self.extract_home_user_info(home_text)

        # 请求用户组织列表。
        org_result = self.post_form(f"{self.base_url}/getUserOrg", {"userId": self.user_info["userId"], "entId": self.user_info["entId"], "pcStatus": "Y"}, home_url)

        # 解析用户组织响应。
        org_data = self.decode_json_response(org_result)

        # 读取默认组织行。
        org_row = (org_data.get("list") or [{}])[0]

        # 合并默认组织信息。
        self.user_info.update(org_row)

        # 写入锁屏状态字段。
        self.user_info["askPassOpen"] = False

        # 提交当前组织上下文到服务端会话。
        set_result = self.post_form(f"{self.base_url}/setUserOrg", self.user_info, home_url)

        # 解析组织设置响应。
        set_data = self.decode_json_response(set_result)

        # 校验组织上下文设置结果。
        if int(set_data.get("result", 0)) != 1:

            # 抛出组织上下文设置失败异常。
            raise RuntimeError(f"{self.company_name} 组织上下文设置失败：{self.extract_message(set_data)}")

    # 定义主页用户上下文提取方法。
    def extract_home_user_info(self, home_text):

        """
        从主页脚本中提取用户上下文。

        :param home_text: 登录后主页 HTML。
        :type home_text: str
        :return: 用户上下文字段。
        :rtype: dict
        """

        # 返回主页脚本中的基础用户上下文。
        return {"userId": self.extract_script_const(home_text, "g_userId"), "entId": self.extract_script_const(home_text, "g_entId"), "userName": self.extract_script_const(home_text, "g_userName"), "userRole": self.extract_script_const(home_text, "g_userRole"), "userType": self.extract_script_const(home_text, "g_userType"), "supeEntId": "", "supeOrgId": "", "mobile": "N", "dateTime": str(int(datetime.now().timestamp() * 1000))}

    # 定义脚本常量提取方法。
    def extract_script_const(self, html_text, const_name):

        """
        从 HTML 脚本文本中提取常量值。

        :param html_text: HTML 文本。
        :type html_text: str
        :param const_name: 常量名称。
        :type const_name: str
        :return: 常量文本值。
        :rtype: str
        """

        # 构建常量赋值前缀。
        marker = f"const {const_name} = '"

        # 查找常量赋值位置。
        start_index = html_text.find(marker)

        # 判断常量是否存在。
        if start_index < 0:

            # 返回空常量值。
            return ""

        # 计算常量值开始位置。
        value_start = start_index + len(marker)

        # 查找常量值结束位置。
        value_end = html_text.find("'", value_start)

        # 判断常量值结束位置是否存在。
        if value_end < 0:

            # 返回空常量值。
            return ""

        # 返回常量值文本。
        return html_text[value_start:value_end]

    # 定义功能设计配置初始化方法。
    def initialize_design_config(self):

        """
        初始化流向查询功能设计配置。

        :return: None。
        :rtype: None
        """

        # 加载查询页面初始化功能上下文。
        self.load_query_page()

        # 构建功能设计请求体。
        payload = [("gnbh", "1093"), ("pcStatus", "true")]

        # 拼接用户上下文字段。
        payload.extend(self.flatten_user_info_fields())

        # 请求功能设计配置。
        result = self.post_form(f"{self.base_url}/getDesignData", payload, self.query_page_url)

        # 解析功能设计响应。
        design_data = self.decode_json_response(result)

        # 读取功能设计文本。
        design_text = str(design_data.get("design") or "")

        # 判断功能设计文本是否有效。
        if not design_text.startswith("'["):

            # 抛出功能设计加载失败异常。
            raise RuntimeError(f"{self.company_name} 功能设计加载失败：{design_text[:300]}")

        # 解析功能设计元素。
        self.design_elements = json.loads(design_text[1:-1])

    # 定义用户上下文字段展开方法。
    def flatten_user_info_fields(self):

        """
        展开用户上下文字段。

        :return: 用户上下文表单字段。
        :rtype: list[tuple[str, str]]
        """

        # 返回 jQuery 表单格式的用户上下文字段。
        return [(f"userInfo[{key}]", self.stringify_form_value(value)) for key, value in self.user_info.items()]


    # 定义查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果数据行。
        :rtype: list[dict]
        """

        # 查询指定业务数据。
        query_result = self.query_sxtzcx_data(data_type, start_date, end_date)

        # 返回查询结果行。
        return query_result["rows"]


    # 定义查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定业务数据。
        query_result = self.query_sxtzcx_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建输出文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建输出文件路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回输出文件路径。
        return output_path

    # 定义业务数据查询方法。
    def query_sxtzcx_data(self, data_type, start_date, end_date):

        """
        查询山西天泽诚信指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 构建查询请求体。
        payload = self.build_query_payload(data_type, start_date, end_date)

        # 提交查询请求。
        result = self.post_form(self.query_api_url, payload, self.query_page_url)

        # 解析查询响应。
        response_data = self.decode_json_response(result)

        # 校验查询响应状态。
        if not self.is_query_success(response_data):

            # 抛出查询失败异常。
            raise RuntimeError(f"{self.company_name} {data_type}查询失败：{self.extract_message(response_data)}")

        # 提取响应数据行。
        rows = self.extract_rows(response_data, data_type)

        # 提取响应表头。
        headers = self.extract_headers(response_data, rows)

        # 打印查询解析日志。
        self.ph.success(self.ph.join(self.ph.kv("查询解析完成", f"{len(rows)} 条", "green"), self.ph.kv("表头数", len(headers), "cyan")))

        # 返回查询数据与表头。
        return {"rows": rows, "headers": headers}

    # 定义查询请求体构建方法。
    def build_query_payload(self, data_type, start_date, end_date):

        """
        构建流向查询表单请求体。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询表单字段序列。
        :rtype: list[tuple[str, str]]
        """

        # 构建主表查询对象。
        dsmain_query = self.build_dsmain_query(start_date, end_date)

        # 构建业务表格查询对象。
        table_query = self.build_table_query(data_type)

        # 初始化查询请求体。
        payload = []

        # 展开主表查询对象。
        payload.extend(self.flatten_form_value("querydata[0]", dsmain_query))

        # 展开业务表格查询对象。
        payload.extend(self.flatten_form_value("querydata[1]", table_query))

        # 拼接用户上下文字段。
        payload.extend(self.flatten_user_info_fields())

        # 返回表单字段序列。
        return payload

    # 定义主表查询对象构建方法。
    def build_dsmain_query(self, start_date, end_date):

        """
        构建主表查询对象。

        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 主表查询对象。
        :rtype: dict
        """

        # 读取主表设计元素。
        stage_element = self.find_design_element("stage")

        # 读取主表字段。
        fields = self.build_dsmain_fields()

        # 构建主表数据行。
        data_row = {"start_rq": start_date or "", "end_rq": end_date or "", "goodscode": "", "goodsname": "", "startDate": self.build_year_start(start_date or end_date or ""), "endDate": self.build_year_end(end_date or start_date or "")}

        # 返回主表查询对象。
        return {"con": (stage_element.get("data") or {}).get("link", "zy"), "elementName": "dsmain", "sqls": self.build_sqls(stage_element), "data": [data_row], "fields": fields, "showData": False, "regUniqueIdentifier": self.build_unique_identifier(), "procName": (stage_element.get("data") or {}).get("procName"), "tables": (stage_element.get("data") or {}).get("tables") or [], "gzid": self.build_unique_identifier()}

    # 定义业务表格查询对象构建方法。
    def build_table_query(self, data_type):

        """
        构建业务表格查询对象。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 业务表格查询对象。
        :rtype: dict
        """

        # 读取目标表格元素。
        table_element = self.find_design_element(self.resolve_table_element_name(data_type))

        # 读取目标表格数据配置。
        table_data = table_element.get("data") or {}

        # 返回业务表格查询对象。
        return {"applyLink": False, "con": table_data.get("link", "SYSLink"), "elementName": table_element.get("thisid"), "sqls": self.build_sqls(table_element), "data": [{}], "showData": True, "fields": self.build_post_fields(table_data.get("fields") or []), "regUniqueIdentifier": self.build_unique_identifier(), "procName": table_data.get("procName"), "tables": table_data.get("tables") or [], "gzid": self.build_unique_identifier()}

    # 定义主表字段构建方法。
    def build_dsmain_fields(self):

        """
        构建主表查询字段。

        :return: 主表查询字段。
        :rtype: list[dict]
        """

        # 初始化主表字段列表。
        fields = []

        # 遍历功能设计元素。
        for element in self.design_elements:

            # 读取元素数据配置。
            element_data = element.get("data") or {}

            # 遍历元素字段配置。
            for field_item in element_data.get("fields") or []:

                # 判断字段是否属于主表查询条件。
                if field_item.get("field") in ["start_rq", "end_rq", "goodscode", "goodsname", "startDate", "endDate"]:

                    # 写入字段配置。
                    fields.append(field_item)

        # 返回前端提交字段格式。
        return self.build_post_fields(fields)

    # 定义前端提交字段构建方法。
    def build_post_fields(self, fields):

        """
        构建前端提交字段格式。

        :param fields: 原始字段配置。
        :type fields: list[dict]
        :return: 提交字段配置。
        :rtype: list[dict]
        """

        # 初始化字段配置列表。
        result = []

        # 遍历原始字段配置。
        for field_item in fields:

            # 复制字段配置。
            post_field = dict(field_item)

            # 写入字段加密默认值。
            post_field["encrypt"] = bool(post_field.get("encrypt", False))

            # 追加字段配置。
            result.append(post_field)

        # 返回字段配置列表。
        return result

    # 定义 SQL 配置构建方法。
    def build_sqls(self, element):

        """
        构建前端提交 SQL 配置。

        :param element: 功能设计元素。
        :type element: dict
        :return: SQL 配置列表。
        :rtype: list[dict]
        """

        # 读取元素数据配置。
        element_data = element.get("data") or {}

        # 读取 SQL 配置列表。
        sqls = element_data.get("sqls")

        # 判断 SQL 配置是否为空。
        if not sqls:

            # 返回空 SQL 配置。
            return [{}]

        # 返回 SQL 配置副本。
        return [dict(sql_item) for sql_item in sqls]

    # 定义表格元素名称解析方法。
    def resolve_table_element_name(self, data_type):

        """
        解析数据类型对应的表格元素名称。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 表格元素名称。
        :rtype: str
        """

        # 定义数据类型与表格元素映射。
        mapping = {"购进流向": "table_2", "配送流向": "table_1", "库存": "table_3"}

        # 返回目标表格元素名称。
        return mapping[data_type]

    # 定义设计元素查找方法。
    def find_design_element(self, element_name):

        """
        查找功能设计元素。

        :param element_name: 元素名称。
        :type element_name: str
        :return: 功能设计元素。
        :rtype: dict
        """

        # 遍历功能设计元素。
        for element in self.design_elements:

            # 判断元素名称是否匹配。
            if element.get("thisid") == element_name:

                # 返回匹配元素。
                return element

        # 抛出元素缺失异常。
        raise KeyError(f"未找到功能设计元素：{element_name}")

    # 定义查询条件字段构建方法。
    def build_query_fields(self, data_type):

        """
        构建流向查询字段定义。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 查询字段定义列表。
        :rtype: list[dict]
        """

        # 返回平台查询条件字段定义。
        return [{"field": "start_rq", "title": "起始日期", "type": "date", "datatype": "varchar", "length": "10"}, {"field": "end_rq", "title": "截止日期", "type": "date", "datatype": "varchar", "length": "10"}, {"field": "goodscode", "title": "商品编号", "type": "text", "datatype": "varchar", "length": "50"}, {"field": "goodsname", "title": "商品名称", "type": "text", "datatype": "varchar", "length": "100"}, {"field": "startDate", "title": "开始日期", "type": "date", "datatype": "varchar", "length": "10"}, {"field": "endDate", "title": "结束日期", "type": "date", "datatype": "varchar", "length": "10"}]

    # 定义查询条件值构建方法。
    def build_query_values(self, data_type, start_date, end_date):

        """
        构建流向查询条件值。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询条件值列表。
        :rtype: list[str]
        """

        # 读取默认开始日期。
        start_text = start_date or ""

        # 读取默认结束日期。
        end_text = end_date or ""

        # 判断当前查询是否为库存。
        if data_type == "库存":

            # 返回库存查询条件值。
            return [start_text, end_text, "", "", self.build_year_start(end_text), self.build_year_end(end_text)]

        # 返回流向查询条件值。
        return [start_text, end_text, "", "", self.build_year_start(start_text), self.build_year_end(end_text)]

    # 定义查询代码构建方法。
    def build_query_code(self, data_type):

        """
        构建业务查询代码。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 业务查询代码。
        :rtype: str
        """

        # 定义数据类型与查询代码映射。
        mapping = {"购进流向": "zy", "配送流向": "zy", "库存": "zy"}

        # 返回查询代码。
        return mapping[data_type]

    # 定义表单请求方法。
    def post_form(self, url, payload, referer):

        """
        提交 AJAX 表单请求。

        :param url: 请求地址。
        :type url: str
        :param payload: 表单字段。
        :type payload: dict | list[tuple[str, str]]
        :param referer: 来源页面地址。
        :type referer: str
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 构建表单请求体。
        body = urlencode(payload).encode("utf-8")

        # 构建表单请求头。
        headers = self.build_form_headers(referer)

        # 发送表单请求。
        return self.http.request("POST", url, headers=headers, data=body, timeout=120)

    # 定义表单请求头构建方法。
    def build_form_headers(self, referer):

        """
        构建 AJAX 表单请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回 AJAX 表单请求头。
        return {"Accept": "*/*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应体文本。
        response_text = result.body.decode("utf-8", errors="replace").strip()

        # 判断响应体是否为空。
        if not response_text:

            # 返回空响应对象。
            return {}

        # 返回 JSON 解析结果。
        return json.loads(response_text)

    # 定义响应数据行提取方法。
    def extract_rows(self, response_data, data_type=None):

        """
        提取查询响应中的数据行。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :param data_type: 数据类型名称。
        :type data_type: str | None
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 尝试按前端结果数组提取目标表格数据。
        table_rows = self.extract_target_table_rows(response_data, data_type)

        # 判断目标表格数据是否存在。
        if table_rows:

            # 返回目标表格数据。
            return table_rows

        # 遍历常见数据行字段。
        for key_name in ["rows", "data", "list", "resultdata", "resultData"]:

            # 读取候选数据。
            candidate = response_data.get(key_name)

            # 判断候选数据是否为字典列表。
            if self.is_dict_list(candidate):

                # 返回候选数据行。
                return candidate

            # 判断候选数据是否为字典。
            if isinstance(candidate, dict):

                # 递归提取候选字典中的数据行。
                rows = self.extract_rows(candidate)

                # 判断递归结果是否包含数据。
                if rows:

                    # 返回递归数据行。
                    return rows

        # 递归扫描整个响应对象。
        rows = self.scan_dict_rows(response_data)

        # 返回扫描到的数据行。
        return rows

    # 定义目标表格数据提取方法。
    def extract_target_table_rows(self, response_data, data_type):

        """
        从前端查询结果数组中提取目标表格数据。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :param data_type: 数据类型名称。
        :type data_type: str | None
        :return: 目标表格数据行。
        :rtype: list[dict]
        """

        # 判断数据类型是否为空。
        if not data_type:

            # 返回空数据行。
            return []

        # 读取目标元素名称。
        target_element = self.resolve_table_element_name(data_type)

        # 读取结果数组。
        result_items = response_data.get("result")

        # 判断结果数组是否为列表。
        if not isinstance(result_items, list):

            # 返回空数据行。
            return []

        # 遍历结果数组成员。
        for item in result_items:

            # 判断当前成员是否为字典。
            if not isinstance(item, dict):

                # 进入下一个结果成员。
                continue

            # 判断当前成员是否为目标表格。
            if item.get("elementName") == target_element or item.get("element") == target_element or item.get("name") == target_element:

                # 判断目标表格是否包含数据行字段。
                if isinstance(item.get("data"), list):

                    # 返回目标表格数据行。
                    return item.get("data") or []

                # 递归提取当前成员数据行。
                return self.scan_dict_rows(item)

        # 返回空数据行。
        return []

    # 定义表头提取方法。
    def extract_headers(self, response_data, rows):

        """
        提取查询响应中的表头。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 按目标响应结果提取字段定义。
        field_headers = self.extract_target_headers(response_data)

        # 判断是否存在字段定义表头。
        if field_headers:

            # 返回字段定义表头。
            return field_headers

        # 根据数据行构建表头。
        return self.build_headers_from_rows(rows)

    # 定义目标响应表头提取方法。
    def extract_target_headers(self, response_data):

        """
        从目标响应结果中提取表头。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 读取结果数组。
        result_items = response_data.get("result")

        # 判断结果数组是否为列表。
        if isinstance(result_items, list):

            # 遍历结果数组成员。
            for item in result_items:

                # 判断当前成员是否为字典。
                if not isinstance(item, dict):

                    # 进入下一个结果成员。
                    continue

                # 读取字段定义。
                columns = item.get("columns")

                # 判断字段定义是否为字典列表。
                if self.is_dict_list(columns):

                    # 返回字段定义表头。
                    return [(str(column.get("title") or column.get("field")), str(column.get("field"))) for column in columns if column.get("field")]

        # 扫描响应中的字段定义。
        return self.scan_field_headers(response_data)

    # 定义字段定义表头扫描方法。
    def scan_field_headers(self, value):

        """
        扫描响应对象中的字段定义表头。

        :param value: 响应对象。
        :type value: object
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 判断当前对象是否为列表。
        if isinstance(value, list):

            # 判断列表是否为字段定义列表。
            if all(isinstance(item, dict) and item.get("field") for item in value):

                # 构建字段定义表头。
                headers = [(str(item.get("title") or item.get("field")), str(item.get("field"))) for item in value if item.get("field")]

                # 判断字段定义表头是否有效。
                if headers:

                    # 返回字段定义表头。
                    return headers

            # 遍历列表成员继续扫描。
            for item in value:

                # 递归扫描当前成员。
                headers = self.scan_field_headers(item)

                # 判断当前成员是否包含表头。
                if headers:

                    # 返回当前成员表头。
                    return headers

        # 判断当前对象是否为字典。
        if isinstance(value, dict):

            # 遍历字典值继续扫描。
            for child_value in value.values():

                # 递归扫描当前字典值。
                headers = self.scan_field_headers(child_value)

                # 判断当前字典值是否包含表头。
                if headers:

                    # 返回当前字典值表头。
                    return headers

        # 返回空表头。
        return []

    # 定义响应数据行扫描方法。
    def scan_dict_rows(self, value):

        """
        递归扫描响应对象中的字典列表。

        :param value: 响应对象。
        :type value: object
        :return: 字典数据行列表。
        :rtype: list[dict]
        """

        # 判断当前对象是否为字典列表。
        if self.is_dict_list(value):

            # 返回当前字典列表。
            return value

        # 判断当前对象是否为列表。
        if isinstance(value, list):

            # 遍历列表成员。
            for item in value:

                # 递归扫描当前成员。
                rows = self.scan_dict_rows(item)

                # 判断当前成员是否包含数据行。
                if rows:

                    # 返回当前成员数据行。
                    return rows

        # 判断当前对象是否为字典。
        if isinstance(value, dict):

            # 遍历字典值。
            for child_value in value.values():

                # 递归扫描当前字典值。
                rows = self.scan_dict_rows(child_value)

                # 判断当前字典值是否包含数据行。
                if rows:

                    # 返回当前字典值数据行。
                    return rows

        # 返回空数据行。
        return []

    # 定义字典列表判断方法。
    def is_dict_list(self, value):

        """
        判断对象是否为字典列表。

        :param value: 待判断对象。
        :type value: object
        :return: 是否为字典列表。
        :rtype: bool
        """

        # 返回字典列表判断结果。
        return isinstance(value, list) and bool(value) and all(isinstance(item, dict) for item in value)

    # 定义数据行表头构建方法。
    def build_headers_from_rows(self, rows):

        """
        根据数据行构建表头。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :return: 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 初始化表头字段列表。
        field_names = []

        # 遍历查询结果行。
        for row in rows:

            # 遍历当前行字段。
            for field_name in row.keys():

                # 判断字段是否已经加入表头。
                if field_name not in field_names:

                    # 追加字段名。
                    field_names.append(field_name)

        # 返回字段名表头。
        return [(field_name, field_name) for field_name in field_names]

    # 定义表单值展开方法。
    def flatten_form_value(self, prefix, value):

        """
        展开嵌套表单值。

        :param prefix: 字段前缀。
        :type prefix: str
        :param value: 字段值。
        :type value: object
        :return: 展开后的表单字段。
        :rtype: list[tuple[str, str]]
        """

        # 判断字段值是否为字典。
        if isinstance(value, dict):

            # 初始化表单字段。
            items = []

            # 遍历字典字段。
            for key, child_value in value.items():

                # 展开字典子字段。
                items.extend(self.flatten_form_value(f"{prefix}[{key}]", child_value))

            # 返回字典展开字段。
            return items

        # 判断字段值是否为列表。
        if isinstance(value, list):

            # 初始化表单字段。
            items = []

            # 遍历列表字段。
            for index, child_value in enumerate(value):

                # 展开列表子字段。
                items.extend(self.flatten_form_value(f"{prefix}[{index}]", child_value))

            # 返回列表展开字段。
            return items

        # 返回普通字段。
        return [(prefix, self.stringify_form_value(value))]

    # 定义表单值文本化方法。
    def stringify_form_value(self, value):

        """
        转换表单字段值为文本。

        :param value: 原始字段值。
        :type value: object
        :return: 表单字段文本。
        :rtype: str
        """

        # 判断字段值是否为布尔值。
        if isinstance(value, bool):

            # 返回小写布尔文本。
            return "true" if value else "false"

        # 判断字段值是否为空。
        if value is None:

            # 返回空文本。
            return ""

        # 返回字符串字段值。
        return str(value)

    # 定义登录成功判断方法。
    def is_login_success(self, response_data):

        """
        判断登录响应是否成功。

        :param response_data: 登录响应数据。
        :type response_data: dict
        :return: 是否登录成功。
        :rtype: bool
        """

        # 返回登录响应成功状态。
        return bool(response_data.get("result")) or str(response_data.get("code", "")) in ["0", "1", "200"]

    # 定义查询成功判断方法。
    def is_query_success(self, response_data):

        """
        判断查询响应是否成功。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :return: 是否查询成功。
        :rtype: bool
        """

        # 判断响应是否显式失败。
        if response_data.get("result") is False:

            # 返回查询失败状态。
            return False

        # 返回查询成功状态。
        return True

    # 定义响应消息提取方法。
    def extract_message(self, response_data):

        """
        提取响应消息文本。

        :param response_data: 响应数据。
        :type response_data: dict
        :return: 响应消息文本。
        :rtype: str
        """

        # 返回常见消息字段或响应摘要。
        return str(response_data.get("message") or response_data.get("msg") or response_data.get("info") or response_data)


    # 定义查询阶段标题打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印查询阶段标题。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 打印查询阶段标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验数据类型是否已接入。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已经接入。
        if data_type not in ["购进流向", "配送流向", "库存"]:

            # 抛出数据类型未接入异常。
            raise ValueError(f"山西天泽诚信暂未接入数据类型：{data_type}")



    # 定义年份开始日期构建方法。
    def build_year_start(self, date_text):

        """
        构建查询年份开始日期。

        :param date_text: 参考日期。
        :type date_text: str
        :return: 年份开始日期。
        :rtype: str
        """

        # 解析参考日期。
        date_value = self.parse_date_or_today(date_text)

        # 返回年份开始日期。
        return f"{date_value.year}-01-01"

    # 定义年份结束日期构建方法。
    def build_year_end(self, date_text):

        """
        构建查询年份结束日期。

        :param date_text: 参考日期。
        :type date_text: str
        :return: 年份结束日期。
        :rtype: str
        """

        # 解析参考日期。
        date_value = self.parse_date_or_today(date_text)

        # 返回年份结束日期。
        return f"{date_value.year}-12-31"

    # 定义日期解析方法。
    def parse_date_or_today(self, date_text):

        """
        解析日期文本或返回当前日期。

        :param date_text: 日期文本。
        :type date_text: str
        :return: 日期对象。
        :rtype: datetime
        """

        # 判断日期文本是否为空。
        if not date_text:

            # 返回当前日期。
            return datetime.now()

        # 返回解析后的日期。
        return datetime.strptime(date_text, "%Y-%m-%d")




    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(field_name, "") if isinstance(row, dict) else "" for _, field_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 读取当前表头名称。
            header_name = header_item[0]

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 捕获文件保存异常。
        try:

            # 保存 Excel 文件。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error


