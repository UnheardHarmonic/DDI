# -*- coding: utf-8 -*-
"""
重庆桐君阁大药房连锁有限责任公司桐君阁供应商流向系统协议适配器。

本模块负责桐君阁供应商流向系统的角色登录、GWT-RPC 分页查询、
标准字段前置、原始槽位保留和 Excel 文件生成。
"""

# 导入抽象语法树解析模块。
import ast

# 导入 Gzip 解压模块。
import gzip

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期时间模块。
from datetime import datetime

# 导入时间增量模块。
from datetime import timedelta

# 导入时区模块。
from datetime import timezone

# 导入 Cookie 容器。
from http.cookiejar import CookieJar

# 导入 URL 错误类型。
from urllib.error import HTTPError

# 导入 Cookie 处理器。
from urllib.request import HTTPCookieProcessor

# 导入 URL 请求对象。
from urllib.request import Request

# 导入 URL 打开器构建函数。
from urllib.request import build_opener

# 导入 Excel 统一样式函数。
from ddi_framework import apply_standard_excel_style

# 导入平台适配器基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义 GWT 长整型日期编码字符表。
GWT_LONG_ALPHABET = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789$_"


# 定义桐君阁供应商流向系统查询配置。
CQTJG_QUERY_DEFINITIONS = {
    "配送流向": {
        "model": "销售流向",
        "date_column": "a.confirmdate",
        "date_field": "销售日期",
        "date_mode": "range",
        "standard_fields": [
            ("发货单位", 26),
            ("销售日期", -5),
            ("客户ID", -8),
            ("客户名称", -9),
            ("客户地址", 32),
            ("商品ID", 23),
            ("品名", 22),
            ("规格", 19),
            ("生产厂家", 25),
            ("数量", 0),
            ("单位", 18),
            ("单价", 20),
            ("金额", 2),
            ("批号", 14),
            ("生产日期", 11),
            ("失效日期", 16),
            ("供应商", 4),
            ("流向提供商ID", 27),
            ("批准文号", -3),
        ],
    },
    "购进流向": {
        "model": "采购流向",
        "date_column": "a.confirmdate",
        "date_field": "采购日期",
        "date_mode": "range",
        "standard_fields": [
            ("购进单位", 27),
            ("采购日期", -5),
            ("供应商ID", 5),
            ("供应商名称", 4),
            ("商品ID", 24),
            ("品名", 23),
            ("规格", 20),
            ("生产厂家", 26),
            ("数量", 0),
            ("单位", 19),
            ("单价", 21),
            ("金额", 2),
            ("批号", 15),
            ("生产日期", 12),
            ("失效日期", 17),
            ("流向提供商ID", 28),
            ("批准文号", -3),
        ],
    },
    "库存": {
        "model": "库存流向",
        "date_column": None,
        "date_field": "库存日期",
        "date_mode": "snapshot",
        "standard_fields": [
            ("库存日期", 1),
            ("商品ID", 25),
            ("品名", 24),
            ("规格", 21),
            ("生产厂家", 27),
            ("库存数量", 22),
            ("单位", 20),
            ("单价", 3),
            ("金额", 5),
            ("批号", 13),
            ("生产日期", 9),
            ("失效日期", 16),
            ("供应商", 28),
            ("流向提供商ID", 29),
            ("批准文号", 31),
        ],
    },
    "纯销": {
        "model": "零售明细",
        "date_column": "useday",
        "date_field": "业务日期",
        "date_mode": "range",
        "page_size": 30,
        "login_each_page": True,
        "standard_fields": [
            ("管理部门", 22),
            ("配送单位", 7),
            ("零售ID", 4),
            ("业务日期", 1),
            ("区域", 29),
            ("商品ID", 19),
            ("品名", 18),
            ("规格", 15),
            ("生产厂家", 21),
            ("数量", 16),
            ("单位", 14),
            ("批号", 10),
            ("失效日期", 12),
            ("配送公司", 6),
            ("流向提供商ID", 23),
        ],
    },
}


# 定义重庆桐君阁平台适配器。
class CqtjgPlatform(BasePlatform):

    """
    重庆桐君阁大药房连锁有限责任公司桐君阁供应商流向系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 平台适配器实例。
    :rtype: CqtjgPlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化桐君阁供应商流向系统协议适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用平台基类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存站点根地址。
        self.site_root_url = account.get("接口根地址", "http://tjgpf.taiji.com:8999").rstrip("/")

        # 保存登录页面地址。
        self.login_page_url = account.get("登录页", f"{self.site_root_url}/gwtapp/")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.site_root_url}/gwtapp/user/login")

        # 保存 GWT 数据接口地址。
        self.data_service_url = account.get("查询接口", f"{self.site_root_url}/gwtapp/saflow/dataService")

        # 保存 GWT 模块根地址。
        self.gwt_module_base = account.get("GWT模块地址", f"{self.site_root_url}/gwtapp/saflow/")

        # 保存 GWT 序列化策略标识。
        self.gwt_policy = account.get("GWT序列化策略", "09DCB1B860EC08D6636FC31B822ED990")

        # 保存 GWT 强名称。
        self.gwt_permutation = account.get("GWT强名称", "97246FD42C3EEA52AB9C2FF29A5E7C25")

        # 保存登录账号。
        self.username = account.get("账号", "")

        # 保存登录密码。
        self.password = account.get("密码", "")

        # 保存登录角色类型。
        self.role_type = str(account.get("角色类型编码", "6"))

        # 保存登录密码前缀。
        self.password_prefix = account.get("密码前缀", "*&^%Ybghs154.")

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 300))

        # 初始化 Cookie 容器。
        self.cookie_jar = CookieJar()

        # 初始化协议请求打开器。
        self.opener = build_opener(HTTPCookieProcessor(self.cookie_jar))

        # 初始化登录令牌。
        self.token = ""

        # 初始化登录状态。
        self.logged_in = False

    # 定义 HTTP 会话重建方法。
    def reset_http_session(self):

        """
        重建当前适配器的 HTTP Cookie 容器和请求打开器。

        :return: None。
        :rtype: None
        """

        # 初始化新的 Cookie 容器。
        self.cookie_jar = CookieJar()

        # 初始化新的协议请求打开器。
        self.opener = build_opener(HTTPCookieProcessor(self.cookie_jar))

        # 清空当前登录令牌。
        self.token = ""

        # 标记当前会话未登录。
        self.logged_in = False

        # 返回空结果。
        return None

    # 定义平台登录方法。
    def login(self):

        """
        登录桐君阁供应商流向系统。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段分隔线。
        self.ph.section("桐君阁供应商流向系统登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 构建登录请求体。
        payload = {
            "comefrom": "web",
            "userType": self.role_type,
            "userId": self.username,
            "passWord": f"{self.password_prefix}{self.password}",
        }

        # 发送登录请求。
        response_data = self.post_login(payload)

        # 判断平台登录状态。
        if not response_data.get("isLogIn"):

            # 抛出登录失败异常。
            raise RuntimeError(f"桐君阁供应商流向系统登录失败：{response_data.get('MessAge') or response_data}")

        # 保存登录令牌。
        self.token = response_data.get("token", "")

        # 标记当前会话已登录。
        self.logged_in = True

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("登录成功", self.username, "green"), self.ph.kv("角色", "查询代理", "cyan")))

        # 返回登录摘要。
        return {"success": True, "token": bool(self.token), "role_type": self.role_type}

    # 定义会话信息保存方法。
    def save_cookie(self, output_dir):

        """
        保存桐君阁供应商流向系统当前登录会话信息。

        :param output_dir: 会话信息输出目录。
        :type output_dir: str
        :return: 会话信息文件路径。
        :rtype: str
        """

        # 创建会话信息输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建会话信息文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 提取当前 Cookie 字段。
        cookie_values = {cookie.name: cookie.value for cookie in self.cookie_jar}

        # 组装当前会话信息。
        cookie_data = {
            "token": self.token,
            "cookies": cookie_values,
            "site_root_url": self.site_root_url,
            "login_api_url": self.login_api_url,
            "data_service_url": self.data_service_url,
            "role_type": self.role_type,
        }

        # 写入会话信息文件。
        with open(cookie_path, "w", encoding="utf-8") as file_obj:

            # 序列化当前会话信息。
            json.dump(cookie_data, file_obj, ensure_ascii=False, indent=2)

        # 打印会话保存日志。
        self.ph.info(self.ph.join(self.ph.kv("会话文件", cookie_path, "cyan")))

        # 返回会话信息文件路径。
        return cookie_path

    # 定义平台查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行列表。
        :rtype: list[dict]
        """

        # 打印数据类型查询分隔线。
        self.print_query_section(data_type, start_date, end_date)

        # 返回协议分页查询结果。
        return self.fetch_protocol_rows(data_type, start_date, end_date)

    # 定义业务导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        导出指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 复用查询落地 Excel 逻辑。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定类型业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 打印数据类型查询分隔线。
        self.print_query_section(data_type, start_date, end_date)

        # 查询协议数据行。
        rows = self.fetch_protocol_rows(data_type, start_date, end_date)

        # 读取输出目录。
        save_dir = output_dir or "downloads"

        # 创建输出目录。
        os.makedirs(save_dir, exist_ok=True)

        # 构建目标文件名。
        target_name = self.build_output_file_name(data_type, start_date, end_date)

        # 拼接目标文件路径。
        target_path = os.path.join(save_dir, target_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(data_type, rows, target_path)

        # 应用统一 Excel 样式。
        apply_standard_excel_style(target_path)

        # 打印 Excel 生成日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(target_path), "green"), self.ph.kv("记录数", f"{len(rows)} 条", "cyan")))

        # 返回 Excel 文件路径。
        return target_path

    # 定义标准输出命名日期解析方法。
    def resolve_output_query_date(self, data_type, start_date=None, end_date=None):

        """
        解析标准输出文件名使用的业务日期。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 标准输出文件名使用的业务日期。
        :rtype: str | None
        """

        # 判断当前数据类型是否为库存类。
        if data_type == "库存":

            # 返回库存类使用的运行当天日期。
            return datetime.now().strftime("%Y-%m-%d")

        # 返回流向类使用的查询开始日期。
        return start_date or end_date

    # 定义数据类型分隔线打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印数据类型查询分隔线。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 打印标准阶段分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 返回空结果。
        return None

    # 定义协议分页查询方法。
    def fetch_protocol_rows(self, data_type, start_date, end_date):

        """
        通过 GWT-RPC 接口分页读取业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 标准字段与原始槽位合并后的数据行。
        :rtype: list[dict]
        """

        # 确保当前会话已经登录。
        self.ensure_logged_in()

        # 读取查询定义。
        query_definition = self.resolve_query_definition(data_type)

        # 初始化分页起始位置。
        start_index = 0

        # 初始化结果行容器。
        result_rows = []

        # 初始化总记录数。
        total_count = None

        # 读取当前业务模型分页大小。
        page_size = int(query_definition.get("page_size", self.page_size))

        # 循环读取分页数据。
        while True:

            # 判断当前模型是否需要按页重建服务端会话。
            if query_definition.get("login_each_page") and start_index > 0:

                # 重建 HTTP 会话以隔离 GWT 分页状态。
                self.reset_http_session()

                # 登录当前分页请求所需的平台会话。
                self.login()

            # 构建当前页 GWT 请求体。
            request_body = self.build_gwt_query_body(query_definition, start_date, end_date, start_index, page_size)

            # 发送 GWT 查询请求。
            response_text = self.post_gwt(request_body)

            # 解析 GWT 响应体。
            parsed_response = self.parse_gwt_response(response_text)

            # 读取当前页原始行。
            page_rows = parsed_response["rows"]

            # 读取接口总记录数。
            total_count = parsed_response["total_count"]

            # 转换当前页业务行。
            converted_rows = [self.convert_gwt_row(data_type, query_definition, row_tokens, parsed_response["strings"], start_index + row_index + 1) for row_index, row_tokens in enumerate(page_rows)]

            # 合并当前页业务行。
            result_rows.extend(converted_rows)

            # 打印分页查询日志。
            self.ph.info(self.ph.join(self.ph.kv("模型", query_definition["model"], "cyan"), self.ph.kv("起始", str(start_index), "cyan"), self.ph.kv("本页", f"{len(page_rows)} 条", "cyan"), self.ph.kv("累计", f"{len(result_rows)} 条", "green")))

            # 判断当前页是否没有返回行。
            if not page_rows:

                # 结束分页循环。
                break

            # 判断是否已经读取到接口总数。
            if total_count is not None and len(result_rows) >= total_count:

                # 结束分页循环。
                break

            # 累加下一页起始位置。
            start_index += len(page_rows)

        # 校验查询日期范围。
        self.validate_date_range(data_type, query_definition, result_rows, start_date, end_date)

        # 返回全部结果行。
        return result_rows

    # 定义查询定义解析方法。
    def resolve_query_definition(self, data_type):

        """
        解析数据类型对应的查询定义。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 查询定义。
        :rtype: dict
        """

        # 校验数据类型是否已配置。
        if data_type not in CQTJG_QUERY_DEFINITIONS:

            # 抛出不支持的数据类型异常。
            raise KeyError(f"桐君阁供应商流向系统暂不支持数据类型：{data_type}")

        # 返回查询定义。
        return CQTJG_QUERY_DEFINITIONS[data_type]

    # 定义登录状态保证方法。
    def ensure_logged_in(self):

        """
        确保当前会话已经登录。

        :return: None。
        :rtype: None
        """

        # 判断当前会话是否尚未登录。
        if not self.logged_in:

            # 执行平台登录。
            self.login()

        # 返回空结果。
        return None

    # 定义登录请求方法。
    def post_login(self, payload):

        """
        发起登录请求并解析响应。

        :param payload: 登录请求体。
        :type payload: dict
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 序列化登录请求体。
        request_body = json.dumps(payload, ensure_ascii=False).encode("utf-8")

        # 构建登录请求头。
        headers = {"Accept": "application/json, text/plain, */*", "Content-Type": "application/json;charset=UTF-8", "Origin": self.site_root_url, "Referer": self.login_page_url, "User-Agent": "Mozilla/5.0"}

        # 构建登录请求对象。
        request = Request(self.login_api_url, data=request_body, headers=headers, method="POST")

        # 发送登录请求并读取响应。
        response_text = self.open_text(request)

        # 返回 JSON 响应。
        return json.loads(response_text)

    # 定义 GWT 请求方法。
    def post_gwt(self, request_body):

        """
        发起 GWT-RPC POST 请求。

        :param request_body: GWT-RPC 请求体。
        :type request_body: str
        :return: 响应文本。
        :rtype: str
        """

        # 构建 GWT 请求头。
        headers = {"Accept": "*/*", "Content-Type": "text/x-gwt-rpc; charset=UTF-8", "Origin": self.site_root_url, "Referer": f"{self.site_root_url}/gwtapp/main.jsp", "X-Gwt-Module-Base": self.gwt_module_base, "X-Gwt-Permutation": self.gwt_permutation, "User-Agent": "Mozilla/5.0"}

        # 构建 GWT 请求对象。
        request = Request(self.data_service_url, data=request_body.encode("utf-8"), headers=headers, method="POST")

        # 发送 GWT 请求并返回响应文本。
        return self.open_text(request)

    # 定义 HTTP 文本请求方法。
    def open_text(self, request):

        """
        发送 HTTP 请求并读取文本响应。

        :param request: HTTP 请求对象。
        :type request: urllib.request.Request
        :return: 响应文本。
        :rtype: str
        """

        # 捕获 HTTP 错误响应。
        try:

            # 发送 HTTP 请求。
            with self.opener.open(request, timeout=45) as response:

                # 读取响应原始字节。
                response_bytes = response.read()

                # 读取响应编码标识。
                content_encoding = response.headers.get("Content-Encoding", "")

        # 处理 HTTP 错误响应。
        except HTTPError as error:

            # 读取错误响应原始字节。
            response_bytes = error.read()

            # 读取错误响应编码标识。
            content_encoding = error.headers.get("Content-Encoding", "")

        # 判断响应体是否使用 Gzip 压缩。
        if content_encoding.lower() == "gzip":

            # 解压 Gzip 响应体。
            response_bytes = gzip.decompress(response_bytes)

        # 返回 UTF-8 响应文本。
        return response_bytes.decode("utf-8", errors="replace")

    # 定义 GWT 查询请求体构建方法。
    def build_gwt_query_body(self, query_definition, start_date, end_date, start_index, page_size):

        """
        构建 GWT-RPC 分页查询请求体。

        :param query_definition: 查询定义。
        :type query_definition: dict
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param start_index: 分页起始位置。
        :type start_index: int
        :param page_size: 分页大小。
        :type page_size: int
        :return: GWT-RPC 请求体。
        :rtype: str
        """

        # 构建模型过滤条件文本。
        filter_text = self.build_filter_text(query_definition, start_date, end_date)

        # 读取平台模型名称。
        model_name = query_definition["model"]

        # 返回 GWT-RPC 请求体。
        return f"7|0|11|{self.gwt_module_base}|{self.gwt_policy}|com.tjg.datasource.client.DataService|getModelPage|java.lang.String/2004016611|Z|com.sencha.gxt.data.shared.loader.PagingLoadConfig|{model_name}|{filter_text}|com.sencha.gxt.data.shared.loader.PagingLoadConfigBean/38458988|java.util.ArrayList/4159755760|1|2|3|4|4|5|6|5|7|8|0|9|10|{page_size}|{start_index}|11|0|"

    # 定义过滤条件文本构建方法。
    def build_filter_text(self, query_definition, start_date, end_date):

        """
        构建平台模型过滤条件文本。

        :param query_definition: 查询定义。
        :type query_definition: dict
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 过滤条件文本。
        :rtype: str
        """

        # 判断当前查询是否为库存快照。
        if query_definition.get("date_mode") == "snapshot":

            # 返回库存快照空过滤条件。
            return '[{"sqlWhereIs":""}]'

        # 校验开始日期是否存在。
        if not start_date:

            # 抛出开始日期缺失异常。
            raise ValueError("桐君阁供应商流向系统区间查询缺少开始日期。")

        # 校验结束日期是否存在。
        if not end_date:

            # 抛出结束日期缺失异常。
            raise ValueError("桐君阁供应商流向系统区间查询缺少结束日期。")

        # 读取平台日期字段。
        date_column = query_definition["date_column"]

        # 构建 SQL 条件文本。
        sql_where = f"{date_column} between to_date('{start_date}','yyyy-mm-dd')  and to_date('{end_date} 23:59:59','yyyy-mm-dd hh24:mi:ss')  "

        # 构建前端过滤器结构。
        filter_items = [
            {"CxIndex": 10, "ColName": date_column, "ColType": "日期", "CxChar": " between ", "Vaule": f"to_date('{start_date}','yyyy-mm-dd') ", "Vaule2": f"to_date('{end_date} 23:59:59','yyyy-mm-dd hh24:mi:ss') "},
            {"sqlWhereIs": sql_where},
        ]

        # 返回平台要求的紧凑 JSON 文本。
        return json.dumps(filter_items, ensure_ascii=False, separators=(",", ":"))

    # 定义 GWT 响应解析方法。
    def parse_gwt_response(self, response_text):

        """
        解析 GWT-RPC 响应文本。

        :param response_text: GWT-RPC 响应文本。
        :type response_text: str
        :return: 解析后的字符串表、行令牌和总记录数。
        :rtype: dict
        """

        # 校验响应前缀是否为成功响应。
        if not response_text.startswith("//OK"):

            # 抛出 GWT 响应异常。
            raise RuntimeError(f"桐君阁供应商流向系统 GWT 响应异常：{response_text[:300]}")

        # 解析 Python 兼容的 GWT 数组文本。
        response_array = ast.literal_eval(response_text[4:])

        # 定位字符串表位置。
        string_table_index = next(index for index, item in enumerate(response_array) if isinstance(item, list))

        # 读取字符串表。
        strings = response_array[string_table_index]

        # 读取字符串表之前的令牌区。
        tokens = response_array[:string_table_index]

        # 拆分行令牌与尾部令牌。
        rows, tail_tokens = self.split_gwt_rows(tokens)

        # 读取总记录数。
        total_count = tail_tokens[2] if len(tail_tokens) >= 3 and isinstance(tail_tokens[2], int) else len(rows)

        # 返回解析结果。
        return {"strings": strings, "rows": rows, "tail": tail_tokens, "total_count": total_count}

    # 定义 GWT 行令牌拆分方法。
    def split_gwt_rows(self, tokens):

        """
        拆分 GWT-RPC 响应中的行令牌和尾部分页令牌。

        :param tokens: 字符串表前的令牌列表。
        :type tokens: list
        :return: 行令牌列表和尾部令牌。
        :rtype: tuple[list[list], list]
        """

        # 初始化行容器。
        rows = []

        # 初始化当前行令牌。
        current_row = []

        # 遍历响应令牌。
        for token in tokens:

            # 将当前令牌加入当前行。
            current_row.append(token)

            # 判断当前令牌是否为对象行结束标记。
            if token == 3 and len(current_row) >= 2 and current_row[-2] == 4:

                # 保存完整行令牌。
                rows.append(current_row)

                # 重置当前行令牌。
                current_row = []

        # 返回行令牌和剩余分页令牌。
        return rows, current_row

    # 定义 GWT 行转换方法。
    def convert_gwt_row(self, data_type, query_definition, row_tokens, strings, row_number):

        """
        将 GWT 行令牌转换为 Excel 行字典。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param query_definition: 查询定义。
        :type query_definition: dict
        :param row_tokens: 当前行 GWT 令牌。
        :type row_tokens: list
        :param strings: GWT 字符串表。
        :type strings: list
        :param row_number: 全局行号。
        :type row_number: int
        :return: Excel 行字典。
        :rtype: dict
        """

        # 解析当前行全部槽位。
        slot_values = self.decode_gwt_slots(row_tokens, strings)

        # 解析纯销明细语义字段。
        retail_values = self.extract_retail_standard_fields(row_tokens, strings) if data_type == "纯销" else {}

        # 初始化输出行。
        output_row = {"数据类型": data_type, "平台模型": query_definition["model"], "页内序号": row_number}

        # 遍历标准字段定义。
        for field_name, slot_index in query_definition["standard_fields"]:

            # 解析标准字段值。
            standard_value = retail_values.get(field_name, self.resolve_slot_value(slot_values, row_tokens, slot_index))

            # 写入标准字段值。
            output_row[field_name] = standard_value

        # 遍历全部 GWT 原始槽位。
        for slot_index in sorted(slot_values.keys()):

            # 构建原始槽位字段名。
            raw_field_name = f"原始GWT槽位{slot_index:03d}"

            # 写入原始槽位值。
            output_row[raw_field_name] = slot_values[slot_index]

        # 返回转换后的输出行。
        return output_row

    # 定义槽位值读取方法。
    def resolve_slot_value(self, slot_values, row_tokens, slot_index):

        """
        按正向或倒序槽位读取标准字段值。

        :param slot_values: 已解码槽位映射。
        :type slot_values: dict[int, object]
        :param row_tokens: 当前行原始令牌。
        :type row_tokens: list
        :param slot_index: 标准字段槽位索引。
        :type slot_index: int
        :return: 标准字段值。
        :rtype: object
        """

        # 判断当前槽位是否为倒序定位。
        if slot_index < 0:

            # 计算倒序槽位对应的正向索引。
            resolved_index = len(row_tokens) + slot_index

            # 返回倒序槽位值。
            return slot_values.get(resolved_index, "")

        # 返回正向槽位值。
        return slot_values.get(slot_index, "")

    # 定义纯销明细语义字段抽取方法。
    def extract_retail_standard_fields(self, row_tokens, strings):

        """
        从零售明细 GWT 行中抽取可空槽位影响较大的标准字段。

        :param row_tokens: 当前行原始令牌。
        :type row_tokens: list
        :param strings: GWT 字符串表。
        :type strings: list
        :return: 纯销标准字段映射。
        :rtype: dict
        """

        # 解码当前行全部顺序值。
        decoded_values = [self.decode_gwt_token(token, row_tokens[index + 1] if index + 1 < len(row_tokens) else None, strings) for index, token in enumerate(row_tokens)]

        # 初始化纯销标准字段。
        retail_values = {}

        # 写入固定前置日期字段。
        retail_values["业务日期"] = decoded_values[1] if len(decoded_values) > 1 else ""

        # 写入固定前置零售标识。
        retail_values["零售ID"] = decoded_values[4] if len(decoded_values) > 4 else ""

        # 写入固定前置配送单位。
        retail_values["配送单位"] = decoded_values[6] if len(decoded_values) > 6 else ""

        # 定位商品规格字段。
        spec_index = self.find_first_value_index(decoded_values, self.is_retail_spec_value)

        # 判断商品规格是否存在。
        if spec_index is not None:

            # 写入单位字段。
            retail_values["单位"] = decoded_values[spec_index - 1] if spec_index >= 1 else ""

            # 写入规格字段。
            retail_values["规格"] = decoded_values[spec_index]

            # 写入数量字段。
            retail_values["数量"] = decoded_values[spec_index + 1] if spec_index + 1 < len(decoded_values) else ""

            # 写入品名字段。
            retail_values["品名"] = decoded_values[spec_index + 3] if spec_index + 3 < len(decoded_values) else ""

            # 写入商品编码字段。
            retail_values["商品ID"] = decoded_values[spec_index + 4] if spec_index + 4 < len(decoded_values) else ""

            # 写入生产厂家字段。
            retail_values["生产厂家"] = decoded_values[spec_index + 6] if spec_index + 6 < len(decoded_values) else ""

            # 写入管理部门字段。
            retail_values["管理部门"] = decoded_values[spec_index + 7] if spec_index + 7 < len(decoded_values) else ""

            # 写入流向提供商标识字段。
            retail_values["流向提供商ID"] = decoded_values[spec_index + 8] if spec_index + 8 < len(decoded_values) else ""

        # 定位批号字段。
        batch_index = self.find_first_value_index(decoded_values, lambda value: isinstance(value, str) and value[:4].isdigit() and any(char.isalpha() for char in value))

        # 写入批号字段。
        retail_values["批号"] = decoded_values[batch_index] if batch_index is not None else ""

        # 读取当前行日期字段列表。
        date_values = [value for value in decoded_values if isinstance(value, str) and len(value) >= 19 and value[:4].isdigit() and value[4] == "-"]

        # 写入失效日期字段。
        retail_values["失效日期"] = date_values[1] if len(date_values) >= 2 else ""

        # 定位区域字段。
        area_index = self.find_first_value_index(decoded_values, lambda value: isinstance(value, str) and value.endswith("片区"))

        # 写入区域字段。
        retail_values["区域"] = decoded_values[area_index] if area_index is not None else ""

        # 写入配送公司字段。
        retail_values["配送公司"] = self.company_name

        # 返回纯销标准字段。
        return retail_values

    # 定义纯销规格字段判断方法。
    def is_retail_spec_value(self, value):

        """
        判断当前顺序值是否为零售明细商品规格。

        :param value: 待判断的顺序值。
        :type value: object
        :return: 是否为商品规格。
        :rtype: bool
        """

        # 判断当前值是否不是文本。
        if not isinstance(value, str):

            # 返回否定结果。
            return False

        # 去除文本首尾空白。
        normalized_value = value.strip()

        # 判断文本是否为空。
        if not normalized_value:

            # 返回否定结果。
            return False

        # 判断文本是否包含常见规格计量单位。
        has_measure_unit = any(unit_text in normalized_value for unit_text in ("IU", "μg", "ug", "mg", "ml", "g", "%"))

        # 判断文本是否包含常见包装单位。
        has_package_unit = any(package_text in normalized_value for package_text in ("/支", "/盒", "x5支", "×5支", "*5支", "支", "盒"))

        # 判断文本是否同时具备规格计量和包装单位。
        return has_measure_unit and has_package_unit

    # 定义顺序值定位方法。
    def find_first_value_index(self, values, predicate):

        """
        返回首个满足条件的顺序值索引。

        :param values: 顺序值列表。
        :type values: list
        :param predicate: 判断函数。
        :type predicate: callable
        :return: 首个满足条件的索引。
        :rtype: int | None
        """

        # 遍历顺序值和索引。
        for index, value in enumerate(values):

            # 判断当前值是否满足条件。
            if predicate(value):

                # 返回当前索引。
                return index

        # 返回空索引。
        return None

    # 定义 GWT 槽位解码方法。
    def decode_gwt_slots(self, row_tokens, strings):

        """
        解码 GWT 行令牌槽位。

        :param row_tokens: 当前行 GWT 令牌。
        :type row_tokens: list
        :param strings: GWT 字符串表。
        :type strings: list
        :return: 槽位索引和值的映射。
        :rtype: dict[int, object]
        """

        # 初始化槽位值映射。
        slot_values = {}

        # 遍历当前行令牌。
        for slot_index, token in enumerate(row_tokens):

            # 读取后继令牌。
            next_token = row_tokens[slot_index + 1] if slot_index + 1 < len(row_tokens) else None

            # 解码当前令牌值。
            slot_values[slot_index] = self.decode_gwt_token(token, next_token, strings)

        # 返回槽位值映射。
        return slot_values

    # 定义单个 GWT 令牌解码方法。
    def decode_gwt_token(self, token, next_token, strings):

        """
        解码单个 GWT 令牌。

        :param token: 当前令牌。
        :type token: object
        :param next_token: 后继令牌。
        :type next_token: object
        :param strings: GWT 字符串表。
        :type strings: list
        :return: 解码后的值。
        :rtype: object
        """

        # 判断当前令牌是否为字符串形式的 GWT 长整型日期。
        if isinstance(token, str):

            # 读取日期解码结果。
            decoded_time = self.decode_gwt_long_datetime(token)

            # 优先返回有效日期文本。
            return decoded_time or token

        # 判断当前令牌是否不是整数。
        if not isinstance(token, int):

            # 返回原始令牌。
            return token

        # 读取后继令牌对应的类型名称。
        next_type = self.resolve_string_reference(next_token, strings)

        # 判断当前令牌是否为 BigDecimal 字符串引用。
        if next_type == "java.math.BigDecimal/8151472":

            # 返回 BigDecimal 的文本值。
            return self.resolve_string_reference(token, strings)

        # 判断当前令牌是否为时间戳类型标记。
        if next_type == "java.sql.Timestamp/3040052672":

            # 返回原始数值。
            return token

        # 判断当前令牌是否为 Integer 数值。
        if next_type == "java.lang.Integer/3438268394":

            # 返回整数值。
            return token

        # 判断当前令牌是否可作为字符串引用。
        if 1 <= token <= len(strings):

            # 返回字符串表引用值。
            return strings[token - 1]

        # 返回原始整数。
        return token

    # 定义字符串表引用解析方法。
    def resolve_string_reference(self, token, strings):

        """
        解析 GWT 字符串表引用。

        :param token: 字符串引用令牌。
        :type token: object
        :param strings: GWT 字符串表。
        :type strings: list
        :return: 字符串引用值。
        :rtype: str
        """

        # 判断令牌是否为有效整数引用。
        if isinstance(token, int) and 1 <= token <= len(strings):

            # 返回引用的字符串值。
            return strings[token - 1]

        # 返回空文本。
        return ""

    # 定义 GWT 长整型日期解码方法。
    def decode_gwt_long_datetime(self, value):

        """
        解码 GWT Base64 长整型毫秒时间。

        :param value: GWT 长整型编码文本。
        :type value: str
        :return: 北京时间日期文本。
        :rtype: str
        """

        # 判断输入值是否为空。
        if not value:

            # 返回空文本。
            return ""

        # 初始化毫秒数。
        milliseconds = 0

        # 遍历编码字符。
        for char in value:

            # 判断字符是否不在 GWT 字符表。
            if char not in GWT_LONG_ALPHABET:

                # 返回空文本。
                return ""

            # 累加当前字符对应的 64 进制数值。
            milliseconds = milliseconds * 64 + GWT_LONG_ALPHABET.index(char)

        # 捕获异常时间戳。
        try:

            # 按北京时间转换毫秒时间戳。
            return datetime.fromtimestamp(milliseconds / 1000, timezone.utc).astimezone(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S")

        # 处理无效时间戳。
        except (OverflowError, OSError, ValueError):

            # 返回空文本。
            return ""

    # 定义日期范围校验方法。
    def validate_date_range(self, data_type, query_definition, rows, start_date, end_date):

        """
        校验流向业务日期是否落在查询区间内。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param query_definition: 查询定义。
        :type query_definition: dict
        :param rows: 查询结果行。
        :type rows: list[dict]
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断库存快照是否不需要区间校验。
        if query_definition.get("date_mode") == "snapshot":

            # 返回空结果。
            return None

        # 判断无数据结果是否不需要区间校验。
        if not rows:

            # 返回空结果。
            return None

        # 读取业务日期字段名。
        date_field = query_definition["date_field"]

        # 读取结果中的业务日期列表。
        row_dates = [str(row.get(date_field, ""))[:10] for row in rows if str(row.get(date_field, ""))[:10]]

        # 判断是否缺少业务日期。
        if not row_dates:

            # 打印缺少日期提示。
            self.ph.warning(f"{data_type} 未解析到可校验的业务日期字段：{date_field}")

            # 返回空结果。
            return None

        # 计算最小业务日期。
        min_date = min(row_dates)

        # 计算最大业务日期。
        max_date = max(row_dates)

        # 判断业务日期是否越过查询区间。
        if (start_date and min_date < start_date) or (end_date and max_date > end_date):

            # 抛出日期范围异常。
            raise RuntimeError(f"{data_type} 业务日期超出查询范围：{min_date} - {max_date}")

        # 打印日期校验日志。
        self.ph.success(self.ph.join(self.ph.kv("日期范围", f"{min_date} - {max_date}", "green"), self.ph.kv("记录数", f"{len(rows)} 条", "cyan")))

        # 返回空结果。
        return None

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, data_type, rows, target_path):

        """
        将查询结果写入 Excel 文件。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果行。
        :type rows: list[dict]
        :param target_path: 目标文件路径。
        :type target_path: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列号转换函数。
        from openpyxl.utils import get_column_letter

        # 构建输出表头。
        headers = self.build_excel_headers(data_type, rows)

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = data_type[:31]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果行。
        for row_data in rows:

            # 按表头顺序写入当前行。
            worksheet.append([row_data.get(header, "") for header in headers])

        # 遍历表头列设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 计算当前列最大文本长度。
            width = max([len(str(header_name))] + [len(str(row_data.get(header_name, ""))) for row_data in rows[:200]])

            # 设置当前列宽。
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(width + 2, 12), 45)

        # 删除已存在的目标文件。
        if os.path.exists(target_path):

            # 移除旧文件。
            os.remove(target_path)

        # 保存 Excel 文件。
        workbook.save(target_path)

        # 返回空结果。
        return None

    # 定义 Excel 表头构建方法。
    def build_excel_headers(self, data_type, rows):

        """
        构建 Excel 输出表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果行。
        :type rows: list[dict]
        :return: Excel 表头列表。
        :rtype: list[str]
        """

        # 读取查询定义。
        query_definition = self.resolve_query_definition(data_type)

        # 构建标准字段表头。
        headers = ["数据类型", "平台模型", "页内序号"] + [field_name for field_name, _ in query_definition["standard_fields"]]

        # 初始化已存在表头集合。
        existing_headers = set(headers)

        # 遍历查询结果行。
        for row_data in rows:

            # 遍历当前行字段。
            for field_name in row_data.keys():

                # 跳过已存在字段。
                if field_name in existing_headers:

                    # 进入下一个字段。
                    continue

                # 追加动态原始字段。
                headers.append(field_name)

                # 标记字段已存在。
                existing_headers.add(field_name)

        # 返回表头列表。
        return headers
