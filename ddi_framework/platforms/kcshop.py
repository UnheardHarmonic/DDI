# -*- coding: utf-8 -*-
"""
KCShop 流向管理系统协议适配器。

本模块负责 KCShop 流向管理系统的验证码登录、分页查询和本地 Excel 生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期时间解析模块。
from datetime import datetime

# 导入正则表达式模块。
import re

# 导入验证码识别器。
from ddi_framework import CaptchaSolver

# 导入下载文件命名工具。
from ddi_framework import DownloadWatcher

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印处理器。
from ddi_framework import PrintHandler


# 定义 KCShop 通用字段中文表头映射。
KCSHOP_COMMON_HEADER_LABELS = {
    "ID": "ID",
    "DJBH": "单据编号",
    "SJJGBH": "所属分公司机构编码",
    "SJJGMC": "所属分公司机构名称",
    "JGBH": "机构编码",
    "JGID": "机构ID",
    "JGMC": "机构名称",
    "GYSBH": "供应商编码",
    "GYSMC": "供应商名称",
    "RQ": "日期",
    "YWLX": "业务类型",
    "SPID": "商品ID",
    "SPBH": "商品编号",
    "SPTYM": "通用名",
    "SPGG": "规格",
    "SCCJ": "生产厂家",
    "SPCD": "产地",
    "BZDW": "单位",
    "FL": "分类",
    "DINGW": "定位",
    "PIHAO": "批号",
    "SXQ": "有效期",
    "SCRQ": "生产日期",
    "PZWH": "批准文号",
    "PCID": "批次ID",
    "PCSHL": "批次数量",
    "KHBH": "客户编码",
    "KHMC": "客户名称",
    "XSEJBH": "下属二级机构编码",
    "XSEJMC": "下属二级机构名称",
    "XSDJ": "销售单价",
    "XSSHL": "销售数量",
    "XSJE": "销售金额",
    "LSHJ": "零售单价",
    "LSHJE": "零售金额",
    "HSCB": "含税成本",
    "HSCHBJE": "含税成本金额",
    "CGDJ": "采购单价",
    "CGSHL": "采购数量",
    "CGJE": "采购金额",
    "FSDJ": "采购单价",
    "FSSHL": "采购数量",
    "FSJE": "采购金额",
    "SLV": "税率",
    "BHSDJ": "不含税单价",
    "CREATED": "创建日期",
    "GYSYJQYMC": "供应商一级企业名称",
    "KHYJQYMC": "客户一级企业名称",
    "CODE": "状态编码",
    "ISDEL": "删除标记",
    "HUOZHU": "货主",
    "SBKH": "社保卡号",
    "BINGZ_CODE": "病种编码",
    "BINGZ": "病种",
    "BEIZHU": "备注",
    "NOTE": "说明",
    "HYKH": "会员卡号",
    "HYMC": "会员名称",
    "GKMC": "顾客名称",
    "SJH": "手机号",
}

# 定义 KCShop 数据类型字段中文表头映射。
KCSHOP_DATA_TYPE_HEADER_LABELS = {
    "配送流向": {
        "RQ": "销售日期",
        "JGBH": "所属分公司机构编码",
        "JGMC": "所属分公司机构名称",
        "PCSHL": "当前库存数量",
    },
    "购进流向": {
        "RQ": "日期",
        "JGBH": "所属分公司机构编码",
        "JGMC": "所属分公司机构名称",
    },
    "总部库存": {
        "PCSHL": "库存数量",
    },
    "门店库存": {
        "PCSHL": "库存数量",
    },
    "纯销流向": {
        "BZDW": "包装单位",
    },
}


# 定义 KCShop 流向管理系统协议适配器。
class KcshopPlatform(BasePlatform):

    """
    KCShop 流向管理系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: KCShop 平台适配器实例。
    :rtype: KcshopPlatform
    """

    # 定义 KCShop 适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化 KCShop 平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存前端入口地址。
        self.login_page_url = account.get("登录页", "https://dongguanguoyao.kcshop.vip/")

        # 保存接口根地址。
        self.api_base_url = account.get("接口根地址", "https://api.dongguanguoyao.kcshop.vip")

        # 保存验证码接口地址。
        self.captcha_api_url = account.get("验证码接口", f"{self.api_base_url}/captcha.html")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.api_base_url}/api/user/login")

        # 保存接口路径映射。
        self.endpoint_mapping = self.build_endpoint_mapping()

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 5000))

        # 保存查询请求超时时间。
        self.query_timeout = int(account.get("查询超时秒数", 30))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.api_base_url)

        # 初始化验证码识别器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录响应。
        self.login_result = {}

        # 初始化登录令牌。
        self.token = ""

        # 初始化登录状态。
        self.logged_in = False

    # 定义 KCShop 登录方法。
    def login(self):

        """
        执行 KCShop 协议登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("KCShop 登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 访问前端入口以初始化站点会话。
        self.http.get(self.login_page_url, referer=self.login_page_url)

        # 执行带验证码重试的登录流程。
        return self.login_with_captcha_retry()

    # 定义带验证码重试的登录方法。
    def login_with_captcha_retry(self, max_retries=8):

        """
        使用验证码重试机制执行登录。

        :param max_retries: 最大验证码刷新次数。
        :type max_retries: int
        :return: 登录响应数据。
        :rtype: dict
        """

        # 初始化最后一次失败原因。
        last_error = ""

        # 循环尝试验证码登录。
        for retry_index in range(max_retries):

            # 打印验证码尝试日志。
            self.ph.info(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{max_retries} 次", "cyan"))

            # 获取验证码图片字节流。
            captcha_bytes = self.fetch_captcha_image()

            # 识别验证码文本。
            captcha_code = self.solve_captcha_code(captcha_bytes)

            # 打印验证码识别日志。
            self.ph.info(self.ph.kv("图片验证码识别完成", f"长度={len(captcha_code)}", "cyan"))

            # 构建登录请求体。
            payload = self.build_login_payload(captcha_code)

            # 提交登录请求并解析响应。
            response_data = self.post_json(self.login_api_url, payload, self.login_page_url)

            # 判断当前登录响应是否成功。
            if self.is_success_response(response_data):

                # 保存登录成功响应。
                self.login_result = response_data

                # 保存登录令牌。
                self.token = self.extract_login_token(response_data)

                # 标记登录状态。
                self.logged_in = True

                # 打印登录成功日志。
                self.ph.success(self.ph.join(self.ph.kv("正式登录成功", response_data.get("msg", "登录成功"), "green"), self.ph.kv("Token", self.token[:8], "cyan")))

                # 返回登录响应数据。
                return response_data

            # 保存当前登录失败原因。
            last_error = str(response_data.get("msg") or response_data)[:200]

            # 打印验证码失败日志。
            self.ph.warning(self.ph.kv("登录校验未通过", last_error, "yellow"))

        # 抛出验证码登录失败异常。
        raise RuntimeError(f"{self.company_name} KCShop 登录失败：{last_error}")

    # 定义验证码图片获取方法。
    def fetch_captcha_image(self):

        """
        获取登录验证码图片。

        :return: 验证码图片字节流。
        :rtype: bytes
        """

        # 构建验证码请求头。
        headers = self.build_browser_headers(self.login_page_url)

        # 请求验证码图片。
        result = self.http.request("GET", self.captcha_api_url, headers=headers, timeout=30)

        # 校验验证码响应状态。
        if result.status_code != 200:

            # 抛出验证码获取异常。
            raise RuntimeError(f"{self.company_name} 验证码获取失败：HTTP {result.status_code}")

        # 返回验证码图片字节流。
        return result.body

    # 定义验证码文本识别方法。
    def solve_captcha_code(self, image_bytes):

        """
        识别验证码文本。

        :param image_bytes: 验证码图片字节流。
        :type image_bytes: bytes
        :return: 验证码文本。
        :rtype: str
        """

        # 识别验证码原始文本。
        raw_text = self.captcha_solver.solve_raw_text_image_with_ranges(image_bytes, "0123456789abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ")

        # 清洗验证码文本。
        captcha_code = re.sub(r"[^0-9A-Za-z]", "", raw_text or "")

        # 校验验证码文本是否为空。
        if not captcha_code:

            # 抛出验证码识别异常。
            raise RuntimeError("图片验证码识别为空")

        # 返回清洗后的验证码文本。
        return captcha_code

    # 定义登录请求体构建方法。
    def build_login_payload(self, captcha_code):

        """
        构建登录请求体。

        :param captcha_code: 图片验证码文本。
        :type captcha_code: str
        :return: 登录请求体。
        :rtype: dict
        """

        # 返回登录请求字段。
        return {"account": self.clean_login_name(), "password": self.clean_login_password(), "captcha": captcha_code}



    # 定义登录令牌提取方法。
    def extract_login_token(self, response_data):

        """
        提取登录令牌。

        :param response_data: 登录响应数据。
        :type response_data: dict
        :return: 登录令牌。
        :rtype: str
        """

        # 读取业务数据字段。
        data = response_data.get("data") if isinstance(response_data.get("data"), dict) else {}

        # 读取用户信息字段。
        userinfo = data.get("userinfo") if isinstance(data.get("userinfo"), dict) else {}

        # 读取用户令牌字段。
        token = str(userinfo.get("token") or "").strip()

        # 校验登录令牌是否存在。
        if not token:

            # 抛出令牌缺失异常。
            raise RuntimeError(f"{self.company_name} 登录响应缺少 token")

        # 返回登录令牌。
        return token


    # 定义业务数据查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询 KCShop 业务数据。

        :param data_type: 业务数据类型。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行。
        :rtype: list[dict]
        """

        # 查询指定业务数据。
        query_result = self.query_kcshop_data(data_type, start_date, end_date)

        # 返回查询结果行。
        return query_result["rows"]


    # 定义查询并生成 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询 KCShop 业务数据并生成 Excel。

        :param data_type: 业务数据类型。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定业务数据。
        query_result = self.query_kcshop_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建输出文件名称。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建输出文件路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回输出文件路径。
        return output_path

    # 定义 KCShop 数据查询方法。
    def query_kcshop_data(self, data_type, start_date, end_date):

        """
        查询 KCShop 指定类型数据。

        :param data_type: 业务数据类型。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 校验日期型数据的开始日期。
        if self.requires_date_range(data_type) and not start_date:

            # 抛出开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验日期型数据的结束日期。
        if self.requires_date_range(data_type) and not end_date:

            # 抛出结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 确保当前会话已登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 初始化全部结果行。
        all_rows = []

        # 初始化表头列表。
        headers = []

        # 初始化当前页码。
        page_index = 1

        # 初始化总记录数。
        total_count = None

        # 循环读取分页数据。
        while True:

            # 构建当前分页请求体。
            payload = self.build_query_payload(data_type, start_date, end_date, page_index)

            # 打印当前分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"page={page_index}", "cyan"), self.ph.kv("pageSize", self.page_size, "cyan")))

            # 提交当前分页查询。
            response_data = self.post_json(self.build_query_api_url(data_type), payload, self.login_page_url, timeout=self.query_timeout)

            # 校验分页响应状态。
            if not self.is_success_response(response_data):

                # 抛出查询失败异常。
                raise RuntimeError(f"{self.company_name} {data_type}查询失败：{response_data.get('msg') or response_data}")

            # 提取当前页数据行。
            current_rows = self.extract_rows(response_data)

            # 更新表头列表。
            headers = self.merge_headers(headers, current_rows)

            # 读取响应总记录数。
            total_count = self.resolve_total_count(response_data, total_count)

            # 合并当前页数据行。
            all_rows.extend(current_rows)

            # 打印当前页完成日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页读取完成", f"{len(current_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count if total_count is not None else "未知", "cyan")))

            # 判断是否已经读取全部记录。
            if total_count is not None and len(all_rows) >= total_count:

                # 结束分页循环。
                break

            # 判断当前页是否已经是最后一页。
            if len(current_rows) < self.page_size:

                # 结束分页循环。
                break

            # 递增页码。
            page_index += 1

        # 返回查询结果与表头。
        return {"rows": all_rows, "headers": headers or self.build_default_headers(data_type)}


    # 定义查询阶段标题打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印查询阶段标题。

        :param data_type: 业务数据类型。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断当前数据类型是否需要日期范围。
        if self.requires_date_range(data_type):

            # 打印带日期范围的查询标题。
            self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date, "cyan"), self.ph.kv("结束日期", end_date, "cyan")))

            # 结束标题打印。
            return

        # 打印无日期范围的查询标题。
        self.ph.section(f"{data_type}查询", "当前接口不依赖日期范围")

    # 定义查询日期规范化方法。
    def normalize_query_date(self, date_text):

        """
        规范化接口查询日期。

        :param date_text: 原始日期文本。
        :type date_text: str | None
        :return: 标准日期文本。
        :rtype: str
        """

        # 判断日期文本是否为空。
        if not date_text:

            # 返回空日期文本。
            return ""

        # 转换日期文本为字符串。
        raw_text = str(date_text).strip()

        # 捕获日期解析异常。
        try:

            # 解析常见短日期格式。
            date_value = datetime.strptime(raw_text, "%Y-%m-%d")

            # 返回补零后的标准日期。
            return date_value.strftime("%Y-%m-%d")

        # 处理日期解析失败。
        except ValueError:

            # 返回原始日期文本。
            return raw_text

    # 定义查询请求体构建方法。
    def build_query_payload(self, data_type, start_date, end_date, page_index):

        """
        构建业务查询请求体。

        :param data_type: 业务数据类型。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param page_index: 当前页码。
        :type page_index: int
        :return: 业务查询请求体。
        :rtype: dict
        """

        # 构建通用查询字段。
        payload = {"page": page_index, "total": self.page_size, "search": "", "ordername": "", "ordertype": "", "stardate": "", "enddate": ""}

        # 判断当前数据类型是否需要日期范围。
        if self.requires_date_range(data_type):

            # 规范化查询开始日期。
            normalized_start_date = self.normalize_query_date(start_date)

            # 规范化查询结束日期。
            normalized_end_date = self.normalize_query_date(end_date)

            # 写入查询开始日期。
            payload["stardate"] = normalized_start_date

            # 写入查询结束日期。
            payload["enddate"] = normalized_end_date

        # 返回查询请求体。
        return payload

    # 定义查询接口地址构建方法。
    def build_query_api_url(self, data_type):

        """
        构建数据类型对应的查询接口地址。

        :param data_type: 业务数据类型。
        :type data_type: str
        :return: 查询接口地址。
        :rtype: str
        """

        # 读取数据类型对应的接口路径。
        endpoint = self.endpoint_mapping[data_type]

        # 返回完整查询接口地址。
        return f"{self.api_base_url}/api/Output/{endpoint}"

    # 定义数据类型日期范围判断方法。
    def requires_date_range(self, data_type):

        """
        判断数据类型是否需要日期范围。

        :param data_type: 业务数据类型。
        :type data_type: str
        :return: 是否需要日期范围。
        :rtype: bool
        """

        # 返回日期型数据类型判断结果。
        return data_type in {"配送流向", "购进流向", "纯销流向"}

    # 定义接口路径映射构建方法。
    def build_endpoint_mapping(self):

        """
        构建业务数据类型与接口路径的映射。

        :return: 接口路径映射。
        :rtype: dict
        """

        # 返回业务数据类型接口映射。
        return {"总部库存": "inventoryDetail", "门店库存": "inventorysDetail", "配送流向": "salesDetail", "购进流向": "procurementDetail", "纯销流向": "retailsDetail"}

    # 定义 JSON POST 请求方法。
    def post_json(self, url, payload, referer, timeout=90):

        """
        提交 JSON 请求。

        :param url: 请求地址。
        :type url: str
        :param payload: 请求字段。
        :type payload: dict
        :param referer: 来源页面地址。
        :type referer: str
        :param timeout: 请求超时秒数。
        :type timeout: int
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 序列化 JSON 请求体。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建 JSON 请求头。
        headers = self.build_json_headers(referer)

        # 发送 POST 请求。
        result = self.http.request("POST", url, headers=headers, data=body, timeout=timeout)

        # 解析 JSON 响应。
        return self.decode_json_response(result)

    # 定义浏览器请求头构建方法。
    def build_browser_headers(self, referer):

        """
        构建浏览器请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 浏览器请求头。
        :rtype: dict
        """

        # 返回浏览器请求头。
        return {"Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Origin": self.login_page_url.rstrip("/"), "Referer": referer, "User-Agent": self.http.user_agent()}

    # 定义 JSON 请求头构建方法。
    def build_json_headers(self, referer):

        """
        构建 JSON 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: JSON 请求头。
        :rtype: dict
        """

        # 构建基础 JSON 请求头。
        headers = {"Accept": "application/json, text/plain, */*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/json", "Origin": self.login_page_url.rstrip("/"), "Referer": referer, "User-Agent": self.http.user_agent()}

        # 判断登录令牌是否存在。
        if self.token:

            # 写入登录令牌请求头。
            headers["token"] = self.token

        # 返回 JSON 请求头。
        return headers

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应文本。
        response_text = result.body.decode("utf-8", errors="replace").strip()

        # 判断响应文本是否为空。
        if not response_text:

            # 返回空响应结构。
            return {}

        # 返回 JSON 响应数据。
        return json.loads(response_text)

    # 定义响应成功判断方法。
    def is_success_response(self, response_data):

        """
        判断接口响应是否成功。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 是否成功。
        :rtype: bool
        """

        # 返回接口状态判断结果。
        return str(response_data.get("code")) == "1"

    # 定义数据行提取方法。
    def extract_rows(self, response_data):

        """
        提取响应数据行。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 读取业务数据字段。
        data = response_data.get("data") if isinstance(response_data.get("data"), dict) else {}

        # 读取列表字段。
        rows = data.get("list") if isinstance(data.get("list"), list) else []

        # 返回字典结构数据行。
        return [row for row in rows if isinstance(row, dict)]

    # 定义总记录数解析方法。
    def resolve_total_count(self, response_data, fallback):

        """
        解析总记录数。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :param fallback: 上一次总记录数。
        :type fallback: int | None
        :return: 总记录数。
        :rtype: int | None
        """

        # 读取业务数据字段。
        data = response_data.get("data") if isinstance(response_data.get("data"), dict) else {}

        # 读取总记录数字段。
        count_value = data.get("count")

        # 尝试转换总记录数。
        try:

            # 返回整数总记录数。
            return int(count_value)

        # 处理总记录数转换异常。
        except (TypeError, ValueError):

            # 返回兜底总记录数。
            return fallback

    # 定义表头合并方法。
    def merge_headers(self, headers, rows):

        """
        合并数据行中的字段顺序。

        :param headers: 已有表头列表。
        :type headers: list[str]
        :param rows: 当前页数据行。
        :type rows: list[dict]
        :return: 合并后的表头列表。
        :rtype: list[str]
        """

        # 复制已有表头列表。
        merged_headers = list(headers)

        # 遍历当前页数据行。
        for row in rows:

            # 遍历当前数据行字段。
            for field_name in row.keys():

                # 判断字段是否尚未写入表头。
                if field_name not in merged_headers:

                    # 追加新字段。
                    merged_headers.append(field_name)

        # 返回合并后的表头列表。
        return merged_headers

    # 定义默认表头构建方法。
    def build_default_headers(self, data_type):

        """
        构建默认表头。

        :param data_type: 业务数据类型。
        :type data_type: str
        :return: 默认表头列表。
        :rtype: list[str]
        """

        # 返回默认表头字段。
        return ["ID", "DJBH", "JGBH", "JGMC", "RQ", "SPBH", "SPTYM", "SPGG", "PIHAO", "XSSHL", "FSSHL", "PCSHL"]

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验 KCShop 数据类型。

        :param data_type: 业务数据类型。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否已接入。
        if data_type not in self.endpoint_mapping:

            # 抛出未接入数据类型异常。
            raise NotImplementedError(f"当前 KCShop 暂未接入：{data_type}")


    # 定义中文表头构建方法。
    def build_display_headers(self, data_type, headers):

        """
        构建 Excel 展示表头列表。

        :param data_type: 业务数据类型。
        :type data_type: str
        :param headers: 接口原始字段列表。
        :type headers: list[str]
        :return: Excel 展示表头列表。
        :rtype: list[str]
        """

        # 返回逐字段翻译后的表头列表。
        return [self.translate_header(data_type, header_name) for header_name in headers]

    # 定义字段中文表头翻译方法。
    def translate_header(self, data_type, header_name):

        """
        翻译接口字段为中文表头。

        :param data_type: 业务数据类型。
        :type data_type: str
        :param header_name: 接口原始字段名称。
        :type header_name: str
        :return: 中文表头名称。
        :rtype: str
        """

        # 读取数据类型专属表头映射。
        data_type_labels = KCSHOP_DATA_TYPE_HEADER_LABELS.get(data_type, {})

        # 判断当前字段是否存在数据类型专属中文表头。
        if header_name in data_type_labels:

            # 返回数据类型专属中文表头。
            return data_type_labels[header_name]

        # 返回通用中文表头或原始字段名称。
        return KCSHOP_COMMON_HEADER_LABELS.get(header_name, header_name)

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头字段列表。
        :type headers: list[str]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿类。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 获取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 构建 Excel 展示表头。
        display_headers = self.build_display_headers(sheet_name, headers)

        # 写入表头行。
        worksheet.append(display_headers)

        # 遍历查询结果写入数据行。
        for row in rows:

            # 按表头顺序组装并写入一行数据。
            worksheet.append([row.get(header_name, "") if isinstance(row, dict) else "" for header_name in headers])

        # 遍历所有展示表头设置列宽。
        for column_index, header_name in enumerate(display_headers, start=1):

            # 获取当前列字母。
            column_letter = get_column_letter(column_index)

            # 设置当前列宽度。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除同名历史文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出业务路径文件占用异常。
            raise PermissionError(f"无法保存文件，请先关闭已打开的 Excel 文件：{output_path}") from error

        # 关闭工作簿对象。
        workbook.close()

