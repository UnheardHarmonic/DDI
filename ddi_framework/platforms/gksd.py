# -*- coding: utf-8 -*-
"""
国控山东数据流向服务系统协议适配器。

本模块负责国控山东数据流向服务系统的 WebForms 验证码登录、查询表单提交、分页抓取和 Excel 文件生成。
"""

# 导入正则表达式模块。
import re

# 导入时间模块。
import time

# 导入操作系统路径模块。
import os

# 导入日期解析对象。
from datetime import datetime

# 导入 HTML 标准解析器。
from html.parser import HTMLParser

# 导入 URL 拼接函数。
from urllib.parse import urljoin

# 导入表单编码函数。
from urllib.parse import urlencode

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入验证码识别器。
from ddi_framework import CaptchaSolver

# 导入日志打印器。
from ddi_framework import PrintHandler

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient


# 定义国控山东 WebForms 表单解析器。
class GksdFormParser(HTMLParser):

    """
    解析国控山东 WebForms 页面表单字段。

    :return: 国控山东 WebForms 表单解析器实例。
    :rtype: GksdFormParser
    """

    # 定义解析器初始化方法。
    def __init__(self):

        """
        初始化表单字段解析器。

        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(convert_charrefs=True)

        # 初始化表单字段字典。
        self.fields = {}

        # 初始化按钮字段字典。
        self.buttons = {}

        # 初始化验证码地址。
        self.captcha_src = ""

        # 初始化当前下拉框名称。
        self.current_select_name = ""

        # 初始化当前下拉框首个选项值。
        self.current_select_first_value = ""

    # 定义开始标签处理方法。
    def handle_starttag(self, tag, attrs):

        """
        处理 HTML 开始标签。

        :param tag: 标签名称。
        :type tag: str
        :param attrs: 标签属性列表。
        :type attrs: list[tuple[str, str | None]]
        :return: None。
        :rtype: None
        """

        # 转换标签属性为字典。
        attr_map = {name: value or "" for name, value in attrs}

        # 转换标签名称。
        tag_name = tag.lower()

        # 处理输入控件。
        if tag_name == "input":

            # 处理输入控件字段。
            self.handle_input_tag(attr_map)

        # 处理图片标签。
        if tag_name == "img":

            # 读取图片地址。
            image_src = attr_map.get("src", "")

            # 判断图片地址是否为验证码。
            if "login_validata" in image_src.lower():

                # 保存验证码地址。
                self.captcha_src = image_src

        # 处理下拉控件。
        if tag_name == "select":

            # 保存当前下拉框名称。
            self.current_select_name = attr_map.get("name", "")

            # 清理当前下拉框首个选项值。
            self.current_select_first_value = ""

        # 处理下拉选项。
        if tag_name == "option" and self.current_select_name:

            # 处理下拉选项字段。
            self.handle_option_tag(attr_map)

    # 定义结束标签处理方法。
    def handle_endtag(self, tag):

        """
        处理 HTML 结束标签。

        :param tag: 标签名称。
        :type tag: str
        :return: None。
        :rtype: None
        """

        # 判断当前标签是否为下拉框。
        if tag.lower() == "select":

            # 判断下拉框是否没有选中值。
            if self.current_select_name and self.current_select_name not in self.fields:

                # 写入下拉框首个选项值。
                self.fields[self.current_select_name] = self.current_select_first_value

            # 清理当前下拉框名称。
            self.current_select_name = ""

            # 清理当前下拉框首个选项值。
            self.current_select_first_value = ""

    # 定义输入控件处理方法。
    def handle_input_tag(self, attr_map):

        """
        处理输入控件字段。

        :param attr_map: 输入控件属性字典。
        :type attr_map: dict
        :return: None。
        :rtype: None
        """

        # 读取输入控件名称。
        field_name = attr_map.get("name", "")

        # 判断字段名称是否为空。
        if not field_name:

            # 结束当前控件处理。
            return

        # 读取输入控件类型。
        field_type = attr_map.get("type", "text").lower()

        # 读取输入控件值。
        field_value = attr_map.get("value", "")

        # 判断输入控件是否为验证码图片按钮。
        if field_type == "image" and "login_validata" in attr_map.get("src", "").lower():

            # 保存验证码地址。
            self.captcha_src = attr_map.get("src", "")

        # 判断输入控件是否为提交按钮。
        if field_type in ("submit", "button", "image"):

            # 保存按钮字段。
            self.buttons[field_name] = field_value or field_type

            # 结束按钮字段处理。
            return

        # 判断单选框或复选框是否未选中。
        if field_type in ("radio", "checkbox") and "checked" not in attr_map:

            # 跳过未选中的选择控件。
            return

        # 保存输入控件字段。
        self.fields[field_name] = field_value

    # 定义下拉选项处理方法。
    def handle_option_tag(self, attr_map):

        """
        处理下拉选项字段。

        :param attr_map: 下拉选项属性字典。
        :type attr_map: dict
        :return: None。
        :rtype: None
        """

        # 读取选项值。
        option_value = attr_map.get("value", "")

        # 判断当前选项是否为第一个选项。
        if self.current_select_first_value == "":

            # 保存第一个选项值。
            self.current_select_first_value = option_value

        # 判断当前选项是否选中。
        if "selected" in attr_map:

            # 保存下拉框选中值。
            self.fields[self.current_select_name] = option_value


# 定义国控山东 HTML 表格解析器。
class GksdTableParser(HTMLParser):

    """
    解析国控山东页面中的 HTML 表格。

    :return: 国控山东 HTML 表格解析器实例。
    :rtype: GksdTableParser
    """

    # 定义解析器初始化方法。
    def __init__(self):

        """
        初始化 HTML 表格解析器。

        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(convert_charrefs=True)

        # 初始化全部表格字典。
        self.tables = {}

        # 初始化当前表格编号。
        self.current_table_id = ""

        # 初始化业务表格嵌套层级。
        self.current_table_depth = 0

        # 初始化当前行。
        self.current_row = None

        # 初始化当前单元格。
        self.current_cell = None

        # 初始化表格嵌套深度。
        self.table_depth = 0

    # 定义开始标签处理方法。
    def handle_starttag(self, tag, attrs):

        """
        处理 HTML 开始标签。

        :param tag: 标签名称。
        :type tag: str
        :param attrs: 标签属性列表。
        :type attrs: list[tuple[str, str | None]]
        :return: None。
        :rtype: None
        """

        # 转换标签属性为字典。
        attr_map = {name: value or "" for name, value in attrs}

        # 转换标签名称。
        tag_name = tag.lower()

        # 处理表格开始标签。
        if tag_name == "table":

            # 增加表格嵌套深度。
            self.table_depth += 1

            # 判断是否进入业务表格。
            if attr_map.get("id") == "MyGridView1_gvmygridview":

                # 保存当前业务表格编号。
                self.current_table_id = attr_map.get("id", "")

                # 保存当前业务表格层级。
                self.current_table_depth = self.table_depth

                # 初始化业务表格行列表。
                self.tables[self.current_table_id] = []

        # 处理业务表格行。
        if tag_name == "tr" and self.current_table_id and self.table_depth >= 1:

            # 初始化当前行。
            self.current_row = []

        # 处理业务表格单元格。
        if tag_name in ("td", "th") and self.current_row is not None:

            # 初始化当前单元格文本列表。
            self.current_cell = []

    # 定义结束标签处理方法。
    def handle_endtag(self, tag):

        """
        处理 HTML 结束标签。

        :param tag: 标签名称。
        :type tag: str
        :return: None。
        :rtype: None
        """

        # 转换标签名称。
        tag_name = tag.lower()

        # 处理单元格结束标签。
        if tag_name in ("td", "th") and self.current_cell is not None and self.current_row is not None:

            # 追加当前单元格文本。
            self.current_row.append(" ".join(self.current_cell).strip())

            # 清理当前单元格。
            self.current_cell = None

        # 处理行结束标签。
        if tag_name == "tr" and self.current_row is not None and self.current_table_id:

            # 判断当前行是否存在有效文本。
            if any(value.strip() for value in self.current_row):

                # 追加当前行。
                self.tables[self.current_table_id].append(self.current_row)

            # 清理当前行。
            self.current_row = None

        # 处理表格结束标签。
        if tag_name == "table":

            # 判断是否离开当前业务表格。
            if self.current_table_id and self.table_depth == self.current_table_depth:

                # 清理当前表格编号。
                self.current_table_id = ""

                # 清理业务表格层级。
                self.current_table_depth = 0

            # 降低表格嵌套深度。
            self.table_depth = max(0, self.table_depth - 1)

    # 定义文本处理方法。
    def handle_data(self, data):

        """
        处理 HTML 文本内容。

        :param data: 文本内容。
        :type data: str
        :return: None。
        :rtype: None
        """

        # 判断当前单元格是否存在。
        if self.current_cell is None:

            # 结束非单元格文本处理。
            return

        # 清理文本内容。
        text = " ".join(str(data).split()).strip()

        # 判断文本是否为空。
        if text:

            # 追加单元格文本。
            self.current_cell.append(text)


# 定义国控山东数据流向服务系统平台适配器。
class GksdPlatform(BasePlatform):

    """
    国控山东数据流向服务系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 国控山东数据流向服务系统协议适配器实例。
    :rtype: GksdPlatform
    """

    # 定义平台初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化平台协议适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存平台根地址。
        self.base_url = account.get("接口根地址", "http://222.173.24.186:9090")

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", f"{self.base_url}/")

        # 保存响应编码。
        self.response_encoding = account.get("响应编码", "gb2312")

        # 保存表单编码。
        self.form_encoding = account.get("表单编码", "gb2312")

        # 保存货主编号。
        self.owner_id = str(account.get("货主编号", "20"))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化验证码识别器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录状态。
        self.logged_in = False

    # 定义平台登录方法。
    def login(self):

        """
        执行平台验证码登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("国控山东登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 执行验证码登录重试。
        login_result = self.login_with_captcha_retry()

        # 标记登录状态。
        self.logged_in = True

        # 返回登录结果。
        return login_result

    # 定义验证码登录重试方法。
    def login_with_captcha_retry(self, max_retries=8):

        """
        使用图片验证码执行登录重试。

        :param max_retries: 最大重试次数。
        :type max_retries: int
        :return: 登录结果。
        :rtype: dict
        """

        # 初始化最近一次错误信息。
        last_error = ""

        # 遍历验证码重试次数。
        for retry_index in range(max_retries):

            # 打印验证码尝试日志。
            self.ph.info(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{max_retries} 次", "cyan"))

            # 获取登录上下文。
            context = self.get_login_context()

            # 识别验证码文本。
            captcha_code = self.captcha_solver.solve_text_image(context["验证码字节"], min_length=4)

            # 打印验证码识别结果长度。
            self.ph.info(self.ph.kv("图片验证码识别完成", f"长度={len(captcha_code)}", "cyan"))

            # 构建登录表单。
            payload = self.build_login_payload(context["表单字段"], captcha_code)

            # 提交登录表单。
            response_text = self.submit_form(self.login_page_url, payload, self.login_page_url)

            # 判断登录是否成功。
            if self.is_login_success(response_text):

                # 打印登录成功日志。
                self.ph.success(self.ph.join(self.ph.kv("登录状态", "成功", "green"), self.ph.kv("Cookie", len(list(self.http.cookie_jar)), "cyan")))

                # 返回登录结果。
                return {"登录状态": "成功", "Cookie数量": len(list(self.http.cookie_jar))}

            # 读取登录失败原因。
            last_error = self.extract_login_error(response_text)

            # 打印登录失败日志。
            self.ph.warning(self.ph.kv("登录校验未通过", last_error, "yellow"))

            # 等待后刷新验证码。
            time.sleep(0.8)

        # 抛出验证码登录失败异常。
        raise RuntimeError(f"{self.company_name} 验证码登录失败：{last_error}")

    # 定义登录上下文获取方法。
    def get_login_context(self):

        """
        获取登录页表单字段与验证码图片。

        :return: 登录上下文字典。
        :rtype: dict
        """

        # 请求登录页。
        page_text = self.get_text(self.login_page_url, self.base_url)

        # 解析登录页表单。
        parser = self.parse_form(page_text)

        # 标准化验证码地址。
        captcha_url = self.normalize_captcha_url(parser.captcha_src)

        # 请求验证码图片。
        captcha_result = self.http.get(captcha_url, referer=self.login_page_url)

        # 返回登录上下文。
        return {"表单字段": parser.fields, "验证码字节": captcha_result.body}

    # 定义登录表单构建方法。
    def build_login_payload(self, fields, captcha_code):

        """
        构建登录提交表单。

        :param fields: 登录页原始表单字段。
        :type fields: dict
        :param captcha_code: 验证码文本。
        :type captcha_code: str
        :return: 登录提交表单。
        :rtype: dict
        """

        # 复制登录页字段。
        payload = dict(fields)

        # 写入登录账号。
        payload["tbaccount"] = self.account.get("账号", "")

        # 写入登录密码。
        payload["tbpassword"] = self.account.get("密码", "")

        # 写入验证码文本。
        payload["tbvalid"] = captcha_code

        # 写入记住登录字段。
        payload["cb_save"] = "on"

        # 写入屏幕宽度字段。
        payload["tbwidth"] = "1920"

        # 写入屏幕高度字段。
        payload["tbheight"] = "1032"

        # 写入登录图片按钮横坐标。
        payload["imb_login.x"] = "46"

        # 写入登录图片按钮纵坐标。
        payload["imb_login.y"] = "8"

        # 返回登录提交表单。
        return payload


    # 定义平台查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定数据类型并返回行数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行。
        :rtype: list[dict]
        """

        # 查询页面数据。
        result = self.query_webforms_pages(data_type, start_date, end_date)

        # 返回结果行。
        return result["rows"]

    # 定义查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定数据类型并生成 Excel 文件。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 校验数据类型。
        self.validate_data_type(data_type)

        # 打印数据类型查询分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 查询页面数据。
        result = self.query_webforms_pages(data_type, start_date, end_date)

        # 解析输出目录。
        target_dir = output_dir or os.getcwd()

        # 构建输出文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建输出文件路径。
        output_path = os.path.join(target_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(result["rows"], result["headers"], output_path, data_type)

        # 打印文件生成日志。
        self.ph.success(self.ph.join(self.ph.kv("生成文件", output_path, "green"), self.ph.kv("记录数", len(result["rows"]), "cyan")))

        # 返回输出文件路径。
        return output_path

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie。

        :param output_dir: Cookie 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建 Cookie 输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件名。
        file_name = f"{self.company_name}_{self.account.get('平台')}_cookie.json"

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, file_name)

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义 WebForms 分页查询方法。
    def query_webforms_pages(self, data_type, start_date, end_date):

        """
        查询 WebForms 页面并抓取全部分页。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果结构。
        :rtype: dict
        """

        # 校验登录状态。
        self.ensure_logged_in()

        # 读取页面配置。
        page_config = self.resolve_page_config(data_type)

        # 构建业务页面地址。
        page_url = urljoin(self.base_url + "/", page_config["路径"].lstrip("/"))

        # 打印业务页加载日志。
        self.ph.info(self.ph.join(self.ph.kv("加载业务页", page_config["页面名称"], "cyan"), self.ph.kv("端点", page_config["路径"], "yellow")))

        # 请求业务页面。
        page_text = self.get_text(page_url, self.login_page_url)

        # 构建查询表单。
        payload = self.build_query_payload(page_text, page_config, start_date, end_date)

        # 打印查询提交日志。
        self.ph.info(self.ph.join(self.ph.kv("提交查询", data_type, "cyan"), self.ph.kv("货主", self.owner_id, "yellow")))

        # 提交查询表单。
        query_text = self.submit_form(page_url, payload, page_url)

        # 解析首页结果。
        rows, headers = self.parse_result_rows(query_text)

        # 解析总页数。
        total_pages = self.extract_total_pages(query_text)

        # 打印首页解析日志。
        self.ph.info(self.ph.join(self.ph.kv("分页", f"1/{total_pages}", "cyan"), self.ph.kv("本页", len(rows), "yellow"), self.ph.kv("累计", len(rows), "green")))

        # 保存当前页面文本。
        current_text = query_text

        # 遍历后续页码。
        for page_index in range(2, total_pages + 1):

            # 构建下一页表单。
            next_payload = self.build_next_page_payload(current_text, page_config, start_date, end_date)

            # 提交下一页表单。
            current_text = self.submit_form(page_url, next_payload, page_url)

            # 解析当前页结果。
            page_rows, page_headers = self.parse_result_rows(current_text)

            # 合并动态表头。
            headers = self.merge_headers(headers, page_headers)

            # 追加当前页数据。
            rows.extend(page_rows)

            # 打印当前页解析日志。
            self.ph.info(self.ph.join(self.ph.kv("分页", f"{page_index}/{total_pages}", "cyan"), self.ph.kv("本页", len(page_rows), "yellow"), self.ph.kv("累计", len(rows), "green")))

        # 验证业务日期范围。
        self.validate_result_dates(data_type, rows, start_date, end_date)

        # 返回查询结果。
        return {"rows": rows, "headers": headers}

    # 定义查询表单构建方法。
    def build_query_payload(self, page_text, page_config, start_date, end_date):

        """
        构建业务查询表单。

        :param page_text: 业务页 HTML。
        :type page_text: str
        :param page_config: 页面配置。
        :type page_config: dict
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询表单字段。
        :rtype: dict
        """

        # 解析业务页表单。
        parser = self.parse_form(page_text)

        # 复制表单字段。
        payload = dict(parser.fields)

        # 写入货主编号。
        payload["ddl_owner"] = self.owner_id

        # 写入模糊查询字段。
        payload["tbbh"] = ""

        # 写入屏幕宽度字段。
        payload["tbwidth"] = "1920"

        # 写入屏幕高度字段。
        payload["tbheight"] = "1032"

        # 写入销售日期口径。
        self.apply_sale_date_basis(payload, page_config)

        # 写入日期字段。
        self.apply_date_fields(payload, page_config, start_date, end_date)

        # 写入查询按钮。
        payload["btchaxun"] = "查询"

        # 返回查询表单字段。
        return payload

    # 定义下一页表单构建方法。
    def build_next_page_payload(self, page_text, page_config, start_date, end_date):

        """
        构建下一页提交表单。

        :param page_text: 当前页 HTML。
        :type page_text: str
        :param page_config: 页面配置。
        :type page_config: dict
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 下一页表单字段。
        :rtype: dict
        """

        # 解析当前页表单。
        parser = self.parse_form(page_text)

        # 复制当前页表单字段。
        payload = dict(parser.fields)

        # 写入货主编号。
        payload["ddl_owner"] = self.owner_id

        # 写入模糊查询字段。
        payload["tbbh"] = ""

        # 写入屏幕宽度字段。
        payload["tbwidth"] = "1920"

        # 写入屏幕高度字段。
        payload["tbheight"] = "1032"

        # 写入销售日期口径。
        self.apply_sale_date_basis(payload, page_config)

        # 写入日期字段。
        self.apply_date_fields(payload, page_config, start_date, end_date)

        # 写入下一页按钮横坐标。
        payload["MyGridView1$ImageButton3.x"] = "8"

        # 写入下一页按钮纵坐标。
        payload["MyGridView1$ImageButton3.y"] = "8"

        # 返回下一页表单字段。
        return payload

    # 定义销售日期口径写入方法。
    def apply_sale_date_basis(self, payload, page_config):

        """
        写入销售页面日期口径字段。

        :param payload: 表单字段。
        :type payload: dict
        :param page_config: 页面配置。
        :type page_config: dict
        :return: None。
        :rtype: None
        """

        # 判断页面是否存在销售日期口径。
        if page_config.get("日期口径"):

            # 写入销售日期口径。
            payload["rb_riqi"] = page_config["日期口径"]

    # 定义日期字段写入方法。
    def apply_date_fields(self, payload, page_config, start_date, end_date):

        """
        写入页面日期查询字段。

        :param payload: 表单字段。
        :type payload: dict
        :param page_config: 页面配置。
        :type page_config: dict
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断当前页面是否需要日期。
        if not page_config.get("需要日期", False):

            # 结束无日期页面处理。
            return

        # 写入开始日期。
        payload["GMDatePicker1$ctl00"] = self.format_page_date(start_date)

        # 写入结束日期。
        payload["GMDatePicker2$ctl00"] = self.format_page_date(end_date)

    # 定义查询结果解析方法。
    def parse_result_rows(self, html_text):

        """
        解析查询响应中的业务表格。

        :param html_text: 查询响应 HTML。
        :type html_text: str
        :return: 查询结果行与表头。
        :rtype: tuple[list[dict], list[str]]
        """

        # 初始化表格解析器。
        parser = GksdTableParser()

        # 输入 HTML 文本。
        parser.feed(html_text)

        # 读取业务表格。
        table_rows = parser.tables.get("MyGridView1_gvmygridview", [])

        # 判断业务表格是否为空。
        if not table_rows:

            # 返回空结果。
            return [], []

        # 解析表头。
        headers = self.deduplicate_headers(table_rows[0])

        # 初始化结果行列表。
        rows = []

        # 遍历业务数据行。
        for raw_row in table_rows[1:]:

            # 判断当前行是否为空。
            if not any(value.strip() for value in raw_row):

                # 跳过空行。
                continue

            # 构建结果行。
            rows.append(self.build_row_data(headers, raw_row))

        # 返回查询结果行和表头。
        return rows, headers

    # 定义结果行构建方法。
    def build_row_data(self, headers, raw_row):

        """
        构建字典形式结果行。

        :param headers: 表头列表。
        :type headers: list[str]
        :param raw_row: 原始表格行。
        :type raw_row: list[str]
        :return: 结果行字典。
        :rtype: dict
        """

        # 初始化结果行字典。
        row_data = {}

        # 遍历表头字段。
        for column_index, header_name in enumerate(headers):

            # 读取单元格值。
            cell_value = raw_row[column_index] if column_index < len(raw_row) else ""

            # 写入字段值。
            row_data[header_name] = cell_value

        # 返回结果行字典。
        return row_data

    # 定义表头合并方法。
    def merge_headers(self, old_headers, new_headers):

        """
        合并分页动态表头。

        :param old_headers: 已有表头。
        :type old_headers: list[str]
        :param new_headers: 新页面表头。
        :type new_headers: list[str]
        :return: 合并后的表头。
        :rtype: list[str]
        """

        # 复制已有表头。
        result = list(old_headers)

        # 遍历新页面表头。
        for header_name in new_headers:

            # 判断表头是否尚不存在。
            if header_name not in result:

                # 追加新表头。
                result.append(header_name)

        # 返回合并表头。
        return result

    # 定义表头去重方法。
    def deduplicate_headers(self, headers):

        """
        去重并补齐表头名称。

        :param headers: 原始表头列表。
        :type headers: list[str]
        :return: 标准表头列表。
        :rtype: list[str]
        """

        # 初始化表头计数字典。
        header_counts = {}

        # 初始化结果表头列表。
        result = []

        # 遍历原始表头。
        for column_index, header_name in enumerate(headers, start=1):

            # 清理表头文本。
            clean_name = str(header_name or "").strip() or f"字段{column_index}"

            # 累加表头出现次数。
            header_counts[clean_name] = header_counts.get(clean_name, 0) + 1

            # 判断表头是否重复。
            if header_counts[clean_name] > 1:

                # 追加重复表头序号。
                result.append(f"{clean_name}_{header_counts[clean_name]}")

            # 处理首次出现表头。
            else:

                # 追加原始表头。
                result.append(clean_name)

        # 返回结果表头。
        return result

    # 定义业务日期验证方法。
    def validate_result_dates(self, data_type, rows, start_date, end_date):

        """
        验证流向查询结果的业务日期范围。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果行。
        :type rows: list[dict]
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断库存是否不需要日期校验。
        if data_type == "库存":

            # 结束库存日期验证。
            return

        # 判断结果是否为空。
        if not rows:

            # 结束空结果验证。
            return

        # 读取页面配置。
        page_config = self.resolve_page_config(data_type)

        # 读取日期字段名称。
        date_field = page_config.get("日期字段", "")

        # 解析起始日期。
        start_obj = self.parse_date(start_date)

        # 解析结束日期。
        end_obj = self.parse_date(end_date)

        # 初始化日期列表。
        values = []

        # 遍历查询结果。
        for row in rows:

            # 读取业务日期文本。
            raw_value = row.get(date_field, "")

            # 判断业务日期是否存在。
            if raw_value:

                # 追加业务日期。
                values.append(self.parse_date(raw_value[:10]))

        # 判断业务日期是否为空。
        if not values:

            # 结束无日期字段验证。
            return

        # 读取最小业务日期。
        min_value = min(values)

        # 读取最大业务日期。
        max_value = max(values)

        # 判断业务日期是否越界。
        if min_value < start_obj or max_value > end_obj:

            # 抛出日期范围异常。
            raise RuntimeError(f"{self.company_name} {data_type} 业务日期越界：{min_value.date()} 至 {max_value.date()}")

        # 打印日期范围日志。
        self.ph.info(self.ph.join(self.ph.kv("业务日期", f"{min_value.date()} 至 {max_value.date()}", "cyan"), self.ph.kv("字段", date_field, "yellow")))

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel 文件。

        :param rows: 查询结果行。
        :type rows: list[dict]
        :param headers: 表头列表。
        :type headers: list[str]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列字母工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 判断表头是否为空。
        if not headers:

            # 写入空结果表头。
            headers = ["查询结果"]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果行。
        for row in rows:

            # 写入当前数据行。
            worksheet.append([row.get(header_name, "") for header_name in headers])

        # 遍历表头字段。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取列字母。
            column_letter = get_column_letter(column_index)

            # 设置列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 8))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 设置自动筛选范围。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 判断同名文件是否存在。
        if os.path.exists(output_path):

            # 删除同名旧文件。
            os.remove(output_path)

        # 保存 Excel 文件。
        workbook.save(output_path)

        # 关闭工作簿。
        workbook.close()

    # 定义文本页面请求方法。
    def get_text(self, url, referer):

        """
        请求页面并按站点编码解码。

        :param url: 请求地址。
        :type url: str
        :param referer: 来源地址。
        :type referer: str
        :return: 页面文本。
        :rtype: str
        """

        # 请求页面。
        result = self.http.get(url, referer=referer)

        # 返回解码文本。
        return self.decode_response(result.body)

    # 定义表单提交方法。
    def submit_form(self, url, payload, referer):

        """
        提交 WebForms 表单并返回响应文本。

        :param url: 请求地址。
        :type url: str
        :param payload: 表单字段。
        :type payload: dict
        :param referer: 来源地址。
        :type referer: str
        :return: 响应文本。
        :rtype: str
        """

        # 构建请求头。
        headers = self.build_form_headers(referer)

        # 编码请求体。
        body = urlencode(payload, encoding=self.form_encoding, errors="ignore").encode("ascii")

        # 提交表单请求。
        result = self.http.request("POST", url, headers=headers, data=body, timeout=90)

        # 返回解码文本。
        return self.decode_response(result.body)

    # 定义请求头构建方法。
    def build_form_headers(self, referer):

        """
        构建 WebForms 表单请求头。

        :param referer: 来源地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回浏览器表单请求头。
        return {"Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8", "Accept-Language": "zh-CN,zh;q=0.9", "Content-Type": "application/x-www-form-urlencoded", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent()}

    # 定义响应解码方法。
    def decode_response(self, body):

        """
        解码站点响应字节。

        :param body: 响应字节。
        :type body: bytes
        :return: 响应文本。
        :rtype: str
        """

        # 返回响应文本。
        return body.decode(self.response_encoding, errors="ignore")

    # 定义表单解析方法。
    def parse_form(self, html_text):

        """
        解析页面表单字段。

        :param html_text: 页面 HTML 文本。
        :type html_text: str
        :return: 表单解析器。
        :rtype: GksdFormParser
        """

        # 初始化表单解析器。
        parser = GksdFormParser()

        # 输入页面文本。
        parser.feed(html_text)

        # 校验 WebForms 状态字段。
        if "__VIEWSTATE" not in parser.fields:

            # 抛出页面状态异常。
            raise RuntimeError(f"{self.company_name} 页面未返回 __VIEWSTATE")

        # 返回表单解析器。
        return parser

    # 定义验证码地址标准化方法。
    def normalize_captcha_url(self, captcha_src):

        """
        标准化验证码图片地址。

        :param captcha_src: 登录页验证码地址。
        :type captcha_src: str
        :return: 绝对验证码地址。
        :rtype: str
        """

        # 清理验证码地址。
        normalized_src = str(captcha_src or "").strip()

        # 判断验证码地址是否包含脚本时间表达式。
        if "Math.round(new Date().getTime() / 1000)" in normalized_src:

            # 替换脚本时间表达式。
            normalized_src = normalized_src.replace("Math.round(new Date().getTime() / 1000)", str(int(time.time())))

        # 返回验证码绝对地址。
        return urljoin(self.base_url + "/", normalized_src.lstrip("/"))

    # 定义登录成功判断方法。
    def is_login_success(self, response_text):

        """
        判断登录响应是否表示登录成功。

        :param response_text: 登录响应文本。
        :type response_text: str
        :return: 是否登录成功。
        :rtype: bool
        """

        # 判断是否存在验证码错误。
        if "验证码错误" in response_text:

            # 返回登录失败状态。
            return False

        # 判断是否存在账号密码错误。
        if "密码" in response_text and "错误" in response_text:

            # 返回登录失败状态。
            return False

        # 返回登录成功状态。
        return len(list(self.http.cookie_jar)) >= 4

    # 定义登录错误提取方法。
    def extract_login_error(self, response_text):

        """
        提取登录失败原因。

        :param response_text: 登录响应文本。
        :type response_text: str
        :return: 登录失败原因。
        :rtype: str
        """

        # 判断是否存在验证码错误。
        if "验证码错误" in response_text:

            # 返回验证码错误。
            return "验证码错误"

        # 判断是否存在登录超时。
        if "登录已超时" in response_text:

            # 返回登录超时。
            return "登录已超时"

        # 返回默认失败原因。
        return "登录响应未进入业务会话"

    # 定义登录状态校验方法。
    def ensure_logged_in(self):

        """
        确保当前适配器已经登录。

        :return: None。
        :rtype: None
        """

        # 判断当前是否未登录。
        if not self.logged_in:

            # 执行平台登录。
            self.login()

    # 定义页面配置解析方法。
    def resolve_page_config(self, data_type):

        """
        解析数据类型页面配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 页面配置。
        :rtype: dict
        """

        # 读取页面映射。
        page_map = self.account.get("页面映射", {})

        # 判断数据类型是否存在。
        if data_type not in page_map:

            # 抛出数据类型异常。
            raise ValueError(f"{self.company_name} 暂未接入数据类型：{data_type}")

        # 返回页面配置。
        return dict(page_map[data_type])

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验平台是否支持指定数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 读取账号配置中的数据类型列表。
        data_types = self.account.get("数据类型列表", [])

        # 判断当前数据类型是否已接入。
        if data_type not in data_types:

            # 抛出数据类型异常。
            raise ValueError(f"{self.company_name} 暂未接入数据类型：{data_type}")

    # 定义总页数解析方法。
    def extract_total_pages(self, html_text):

        """
        解析查询结果总页数。

        :param html_text: 查询响应 HTML。
        :type html_text: str
        :return: 总页数。
        :rtype: int
        """

        # 匹配分页标签文本。
        match_obj = re.search(r"第\s*\d+\s*页\s*/\s*共\s*(\d+)\s*页", html_text)

        # 判断分页标签是否存在。
        if not match_obj:

            # 返回默认页数。
            return 1

        # 返回总页数。
        return max(1, int(match_obj.group(1)))

    # 定义页面日期格式化方法。
    def format_page_date(self, date_value):

        """
        格式化页面日期文本。

        :param date_value: 输入日期。
        :type date_value: str | None
        :return: 页面日期文本。
        :rtype: str
        """

        # 解析日期对象。
        date_obj = self.parse_date(date_value)

        # 返回页面日期文本。
        return f"{date_obj.year}/{date_obj.month:02d}/{date_obj.day:02d}"

    # 定义日期解析方法。
    def parse_date(self, date_value):

        """
        解析常见日期文本。

        :param date_value: 日期文本。
        :type date_value: str | None
        :return: 日期对象。
        :rtype: datetime
        """

        # 判断日期文本是否为空。
        if not date_value:

            # 返回当前日期。
            return datetime.now()

        # 清理日期文本。
        clean_value = str(date_value).strip().replace("/", "-")

        # 截取日期部分。
        clean_value = clean_value[:10]

        # 返回日期对象。
        return datetime.strptime(clean_value, "%Y-%m-%d")
