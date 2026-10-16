# -*- coding: utf-8 -*-
"""
太极大药房流向查询系统协议适配器测试。

本模块使用固定 HTML 表格验证页面解析、字段完整性和 Excel 落地逻辑。

:return: 模块提供太极平台协议适配器测试。
:rtype: None
"""

# 导入临时目录模块。
import tempfile

# 导入单元测试模块。
import unittest

# 导入 Excel 工作簿读取函数。
from openpyxl import load_workbook

# 导入太极平台适配器。
from ddi_framework.platforms.taiji import TaijiPlatform


# 定义太极平台适配器测试类。
class TaijiPlatformTests(unittest.TestCase):

    """
    验证太极平台 HTML 表格解析和 Excel 原始字段保留。

    :return: None。
    :rtype: None
    """

    # 定义 HTML 表格解析与 Excel 写入测试方法。
    def test_parser_preserves_raw_headers_and_rows_in_excel(self):

        """
        验证页面表头和记录会按真实文本与顺序写入正式原始 Excel。

        :return: None。
        :rtype: None
        """

        # 构建模拟页面 HTML 表格。
        html_text = "<table><tr><td>日期</td><td>供应商id</td><td>货品id</td><td>品名</td><td>数量</td></tr><tr><td>2026-03-03</td><td>S01</td><td>134412</td><td>测试药品</td><td>5</td></tr></table>"

        # 构建最小平台账号配置。
        account = {"平台": "taiji", "页面映射": {"购进流向": {"查询页": "show_su.jsp", "详情页": "show_su_detail.jsp", "需要日期": True}}}

        # 创建待测平台适配器。
        platform = TaijiPlatform("测试商业", account)

        # 解析模拟页面表格。
        headers, rows = platform.parse_result_table(html_text)

        # 断言原始表头完整保留。
        self.assertEqual(headers, ["日期", "供应商id", "货品id", "品名", "数量"])

        # 断言页面记录完整解析。
        self.assertEqual(rows, [{"日期": "2026-03-03", "供应商id": "S01", "货品id": "134412", "品名": "测试药品", "数量": "5"}])

        # 创建临时 Excel 输出目录。
        with tempfile.TemporaryDirectory() as output_dir:

            # 构建临时 Excel 文件路径。
            output_path = f"{output_dir}/taiji.xlsx"

            # 写入测试 Excel 文件。
            platform.write_rows_to_excel(rows, headers, "购进流向", output_path)

            # 读取生成后的 Excel 工作簿。
            workbook = load_workbook(output_path, read_only=True)

            # 读取当前工作表。
            worksheet = workbook.active

            # 读取 Excel 表头。
            excel_headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]

            # 读取 Excel 首行记录。
            excel_row = [cell.value for cell in next(worksheet.iter_rows(min_row=2, max_row=2))]

            # 关闭 Excel 工作簿。
            workbook.close()

        # 断言 Excel 表头严格保持页面原始字段顺序。
        self.assertEqual(excel_headers, headers)

        # 断言 Excel 数据严格保持页面原始记录顺序。
        self.assertEqual(excel_row, ["2026-03-03", "S01", "134412", "测试药品", "5"])

    # 定义库存查询请求体测试方法。
    def test_stock_query_payload_omits_date_parameters(self):

        """
        验证库存查询仅提交货品和查询动作字段。

        :return: None。
        :rtype: None
        """

        # 构建包含库存页面配置的最小账号配置。
        account = {"平台": "taiji", "货品ID": "134412", "页面映射": {"库存": {"查询页": "show_st.jsp", "详情页": "show_st_detail.jsp", "需要日期": False}}}

        # 创建待测平台适配器。
        platform = TaijiPlatform("测试商业", account)

        # 构建库存查询请求体。
        payload = platform.build_query_payload(account["页面映射"]["库存"], "2026-03-03", "2026-03-03")

        # 断言库存请求包含货品 ID。
        self.assertEqual(payload["goodsid"], "134412")

        # 断言库存请求不包含开始日期。
        self.assertNotIn("startdate", payload)

        # 断言库存请求不包含结束日期。
        self.assertNotIn("enddate", payload)
