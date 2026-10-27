# -*- coding: utf-8 -*-
"""
DDI 公共单商业协议执行服务。

本模块承载登录、查询、导出、Excel 落地、复制用 Excel 和 Excel 任务日志等业务行为。
终端入口与图形界面仅通过本模块调用同一份平台适配器能力。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期时间模块。
from datetime import datetime

# 导入 Excel 标准样式函数。
from ddi_framework.excel_style import apply_standard_excel_style

# 导入复制用流向 Excel 生成函数。
from ddi_framework.flow_cleaner import generate_copy_flow_excel

# 导入平台适配器工厂函数。
from ddi_framework.platform_registry import create_platform

# 导入全局任务配置读取函数。
from ddi_framework.task_config import get_task_config


# 定义数据类型解析函数。
def resolve_data_types(data_type, account=None):

    """
    解析单商业任务需要执行的数据类型列表。

    :param data_type: 请求的数据类型文本。
    :type data_type: str
    :param account: 当前商业账号配置。
    :type account: dict | None
    :return: 有序数据类型列表。
    :rtype: list[str]
    """

    # 判断请求是否为全部数据类型。
    if str(data_type) == "全部":

        # 读取商业专属数据类型配置。
        account_data_types = (account or {}).get("数据类型列表")

        # 判断商业是否配置专属数据类型。
        if account_data_types:

            # 返回商业专属数据类型顺序。
            return list(account_data_types)

        # 返回全局默认数据类型顺序。
        return list(get_task_config()["数据类型列表"])

    # 返回逗号分隔的非空数据类型。
    return [item.strip() for item in str(data_type).split(",") if item.strip()]


# 定义输出路径标准化函数。
def normalize_output_paths(output_value):

    """
    将平台适配器返回的单个或多个文件路径统一转换为列表。

    :param output_value: 平台适配器返回的文件路径或文件路径集合。
    :type output_value: str | list[str] | tuple[str] | None
    :return: 有效文件路径列表。
    :rtype: list[str]
    """

    # 判断平台是否未返回输出路径。
    if not output_value:

        # 返回空文件路径列表。
        return []

    # 判断输出是否为集合类型。
    if isinstance(output_value, (list, tuple)):

        # 返回去除空项后的文件路径列表。
        return [str(file_path) for file_path in output_value if str(file_path).strip()]

    # 返回单文件路径列表。
    return [str(output_value)]


# 定义异常阶段识别函数。
def resolve_error_stage(error):

    """
    根据异常文本归类 Excel 任务日志中的失败阶段。

    :param error: 当前异常对象。
    :type error: Exception
    :return: 标准失败阶段名称。
    :rtype: str
    """

    # 转换异常为文本。
    error_text = str(error)

    # 判断是否为验证码异常。
    if "验证码" in error_text or "滑块" in error_text:

        # 返回验证码阶段。
        return "验证码"

    # 判断是否为登录异常。
    if "登录" in error_text or "账号" in error_text or "密码" in error_text:

        # 返回登录阶段。
        return "登录"

    # 判断是否为文件异常。
    if "Excel" in error_text or "覆盖文件" in error_text or "文件" in error_text:

        # 返回文件保存阶段。
        return "文件保存"

    # 判断是否为协议查询异常。
    if "查询" in error_text or "HTTP" in error_text or "请求" in error_text:

        # 返回查询阶段。
        return "查询"

    # 返回通用业务阶段。
    return "业务执行"


# 定义 Excel 记录数量读取函数。
def read_excel_record_count(excel_path):

    """
    读取 Excel 活动工作表中去除表头后的记录数量。

    :param excel_path: Excel 文件绝对路径。
    :type excel_path: str
    :return: 记录数量或空文本。
    :rtype: int | str
    """

    # 捕获 Excel 文件读取异常。
    try:

        # 导入只读工作簿加载函数。
        from openpyxl import load_workbook

        # 以只读模式打开 Excel 文件。
        workbook = load_workbook(excel_path, read_only=True, data_only=True)

        # 读取活动工作表。
        worksheet = workbook.active

        # 计算去除表头后的记录数量。
        record_count = max(0, int(worksheet.max_row or 1) - 1)

        # 关闭工作簿。
        workbook.close()

        # 返回读取到的记录数量。
        return record_count

    # 处理 Excel 文件读取异常。
    except Exception:

        # 返回未知记录数量。
        return ""


# 定义任务日志写入函数。
def record_task_log(task_logger, company_name, platform_name, data_type, stage, status, message, record_count="", file_path="", cookie_path="", error_detail="", started_at=None):

    """
    向可选的 Excel 任务日志记录器写入一条结构化日志。

    :param task_logger: Excel 任务日志记录器。
    :type task_logger: object | None
    :param company_name: 商业名称。
    :type company_name: str
    :param platform_name: 平台标识。
    :type platform_name: str
    :param data_type: 数据类型名称。
    :type data_type: str
    :param stage: 执行阶段。
    :type stage: str
    :param status: 执行状态。
    :type status: str
    :param message: 执行说明。
    :type message: str
    :param record_count: 数据记录数量。
    :type record_count: int | str
    :param file_path: 输出文件路径。
    :type file_path: str
    :param cookie_path: Cookie 文件路径。
    :type cookie_path: str
    :param error_detail: 异常详情。
    :type error_detail: object
    :param started_at: 阶段开始时间。
    :type started_at: datetime | None
    :return: None。
    :rtype: None
    """

    # 判断任务日志记录器是否不存在。
    if task_logger is None:

        # 结束日志写入处理。
        return

    # 写入结构化 Excel 任务日志。
    task_logger.record(company_name=company_name, platform_name=platform_name, data_type=data_type, stage=stage, status=status, message=message, record_count=record_count, file_path=file_path, cookie_path=cookie_path, error_detail=error_detail, started_at=started_at)


# 定义查询结果 JSON 保存函数。
def save_query_result(company_name, data_type, start_date, end_date, output_dir, rows):

    """
    将协议查询返回的原始记录保存为 UTF-8 JSON 文件。

    :param company_name: 商业名称。
    :type company_name: str
    :param data_type: 数据类型名称。
    :type data_type: str
    :param start_date: 查询开始日期。
    :type start_date: str
    :param end_date: 查询结束日期。
    :type end_date: str
    :param output_dir: JSON 输出目录。
    :type output_dir: str
    :param rows: 协议查询原始记录。
    :type rows: list[dict]
    :return: JSON 文件路径。
    :rtype: str
    """

    # 创建 JSON 输出目录。
    os.makedirs(output_dir, exist_ok=True)

    # 构建 JSON 输出文件名。
    file_name = f"{company_name}_{start_date}_至_{end_date}_{data_type}.json"

    # 拼接 JSON 输出绝对路径。
    output_path = os.path.join(output_dir, file_name)

    # 打开 UTF-8 JSON 输出文件。
    with open(output_path, "w", encoding="utf-8") as file_obj:

        # 写入未裁剪的查询原始记录。
        json.dump(rows, file_obj, ensure_ascii=False, indent=2)

    # 返回 JSON 输出文件路径。
    return output_path


# 定义复制用 Excel 生成函数。
def generate_copy_excel_if_needed(company_name, account, data_type, source_path, output_dir, copy_excel_enabled, printer, task_logger, cookie_path):

    """
    按配置为流向文件生成复制用标准 Excel。

    :param company_name: 商业名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param data_type: 数据类型名称。
    :type data_type: str
    :param source_path: 原始 Excel 文件路径。
    :type source_path: str
    :param output_dir: 输出目录。
    :type output_dir: str
    :param copy_excel_enabled: 是否启用复制用 Excel。
    :type copy_excel_enabled: bool
    :param printer: 公共日志输出器。
    :type printer: object
    :param task_logger: Excel 任务日志记录器。
    :type task_logger: object | None
    :param cookie_path: Cookie 文件路径。
    :type cookie_path: str
    :return: 复制用 Excel 文件路径。
    :rtype: str
    """

    # 判断复制用 Excel 是否启用。
    if not copy_excel_enabled:

        # 返回空文件路径。
        return ""

    # 记录复制用 Excel 处理开始时间。
    copy_started_at = datetime.now()

    # 捕获复制用 Excel 生成异常。
    try:

        # 生成复制用标准 Excel。
        copy_result = generate_copy_flow_excel(company_name, account.get("平台"), data_type, source_path, output_dir)

        # 判断当前数据类型是否不需要复制用文件。
        if copy_result.get("跳过原因"):

            # 返回空文件路径。
            return ""

        # 读取复制用 Excel 文件路径。
        copy_path = copy_result.get("输出文件", "")

        # 读取复制用文件缺失字段列表。
        missing_fields = copy_result.get("缺失字段", [])

        # 读取复制用文件日志状态。
        log_status = "警告" if missing_fields else "成功"

        # 构建复制用文件日志说明。
        log_message = "复制用 Excel 已生成" if not missing_fields else f"复制用 Excel 已生成，缺失字段：{','.join(missing_fields)}"

        # 记录复制用 Excel 日志。
        record_task_log(task_logger, company_name, account.get("平台"), data_type, "复制用Excel", log_status, log_message, record_count=copy_result.get("记录数量", 0), file_path=copy_path, cookie_path=cookie_path, started_at=copy_started_at)

        # 打印复制用 Excel 生成日志。
        printer.success(f"复制用Excel已生成: {printer.path(copy_path)}")

        # 返回复制用 Excel 文件路径。
        return copy_path

    # 处理复制用 Excel 生成异常。
    except Exception as error:

        # 记录复制用 Excel 生成失败日志。
        record_task_log(task_logger, company_name, account.get("平台"), data_type, "复制用Excel", "失败", "复制用 Excel 生成失败", file_path=source_path, cookie_path=cookie_path, error_detail=error, started_at=copy_started_at)

        # 打印复制用 Excel 生成失败日志。
        printer.warning(f"复制用Excel生成失败: {error}")

        # 返回空文件路径。
        return ""


# 定义单商业公共执行函数。
def execute_company(company_name, account, cookie_dir, action, data_type, start_date, end_date, output_dir, printer, task_logger=None, copy_excel_enabled=True):

    """
    使用平台适配器完成单商业的登录、查询、导出与日志记录。

    :param company_name: 商业名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param cookie_dir: Cookie 输出目录。
    :type cookie_dir: str
    :param action: 任务动作。
    :type action: str
    :param data_type: 数据类型文本。
    :type data_type: str
    :param start_date: 查询开始日期。
    :type start_date: str
    :param end_date: 查询结束日期。
    :type end_date: str
    :param output_dir: 查询和导出文件目录。
    :type output_dir: str
    :param printer: 公共日志输出器。
    :type printer: object
    :param task_logger: Excel 任务日志记录器。
    :type task_logger: object | None
    :param copy_excel_enabled: 是否生成复制用 Excel。
    :type copy_excel_enabled: bool
    :return: 单商业任务执行结果。
    :rtype: dict
    """

    # 读取空密码许可状态。
    allow_empty_password = bool(account.get("允许空密码", False))

    # 校验商业账号是否已配置。
    if not account.get("账号"):

        # 抛出账号缺失异常。
        raise RuntimeError(f"{company_name} 缺少账号，请在 account_config.py 中配置")

    # 校验商业密码是否已配置。
    if not allow_empty_password and not account.get("密码"):

        # 抛出密码缺失异常。
        raise RuntimeError(f"{company_name} 缺少密码，请在 account_config.py 中配置")

    # 打印单商业任务开始分隔线。
    printer.section("任务开始", printer.join(printer.kv("商业", company_name, "cyan"), printer.kv("平台", account.get("平台"), "yellow")))

    # 创建商业上下文日志输出器。
    company_printer = printer.for_company(company_name)

    # 记录商业任务开始时间。
    company_started_at = datetime.now()

    # 记录商业任务开始日志。
    record_task_log(task_logger, company_name, account.get("平台"), "", "商业任务", "开始", "商业任务开始执行", started_at=company_started_at)

    # 创建当前商业对应的平台协议适配器。
    platform_adapter = create_platform(company_name, account, printer=company_printer)

    # 记录平台登录开始时间。
    login_started_at = datetime.now()

    # 捕获平台登录异常。
    try:

        # 执行平台登录协议。
        login_result = platform_adapter.login()

    # 处理平台登录异常。
    except Exception as error:

        # 记录平台登录失败日志。
        record_task_log(task_logger, company_name, account.get("平台"), "", resolve_error_stage(error), "失败", "登录阶段执行失败", error_detail=error, started_at=login_started_at)

        # 继续抛出登录异常。
        raise

    # 保存当前平台会话 Cookie。
    cookie_path = platform_adapter.save_cookie(cookie_dir)

    # 记录平台登录成功日志。
    record_task_log(task_logger, company_name, account.get("平台"), "", "登录", "成功", "登录成功并保存 Cookie", cookie_path=cookie_path, started_at=login_started_at)

    # 初始化单商业任务结果字典。
    task_result = {"商业名称": company_name, "平台": account.get("平台"), "登录结果": login_result, "Cookie文件": cookie_path}

    # 遍历当前商业需要执行的数据类型。
    for current_data_type in resolve_data_types(data_type, account):

        # 记录当前数据类型开始时间。
        data_type_started_at = datetime.now()

        # 执行仅查询并保存 JSON 的动作。
        if action == "query":

            # 捕获平台查询协议异常。
            try:

                # 执行平台协议查询。
                query_rows = platform_adapter.query_flow(current_data_type, start_date=start_date, end_date=end_date)

            # 处理平台查询协议异常。
            except Exception as error:

                # 记录当前数据类型查询失败日志。
                record_task_log(task_logger, company_name, account.get("平台"), current_data_type, resolve_error_stage(error), "失败", "查询任务执行失败", cookie_path=cookie_path, error_detail=error, started_at=data_type_started_at)

                # 继续抛出当前查询异常。
                raise

            # 保存协议查询原始记录。
            query_path = save_query_result(company_name, current_data_type, start_date, end_date, output_dir, query_rows)

            # 保存查询记录数量。
            task_result[f"{current_data_type}查询条数"] = len(query_rows)

            # 保存查询 JSON 文件路径。
            task_result[f"{current_data_type}查询文件"] = query_path

            # 记录查询成功日志。
            record_task_log(task_logger, company_name, account.get("平台"), current_data_type, "查询", "成功", "查询完成并保存 JSON", record_count=len(query_rows), file_path=query_path, cookie_path=cookie_path, started_at=data_type_started_at)

            # 打印查询完成日志。
            company_printer.success(printer.join(printer.kv("查询任务完成", current_data_type, "green"), printer.kv("记录数", f"{len(query_rows)} 条", "cyan")))

        # 执行查询并生成 Excel 的动作。
        if action == "query_excel":

            # 捕获平台查询与 Excel 生成异常。
            try:

                # 执行平台查询并生成 Excel 文件。
                excel_result = platform_adapter.query_flow_to_excel(current_data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

            # 处理平台查询与 Excel 生成异常。
            except Exception as error:

                # 记录当前数据类型 Excel 失败日志。
                record_task_log(task_logger, company_name, account.get("平台"), current_data_type, resolve_error_stage(error), "失败", "查询或 Excel 生成失败", cookie_path=cookie_path, error_detail=error, started_at=data_type_started_at)

                # 继续抛出当前 Excel 异常。
                raise

            # 标准化平台返回的 Excel 文件路径。
            excel_paths = normalize_output_paths(excel_result)

            # 校验当前数据类型是否生成 Excel 文件。
            if not excel_paths:

                # 抛出 Excel 文件缺失异常。
                raise RuntimeError(f"{current_data_type} 未生成 Excel 文件")

            # 保存当前数据类型 Excel 文件路径。
            task_result[f"{current_data_type}Excel文件"] = ";".join(excel_paths)

            # 初始化复制用 Excel 文件路径集合。
            copy_excel_paths = []

            # 遍历当前数据类型的全部 Excel 文件。
            for excel_path in excel_paths:

                # 应用统一 Excel 输出样式。
                apply_standard_excel_style(excel_path)

                # 读取当前 Excel 记录数量。
                row_count = read_excel_record_count(excel_path)

                # 记录查询并生成 Excel 的成功日志。
                record_task_log(task_logger, company_name, account.get("平台"), current_data_type, "查询并生成Excel", "成功", "查询完成并生成 Excel", record_count=row_count, file_path=excel_path, cookie_path=cookie_path, started_at=data_type_started_at)

                # 生成当前 Excel 对应的复制用文件。
                copy_excel_path = generate_copy_excel_if_needed(company_name, account, current_data_type, excel_path, output_dir, copy_excel_enabled, company_printer, task_logger, cookie_path)

                # 判断复制用 Excel 文件是否已生成。
                if copy_excel_path:

                    # 追加复制用 Excel 文件路径。
                    copy_excel_paths.append(copy_excel_path)

                # 打印 Excel 任务完成日志。
                company_printer.success(printer.join(printer.kv("Excel任务完成", current_data_type, "green"), printer.kv("文件", printer.path(excel_path), "cyan")))

            # 保存复制用 Excel 文件路径集合。
            task_result[f"{current_data_type}复制用Excel文件"] = ";".join(copy_excel_paths)

        # 执行平台原始导出动作。
        if action == "export":

            # 捕获平台原始导出协议异常。
            try:

                # 执行平台原始导出协议。
                export_path = platform_adapter.export_flow(current_data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

            # 处理平台原始导出协议异常。
            except Exception as error:

                # 记录当前数据类型导出失败日志。
                record_task_log(task_logger, company_name, account.get("平台"), current_data_type, resolve_error_stage(error), "失败", "导出任务执行失败", cookie_path=cookie_path, error_detail=error, started_at=data_type_started_at)

                # 继续抛出当前导出异常。
                raise

            # 应用统一 Excel 输出样式。
            apply_standard_excel_style(export_path)

            # 保存当前数据类型导出文件路径。
            task_result[f"{current_data_type}导出文件"] = export_path

            # 记录导出成功日志。
            record_task_log(task_logger, company_name, account.get("平台"), current_data_type, "导出", "成功", "平台导出接口执行完成", file_path=export_path, cookie_path=cookie_path, started_at=data_type_started_at)

            # 生成当前导出文件对应的复制用文件。
            copy_excel_path = generate_copy_excel_if_needed(company_name, account, current_data_type, export_path, output_dir, copy_excel_enabled, company_printer, task_logger, cookie_path)

            # 保存复制用 Excel 文件路径。
            task_result[f"{current_data_type}复制用Excel文件"] = copy_excel_path

            # 打印导出任务完成日志。
            company_printer.success(printer.join(printer.kv("导出任务完成", current_data_type, "green"), printer.kv("文件", printer.path(export_path), "cyan")))

    # 打印商业任务成功日志。
    company_printer.success(printer.join(printer.kv("任务执行成功", company_name, "green"), printer.kv("Cookie", printer.path(cookie_path), "cyan")))

    # 记录商业任务成功日志。
    record_task_log(task_logger, company_name, account.get("平台"), "", "商业任务", "成功", "商业任务全部执行完成", cookie_path=cookie_path, started_at=company_started_at)

    # 返回单商业任务结果。
    return task_result
