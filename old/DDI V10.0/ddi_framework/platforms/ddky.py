# -*- coding: utf-8 -*-
"""
叮当快药 DTP 流向系统协议适配器。

本模块负责叮当快药 DTP 流向系统的验证码登录、令牌保存、分页查询和本地 Excel 生成。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入随机数模块。
import random

# 导入时间模块。
import time

# 导入日期时间解析对象。
from datetime import datetime

# 导入 URL 表单编码函数。
from urllib.parse import urlencode

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入验证码识别器。
from ddi_framework import CaptchaSolver

# 导入平台适配器基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义业务类型与 DTP 视图编码映射。
DATA_TYPE_APIS = {
    "购进流向": "V_JXXGYY_FLOW_PUR",
    "配送流向": "V_JXXGYY_FLOW_DIS",
    "库存": "V_JXXGYY_FLOW_STK",
}


# 定义购进流向输出表头。
PURCHASE_HEADERS = [
    ("日期", "日期"),
    ("销售方代码", "销售方代码"),
    ("销售方名称", "销售方名称"),
    ("采购方代码", "采购方代码"),
    ("采购方名称", "采购方名称"),
    ("产品代码", "产品代码"),
    ("产品名称", "产品名称"),
    ("产品规格", "产品规格"),
    ("批号", "批号"),
    ("数量", "数量"),
    ("单位", "单位"),
    ("单价", "单价"),
    ("金额", "金额"),
    ("生产厂家", "生产厂家"),
    ("有效期", "有效期"),
    ("采购类型", "采购类型"),
    ("订单编号", "订单编号"),
]


# 定义配送流向输出表头。
DELIVERY_HEADERS = [
    ("销售日期", "日期"),
    ("销售方代码", "销售方代码"),
    ("销售方名称", "销售方名称"),
    ("采购方代码", "采购方代码"),
    ("买方名称", "采购方名称"),
    ("产品代码", "产品代码"),
    ("产品名称", "产品名称"),
    ("产品规格", "产品规格"),
    ("批号", "批号"),
    ("交易数量", "数量"),
    ("单位", "单位"),
    ("销售单价", "单价"),
    ("金额", "金额"),
    ("销售单类型", "销售单类型"),
    ("单据编号", "单据编号"),
    ("有效期", "有效期"),
]


# 定义库存权限未开通输出表头。
STOCK_PASS_HEADERS = [
    ("状态", "状态"),
    ("说明", "说明"),
]


# 定义业务类型与输出表头映射。
DATA_TYPE_HEADERS = {
    "购进流向": PURCHASE_HEADERS,
    "配送流向": DELIVERY_HEADERS,
    "库存": STOCK_PASS_HEADERS,
}


class DdkyPlatform(BasePlatform):
    """
    叮当快药 DTP 流向系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 叮当快药 DTP 平台适配器实例。
    :rtype: DdkyPlatform
    """

    def __init__(self, company_name, account, printer=None):
        """
        初始化叮当快药 DTP 平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用平台基类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "https://flowsails.ddky.com:11443/user/login")

        # 保存接口根地址。
        self.api_base_url = account.get("接口根地址", "https://flowsails.ddky.com:11443").rstrip("/")

        # 保存站点源地址。
        self.origin_url = self.build_origin(self.api_base_url)

        # 保存验证码接口地址。
        self.captcha_api_url = account.get("验证码接口", f"{self.api_base_url}/api/kcone/auth/captcha")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.api_base_url}/api/kcone/auth/login")

        # 保存系统初始化接口地址。
        self.init_api_url = account.get("初始化接口", f"{self.api_base_url}/api/kcone/system/init")

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 100))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.api_base_url)

        # 初始化验证码识别器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录令牌。
        self.login_token = ""

        # 初始化登录响应。
        self.login_result = {}

        # 初始化登录状态。
        self.logged_in = False

    def login(self):
        """
        执行叮当快药 DTP 平台登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 执行带验证码重试的登录流程。
        return self.login_with_captcha_retry()

    def login_with_captcha_retry(self, max_retries=10):
        """
        使用图片验证码重试机制执行登录。

        :param max_retries: 最大验证码刷新次数。
        :type max_retries: int
        :return: 登录响应数据。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("叮当快药DTP登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.join(self.ph.kv("商业", self.company_name, "cyan"), self.ph.kv("访问登录页", self.login_page_url, "cyan")))

        # 请求登录页初始化站点上下文。
        self.http.get(self.login_page_url)

        # 初始化最后一次登录错误。
        last_error = None

        # 循环获取验证码并提交登录。
        for retry_index in range(max_retries):

            # 打印验证码尝试次数。
            self.ph.info(self.ph.join(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{max_retries} 次", "cyan")))

            # 捕获验证码识别与登录提交异常。
            try:

                # 生成验证码票据。
                ticket = str(random.random())

                # 获取验证码图片字节。
                image_bytes = self.get_captcha_image(ticket)

                # 识别验证码文本。
                captcha_code = self.solve_captcha_code(image_bytes)

                # 打印验证码识别日志。
                self.ph.info(self.ph.join(self.ph.kv("验证码识别完成", f"长度={len(captcha_code)}", "cyan")))

                # 构建登录请求体。
                payload = self.build_login_payload(ticket, captcha_code)

                # 提交登录表单。
                result = self.post_form(self.login_api_url, payload, self.login_page_url)

                # 解析登录响应。
                login_result = self.decode_json_response(result)

                # 判断登录响应是否成功。
                if self.is_success_response(login_result):

                    # 保存登录响应。
                    self.login_result = login_result

                    # 保存登录令牌。
                    self.login_token = str((login_result.get("data") or {}).get("token") or "")

                    # 校验登录令牌。
                    if not self.login_token:

                        # 抛出令牌缺失异常。
                        raise RuntimeError(f"{self.company_name} 登录响应缺少 token")

                    # 保存登录状态。
                    self.logged_in = True

                    # 调用系统初始化接口加载菜单上下文。
                    self.load_system_init()

                    # 打印登录成功日志。
                    self.ph.success(self.ph.join(self.ph.kv("正式登录成功", (login_result.get("data") or {}).get("name", ""), "green"), self.ph.kv("Token", self.login_token[:8], "cyan")))

                    # 返回登录响应。
                    return login_result

                # 保存当前登录失败响应。
                last_error = login_result

                # 打印登录失败重试日志。
                self.ph.warning(self.ph.join(self.ph.kv("登录校验未通过", login_result.get("message", ""), "yellow")))

            # 处理验证码识别或登录提交异常。
            except Exception as error:

                # 保存当前异常。
                last_error = error

                # 打印验证码处理失败日志。
                self.ph.warning(self.ph.join(self.ph.kv("验证码或登录校验未通过", str(error), "yellow")))

            # 判断是否已经达到最后一次重试。
            if retry_index >= max_retries - 1:

                # 结束验证码重试循环。
                break

            # 计算失败后的等待秒数。
            delay_seconds = 0.8 + random.random() * 0.8

            # 打印失败等待日志。
            self.ph.wait(self.ph.kv("验证码失败后等待", f"{delay_seconds:.1f} 秒", "cyan"))

            # 延迟后进入下一次验证码识别。
            time.sleep(delay_seconds)

        # 抛出登录失败异常。
        raise RuntimeError(f"{self.company_name} 图片验证码登录失败：{last_error}")

    def get_captcha_image(self, ticket):
        """
        获取登录验证码图片。

        :param ticket: 验证码票据。
        :type ticket: str
        :return: 验证码图片字节。
        :rtype: bytes
        """

        # 构建验证码请求地址。
        captcha_url = f"{self.captcha_api_url}?ticket={ticket}"

        # 构建验证码请求头。
        headers = self.build_image_headers(self.login_page_url)

        # 请求验证码图片。
        result = self.http.request("GET", captcha_url, headers=headers)

        # 校验验证码响应内容。
        if not result.body:

            # 抛出验证码图片缺失异常。
            raise RuntimeError("验证码接口未返回图片内容")

        # 返回验证码图片字节。
        return result.body

    def solve_captcha_code(self, image_bytes):
        """
        识别图片验证码文本。

        :param image_bytes: 验证码图片字节。
        :type image_bytes: bytes
        :return: 验证码文本。
        :rtype: str
        """

        # 按字母数字字符范围识别验证码。
        raw_text = self.captcha_solver.solve_raw_text_image_with_ranges(image_bytes, "0123456789abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ")

        # 清洗验证码文本。
        captcha_code = "".join([char for char in str(raw_text or "") if char.isalnum()])

        # 判断验证码识别结果是否过短。
        if len(captcha_code) < 4:

            # 使用通用文本验证码识别作为备用。
            captcha_code = self.captcha_solver.solve_text_image(image_bytes, min_length=4)

        # 清洗备用识别结果。
        captcha_code = "".join([char for char in str(captcha_code or "") if char.isalnum()])

        # 校验最终验证码文本。
        if not captcha_code:

            # 抛出验证码识别失败异常。
            raise RuntimeError("图片验证码识别为空")

        # 返回验证码文本。
        return captcha_code

    def build_login_payload(self, ticket, captcha_code):
        """
        构建登录表单请求体。

        :param ticket: 验证码票据。
        :type ticket: str
        :param captcha_code: 验证码文本。
        :type captcha_code: str
        :return: 登录表单请求体。
        :rtype: dict
        """

        # 返回登录表单字段。
        return {"username": self.clean_login_name(), "password": self.clean_login_password(), "verify": captcha_code, "ticket": ticket, "remember": "1"}

    def clean_login_name(self):
        """
        读取登录账号。

        :return: 登录账号。
        :rtype: str
        """

        # 返回清洗后的账号。
        return str(self.account.get("账号", "")).strip()

    def clean_login_password(self):
        """
        读取登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回清洗后的密码。
        return str(self.account.get("密码", "")).strip()

    def load_system_init(self):
        """
        调用系统初始化接口加载菜单上下文。

        :return: 初始化响应数据。
        :rtype: dict
        """

        # 构建空初始化请求体。
        payload = {}

        # 提交初始化请求。
        result = self.post_json(self.init_api_url, payload, self.login_page_url)

        # 解析初始化响应。
        response_data = self.decode_json_response(result)

        # 返回初始化响应。
        return response_data

    def save_cookie(self, output_dir):
        """
        保存登录令牌到 Cookie 文件目录。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 构建令牌保存数据。
        token_data = {"token": self.login_token, "authorization": f"Bearer {self.login_token}", "login": self.login_result}

        # 写入令牌文件。
        with open(cookie_path, "w", encoding="utf-8") as file_obj:
            json.dump(token_data, file_obj, ensure_ascii=False, indent=2)

        # 打印 Cookie 保存日志。
        self.ph.success(self.ph.join(self.ph.kv("Token已保存", self.ph.path(cookie_path), "green")))

        # 返回 Cookie 文件路径。
        return cookie_path

    def query_flow(self, data_type, start_date=None, end_date=None):
        """
        查询指定业务类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行。
        :rtype: list[dict]
        """

        # 查询业务数据并读取行集合。
        return self.query_ddky_data(data_type, start_date, end_date)["rows"]

    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):
        """
        导出指定业务类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 使用查询结果生成本地 Excel 文件。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):
        """
        查询指定业务类型数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询业务数据。
        query_result = self.query_ddky_data(data_type, start_date, end_date)

        # 使用当前工作目录作为默认输出目录。
        output_dir = output_dir or os.getcwd()

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    def query_ddky_data(self, data_type, start_date, end_date):
        """
        查询叮当快药 DTP 指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型。
        self.validate_data_type(data_type)

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 判断库存是否使用占位输出。
        if data_type == "库存":

            # 返回库存权限未开通占位数据。
            return {"rows": [{"状态": "pass", "说明": "库存权限未开通"}], "headers": STOCK_PASS_HEADERS}

        # 校验流向查询开始日期。
        if not start_date:

            # 抛出查询开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验流向查询结束日期。
        if not end_date:

            # 抛出查询结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 初始化查询结果行。
        all_rows = []

        # 初始化当前页码。
        page_index = 1

        # 初始化总记录数。
        total_count = None

        # 循环读取分页数据。
        while True:

            # 构建当前分页请求体。
            payload = self.build_query_payload(data_type, start_date, end_date, page_index)

            # 构建当前分页接口地址。
            query_url = self.build_query_url(data_type)

            # 构建当前分页来源地址。
            referer = self.build_query_referer(data_type)

            # 打印当前分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"pageNo={page_index}", "cyan"), self.ph.kv("pageSize", self.page_size, "cyan")))

            # 发送当前分页请求。
            result = self.post_json(query_url, payload, referer)

            # 解析当前分页响应。
            response_data = self.decode_json_response(result)

            # 校验当前分页响应状态。
            if not self.is_success_response(response_data):

                # 抛出查询失败异常。
                raise RuntimeError(f"{self.company_name} {data_type}查询失败：{response_data.get('message') or response_data}")

            # 提取响应表格对象。
            table_data = self.extract_table_data(response_data)

            # 提取当前页数据行。
            current_rows = self.extract_rows(table_data)

            # 读取当前页总记录数。
            total_count = self.resolve_total_count(table_data, total_count)

            # 合并当前页数据行。
            all_rows.extend(current_rows)

            # 打印当前页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(current_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count if total_count is not None else "未知", "cyan")))

            # 判断总数是否已经取完。
            if total_count is not None and len(all_rows) >= total_count:

                # 结束分页循环。
                break

            # 判断当前页是否已经不足一页。
            if len(current_rows) < self.page_size:

                # 结束分页循环。
                break

            # 递增页码。
            page_index += 1

        # 返回查询数据与表头。
        return {"rows": all_rows, "headers": DATA_TYPE_HEADERS[data_type]}

    def ensure_login(self):
        """
        确保当前适配器已经完成登录。

        :return: None。
        :rtype: None
        """

        # 判断登录状态是否已经存在。
        if self.logged_in:

            # 登录状态存在时直接返回。
            return

        # 执行登录流程。
        self.login()

    def print_query_section(self, data_type, start_date, end_date):
        """
        打印查询阶段标题。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断当前查询是否为库存。
        if data_type == "库存":

            # 打印库存占位查询标题。
            self.ph.section(f"{data_type}查询", self.ph.kv("处理方式", "pass", "cyan"))

            # 结束标题打印。
            return

        # 打印流向查询标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

    def build_query_payload(self, data_type, start_date, end_date, page_index):
        """
        构建分页查询请求体。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :param page_index: 当前页码。
        :type page_index: int
        :return: 分页查询请求体。
        :rtype: dict
        """

        # 读取业务视图编码。
        api_code = DATA_TYPE_APIS[data_type]

        # 构建业务 API 路径。
        api_path = f"/dtp/{api_code}/list"

        # 返回分页查询请求体。
        return {"api": api_path, "pageNo": page_index, "pageSize": self.page_size, "type": "table", "extra": {}, "sort": [], "where": [{"key": "日期", "op": "BETWEEN", "value": [start_date, end_date]}]}

    def build_query_url(self, data_type):
        """
        构建业务查询接口地址。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 业务查询接口地址。
        :rtype: str
        """

        # 读取业务视图编码。
        api_code = DATA_TYPE_APIS[data_type]

        # 返回业务查询接口地址。
        return f"{self.api_base_url}/api/dtp/{api_code}/list"

    def build_query_referer(self, data_type):
        """
        构建业务查询来源页面地址。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 来源页面地址。
        :rtype: str
        """

        # 读取业务视图编码。
        api_code = DATA_TYPE_APIS[data_type]

        # 返回来源页面地址。
        return f"{self.api_base_url}/dtp/{api_code}/list"

    def extract_table_data(self, response_data):
        """
        提取响应表格对象。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 表格对象。
        :rtype: dict
        """

        # 读取第一层业务数据。
        outer_data = response_data.get("data") or {}

        # 读取表格数据对象。
        table_data = outer_data.get("data") or {}

        # 返回表格数据对象。
        return table_data

    def extract_rows(self, table_data):
        """
        提取并过滤业务数据行。

        :param table_data: 表格对象。
        :type table_data: dict
        :return: 业务数据行。
        :rtype: list[dict]
        """

        # 读取原始记录列表。
        records = table_data.get("records") or []

        # 初始化清洗后的记录列表。
        clean_rows = []

        # 遍历原始记录列表。
        for row in records:

            # 判断当前记录是否为字典。
            if not isinstance(row, dict):

                # 跳过非字典记录。
                continue

            # 判断当前记录是否为空占位行。
            if self.is_blank_row(row):

                # 跳过空占位行。
                continue

            # 保存有效业务记录。
            clean_rows.append(row)

        # 返回清洗后的记录列表。
        return clean_rows

    def is_blank_row(self, row):
        """
        判断记录是否为空占位行。

        :param row: 业务记录。
        :type row: dict
        :return: 是否为空占位行。
        :rtype: bool
        """

        # 返回所有字段值均为空的判断结果。
        return all(str(value or "").strip() == "" for value in row.values())

    def resolve_total_count(self, table_data, current_total):
        """
        解析分页总记录数。

        :param table_data: 表格对象。
        :type table_data: dict
        :param current_total: 当前已知总数。
        :type current_total: int | None
        :return: 总记录数。
        :rtype: int | None
        """

        # 读取响应中的总数字段。
        total_value = table_data.get("total")

        # 判断总数字段是否存在。
        if total_value is None:

            # 返回当前已知总数。
            return current_total

        # 捕获数字转换异常。
        try:

            # 返回整数总数。
            return int(total_value)

        # 处理非数字总数字段。
        except (TypeError, ValueError):

            # 返回当前已知总数。
            return current_total

    def validate_data_type(self, data_type):
        """
        校验数据类型是否已接入。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断数据类型是否存在映射。
        if data_type not in DATA_TYPE_APIS:

            # 抛出数据类型未接入异常。
            raise ValueError(f"叮当快药DTP平台暂不支持数据类型：{data_type}")

    def post_form(self, url, payload, referer):
        """
        发送表单 POST 请求。

        :param url: 请求地址。
        :type url: str
        :param payload: 表单字段。
        :type payload: dict
        :param referer: 来源页面地址。
        :type referer: str
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 构建表单请求头。
        headers = self.build_form_headers(referer)

        # 编码表单请求体。
        data = urlencode(payload).encode("utf-8")

        # 发送表单请求。
        return self.http.request("POST", url, headers=headers, data=data)

    def post_json(self, url, payload, referer):
        """
        发送 JSON POST 请求。

        :param url: 请求地址。
        :type url: str
        :param payload: JSON 字段。
        :type payload: dict
        :param referer: 来源页面地址。
        :type referer: str
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 构建 JSON 请求头。
        headers = self.build_json_headers(referer)

        # 编码 JSON 请求体。
        data = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 发送 JSON 请求。
        return self.http.request("POST", url, headers=headers, data=data)

    def build_image_headers(self, referer):
        """
        构建验证码图片请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回验证码图片请求头。
        return {"Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Referer": referer, "User-Agent": self.user_agent()}

    def build_form_headers(self, referer):
        """
        构建登录表单请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回表单请求头。
        return {"Accept": "*/*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8", "Origin": self.origin_url, "Referer": referer, "User-Agent": self.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    def build_json_headers(self, referer):
        """
        构建授权 JSON 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回 JSON 请求头。
        return {"Accept": "*/*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Authorization": f"Bearer {self.login_token}", "Connection": "keep-alive", "Content-Type": "application/json;charset=UTF-8", "Origin": self.origin_url, "Referer": referer, "timestamp": str(round(time.time())), "User-Agent": self.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    def decode_json_response(self, result):
        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应体文本。
        text = result.body.decode("utf-8", errors="replace")

        # 返回 JSON 解析结果。
        return json.loads(text)

    def is_success_response(self, response_data):
        """
        判断接口响应是否成功。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 是否成功。
        :rtype: bool
        """

        # 返回状态码成功判断结果。
        return str(response_data.get("code")) == "0"

    def build_output_file_name(self, data_type, start_date, end_date):
        """
        构建输出文件名。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 输出文件名。
        :rtype: str
        """

        # Delegate standard naming to the base platform.
        return super().build_output_file_name(data_type, start_date, end_date)

    def build_month_label(self, start_date):
        """
        构建流向文件月份标签。

        :param start_date: 查询开始日期。
        :type start_date: str
        :return: 月份标签。
        :rtype: str
        """

        # 解析开始日期。
        date_value = datetime.strptime(start_date, "%Y-%m-%d")

        # 返回 yy.m月 标签。
        return f"{date_value.year % 100}.{date_value.month}月"

    def build_stock_date_label(self):
        """
        构建库存导出日期标签。

        :return: 日期标签。
        :rtype: str
        """

        # 读取当前日期。
        date_value = datetime.now()

        # 返回 yy.m.d 标签。
        return f"{date_value.year % 100}.{date_value.month}.{date_value.day}"

    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):
        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: Excel 输出路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 删除已存在的目标文件。
        self.remove_existing_file(output_path)

        # 创建 Excel 工作簿。
        workbook = Workbook(write_only=False)

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 写入当前数据行。
            worksheet.append([row.get(field_key, "") for _, field_key in headers])

        # 遍历表头列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取列字母。
            column_letter = get_column_letter(column_index)

            # 设置基础列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(28, len(str(header_item[0])) + 6))

        # 保存工作簿。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

    def remove_existing_file(self, output_path):
        """
        删除已存在的同名输出文件。

        :param output_path: 输出文件路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 判断输出文件是否已经存在。
        if os.path.exists(output_path):

            # 删除已存在的输出文件。
            os.remove(output_path)

    def build_origin(self, url):
        """
        构建站点源地址。

        :param url: 任意站点地址。
        :type url: str
        :return: 源地址。
        :rtype: str
        """

        # 解析 URL 地址。
        parsed_url = urlparse(url)

        # 返回协议与主机地址。
        return f"{parsed_url.scheme}://{parsed_url.netloc}"

    def user_agent(self):
        """
        返回浏览器 User-Agent。

        :return: User-Agent 字符串。
        :rtype: str
        """

        # 返回 Edge 浏览器 User-Agent。
        return "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36 Edg/148.0.0.0"
