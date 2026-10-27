# -*- coding: utf-8 -*-
"""
DDI 流向复制用 Excel 生成工具。

本模块负责从各商业原始流向 Excel 中提取统一字段，并生成便于复制到内部 Excel 的独立工作簿。
"""

# 导入操作系统路径模块。
import os

# 导入正则表达式模块。
import re

# 定义复制用流向支持的数据类型集合。
SUPPORTED_COPY_DATA_TYPES = {"配送流向", "纯销流向"}

# 定义复制用流向标准字段顺序。
STANDARD_COPY_FIELDS = ["销售日期", "产品规格", "产品名称", "产品批号", "交易数量", "买方名称", "销售单价"]

# 定义通用字段别名映射。
COMMON_FIELD_ALIASES = {
    "销售日期": ["销售日期", "销货日期", "业务日期", "单据日期", "开单日期", "开票日期", "出库日期", "发货日期", "日期", "bill_date", "saleDate", "sale_date", "xsrq", "rq"],
    "产品规格": ["产品规格", "商品规格", "规格", "品规", "药品规格", "goodsSpec", "prodSpec", "spec", "gg"],
    "产品名称": ["产品名称", "商品名称", "药品名称", "物料名称", "货品名称", "goodsName", "prodName", "productName", "spmc", "ypmc"],
    "产品批号": ["产品批号", "商品批号", "批号", "生产批号", "batchNo", "batch_no", "ph"],
    "交易数量": ["交易数量", "销售数量", "数量", "实发数量", "出库数量", "开票数量", "qty", "quantity", "saleQty", "sl"],
    "买方名称": ["买方名称", "客户名称", "购方名称", "购货单位", "收货单位", "终端名称", "门店名称", "医院名称", "buyerName", "customerName", "khmc"],
    "销售单价": ["销售单价", "单价", "含税单价", "开票单价", "价格", "price", "unitPrice", "hsdj", "dj"],
}

# 定义商业专属字段映射预留表。
COMPANY_FIELD_MAPPINGS = {}

# 定义平台专属字段映射预留表。
PLATFORM_FIELD_MAPPINGS = {
    "zztj": {
        "配送流向": {
            "买方名称": ["单位名称"],
        },
    },
    "kcshop": {
        "配送流向": {
            "销售日期": ["RQ", "销售日期", "日期"],
            "产品规格": ["SPGG", "规格"],
            "产品名称": ["SPTYM", "通用名"],
            "产品批号": ["PIHAO", "批号"],
            "交易数量": ["XSSHL", "销售数量"],
            "买方名称": ["KHMC", "XSEJMC", "客户名称", "下属二级机构名称"],
            "销售单价": ["XSDJ", "销售单价"],
        },
        "纯销流向": {
            "销售日期": ["RQ", "日期"],
            "产品规格": ["SPGG", "规格"],
            "产品名称": ["SPTYM", "通用名"],
            "产品批号": ["PIHAO", "批号"],
            "交易数量": ["XSSHL", "销售数量"],
            "买方名称": ["JGMC", "GKMC", "机构名称", "顾客名称"],
            "销售单价": ["XSDJ", "LSHJ", "销售单价", "零售单价"],
        },
    },
    "hx": {
        "配送流向": {
            "销售日期": ["billdate"],
            "产品规格": ["standard"],
            "产品名称": ["name"],
            "产品批号": ["batchcode"],
            "交易数量": ["qty"],
            "买方名称": ["clientname"],
            "销售单价": ["price", "taxprice"],
        },
    },
}

# 定义商业专属扩展字段预留表。
COMPANY_EXTRA_FIELDS = {}


# 定义复制用流向 Excel 生成函数。
def generate_copy_flow_excel(company_name, platform_name, data_type, source_path, output_dir):

    """
    生成复制用标准流向 Excel 文件。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param platform_name: 平台标识。
    :type platform_name: str
    :param data_type: 业务数据类型。
    :type data_type: str
    :param source_path: 原始流向 Excel 文件路径。
    :type source_path: str
    :param output_dir: 复制用 Excel 输出目录。
    :type output_dir: str
    :return: 复制用 Excel 生成结果。
    :rtype: dict
    """

    # 判断当前数据类型是否支持复制用流向生成。
    if data_type not in SUPPORTED_COPY_DATA_TYPES:

        # 返回数据类型跳过结果。
        return {"输出文件": "", "记录数量": 0, "缺失字段": [], "跳过原因": "当前数据类型不需要生成复制用 Excel"}

    # 判断原始文件是否可以读取。
    if not is_readable_excel_path(source_path):

        # 返回原始文件跳过结果。
        return {"输出文件": "", "记录数量": 0, "缺失字段": [], "跳过原因": "原始 Excel 文件不可读取"}

    # 导入 Excel 工作簿加载函数。
    from openpyxl import load_workbook

    # 加载原始流向工作簿。
    source_workbook = load_workbook(source_path, read_only=True, data_only=True)

    # 读取原始流向活动工作表。
    source_worksheet = source_workbook.active

    # 读取原始流向表头列表。
    source_headers = read_source_headers(source_worksheet)

    # 构建当前商业的字段映射。
    field_mapping = build_field_mapping(company_name, platform_name, data_type, source_headers)

    # 构建当前商业的输出字段列表。
    output_fields = build_output_fields(company_name, data_type)

    # 构建复制用数据行。
    output_rows = build_output_rows(source_worksheet, source_headers, field_mapping, output_fields)

    # 关闭原始流向工作簿。
    source_workbook.close()

    # 创建复制用输出目录。
    os.makedirs(output_dir, exist_ok=True)

    # 构建复制用输出路径。
    output_path = build_copy_output_path(source_path, output_dir)

    # 写入复制用 Excel 工作簿。
    write_copy_workbook(output_path, company_name, platform_name, data_type, source_path, output_fields, output_rows, field_mapping)

    # 读取缺失标准字段列表。
    missing_fields = [field_name for field_name in output_fields if not field_mapping.get(field_name)]

    # 返回复制用 Excel 生成结果。
    return {"输出文件": output_path, "记录数量": len(output_rows), "缺失字段": missing_fields, "跳过原因": ""}


# 定义 Excel 路径可读性判断函数。
def is_readable_excel_path(source_path):

    """
    判断原始文件路径是否为可读取的 xlsx 文件。

    :param source_path: 原始文件路径。
    :type source_path: str
    :return: 是否为可读取的 xlsx 文件。
    :rtype: bool
    """

    # 判断原始路径是否为空。
    if not source_path:

        # 返回不可读取状态。
        return False

    # 判断原始路径是否为 xlsx 文件。
    if not str(source_path).lower().endswith(".xlsx"):

        # 返回不可读取状态。
        return False

    # 返回文件存在状态。
    return os.path.exists(source_path)


# 定义原始表头读取函数。
def read_source_headers(worksheet):

    """
    读取原始 Excel 工作表表头。

    :param worksheet: 原始 Excel 工作表对象。
    :type worksheet: openpyxl.worksheet.worksheet.Worksheet
    :return: 原始表头列表。
    :rtype: list[str]
    """

    # 读取工作表第一行单元格。
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), [])

    # 返回清洗后的表头列表。
    return [stringify_cell_value(cell_value) for cell_value in header_row]


# 定义字段映射构建函数。
def build_field_mapping(company_name, platform_name, data_type, source_headers):

    """
    构建标准字段与原始字段的映射关系。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param platform_name: 平台标识。
    :type platform_name: str
    :param data_type: 业务数据类型。
    :type data_type: str
    :param source_headers: 原始表头列表。
    :type source_headers: list[str]
    :return: 标准字段到原始字段的映射关系。
    :rtype: dict
    """

    # 初始化字段映射结果。
    field_mapping = {}

    # 遍历当前输出字段列表。
    for standard_field in build_output_fields(company_name, data_type):

        # 读取当前字段候选别名列表。
        aliases = build_field_aliases(company_name, platform_name, data_type, standard_field)

        # 匹配原始表头中的字段名称。
        field_mapping[standard_field] = find_source_header(source_headers, aliases)

    # 返回字段映射结果。
    return field_mapping


# 定义字段别名列表构建函数。
def build_field_aliases(company_name, platform_name, data_type, standard_field):

    """
    构建字段匹配候选别名列表。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param platform_name: 平台标识。
    :type platform_name: str
    :param data_type: 业务数据类型。
    :type data_type: str
    :param standard_field: 标准字段名称。
    :type standard_field: str
    :return: 字段匹配候选别名列表。
    :rtype: list[str]
    """

    # 初始化候选别名列表。
    aliases = [standard_field]

    # 追加通用字段别名。
    aliases.extend(COMMON_FIELD_ALIASES.get(standard_field, []))

    # 追加平台专属字段别名。
    aliases.extend(resolve_mapping_aliases(PLATFORM_FIELD_MAPPINGS.get(platform_name, {}), data_type, standard_field))

    # 追加商业专属字段别名。
    aliases.extend(resolve_mapping_aliases(COMPANY_FIELD_MAPPINGS.get(company_name, {}), data_type, standard_field))

    # 返回去重后的候选别名列表。
    return deduplicate_texts(aliases)


# 定义映射表别名解析函数。
def resolve_mapping_aliases(mapping_config, data_type, standard_field):

    """
    解析字段映射表中的候选别名。

    :param mapping_config: 字段映射配置。
    :type mapping_config: dict
    :param data_type: 业务数据类型。
    :type data_type: str
    :param standard_field: 标准字段名称。
    :type standard_field: str
    :return: 字段候选别名列表。
    :rtype: list[str]
    """

    # 初始化候选别名列表。
    aliases = []

    # 追加通用映射值。
    aliases.extend(normalize_alias_value(mapping_config.get(standard_field)))

    # 读取数据类型专属映射字典。
    data_type_mapping = mapping_config.get(data_type, {})

    # 判断数据类型映射是否为字典。
    if isinstance(data_type_mapping, dict):

        # 追加数据类型专属映射值。
        aliases.extend(normalize_alias_value(data_type_mapping.get(standard_field)))

    # 返回候选别名列表。
    return aliases


# 定义别名值规范化函数。
def normalize_alias_value(alias_value):

    """
    规范化映射配置中的字段别名值。

    :param alias_value: 原始字段别名值。
    :type alias_value: object
    :return: 字段别名列表。
    :rtype: list[str]
    """

    # 判断别名值是否为空。
    if not alias_value:

        # 返回空列表。
        return []

    # 判断别名值是否为列表或元组。
    if isinstance(alias_value, (list, tuple)):

        # 返回文本化后的别名列表。
        return [stringify_cell_value(item) for item in alias_value]

    # 返回单个文本别名列表。
    return [stringify_cell_value(alias_value)]


# 定义输出字段列表构建函数。
def build_output_fields(company_name, data_type):

    """
    构建复制用 Excel 输出字段列表。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param data_type: 业务数据类型。
    :type data_type: str
    :return: 输出字段列表。
    :rtype: list[str]
    """

    # 初始化标准输出字段列表。
    output_fields = list(STANDARD_COPY_FIELDS)

    # 读取商业扩展字段配置。
    extra_config = COMPANY_EXTRA_FIELDS.get(company_name, {})

    # 读取当前数据类型的扩展字段列表。
    extra_fields = extra_config.get(data_type, extra_config.get("全部", [])) if isinstance(extra_config, dict) else []

    # 追加商业扩展字段列表。
    output_fields.extend([field_name for field_name in extra_fields if field_name not in output_fields])

    # 返回输出字段列表。
    return output_fields


# 定义原始表头匹配函数。
def find_source_header(source_headers, aliases):

    """
    根据候选别名匹配原始表头。

    :param source_headers: 原始表头列表。
    :type source_headers: list[str]
    :param aliases: 候选别名列表。
    :type aliases: list[str]
    :return: 匹配到的原始表头。
    :rtype: str
    """

    # 构建规范化原始表头映射。
    normalized_headers = {normalize_header(header): header for header in source_headers if header}

    # 遍历候选别名列表。
    for alias in aliases:

        # 规范化当前候选别名。
        normalized_alias = normalize_header(alias)

        # 判断候选别名是否直接命中表头。
        if normalized_alias in normalized_headers:

            # 返回命中的原始表头。
            return normalized_headers[normalized_alias]

    # 遍历候选别名列表进行包含匹配。
    for alias in aliases:

        # 读取包含匹配结果。
        source_header = find_contains_header(normalized_headers, alias)

        # 判断包含匹配结果是否存在。
        if source_header:

            # 返回命中的原始表头。
            return source_header

    # 返回空映射结果。
    return ""


# 定义包含式表头匹配函数。
def find_contains_header(normalized_headers, alias):

    """
    使用包含关系匹配原始表头。

    :param normalized_headers: 规范化表头映射。
    :type normalized_headers: dict
    :param alias: 候选别名。
    :type alias: str
    :return: 匹配到的原始表头。
    :rtype: str
    """

    # 规范化当前候选别名。
    normalized_alias = normalize_header(alias)

    # 跳过空候选别名。
    if not normalized_alias:

        # 返回空匹配结果。
        return ""

    # 遍历规范化原始表头映射。
    for normalized_header, source_header in normalized_headers.items():

        # 判断候选别名是否包含在原始表头中。
        if normalized_alias in normalized_header:

            # 返回命中的原始表头。
            return source_header

    # 返回空匹配结果。
    return ""


# 定义输出数据行构建函数。
def build_output_rows(worksheet, source_headers, field_mapping, output_fields):

    """
    根据字段映射构建复制用数据行。

    :param worksheet: 原始 Excel 工作表对象。
    :type worksheet: openpyxl.worksheet.worksheet.Worksheet
    :param source_headers: 原始表头列表。
    :type source_headers: list[str]
    :param field_mapping: 标准字段到原始字段的映射关系。
    :type field_mapping: dict
    :param output_fields: 输出字段列表。
    :type output_fields: list[str]
    :return: 复制用数据行列表。
    :rtype: list[list]
    """

    # 构建原始表头到列序号的映射。
    header_indexes = {header: index for index, header in enumerate(source_headers)}

    # 初始化输出数据行列表。
    output_rows = []

    # 遍历原始工作表数据行。
    for row_values in worksheet.iter_rows(min_row=2, values_only=True):

        # 构建当前输出行。
        output_row = [resolve_row_field_value(row_values, header_indexes, field_mapping.get(field_name)) for field_name in output_fields]

        # 判断当前输出行是否存在有效数据。
        if any(stringify_cell_value(value) for value in output_row):

            # 追加当前输出行。
            output_rows.append(output_row)

    # 返回输出数据行列表。
    return output_rows


# 定义行字段值读取函数。
def resolve_row_field_value(row_values, header_indexes, source_header):

    """
    读取原始行中的指定字段值。

    :param row_values: 原始行值列表。
    :type row_values: tuple
    :param header_indexes: 原始表头到列序号的映射。
    :type header_indexes: dict
    :param source_header: 原始字段名称。
    :type source_header: str
    :return: 原始字段值。
    :rtype: object
    """

    # 判断原始字段名称是否为空。
    if not source_header:

        # 返回空字符串。
        return ""

    # 读取原始字段列序号。
    column_index = header_indexes.get(source_header)

    # 判断列序号是否存在。
    if column_index is None:

        # 返回空字符串。
        return ""

    # 判断列序号是否超出当前行范围。
    if column_index >= len(row_values):

        # 返回空字符串。
        return ""

    # 返回原始字段值。
    return row_values[column_index]


# 定义复制用输出路径构建函数。
def build_copy_output_path(source_path, output_dir):

    """
    构建复制用 Excel 输出路径。

    :param source_path: 原始流向 Excel 文件路径。
    :type source_path: str
    :param output_dir: 复制用 Excel 输出目录。
    :type output_dir: str
    :return: 复制用 Excel 输出路径。
    :rtype: str
    """

    # 读取原始文件主名称。
    source_stem = os.path.splitext(os.path.basename(source_path))[0]

    # 构建复制用文件名。
    file_name = f"{source_stem}_复制用.xlsx"

    # 返回复制用输出完整路径。
    return os.path.join(output_dir, file_name)


# 定义复制用工作簿写入函数。
def write_copy_workbook(output_path, company_name, platform_name, data_type, source_path, output_fields, output_rows, field_mapping):

    """
    写入复制用 Excel 工作簿。

    :param output_path: 复制用 Excel 输出路径。
    :type output_path: str
    :param company_name: 商业公司名称。
    :type company_name: str
    :param platform_name: 平台标识。
    :type platform_name: str
    :param data_type: 业务数据类型。
    :type data_type: str
    :param source_path: 原始流向 Excel 文件路径。
    :type source_path: str
    :param output_fields: 输出字段列表。
    :type output_fields: list[str]
    :param output_rows: 输出数据行列表。
    :type output_rows: list[list]
    :param field_mapping: 标准字段到原始字段的映射关系。
    :type field_mapping: dict
    :return: None。
    :rtype: None
    """

    # 导入 Excel 工作簿对象。
    from openpyxl import Workbook

    # 创建复制用工作簿。
    workbook = Workbook()

    # 读取复制用数据工作表。
    copy_sheet = workbook.active

    # 设置复制用数据工作表名称。
    copy_sheet.title = "复制用流向"

    # 写入复制用数据表头。
    copy_sheet.append(output_fields)

    # 遍历复制用数据行。
    for output_row in output_rows:

        # 写入复制用数据行。
        copy_sheet.append(output_row)

    # 设置复制用数据工作表样式。
    style_simple_sheet(copy_sheet)

    # 创建字段映射说明工作表。
    mapping_sheet = workbook.create_sheet("字段映射")

    # 写入字段映射说明。
    write_mapping_sheet(mapping_sheet, company_name, platform_name, data_type, source_path, output_fields, field_mapping)

    # 保存复制用工作簿。
    workbook.save(output_path)

    # 关闭复制用工作簿。
    workbook.close()

    # 复制用工作簿已在写入阶段完成基础样式设置。
    return


# 定义简单工作表样式函数。
def style_simple_sheet(worksheet):

    """
    设置复制用数据工作表基础样式。

    :param worksheet: Excel 工作表对象。
    :type worksheet: openpyxl.worksheet.worksheet.Worksheet
    :return: None。
    :rtype: None
    """

    # 导入单元格对齐样式。
    from openpyxl.styles import Alignment

    # 导入单元格字体样式。
    from openpyxl.styles import Font

    # 导入单元格填充样式。
    from openpyxl.styles import PatternFill

    # 冻结表头行。
    worksheet.freeze_panes = "A2"

    # 启用自动筛选。
    worksheet.auto_filter.ref = worksheet.dimensions

    # 创建表头填充样式。
    header_fill = PatternFill("solid", fgColor="1F4E78")

    # 遍历表头列。
    for column_index in range(1, worksheet.max_column + 1):

        # 读取表头单元格。
        cell = worksheet.cell(row=1, column=column_index)

        # 设置表头单元格填充。
        cell.fill = header_fill

        # 设置表头单元格字体。
        cell.font = Font(name="等线", color="FFFFFF", bold=True, size=11)

        # 设置表头单元格对齐。
        cell.alignment = Alignment(horizontal="center", vertical="center")

        # 设置当前列宽。
        worksheet.column_dimensions[cell.column_letter].width = 18


# 定义字段映射工作表写入函数。
def write_mapping_sheet(worksheet, company_name, platform_name, data_type, source_path, output_fields, field_mapping):

    """
    写入字段映射说明工作表。

    :param worksheet: 字段映射工作表对象。
    :type worksheet: openpyxl.worksheet.worksheet.Worksheet
    :param company_name: 商业公司名称。
    :type company_name: str
    :param platform_name: 平台标识。
    :type platform_name: str
    :param data_type: 业务数据类型。
    :type data_type: str
    :param source_path: 原始流向 Excel 文件路径。
    :type source_path: str
    :param output_fields: 输出字段列表。
    :type output_fields: list[str]
    :param field_mapping: 标准字段到原始字段的映射关系。
    :type field_mapping: dict
    :return: None。
    :rtype: None
    """

    # 写入映射元信息标题。
    worksheet.append(["项目", "内容"])

    # 写入商业名称。
    worksheet.append(["商业名称", company_name])

    # 写入平台标识。
    worksheet.append(["平台", platform_name])

    # 写入数据类型。
    worksheet.append(["数据类型", data_type])

    # 写入原始文件路径。
    worksheet.append(["原始文件", os.path.abspath(source_path)])

    # 写入空行。
    worksheet.append([])

    # 写入字段映射表头。
    worksheet.append(["标准字段", "原始字段", "匹配状态"])

    # 遍历输出字段列表。
    for field_name in output_fields:

        # 读取原始字段名称。
        source_header = field_mapping.get(field_name, "")

        # 写入字段映射行。
        worksheet.append([field_name, source_header, "已匹配" if source_header else "未匹配"])

    # 设置字段映射工作表样式。
    style_simple_sheet(worksheet)

    # 设置第一列宽。
    worksheet.column_dimensions["A"].width = 18

    # 设置第二列宽。
    worksheet.column_dimensions["B"].width = 60

    # 设置第三列宽。
    worksheet.column_dimensions["C"].width = 12


# 定义表头规范化函数。
def normalize_header(value):

    """
    规范化表头文本。

    :param value: 原始表头文本。
    :type value: object
    :return: 规范化表头文本。
    :rtype: str
    """

    # 转换表头为文本。
    text = stringify_cell_value(value)

    # 转换表头为小写文本。
    text = text.lower()

    # 移除空白和常见分隔符。
    return re.sub(r"[\s_（）()【】\[\]\-:/：]+", "", text)


# 定义单元格文本化函数。
def stringify_cell_value(value):

    """
    转换单元格值为文本。

    :param value: 原始单元格值。
    :type value: object
    :return: 单元格文本。
    :rtype: str
    """

    # 判断值是否为空。
    if value is None:

        # 返回空字符串。
        return ""

    # 返回去除首尾空白后的文本。
    return str(value).strip()


# 定义文本列表去重函数。
def deduplicate_texts(values):

    """
    按顺序去重文本列表。

    :param values: 原始文本列表。
    :type values: list[str]
    :return: 去重后的文本列表。
    :rtype: list[str]
    """

    # 初始化已存在文本集合。
    exists = set()

    # 初始化去重结果列表。
    result = []

    # 遍历原始文本列表。
    for value in values:

        # 转换当前值为文本。
        text = stringify_cell_value(value)

        # 构建规范化文本键。
        key = normalize_header(text)

        # 判断当前文本是否为空或重复。
        if not key or key in exists:

            # 继续处理下一个文本。
            continue

        # 记录当前文本键。
        exists.add(key)

        # 追加当前文本。
        result.append(text)

    # 返回去重后的文本列表。
    return result
