# -*- coding: utf-8 -*-
"""
DDI Excel 输出样式工具。

本模块负责统一业务导出文件的字体、行高、冻结窗格、表头样式和内容行样式。
"""

# 导入操作系统路径模块。
import os


# 定义业务 Excel 默认样式配置。
DEFAULT_FONT_NAME = "等线"

# 定义业务 Excel 默认字号。
DEFAULT_FONT_SIZE = 11

# 定义业务 Excel 默认行高。
DEFAULT_ROW_HEIGHT = 30


# 定义业务 Excel 样式应用函数。
def apply_standard_excel_style(file_path):

    """
    对业务 Excel 文件应用统一样式。

    :param file_path: Excel 文件路径。
    :type file_path: str
    :return: Excel 文件路径。
    :rtype: str
    """

    # 判断文件路径是否为空。
    if not file_path:

        # 返回原始文件路径。
        return file_path

    # 判断文件是否为 xlsx 文件。
    if not str(file_path).lower().endswith(".xlsx"):

        # 返回原始文件路径。
        return file_path

    # 判断文件是否存在。
    if not os.path.exists(file_path):

        # 返回原始文件路径。
        return file_path

    # 导入 Excel 工作簿加载函数。
    from openpyxl import load_workbook

    # 加载 Excel 工作簿。
    workbook = load_workbook(file_path)

    # 遍历工作簿中的所有工作表。
    for worksheet in workbook.worksheets:

        # 应用当前工作表样式。
        apply_standard_worksheet_style(worksheet)

    # 保存样式调整后的工作簿。
    workbook.save(file_path)

    # 关闭工作簿。
    workbook.close()

    # 返回 Excel 文件路径。
    return file_path


# 定义工作表统一样式应用函数。
def apply_standard_worksheet_style(worksheet):

    """
    对单个工作表应用统一样式。

    :param worksheet: Excel 工作表对象。
    :type worksheet: openpyxl.worksheet.worksheet.Worksheet
    :return: None。
    :rtype: None
    """

    # 导入单元格字体样式。
    from openpyxl.styles import Font

    # 导入单元格填充样式。
    from openpyxl.styles import PatternFill

    # 导入单元格对齐样式。
    from openpyxl.styles import Alignment

    # 导入单元格边框样式。
    from openpyxl.styles import Border

    # 导入边框线条样式。
    from openpyxl.styles import Side

    # 构建标题行字体。
    header_font = Font(name=DEFAULT_FONT_NAME, size=DEFAULT_FONT_SIZE, color="FFFFFF", bold=True)

    # 构建内容行字体。
    body_font = Font(name=DEFAULT_FONT_NAME, size=DEFAULT_FONT_SIZE, color="1F2933")

    # 构建标题行填充色。
    header_fill = PatternFill("solid", fgColor="1F4E78")

    # 构建奇数内容行填充色。
    odd_fill = PatternFill("solid", fgColor="FFFFFF")

    # 构建偶数内容行填充色。
    even_fill = PatternFill("solid", fgColor="F7FBFF")

    # 构建细边框线条。
    thin_side = Side(style="thin", color="D9E2F3")

    # 构建表格边框。
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # 设置首行冻结窗格。
    worksheet.freeze_panes = "A2"

    # 设置自动筛选区域。
    worksheet.auto_filter.ref = worksheet.dimensions

    # 遍历工作表所有行。
    for row_index, row_cells in enumerate(worksheet.iter_rows(), start=1):

        # 设置当前行高度。
        worksheet.row_dimensions[row_index].height = DEFAULT_ROW_HEIGHT

        # 判断当前行是否为标题行。
        is_header_row = row_index == 1

        # 遍历当前行单元格。
        for cell in row_cells:

            # 设置单元格字体。
            cell.font = header_font if is_header_row else body_font

            # 设置单元格填充。
            cell.fill = resolve_row_fill(row_index, header_fill, odd_fill, even_fill)

            # 设置单元格对齐。
            cell.alignment = resolve_cell_alignment(is_header_row)

            # 设置单元格边框。
            cell.border = table_border


# 定义行填充样式解析函数。
def resolve_row_fill(row_index, header_fill, odd_fill, even_fill):

    """
    根据行号解析填充样式。

    :param row_index: Excel 行号。
    :type row_index: int
    :param header_fill: 标题行填充样式。
    :type header_fill: openpyxl.styles.PatternFill
    :param odd_fill: 奇数内容行填充样式。
    :type odd_fill: openpyxl.styles.PatternFill
    :param even_fill: 偶数内容行填充样式。
    :type even_fill: openpyxl.styles.PatternFill
    :return: 当前行填充样式。
    :rtype: openpyxl.styles.PatternFill
    """

    # 判断当前行是否为标题行。
    if row_index == 1:

        # 返回标题行填充样式。
        return header_fill

    # 判断当前内容行是否为偶数行。
    if row_index % 2 == 0:

        # 返回偶数内容行填充样式。
        return even_fill

    # 返回奇数内容行填充样式。
    return odd_fill


# 定义单元格对齐样式解析函数。
def resolve_cell_alignment(is_header_row):

    """
    根据行类型解析单元格对齐样式。

    :param is_header_row: 是否为标题行。
    :type is_header_row: bool
    :return: 单元格对齐样式。
    :rtype: openpyxl.styles.Alignment
    """

    # 导入单元格对齐样式。
    from openpyxl.styles import Alignment

    # 判断当前单元格是否属于标题行。
    if is_header_row:

        # 返回标题行居中对齐样式。
        return Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 返回内容行垂直居中样式。
    return Alignment(vertical="center", wrap_text=True)
