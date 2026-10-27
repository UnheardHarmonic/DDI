# -*- coding: utf-8 -*-
"""
燕喜堂 SRM 系统协议适配器。

本模块负责燕喜堂 SRM 系统的验证码登录、会话 Cookie 保存和后续报表协议查询入口。
"""

# 导入 JSON 编解码模块。
import json

# 导入 MD5 摘要模块。
import hashlib

# 导入操作系统路径模块。
import os

# 导入随机数模块。
import random

# 导入时间模块。
import time

# 导入日期时间对象。
from datetime import datetime, timedelta

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入平台适配器基类。
from ddi_framework import BasePlatform

# 导入验证码识别器。
from ddi_framework import CaptchaSolver

# 导入日志打印器。
from ddi_framework import PrintHandler

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient


# 定义燕喜堂 SRM 系统协议适配器。
class YanxitangPlatform(BasePlatform):

    """
    燕喜堂 SRM 系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 燕喜堂 SRM 系统协议适配器实例。
    :rtype: YanxitangPlatform
    """

    # 定义燕喜堂适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化燕喜堂 SRM 系统协议适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用平台基类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "https://srm.sdyanxitang.com/ELSServer_YXT/login/login.html")

        # 保存应用根地址。
        self.app_base_url = account.get("接口根地址", "https://srm.sdyanxitang.com/ELSServer_YXT").rstrip("/")

        # 保存站点根地址。
        self.base_url = self.build_origin(self.app_base_url)

        # 保存验证码图片地址。
        self.captcha_url = account.get("验证码接口", f"{self.app_base_url}/image/identifyingCode.img")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", f"{self.app_base_url}/rest/AccountService/login")

        # 保存报表查询接口地址。
        self.report_query_url = account.get("报表查询接口", f"{self.app_base_url}/rest/reportForm/ReportCommonService/queryReportByType")

        # 保存分页大小。
        self.page_size = int(account.get("分页大小", 50))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.app_base_url)

        # 初始化验证码识别器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录响应。
        self.login_result = {}

        # 初始化登录状态。
        self.logged_in = False

        # 保存验证码最大重试次数。
        self.captcha_max_retries = int(account.get("验证码重试次数", 8))

    # 定义平台登录方法。
    def login(self):

        """
        执行燕喜堂 SRM 系统协议登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("燕喜堂 SRM 登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 访问登录页以建立验证码会话。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 执行验证码登录流程。
        return self.login_with_captcha_retry()

    # 定义带验证码重试的登录方法。
    def login_with_captcha_retry(self):

        """
        使用图片验证码重试机制执行登录。

        :return: 登录响应数据。
        :rtype: dict
        """

        # 初始化最后一次失败原因。
        last_error = ""

        # 按最大重试次数循环提交登录。
        for retry_index in range(self.captcha_max_retries):

            # 打印验证码尝试日志。
            self.ph.info(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{self.captcha_max_retries} 次", "cyan"))

            # 捕获单次验证码登录异常。
            try:

                # 获取验证码图片字节。
                image_bytes = self.fetch_captcha_image()

                # 识别验证码文本。
                captcha_code = self.solve_captcha_code(image_bytes)

                # 打印验证码识别完成日志。
                self.ph.info(self.ph.kv("图片验证码识别完成", f"长度={len(captcha_code)}", "cyan"))

                # 构建登录请求体。
                payload = self.build_login_payload(captcha_code)

                # 打印正式登录请求日志。
                self.ph.info("提交正式登录请求: /rest/AccountService/login")

                # 提交登录请求。
                result = self.post_login(payload)

                # 解析登录响应。
                login_result = self.decode_json_response(result)

                # 判断登录是否成功。
                if self.is_login_success(login_result):

                    # 保存登录响应。
                    self.login_result = login_result

                    # 标记登录状态。
                    self.logged_in = True

                    # 打印登录成功日志。
                    self.ph.success(self.ph.join(self.ph.kv("正式登录成功", self.extract_login_name(login_result), "green"), self.ph.kv("Cookie数", len(list(self.http.cookie_jar)), "cyan")))

                    # 返回登录响应。
                    return login_result

                # 保存业务失败原因。
                last_error = self.extract_response_message(login_result)

                # 打印验证码校验失败日志。
                self.ph.warning(self.ph.kv("登录校验未通过，准备刷新验证码", last_error, "yellow"))

            # 处理验证码识别或登录提交异常。
            except Exception as error:

                # 保存异常失败原因。
                last_error = str(error)

                # 打印验证码处理失败日志。
                self.ph.warning(self.ph.kv("图片验证码处理失败，准备刷新验证码", last_error, "yellow"))

            # 判断是否已经到达最后一次重试。
            if retry_index >= self.captcha_max_retries - 1:

                # 结束重试等待。
                continue

            # 计算验证码刷新间隔。
            delay_seconds = random.uniform(0.8, 1.6)

            # 等待下一次验证码刷新。
            time.sleep(delay_seconds)

        # 抛出登录失败异常。
        raise RuntimeError(f"{self.company_name} 登录失败：{last_error}")

    # 定义验证码图片获取方法。
    def fetch_captcha_image(self):

        """
        获取当前会话的验证码图片。

        :return: 验证码图片字节。
        :rtype: bytes
        """

        # 拼接防缓存时间戳。
        request_url = f"{self.captcha_url}?t={int(time.time() * 1000)}"

        # 构建验证码图片请求头。
        headers = self.build_captcha_headers()

        # 请求验证码图片。
        result = self.http.request("GET", request_url, headers=headers, timeout=30)

        # 校验验证码图片响应体。
        if not result.body:

            # 抛出验证码图片为空异常。
            raise RuntimeError("验证码图片为空")

        # 返回验证码图片字节。
        return result.body

    # 定义验证码识别方法。
    def solve_captcha_code(self, image_bytes):

        """
        识别验证码图片文本。

        :param image_bytes: 验证码图片字节。
        :type image_bytes: bytes
        :return: 验证码文本。
        :rtype: str
        """

        # 使用通用 OCR 识别验证码。
        captcha_code = self.captcha_solver.solve_text_image(image_bytes, min_length=4)

        # 清洗验证码识别结果。
        captcha_code = "".join([char for char in str(captcha_code) if char.isalnum()])

        # 校验验证码长度。
        if len(captcha_code) != 4:

            # 抛出验证码长度异常。
            raise RuntimeError(f"验证码识别长度异常：{len(captcha_code)}")

        # 返回验证码文本。
        return captcha_code

    # 定义登录请求体构建方法。
    def build_login_payload(self, captcha_code):

        """
        构建登录请求体。

        :param captcha_code: 验证码文本。
        :type captcha_code: str
        :return: 登录请求体。
        :rtype: dict
        """

        # 返回燕喜堂登录 JSON 字段。
        return {"language": "zh_cn", "elsAccount": self.get_els_account(), "elsSubAccount": self.get_login_name(), "elsSubAccountPassword": self.build_password_digest(), "identifyCode": captcha_code, "loginType": self.get_login_type(), "clientType": "pc", "clientOs": "Win10", "browser": self.account.get("浏览器标识", "chrome/150.0.0.0")}

    # 定义登录提交方法。
    def post_login(self, payload):

        """
        提交登录请求。

        :param payload: 登录请求体。
        :type payload: dict
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 序列化登录请求体。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建登录请求头。
        headers = self.build_login_headers()

        # 发送登录请求。
        return self.http.request("POST", self.login_api_url, headers=headers, data=body, timeout=30)

    # 定义验证码请求头构建方法。
    def build_captcha_headers(self):

        """
        构建验证码图片请求头。

        :return: 请求头字典。
        :rtype: dict
        """

        # 返回验证码图片请求头。
        return {"Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Referer": self.login_page_url, "User-Agent": self.http.user_agent()}

    # 定义登录请求头构建方法。
    def build_login_headers(self):

        """
        构建登录接口请求头。

        :return: 请求头字典。
        :rtype: dict
        """

        # 返回浏览器 JSON 登录请求头。
        return {"Accept": "application/json, text/plain, */*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/json;charset=UTF-8", "Origin": self.base_url, "Referer": self.login_page_url, "User-Agent": self.http.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    # 定义 JSON 响应解析方法。
    def decode_json_response(self, result):

        """
        解析 JSON 响应体。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: JSON 响应数据。
        :rtype: dict
        """

        # 解码响应文本。
        response_text = result.body.decode("utf-8", errors="replace").strip()

        # 判断响应文本是否为空。
        if not response_text:

            # 返回空响应错误结构。
            return {"statusCode": "-1", "message": "接口返回空响应"}

        # 返回 JSON 响应数据。
        return json.loads(response_text)

    # 定义登录成功判断方法。
    def is_login_success(self, login_result):

        """
        判断登录响应是否成功。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 是否登录成功。
        :rtype: bool
        """

        # 返回状态码判断结果。
        return str(login_result.get("statusCode", "")) == "200"

    # 定义响应消息提取方法。
    def extract_response_message(self, response_data):

        """
        提取接口响应消息。

        :param response_data: 接口响应数据。
        :type response_data: dict
        :return: 响应消息。
        :rtype: str
        """

        # 返回消息字段或响应摘要。
        return str(response_data.get("message") or response_data.get("errMessage") or response_data)

    # 定义登录用户名称提取方法。
    def extract_login_name(self, login_result):

        """
        提取登录响应中的账号名称。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 账号名称。
        :rtype: str
        """

        # 返回登录响应中的机构或账号名称。
        return str(login_result.get("orgname") or login_result.get("companyShortName") or login_result.get("elsSubAccountName") or self.company_name)

    # 定义登录账号读取方法。
    def get_login_name(self):

        """
        读取个人登录账号。

        :return: 个人登录账号。
        :rtype: str
        """

        # 返回账号配置中的个人账号。
        return str(self.account.get("账号", "")).strip()

    # 定义登录密码读取方法。
    def get_login_password(self):

        """
        读取登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回账号配置中的登录密码。
        return str(self.account.get("密码", "")).strip()

    # 定义 ELS 企业号读取方法。
    def get_els_account(self):

        """
        读取 ELS 企业号。

        :return: ELS 企业号。
        :rtype: str
        """

        # 返回账号配置中的 ELS 企业号。
        return str(self.account.get("ELS账号", "")).strip()

    # 定义登录类型读取方法。
    def get_login_type(self):

        """
        读取平台登录类型。

        :return: 登录类型。
        :rtype: str
        """

        # 返回账号配置中的登录类型。
        return str(self.account.get("登录类型", "els")).strip()

    # 定义密码摘要构建方法。
    def build_password_digest(self):

        """
        构建登录密码 MD5 摘要。

        :return: MD5 密码摘要。
        :rtype: str
        """

        # 编码明文密码。
        password_bytes = self.get_login_password().encode("utf-8")

        # 返回 MD5 十六进制摘要。
        return hashlib.md5(password_bytes).hexdigest()

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie 和登录响应。

        :param output_dir: Cookie 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 保存登录响应文件。
        self.save_login_result_file(output_dir)

        # 打印 Cookie 保存日志。
        self.ph.success(self.ph.kv("登录 Cookie 已保存", self.ph.path(cookie_path), "green"))

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义登录响应保存方法。
    def save_login_result_file(self, output_dir):

        """
        保存登录响应到本地文件。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: 登录响应文件路径。
        :rtype: str
        """

        # 构建登录响应文件路径。
        result_path = os.path.join(output_dir, f"{self.company_name}_login_result.json")

        # 写入登录响应文件。
        with open(result_path, "w", encoding="utf-8") as file_obj:

            # 序列化登录响应数据。
            json.dump(self.login_result, file_obj, ensure_ascii=False, indent=2)

        # 打印登录响应保存日志。
        self.ph.success(self.ph.kv("登录响应已保存", self.ph.path(result_path), "green"))

        # 返回登录响应文件路径。
        return result_path

    # 定义平台查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 返回报表分页查询结果。
        return self.fetch_report_rows(data_type, start_date, end_date)

    # 定义平台导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        导出指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: 导出结果。
        :rtype: str
        """

        # 使用分页查询结果生成本地 Excel 文件。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义平台查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定类型业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定类型报表数据。
        rows = self.fetch_report_rows(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(data_type, rows, output_path)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(rows)} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义报表分页查询方法。
    def fetch_report_rows(self, data_type, start_date, end_date):

        """
        查询指定类型的报表明细行。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 报表明细行列表。
        :rtype: list[dict]
        """

        # 校验当前数据类型是否已配置。
        self.validate_data_type(data_type)

        # 确保当前会话已经登录。
        self.ensure_login()

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 初始化当前页码。
        current_page = 1

        # 初始化累计明细行。
        all_rows = []

        # 初始化上一页原始行键集合。
        previous_page_keys = set()

        # 初始化分页边界重叠行数。
        overlap_count = 0

        # 初始化接口总数。
        total_count = None

        # 循环请求全部分页。
        while True:

            # 构建当前页查询载荷。
            payload = self.build_report_payload(data_type, start_date, end_date, current_page)

            # 打印分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv("请求报表分页", f"{data_type} page={current_page}", "cyan"), self.ph.kv("pageSize", payload["pageSize"], "yellow")))

            # 请求当前页报表接口。
            result = self.post_report_page(payload)

            # 解析当前页报表响应。
            page_result = self.decode_report_page(result)

            # 读取当前页原始行。
            raw_rows = page_result["rows"]

            # 保存接口总数。
            total_count = page_result["total"]

            # 初始化当前页有效行。
            current_rows = []

            # 初始化当前页原始行键集合。
            current_page_keys = set()

            # 初始化分页边界重叠状态。
            boundary_overlap_open = current_page > 1

            # 遍历当前页原始行。
            for row_data in raw_rows:

                # 构建当前原始行去重键。
                row_key = self.build_report_row_key(row_data)

                # 记录当前页原始行键。
                current_page_keys.add(row_key)

                # 判断当前行是否为页首边界重叠行。
                if boundary_overlap_open and row_key in previous_page_keys:

                    # 累加分页边界重叠行数量。
                    overlap_count += 1

                    # 关闭当前页页首重叠判断。
                    boundary_overlap_open = False

                    # 跳过分页边界重叠行。
                    continue

                # 关闭当前页页首重叠判断。
                boundary_overlap_open = False

                # 转换并追加当前页有效行。
                current_rows.append(self.convert_report_row(data_type, row_data))

            # 追加当前页有效行。
            all_rows.extend(current_rows)

            # 保存当前页原始行键集合。
            previous_page_keys = current_page_keys

            # 打印当前分页统计。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(raw_rows)} 条", "green"), self.ph.kv("新增", f"{len(current_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("分页重叠", f"{overlap_count} 条", "yellow"), self.ph.kv("总数", total_count, "yellow")))

            # 判断是否已经完成分页。
            if not raw_rows or not current_rows or len(all_rows) >= total_count:

                # 结束分页循环。
                break

            # 递增页码。
            current_page += 1

        # 判断累计行数是否超过接口总数。
        if total_count is not None and len(all_rows) > total_count:

            # 按接口总数截取有效明细行。
            all_rows = all_rows[:total_count]

        # 返回累计明细行。
        return all_rows

    # 定义登录状态保障方法。
    def ensure_login(self):

        """
        确保当前适配器已经完成登录。

        :return: None。
        :rtype: None
        """

        # 判断当前登录状态是否有效。
        if self.logged_in:

            # 已登录时直接返回。
            return None

        # 执行登录流程。
        self.login()

        # 返回空结果。
        return None

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验报表数据类型是否已配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 读取报表类型映射。
        report_mapping = self.account.get("报表类型映射") or {}

        # 判断数据类型是否存在。
        if data_type not in report_mapping:

            # 抛出数据类型异常。
            raise NotImplementedError(f"燕喜堂 SRM 暂未接入数据类型：{data_type}")

        # 返回空结果。
        return None

    # 定义查询分隔线打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印报表查询阶段分隔线。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 打印标准阶段分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 返回空结果。
        return None

    # 定义报表查询载荷构建方法。
    def build_report_payload(self, data_type, start_date, end_date, current_page):

        """
        构建报表分页查询载荷。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param current_page: 当前页码。
        :type current_page: int
        :return: 报表查询载荷。
        :rtype: dict
        """

        # 读取报表类型编码。
        report_type = str((self.account.get("报表类型映射") or {})[data_type])

        # 构建通用报表载荷。
        payload = {"elsAccount": self.get_els_account(), "elsSubAccount": self.get_login_name(), "currentPage": int(current_page), "pageSize": str(self.page_size), "type": report_type, "wareCode": ""}

        # 判断当前查询是否为库存类型。
        if data_type in self.get_stock_data_types():

            # 写入库存类型参数。
            payload["storeType"] = str((self.account.get("库存类型映射") or {})[data_type])

            # 写入库存页面固定参数。
            payload["undefined"] = "0"

        # 处理流向类日期参数。
        else:

            # 写入查询开始日期。
            payload["startDate"] = start_date

            # 写入查询结束日期。
            payload["endDate"] = end_date

            # 写入流向页面日期跨度参数。
            self.apply_flow_days_param(payload, report_type)

        # 返回报表查询载荷。
        return payload

    # 定义库存数据类型读取方法。
    def get_stock_data_types(self):

        """
        读取库存数据类型集合。

        :return: 库存数据类型集合。
        :rtype: set[str]
        """

        # 返回库存类型映射中的数据类型。
        return set((self.account.get("库存类型映射") or {}).keys())

    # 定义流向日期跨度参数写入方法。
    def apply_flow_days_param(self, payload, report_type):

        """
        写入流向报表页面附加日期参数。

        :param payload: 报表查询载荷。
        :type payload: dict
        :param report_type: 报表类型编码。
        :type report_type: str
        :return: None。
        :rtype: None
        """

        # 判断是否为销售退货页面。
        if report_type == "2":

            # 写入销售退货页面使用的日期跨度字段。
            payload["days"] = "90"

        # 处理其他流向页面。
        else:

            # 写入入库退货和零售页面使用的日期跨度字段。
            payload["undefined"] = "90"

        # 返回空结果。
        return None

    # 定义报表单页提交方法。
    def post_report_page(self, payload):

        """
        提交单页报表查询请求。

        :param payload: 报表查询载荷。
        :type payload: dict
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 序列化查询载荷。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建报表请求头。
        headers = self.build_report_headers()

        # 发送报表查询请求。
        return self.http.request("POST", self.report_query_url, headers=headers, data=body, timeout=90)

    # 定义报表请求头构建方法。
    def build_report_headers(self):

        """
        构建报表查询请求头。

        :return: 请求头字典。
        :rtype: dict
        """

        # 返回报表 JSON 请求头。
        return {"Accept": "application/json, text/plain, */*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/json;charset=UTF-8", "Origin": self.base_url, "Referer": f"{self.app_base_url}/default2.jsp", "User-Agent": self.http.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    # 定义报表响应解析方法。
    def decode_report_page(self, result):

        """
        解析报表分页响应。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 标准分页结果。
        :rtype: dict
        """

        # 解析 JSON 响应。
        response_data = self.decode_json_response(result)

        # 校验业务状态码。
        if str(response_data.get("statusCode", "")) != "200":

            # 抛出报表接口异常。
            raise RuntimeError(f"报表接口返回异常：{self.extract_response_message(response_data)}")

        # 读取当前页数据行。
        rows = response_data.get("rows") or []

        # 读取接口总记录数。
        total = int(response_data.get("total") or len(rows))

        # 返回标准分页结果。
        return {"rows": rows, "total": total}

    # 定义报表原始行去重键构建方法。
    def build_report_row_key(self, raw_row):

        """
        构建报表分页原始行去重键。

        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :return: 原始行去重键。
        :rtype: str
        """

        # 捕获 JSON 序列化异常。
        try:

            # 返回稳定排序后的 JSON 文本。
            return json.dumps(raw_row, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":"))

        # 处理无法 JSON 序列化的原始行。
        except TypeError:

            # 返回原始行文本表示。
            return str(raw_row)

    # 定义报表行转换方法。
    def convert_report_row(self, data_type, raw_row):

        """
        转换报表原始行并保留全部原始字段。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :return: 标准字段前置的数据行。
        :rtype: dict
        """

        # 构建标准字段数据。
        row_data = self.build_standard_row(data_type, raw_row)

        # 遍历原始字段并追加到结果行。
        for field_name, field_value in raw_row.items():

            # 判断字段是否已经存在。
            if field_name in row_data:

                # 跳过已经写入的字段。
                continue

            # 写入标准化后的原始字段值。
            row_data[field_name] = self.normalize_cell_value(field_name, field_value)

        # 返回转换后的数据行。
        return row_data

    # 定义标准字段行构建方法。
    def build_standard_row(self, data_type, raw_row):

        """
        构建指定数据类型的标准字段行。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :return: 标准字段行。
        :rtype: dict
        """

        # 判断是否为配送流向。
        if data_type == "配送流向":

            # 返回配送流向标准字段。
            return self.build_distribution_row(raw_row)

        # 判断是否为购进流向。
        if data_type == "购进流向":

            # 返回购进流向标准字段。
            return self.build_purchase_row(raw_row)

        # 判断是否为纯销流向。
        if data_type == "纯销流向":

            # 返回纯销流向标准字段。
            return self.build_retail_row(raw_row)

        # 返回库存标准字段。
        return self.build_stock_row(raw_row)

    # 定义配送流向标准字段构建方法。
    def build_distribution_row(self, raw_row):

        """
        构建配送流向标准字段行。

        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :return: 标准字段行。
        :rtype: dict
        """

        # 返回配送流向字段映射。
        return self.build_row_from_mapping(raw_row, [("往来帐日期", "execDate"), ("客户名称", "orgname"), ("商品编码", "wareCode"), ("商品名称", "wareName"), ("商品规格", "wareSpec"), ("生产厂家", "factoryName"), ("单位", "wareUnit"), ("批号", "maxkNo"), ("数量", "wareqty"), ("区域", "commname")])

    # 定义购进流向标准字段构建方法。
    def build_purchase_row(self, raw_row):

        """
        构建购进流向标准字段行。

        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :return: 标准字段行。
        :rtype: dict
        """

        # 返回购进流向字段映射。
        return self.build_row_from_mapping(raw_row, [("商品编号", "wareCode"), ("供应商名称", "vencusName"), ("商品编码", "wareCode"), ("商品名称", "wareName"), ("商品规格", "wareSpec"), ("生产厂家", "factoryName"), ("单位", "wareUnit"), ("批号", "maxkNo"), ("数量", "wareqty")])

    # 定义纯销流向标准字段构建方法。
    def build_retail_row(self, raw_row):

        """
        构建纯销流向标准字段行。

        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :return: 标准字段行。
        :rtype: dict
        """

        # 返回纯销流向字段映射。
        return self.build_row_from_mapping(raw_row, [("会计日", "accdate"), ("门店名称", "orgname"), ("业务机构群", "commname"), ("商品编码", "wareCode"), ("商品名称", "wareName"), ("商品规格", "wareSpec"), ("生产厂家", "factoryName"), ("单位", "wareUnit"), ("批号", "maxkNo"), ("数量", "wareqty")])

    # 定义库存标准字段构建方法。
    def build_stock_row(self, raw_row):

        """
        构建库存标准字段行。

        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :return: 标准字段行。
        :rtype: dict
        """

        # 返回库存字段映射。
        return self.build_row_from_mapping(raw_row, [("业务机构", "orgname"), ("业务机构群", "commname"), ("商品编码", "wareCode"), ("商品名称", "wareName"), ("商品规格", "wareSpec"), ("生产厂家", "factoryName"), ("单位", "wareUnit"), ("批号", "maxkNo"), ("有效期", "invalidate"), ("数量", "wareqty")])

    # 定义字段映射行构建方法。
    def build_row_from_mapping(self, raw_row, mapping):

        """
        按字段映射构建标准字段行。

        :param raw_row: 接口原始数据行。
        :type raw_row: dict
        :param mapping: 标准字段和原始字段映射。
        :type mapping: list[tuple[str, str]]
        :return: 标准字段行。
        :rtype: dict
        """

        # 初始化标准字段行。
        row_data = {}

        # 遍历字段映射。
        for standard_name, raw_name in mapping:

            # 写入标准化字段值。
            row_data[standard_name] = self.normalize_cell_value(raw_name, raw_row.get(raw_name))

        # 返回标准字段行。
        return row_data

    # 定义单元格值标准化方法。
    def normalize_cell_value(self, field_name, field_value):

        """
        标准化 Excel 单元格值。

        :param field_name: 字段名称。
        :type field_name: str
        :param field_value: 原始字段值。
        :type field_value: object
        :return: 标准化字段值。
        :rtype: object
        """

        # 判断空值。
        if field_value is None:

            # 返回空字符串。
            return ""

        # 判断是否为毫秒时间戳日期字段。
        if field_name in {"execDate", "accdate"} and isinstance(field_value, (int, float)):

            # 返回平台业务日期。
            return self.format_business_timestamp(field_value)

        # 判断是否为日期文本字段。
        if field_name in {"invalidate", "makedate", "accdate", "startDate", "endDate"} and isinstance(field_value, str):

            # 返回日期文本前十位。
            return field_value[:10]

        # 判断是否为复杂结构字段。
        if isinstance(field_value, (dict, list)):

            # 返回 JSON 文本以保留原始结构。
            return json.dumps(field_value, ensure_ascii=False, default=str, separators=(",", ":"))

        # 判断是否为普通文本。
        if isinstance(field_value, str):

            # 返回去除首尾空白后的文本。
            return field_value.strip()

        # 返回原始字段值。
        return field_value

    # 定义业务时间戳格式化方法。
    def format_business_timestamp(self, field_value):

        """
        将接口毫秒时间戳转换为业务日期。

        :param field_value: 毫秒时间戳。
        :type field_value: int | float
        :return: 业务日期文本。
        :rtype: str
        """

        # 转换为 UTC 时间。
        date_value = datetime.utcfromtimestamp(float(field_value) / 1000.0)

        # 转换为平台页面使用的北京时间日期。
        date_value = date_value + timedelta(hours=8)

        # 返回日期文本。
        return date_value.strftime("%Y-%m-%d")

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, data_type, rows, output_path):

        """
        将查询结果写入 Excel 文件。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果行列表。
        :type rows: list[dict]
        :param output_path: 输出文件路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 构建输出表头。
        headers = self.build_excel_headers(data_type, rows)

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = data_type[:31]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历数据行。
        for row_data in rows:

            # 写入当前数据行。
            worksheet.append([row_data.get(header_name, "") for header_name in headers])

        # 遍历表头设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 8))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        workbook.save(output_path)

        # 关闭工作簿。
        workbook.close()

        # 返回空结果。
        return None

    # 定义 Excel 表头构建方法。
    def build_excel_headers(self, data_type, rows):

        """
        构建 Excel 输出表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果行列表。
        :type rows: list[dict]
        :return: Excel 表头列表。
        :rtype: list[str]
        """

        # 初始化标准表头。
        headers = list(self.get_standard_headers(data_type))

        # 遍历结果行补充原始字段。
        for row_data in rows:

            # 遍历当前行字段。
            for field_name in row_data.keys():

                # 判断字段是否已存在。
                if field_name not in headers:

                    # 追加动态字段。
                    headers.append(field_name)

        # 返回 Excel 表头。
        return headers

    # 定义标准表头读取方法。
    def get_standard_headers(self, data_type):

        """
        读取指定数据类型的标准表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 标准表头列表。
        :rtype: list[str]
        """

        # 定义标准表头映射。
        header_mapping = {"配送流向": ["往来帐日期", "客户名称", "商品编码", "商品名称", "商品规格", "生产厂家", "单位", "批号", "数量", "区域"], "购进流向": ["商品编号", "供应商名称", "商品编码", "商品名称", "商品规格", "生产厂家", "单位", "批号", "数量"], "纯销流向": ["会计日", "门店名称", "业务机构群", "商品编码", "商品名称", "商品规格", "生产厂家", "单位", "批号", "数量"], "大仓库存": ["业务机构", "业务机构群", "商品编码", "商品名称", "商品规格", "生产厂家", "单位", "批号", "有效期", "数量"], "门店库存": ["业务机构", "业务机构群", "商品编码", "商品名称", "商品规格", "生产厂家", "单位", "批号", "有效期", "数量"]}

        # 返回当前数据类型表头。
        return header_mapping.get(data_type, [])

    # 定义同名文件删除方法。
    def remove_existing_file(self, output_path):

        """
        删除已有同名输出文件。

        :param output_path: 输出文件路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 判断目标文件是否存在。
        if not os.path.exists(output_path):

            # 无需删除不存在的文件。
            return None

        # 删除已有目标文件。
        os.remove(output_path)

        # 返回空结果。
        return None

    # 定义站点根地址构建方法。
    def build_origin(self, url):

        """
        构建 URL 的站点根地址。

        :param url: 原始 URL。
        :type url: str
        :return: 站点根地址。
        :rtype: str
        """

        # 解析 URL 结构。
        parsed_url = urlparse(url)

        # 返回协议与主机组合。
        return f"{parsed_url.scheme}://{parsed_url.netloc}"
