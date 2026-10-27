# -*- coding: utf-8 -*-
"""
商业录入明细聚合工具。

本模块从当前商业本轮生成的原始流向 Excel 中读取明确配置的字段，
按商业独立 JSON 配置生成单个“录入明细”工作簿。
"""

# 导入日期时间处理模块。
from datetime import date

# 导入日期时间处理模块。
from datetime import datetime

# 导入 JSON 编解码模块。
import json

# 导入数学工具模块。
import math

# 导入操作系统路径模块。
import os

# 导入路径处理工具。
from pathlib import Path

# 导入正则表达式模块。
import re

# 导入统一 Excel 样式函数。
from ddi_framework.excel_style import apply_standard_excel_style


# 定义录入明细标准字段顺序。
ENTRY_DETAIL_FIELDS = [
    "规格",
    "流向日期",
    "买方-原始简称",
    "销售数量（支）",
    "销售数量（盒）",
    "供货单价",
    "批号",
    "来源表",
]

# 定义需要从配置中解析的业务字段集合。
ENTRY_DETAIL_DATA_FIELDS = ENTRY_DETAIL_FIELDS[:-1]

# 定义当前录入明细配置版本。
ENTRY_DETAIL_SCHEMA_VERSION = 1

# 定义首版默认处理的数据类型。
DEFAULT_ENTRY_DETAIL_DATA_TYPES = ["配送流向"]

# 定义允许的字段计算操作。
ALLOWED_FIELD_OPERATIONS = {
    "direct",
    "coalesce",
    "sum",
    "difference",
    "signed",
    "multiply",
    "abs",
    "trim",
    "to_number",
    "to_date",
}

# 定义允许的行级规则类型。
ALLOWED_ROW_RULES = {"sign_by_field", "replace"}

# 定义配置文件目录名称。
MAPPING_DIRECTORY_NAME = "copy_mappings"


# 定义配置目录解析函数。
def resolve_mapping_directory(mapping_dir=None):

    """
    解析商业录入明细 JSON 配置目录。

    :param mapping_dir: 可选的外部配置目录。
    :type mapping_dir: str | os.PathLike | None
    :return: 配置目录绝对路径。
    :rtype: str
    """

    # 优先使用调用方明确传入的配置目录。
    if mapping_dir:

        # 返回调用方配置目录的绝对路径。
        return os.path.abspath(os.fspath(mapping_dir))

    # 解析项目根目录。
    project_root = Path(__file__).resolve().parent.parent

    # 返回项目默认配置目录。
    return str(project_root / "config" / MAPPING_DIRECTORY_NAME)


# 定义商业配置文件路径解析函数。
def resolve_mapping_path(company_name, mapping_dir=None):

    """
    构建指定商业的独立 JSON 配置路径。

    :param company_name: 商业名称。
    :type company_name: str
    :param mapping_dir: 可选的外部配置目录。
    :type mapping_dir: str | os.PathLike | None
    :return: 商业 JSON 配置路径。
    :rtype: str
    """

    # 读取商业配置目录。
    directory = resolve_mapping_directory(mapping_dir)

    # 返回商业名称对应的 JSON 文件路径。
    return os.path.join(directory, f"{str(company_name).strip()}.json")


# 定义空白商业配置构建函数。
def build_empty_mapping(company_name):

    """
    构建未配置商业的安全空白映射。

    :param company_name: 商业名称。
    :type company_name: str
    :return: 未启用的空白映射配置。
    :rtype: dict
    """

    # 构建空白字段配置。
    fields = {field_name: {} for field_name in ENTRY_DETAIL_DATA_FIELDS}

    # 返回未启用的商业映射配置。
    return {
        "schema_version": ENTRY_DETAIL_SCHEMA_VERSION,
        "company_name": str(company_name),
        "enabled": False,
        "data_types": {
            "配送流向": {
                "fields": fields,
                "rules": [],
            },
        },
    }


# 定义商业映射加载函数。
def load_entry_detail_mapping(company_name, mapping_dir=None):

    """
    加载指定商业的独立录入明细 JSON 配置。

    :param company_name: 商业名称。
    :type company_name: str
    :param mapping_dir: 可选的外部配置目录。
    :type mapping_dir: str | os.PathLike | None
    :return: 映射配置及加载状态。
    :rtype: dict
    """

    # 解析商业配置文件路径。
    mapping_path = resolve_mapping_path(company_name, mapping_dir)

    # 判断商业配置文件是否存在。
    if not os.path.isfile(mapping_path):

        # 返回安全空白配置并标记文件缺失。
        mapping = build_empty_mapping(company_name)

        # 保存配置文件路径供日志和结果追踪。
        mapping["_mapping_path"] = mapping_path

        # 保存缺失配置状态。
        mapping["_load_error"] = "商业映射 JSON 不存在"

        # 返回缺失配置结果。
        return mapping

    # 捕获 JSON 文件读取异常。
    try:

        # 以 UTF-8 编码读取商业映射文件。
        with open(mapping_path, "r", encoding="utf-8") as mapping_file:

            # 解析 JSON 配置对象。
            mapping = json.load(mapping_file)

    # 将读取异常转换为可记录的配置结果。
    except Exception as error:

        # 返回结构化 JSON 读取错误。
        mapping = build_empty_mapping(company_name)

        # 保存配置文件路径供日志和结果追踪。
        mapping["_mapping_path"] = mapping_path

        # 保存读取异常文本。
        mapping["_load_error"] = f"商业映射 JSON 读取失败：{error}"

        # 返回读取失败结果。
        return mapping

    # 判断 JSON 根节点是否为对象。
    if not isinstance(mapping, dict):

        # 返回结构化根节点错误。
        mapping = build_empty_mapping(company_name)

        # 保存配置文件路径供日志和结果追踪。
        mapping["_mapping_path"] = mapping_path

        # 保存根节点类型错误。
        mapping["_load_error"] = "商业映射 JSON 根节点必须是对象"

        # 返回根节点错误结果。
        return mapping

    # 保存配置文件路径供日志和结果追踪。
    mapping["_mapping_path"] = mapping_path

    # 返回已加载的商业配置。
    return mapping


# 定义配置错误追加函数。
def append_unique_message(messages, message):

    """
    向消息列表追加不重复的文本。

    :param messages: 目标消息列表。
    :type messages: list[str]
    :param message: 待追加消息。
    :type message: str
    :return: None。
    :rtype: None
    """

    # 标准化待追加消息。
    normalized_message = str(message).strip()

    # 忽略空消息和重复消息。
    if not normalized_message or normalized_message in messages:

        # 结束消息追加处理。
        return

    # 追加唯一消息。
    messages.append(normalized_message)


# 定义字段配置来源读取函数。
def resolve_definition_sources(definition):

    """
    从字段定义中读取原始表头候选列表。

    :param definition: 标准字段配置。
    :type definition: object
    :return: 原始表头候选列表。
    :rtype: list[str]
    """

    # 将字符串定义转换为单字段来源列表。
    if isinstance(definition, str):

        # 返回单个原始字段。
        return [definition]

    # 将列表定义转换为候选字段列表。
    if isinstance(definition, (list, tuple)):

        # 返回非空候选字段文本。
        return [str(item).strip() for item in definition if str(item).strip()]

    # 判断当前定义是否为对象。
    if not isinstance(definition, dict):

        # 返回空来源列表。
        return []

    # 优先读取单字段来源配置。
    source = definition.get("source")

    # 读取多字段来源配置。
    sources = definition.get("sources")

    # 将单字段来源放入候选列表。
    resolved = [str(source).strip()] if source else []

    # 追加多字段来源列表。
    if isinstance(sources, (list, tuple)):

        # 追加非空来源字段。
        resolved.extend(str(item).strip() for item in sources if str(item).strip())

    # 返回去重后的来源字段。
    return deduplicate_texts(resolved)


# 定义映射配置校验函数。
def validate_entry_detail_mapping(mapping_config):

    """
    校验商业录入明细 JSON 配置结构。

    :param mapping_config: 商业映射配置。
    :type mapping_config: dict
    :return: 配置错误列表。
    :rtype: list[str]
    """

    # 初始化配置错误列表。
    errors = []

    # 判断配置根节点是否为对象。
    if not isinstance(mapping_config, dict):

        # 返回根节点结构错误。
        return ["映射配置根节点必须是对象"]

    # 校验配置版本字段。
    if mapping_config.get("schema_version") != ENTRY_DETAIL_SCHEMA_VERSION:

        # 记录不支持的配置版本。
        append_unique_message(errors, "映射配置 schema_version 必须为 1")

    # 校验商业名称字段。
    if not str(mapping_config.get("company_name", "")).strip():

        # 记录缺少商业名称错误。
        append_unique_message(errors, "映射配置缺少 company_name")

    # 读取数据类型配置。
    data_types = mapping_config.get("data_types")

    # 校验数据类型配置对象。
    if not isinstance(data_types, dict):

        # 记录数据类型结构错误。
        append_unique_message(errors, "映射配置 data_types 必须是对象")

        # 返回已发现的配置错误。
        return errors

    # 遍历每个数据类型配置。
    for data_type, data_type_config in data_types.items():

        # 校验数据类型配置对象。
        if not isinstance(data_type_config, dict):

            # 记录数据类型结构错误。
            append_unique_message(errors, f"数据类型 {data_type} 配置必须是对象")

            # 继续检查其他数据类型。
            continue

        # 读取标准字段映射对象。
        fields = data_type_config.get("fields")

        # 校验标准字段映射对象。
        if not isinstance(fields, dict):

            # 记录标准字段结构错误。
            append_unique_message(errors, f"数据类型 {data_type} 的 fields 必须是对象")

        # 读取行级规则配置。
        rules = data_type_config.get("rules", [])

        # 校验行级规则列表。
        if not isinstance(rules, list):

            # 记录行级规则结构错误。
            append_unique_message(errors, f"数据类型 {data_type} 的 rules 必须是列表")

        # 遍历字段定义。
        for field_name, definition in (fields or {}).items():

            # 校验标准字段名称是否属于受支持字段。
            if field_name not in ENTRY_DETAIL_DATA_FIELDS:

                # 记录未知标准字段错误。
                append_unique_message(errors, f"数据类型 {data_type} 存在未知标准字段：{field_name}")

            # 校验字段定义结构。
            if definition and not isinstance(definition, (str, list, tuple, dict)):

                # 记录字段定义类型错误。
                append_unique_message(errors, f"字段 {data_type}.{field_name} 定义类型不受支持")

            # 读取字段操作名称。
            operation = definition.get("operation", "direct") if isinstance(definition, dict) else "direct"

            # 校验字段操作白名单。
            if operation not in ALLOWED_FIELD_OPERATIONS:

                # 记录未知字段操作错误。
                append_unique_message(errors, f"字段 {data_type}.{field_name} 操作不受支持：{operation}")

            # 校验字段操作来源是否存在。
            if definition and operation in {"direct", "coalesce", "sum", "abs", "trim", "to_number", "to_date"} and not resolve_definition_sources(definition):

                # 记录字段来源缺失错误。
                append_unique_message(errors, f"字段 {data_type}.{field_name} 缺少 source 或 sources")

        # 遍历行级规则。
        for rule in rules if isinstance(rules, list) else []:

            # 校验规则对象结构。
            if not isinstance(rule, dict):

                # 记录行级规则类型错误。
                append_unique_message(errors, f"数据类型 {data_type} 存在非对象规则")

                # 继续检查下一条规则。
                continue

            # 读取规则类型。
            rule_type = rule.get("type")

            # 校验规则类型白名单。
            if rule_type not in ALLOWED_ROW_RULES:

                # 记录未知规则类型错误。
                append_unique_message(errors, f"数据类型 {data_type} 规则不受支持：{rule_type}")

    # 返回全部配置错误。
    return errors


# 定义表头规范化函数。
def normalize_header(value):

    """
    规范化原始表头用于精确匹配。

    :param value: 原始表头。
    :type value: object
    :return: 规范化表头文本。
    :rtype: str
    """

    # 转换表头为文本。
    text = "" if value is None else str(value).strip()

    # 去除字段周围空白并统一大小写。
    return re.sub(r"\s+", "", text).lower()


# 定义来源表头索引构建函数。
def build_header_indexes(headers):

    """
    构建原始表头到列序号的精确索引。

    :param headers: 原始表头列表。
    :type headers: list[str]
    :return: 规范化表头索引。
    :rtype: dict[str, int]
    """

    # 初始化表头索引。
    indexes = {}

    # 遍历原始表头及其列序号。
    for index, header in enumerate(headers):

        # 规范化当前表头。
        normalized_header = normalize_header(header)

        # 仅保留首个同名列，避免重复表头被静默覆盖。
        if normalized_header and normalized_header not in indexes:

            # 保存首个同名列序号。
            indexes[normalized_header] = index

    # 返回表头索引。
    return indexes


# 定义来源字段读取函数。
def read_source_value(row_values, headers, source_name):

    """
    按精确表头读取原始行字段值。

    :param row_values: 原始行值。
    :type row_values: tuple | list
    :param headers: 原始表头列表。
    :type headers: list[str]
    :param source_name: 原始字段名称。
    :type source_name: str
    :return: 字段值和是否命中状态。
    :rtype: tuple[object, bool]
    """

    # 构建当前表头索引。
    header_indexes = build_header_indexes(headers)

    # 读取来源字段列序号。
    column_index = header_indexes.get(normalize_header(source_name))

    # 判断来源字段是否存在。
    if column_index is None:

        # 返回未命中状态。
        return "", False

    # 判断当前行是否包含来源列。
    if column_index >= len(row_values):

        # 返回空值和命中状态。
        return "", True

    # 返回来源字段值和命中状态。
    return row_values[column_index], True


# 定义字符串列表去重函数。
def deduplicate_texts(values):

    """
    按规范化文本顺序去重。

    :param values: 待去重文本列表。
    :type values: list[object]
    :return: 去重后的文本列表。
    :rtype: list[str]
    """

    # 初始化已存在值集合。
    exists = set()

    # 初始化去重结果。
    result = []

    # 遍历待去重值。
    for value in values:

        # 转换当前值为文本。
        text = str(value).strip()

        # 构建规范化键。
        key = normalize_header(text)

        # 跳过空值和重复值。
        if not key or key in exists:

            # 继续处理下一个值。
            continue

        # 记录规范化键。
        exists.add(key)

        # 追加原始文本。
        result.append(text)

    # 返回去重结果。
    return result


# 定义单元格空值判断函数。
def is_blank(value):

    """
    判断单元格值是否为空。

    :param value: 单元格值。
    :type value: object
    :return: 是否为空。
    :rtype: bool
    """

    # 判断空值对象。
    if value is None:

        # 返回空状态。
        return True

    # 判断空文本。
    return not str(value).strip()


# 定义数值解析函数。
def coerce_number(value):

    """
    将原始值转换为数值。

    :param value: 原始数值文本或数字。
    :type value: object
    :return: 数值或空值。
    :rtype: int | float | None
    """

    # 保留空值为 None。
    if is_blank(value):

        # 返回空数值。
        return None

    # 直接保留有限数字。
    if isinstance(value, (int, float)) and not isinstance(value, bool):

        # 返回有限数字。
        return value if math.isfinite(float(value)) else None

    # 清理数值文本中的千分位和空白。
    text = str(value).strip().replace(",", "").replace("，", "")

    # 将括号数值转换为负数文本。
    if text.startswith("(") and text.endswith(")"):

        # 添加负号并去除括号。
        text = "-" + text[1:-1]

    # 尝试解析浮点数。
    try:

        # 解析浮点数。
        number = float(text)

    # 将非法数值转换为空值。
    except (TypeError, ValueError):

        # 返回空数值。
        return None

    # 将整数浮点数还原为整数。
    if number.is_integer():

        # 返回整数结果。
        return int(number)

    # 返回浮点结果。
    return number


# 定义日期标准化函数。
def normalize_date_value(value):

    """
    将常见日期值标准化为 ISO 日期文本。

    :param value: 原始日期值。
    :type value: object
    :return: ISO 日期文本或原始文本。
    :rtype: str
    """

    # 处理空日期值。
    if is_blank(value):

        # 返回空日期。
        return ""

    # 处理日期时间对象。
    if isinstance(value, datetime):

        # 返回日期部分文本。
        return value.strftime("%Y-%m-%d")

    # 处理日期对象。
    if isinstance(value, date):

        # 返回日期对象文本。
        return value.strftime("%Y-%m-%d")

    # 清理日期文本。
    text = str(value).strip()

    # 定义允许的日期格式。
    formats = ["%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S"]

    # 遍历允许的日期格式。
    for date_format in formats:

        # 尝试按当前格式解析日期。
        try:

            # 返回标准日期文本。
            return datetime.strptime(text, date_format).strftime("%Y-%m-%d")

        # 当前格式不匹配时继续尝试。
        except ValueError:

            # 继续尝试下一种格式。
            continue

    # 返回无法解析的原始文本，交由调用方记录警告。
    return text


# 定义数值列表求和函数。
def sum_numbers(values):

    """
    对可解析数值求和。

    :param values: 数值列表。
    :type values: list[object]
    :return: 求和结果或空值。
    :rtype: int | float | None
    """

    # 初始化求和结果。
    total = 0

    # 初始化是否命中有效数值。
    found = False

    # 遍历待求和字段。
    for value in values:

        # 解析当前字段数值。
        number = coerce_number(value)

        # 跳过空数值。
        if number is None:

            # 继续读取下一个字段。
            continue

        # 累加当前数值。
        total += number

        # 标记已经命中有效数值。
        found = True

    # 返回空字段集合的空值。
    return total if found else None


# 定义字段来源值读取函数。
def read_definition_values(definition, row_values, headers):

    """
    读取字段定义中的全部原始来源值。

    :param definition: 字段定义。
    :type definition: object
    :param row_values: 原始行值。
    :type row_values: tuple | list
    :param headers: 原始表头列表。
    :type headers: list[str]
    :return: 来源值、命中字段和未命中字段。
    :rtype: tuple[list[object], list[str], list[str]]
    """

    # 读取字段候选来源。
    sources = resolve_definition_sources(definition)

    # 初始化来源值集合。
    values = []

    # 初始化命中字段集合。
    matched = []

    # 初始化未命中字段集合。
    missing = []

    # 遍历字段候选来源。
    for source_name in sources:

        # 读取原始字段值。
        value, found = read_source_value(row_values, headers, source_name)

        # 记录命中字段。
        if found:

            # 保存命中字段名称。
            matched.append(source_name)

            # 保存原始字段值。
            values.append(value)

        # 记录未命中字段。
        else:

            # 保存未命中字段名称。
            missing.append(source_name)

    # 返回来源读取结果。
    return values, matched, missing


# 定义字段值变换函数。
def apply_transforms(value, transforms, warnings, context):

    """
    按白名单顺序变换字段值。

    :param value: 当前字段值。
    :type value: object
    :param transforms: 变换操作列表。
    :type transforms: list[str]
    :param warnings: 警告列表。
    :type warnings: list[str]
    :param context: 当前数据上下文。
    :type context: str
    :return: 变换后的字段值。
    :rtype: object
    """

    # 遍历字段变换操作。
    for transform in transforms or []:

        # 读取变换名称。
        operation = transform.get("type") if isinstance(transform, dict) else str(transform)

        # 校验变换操作白名单。
        if operation not in ALLOWED_FIELD_OPERATIONS:

            # 记录未知变换警告。
            append_unique_message(warnings, f"{context} 使用了未支持的变换：{operation}")

            # 跳过未知变换。
            continue

        # 执行文本清洗。
        if operation == "trim":

            # 清理当前值两侧空白。
            value = "" if value is None else str(value).strip()

        # 执行数值转换。
        elif operation == "to_number":

            # 转换当前值为数值。
            value = coerce_number(value)

        # 执行日期转换。
        elif operation == "to_date":

            # 标准化当前值为日期文本。
            value = normalize_date_value(value)

        # 执行绝对值转换。
        elif operation == "abs":

            # 解析当前值数值。
            number = coerce_number(value)

            # 将有效数值转换为绝对值。
            value = abs(number) if number is not None else value

    # 返回变换后的字段值。
    return value


# 定义字段值解析函数。
def resolve_field_value(field_name, definition, row_values, headers, warnings, context, unmatched_fields=None):

    """
    按字段定义解析一个标准字段。

    :param field_name: 标准字段名称。
    :type field_name: str
    :param definition: 字段定义。
    :type definition: object
    :param row_values: 原始行值。
    :type row_values: tuple | list
    :param headers: 原始表头列表。
    :type headers: list[str]
    :param warnings: 警告列表。
    :type warnings: list[str]
    :param context: 当前行上下文。
    :type context: str
    :param unmatched_fields: 未匹配标准字段列表。
    :type unmatched_fields: list[str] | None
    :return: 标准字段值。
    :rtype: object
    """

    # 跳过空字段定义。
    if not definition:

        # 返回空字段值。
        return ""

    # 将字符串和列表定义转换为直接映射定义。
    if isinstance(definition, (str, list, tuple)):

        # 构建直接字段定义。
        definition = {"operation": "direct", "sources": list(definition) if not isinstance(definition, str) else [definition]}

    # 读取字段操作名称。
    operation = definition.get("operation", "direct")

    # 读取来源值集合。
    values, matched, missing = read_definition_values(definition, row_values, headers)

    # 记录缺失来源字段警告。
    if missing and not matched:

        # 记录当前字段完全未匹配警告。
        append_unique_message(warnings, f"{context} 字段 {field_name} 未匹配：{','.join(missing)}")

        # 记录未匹配标准字段。
        if unmatched_fields is not None:

            # 追加当前未匹配字段。
            append_unique_message(unmatched_fields, field_name)

    # 解析直接字段值。
    if operation == "direct":

        # 读取首个来源值。
        value = values[0] if values else ""

    # 解析候选字段值。
    elif operation == "coalesce":

        # 读取首个非空来源值。
        value = next((item for item in values if not is_blank(item)), "")

    # 解析多字段求和。
    elif operation == "sum":

        # 汇总全部来源数值。
        value = sum_numbers(values)

    # 解析正负列差值。
    elif operation == "difference":

        # 读取正数来源字段。
        positive_values, positive_missing = read_named_values_with_status(definition.get("positive_sources", definition.get("positive_source", [])), row_values, headers)

        # 读取负数来源字段。
        negative_values, negative_missing = read_named_values_with_status(definition.get("negative_sources", definition.get("negative_source", [])), row_values, headers)

        # 记录正负列未匹配警告。
        for missing_source in positive_missing + negative_missing:

            # 记录当前差值来源字段。
            append_unique_message(warnings, f"{context} difference 来源未匹配：{missing_source}")

        # 计算正数减负数。
        value = safe_difference(positive_values, negative_values)

    # 解析按方向加符号的字段。
    elif operation == "signed":

        # 读取基础数量值。
        value = sum_numbers(values) if len(values) > 1 else (values[0] if values else "")

        # 读取方向字段名称。
        direction_source = definition.get("direction_source", definition.get("direction_field", ""))

        # 读取方向字段值。
        direction_value, direction_found = read_source_value(row_values, headers, direction_source)

        # 读取方向规则集合。
        negative_values = definition.get("negative_values", ["退货"])

        # 读取正向规则集合。
        positive_values = definition.get("positive_values", [])

        # 规范化方向文本。
        normalized_direction = str(direction_value).strip()

        # 解析数量绝对值。
        number = coerce_number(value)

        # 判断方向字段未命中。
        if not direction_found:

            # 记录方向字段缺失警告。
            append_unique_message(warnings, f"{context} signed 字段方向来源未匹配：{direction_source}")

        # 判断当前方向属于负向类型。
        if normalized_direction in [str(item).strip() for item in negative_values]:

            # 将数量设为负绝对值。
            value = -abs(number) if number is not None else value

        # 判断当前方向属于正向类型。
        elif normalized_direction in [str(item).strip() for item in positive_values]:

            # 将数量设为正绝对值。
            value = abs(number) if number is not None else value

        # 处理未识别方向。
        else:

            # 记录未识别方向警告。
            if normalized_direction:

                # 保存未识别方向信息。
                append_unique_message(warnings, f"{context} signed 字段未识别方向：{normalized_direction}")

            # 保留原始符号。
            value = value

    # 解析乘法换算字段。
    elif operation == "multiply":

        # 读取基础值。
        base_value = sum_numbers(values) if len(values) > 1 else (values[0] if values else "")

        # 读取固定换算系数。
        factor = definition.get("factor")

        # 优先从来源字段读取换算系数。
        if definition.get("factor_source"):

            # 读取动态换算系数。
            factor_value, factor_found = read_source_value(row_values, headers, definition.get("factor_source"))

            # 记录换算系数缺失。
            if not factor_found:

                # 记录换算系数警告。
                append_unique_message(warnings, f"{context} multiply 换算系数未匹配：{definition.get('factor_source')}")

            # 使用动态换算系数。
            factor = factor_value

        # 记录乘法来源未匹配警告。
        if not values and not matched:

            # 记录基础值来源缺失。
            append_unique_message(warnings, f"{context} multiply 基础来源未匹配")

        # 解析基础值和换算系数。
        base_number = coerce_number(base_value)

        # 解析换算系数数值。
        factor_number = coerce_number(factor)

        # 计算换算结果。
        value = base_number * factor_number if base_number is not None and factor_number is not None else ""

    # 处理独立变换操作。
    elif operation in {"abs", "trim", "to_number", "to_date"}:

        # 读取首个来源值。
        value = values[0] if values else ""

        # 执行当前字段操作。
        value = apply_transforms(value, [{"type": operation}], warnings, context)

    # 处理未知操作。
    else:

        # 记录未知操作警告。
        append_unique_message(warnings, f"{context} 字段操作未支持：{operation}")

        # 返回空字段值。
        value = ""

    # 读取字段后置变换列表。
    transforms = definition.get("transforms", definition.get("transform", [])) if isinstance(definition, dict) else []

    # 将单个变换转换为列表。
    if isinstance(transforms, str):

        # 构建单个变换列表。
        transforms = [transforms]

    # 执行字段后置变换。
    return apply_transforms(value, transforms, warnings, context)


# 定义指定来源读取函数。
def read_named_values(sources, row_values, headers):

    """
    读取指定来源字段的值列表。

    :param sources: 来源字段或来源字段列表。
    :type sources: str | list[str] | tuple[str]
    :param row_values: 原始行值。
    :type row_values: tuple | list
    :param headers: 原始表头列表。
    :type headers: list[str]
    :return: 来源字段值列表。
    :rtype: list[object]
    """

    # 标准化来源字段列表。
    source_list = [sources] if isinstance(sources, str) else list(sources or [])

    # 初始化来源字段值列表。
    values = []

    # 遍历来源字段。
    for source_name in source_list:

        # 读取来源字段值。
        value, found = read_source_value(row_values, headers, source_name)

        # 追加命中或空值。
        values.append(value if found else "")

    # 返回来源字段值。
    return values


# 定义指定来源读取函数并返回未命中字段。
def read_named_values_with_status(sources, row_values, headers):

    """
    读取来源字段值并返回未命中来源。

    :param sources: 来源字段或来源字段列表。
    :type sources: str | list[str] | tuple[str]
    :param row_values: 原始行值。
    :type row_values: tuple | list
    :param headers: 原始表头列表。
    :type headers: list[str]
    :return: 来源值列表和未命中字段列表。
    :rtype: tuple[list[object], list[str]]
    """

    # 标准化来源字段列表。
    source_list = [sources] if isinstance(sources, str) else list(sources or [])

    # 初始化来源字段值列表。
    values = []

    # 初始化未命中字段列表。
    missing = []

    # 遍历来源字段。
    for source_name in source_list:

        # 读取来源字段值。
        value, found = read_source_value(row_values, headers, source_name)

        # 保存命中值或未命中字段。
        if found:

            # 追加来源字段值。
            values.append(value)

        else:

            # 追加未命中字段名称。
            missing.append(str(source_name))

    # 返回来源值和未命中字段。
    return values, missing


# 定义安全差值计算函数。
def safe_difference(positive_values, negative_values):

    """
    计算正数来源与负数来源的差值。

    :param positive_values: 正数来源值。
    :type positive_values: list[object]
    :param negative_values: 负数来源值。
    :type negative_values: list[object]
    :return: 差值或空值。
    :rtype: int | float | None
    """

    # 汇总正数来源。
    positive_total = sum_numbers(positive_values)

    # 汇总负数来源。
    negative_total = sum_numbers(negative_values)

    # 判断两侧均无有效值。
    if positive_total is None and negative_total is None:

        # 返回空差值。
        return None

    # 将缺失的一侧按零处理。
    positive_total = positive_total or 0

    # 将缺失的一侧按零处理。
    negative_total = negative_total or 0

    # 返回正数减负数结果。
    return positive_total - negative_total


# 定义行级规则应用函数。
def apply_row_rules(row_data, rules, raw_row, headers, warnings, context):

    """
    对标准字段行执行配置声明的行级规则。

    :param row_data: 标准字段行。
    :type row_data: dict[str, object]
    :param rules: 行级规则列表。
    :type rules: list[dict]
    :param raw_row: 原始行值。
    :type raw_row: tuple | list
    :param headers: 原始表头列表。
    :type headers: list[str]
    :param warnings: 警告列表。
    :type warnings: list[str]
    :param context: 当前行上下文。
    :type context: str
    :return: None。
    :rtype: None
    """

    # 遍历行级规则。
    for rule in rules or []:

        # 跳过非对象规则。
        if not isinstance(rule, dict):

            # 继续处理下一条规则。
            continue

        # 读取规则类型。
        rule_type = rule.get("type")

        # 处理按方向添加符号规则。
        if rule_type == "sign_by_field":

            # 读取目标标准字段。
            target_field = rule.get("target_field")

            # 读取方向字段名称。
            direction_source = rule.get("direction_source", rule.get("direction_field", ""))

            # 读取方向字段值。
            direction_value, direction_found = read_source_value(raw_row, headers, direction_source)

            # 标准化方向文本。
            direction_text = str(direction_value).strip()

            # 读取目标当前数值。
            current_number = coerce_number(row_data.get(target_field))

            # 判断目标字段不存在。
            if target_field not in row_data:

                # 记录目标字段错误。
                append_unique_message(warnings, f"{context} sign_by_field 目标字段不存在：{target_field}")

                # 继续处理下一条规则。
                continue

            # 判断方向字段未命中。
            if not direction_found:

                # 记录方向字段警告。
                append_unique_message(warnings, f"{context} sign_by_field 方向来源未匹配：{direction_source}")

            # 读取负向值集合。
            negative_values = [str(item).strip() for item in rule.get("negative_values", ["退货"])]

            # 读取正向值集合。
            positive_values = [str(item).strip() for item in rule.get("positive_values", [])]

            # 负向业务设置负绝对值。
            if direction_text in negative_values and current_number is not None:

                # 写入负绝对值。
                row_data[target_field] = -abs(current_number)

            # 正向业务设置正绝对值。
            elif direction_text in positive_values and current_number is not None:

                # 写入正绝对值。
                row_data[target_field] = abs(current_number)

            # 记录未识别方向。
            elif direction_text:

                # 记录未知方向警告并保留当前符号。
                append_unique_message(warnings, f"{context} sign_by_field 未识别方向：{direction_text}")

        # 处理字段替换规则。
        elif rule_type == "replace":

            # 读取目标标准字段。
            target_field = rule.get("target_field")

            # 读取替换字典。
            replacements = rule.get("replacements", {})

            # 读取当前字段值。
            current_value = row_data.get(target_field, "")

            # 按精确文本执行替换。
            if isinstance(replacements, dict) and str(current_value) in replacements:

                # 写入替换后的字段值。
                row_data[target_field] = replacements[str(current_value)]


# 定义数据类型解析函数。
def resolve_mapping_data_types(mapping_config):

    """
    解析映射配置声明的数据类型。

    :param mapping_config: 商业映射配置。
    :type mapping_config: dict
    :return: 有序数据类型列表。
    :rtype: list[str]
    """

    # 读取映射数据类型对象。
    data_types = mapping_config.get("data_types", {})

    # 返回配置声明的数据类型顺序。
    return [str(data_type) for data_type in data_types if str(data_type).strip()]


# 定义源文件标准化函数。
def normalize_source_file(source_file, default_data_type="配送流向"):

    """
    将源文件描述标准化为路径和数据类型。

    :param source_file: 源文件路径或源文件对象。
    :type source_file: str | os.PathLike | dict
    :param default_data_type: 缺省数据类型。
    :type default_data_type: str
    :return: 标准源文件对象。
    :rtype: dict
    """

    # 处理对象形式的源文件。
    if isinstance(source_file, dict):

        # 返回对象形式源文件描述。
        return {"path": str(source_file.get("path", "")), "data_type": str(source_file.get("data_type", default_data_type))}

    # 返回路径形式源文件描述。
    return {"path": os.fspath(source_file), "data_type": default_data_type}


# 定义原始流向行采集函数。
def collect_entry_detail_rows(company_name, source_files, mapping_config):

    """
    从多个原始流向 Excel 中采集录入明细行。

    :param company_name: 商业名称。
    :type company_name: str
    :param source_files: 原始 Excel 路径或路径描述列表。
    :type source_files: list[str | os.PathLike | dict]
    :param mapping_config: 商业映射配置。
    :type mapping_config: dict
    :return: 采集结果字典。
    :rtype: dict
    """

    # 初始化采集结果。
    result = {"行": [], "记录数量": 0, "来源文件": [], "字段映射": [], "未匹配字段": [], "跳过行数": 0, "警告": [], "配置错误": []}

    # 读取配置加载错误。
    if mapping_config.get("_load_error"):

        # 记录配置加载错误。
        result["配置错误"].append(mapping_config["_load_error"])

        # 返回配置错误结果。
        return result

    # 校验映射配置所属商业。
    if str(mapping_config.get("company_name", "")).strip() != str(company_name).strip():

        # 记录商业名称不一致错误。
        result["配置错误"].append("映射配置 company_name 与任务商业不一致")

        # 返回商业名称错误结果。
        return result

    # 校验映射配置结构。
    validation_errors = validate_entry_detail_mapping(mapping_config)

    # 保存结构校验错误。
    result["配置错误"].extend(validation_errors)

    # 判断商业映射是否启用。
    if not mapping_config.get("enabled", False):

        # 记录映射未启用状态。
        result["警告"].append("商业录入明细映射未启用")

        # 返回未启用结果。
        return result

    # 判断配置结构是否有效。
    if validation_errors:

        # 返回结构错误结果。
        return result

    # 读取数据类型映射。
    data_type_configs = mapping_config.get("data_types", {})

    # 遍历源文件描述。
    for source_file in source_files or []:

        # 标准化源文件描述。
        source = normalize_source_file(source_file, next(iter(data_type_configs), "配送流向"))

        # 读取源文件路径。
        source_path = source.get("path", "")

        # 读取源文件数据类型。
        data_type = source.get("data_type", "")

        # 跳过未配置的数据类型。
        if data_type not in data_type_configs:

            # 记录未配置数据类型警告。
            append_unique_message(result["警告"], f"源文件 {os.path.basename(source_path)} 的数据类型未配置：{data_type}")

            # 继续处理下一个源文件。
            continue

        # 判断源文件是否存在。
        if not os.path.isfile(source_path):

            # 记录源文件缺失警告。
            append_unique_message(result["警告"], f"源文件不存在：{source_path}")

            # 继续处理下一个源文件。
            continue

        # 导入 Excel 工作簿读取函数。
        from openpyxl import load_workbook

        # 捕获源工作簿读取异常。
        try:

            # 以只读模式加载原始工作簿。
            workbook = load_workbook(source_path, read_only=True, data_only=True)

            # 读取活动工作表。
            worksheet = workbook.active

            # 读取第一行原始表头。
            header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())

            # 转换原始表头文本。
            headers = ["" if value is None else str(value).strip() for value in header_row]

            # 读取当前数据类型字段配置。
            data_type_config = data_type_configs[data_type]

            # 读取当前字段配置。
            field_definitions = data_type_config.get("fields", {})

            # 保存字段映射说明。
            result["字段映射"].append({"数据类型": data_type, "来源表": os.path.basename(source_path), "字段": field_definitions})

            # 遍历原始数据行。
            for row_number, row_values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):

                # 构建当前行上下文。
                context = f"{os.path.basename(source_path)} 第{row_number}行"

                # 跳过整行空白数据。
                if not any(not is_blank(value) for value in row_values):

                    # 增加跳过行计数。
                    result["跳过行数"] += 1

                    # 继续读取下一行。
                    continue

                # 初始化当前标准字段行。
                row_data = {}

                # 解析全部标准字段。
                for field_name in ENTRY_DETAIL_DATA_FIELDS:

                    # 读取字段定义。
                    definition = field_definitions.get(field_name, {})

                    # 解析标准字段值。
                    row_data[field_name] = resolve_field_value(field_name, definition, row_values, headers, result["警告"], context, result["未匹配字段"])

                # 应用当前数据类型的行级规则。
                apply_row_rules(row_data, data_type_config.get("rules", []), row_values, headers, result["警告"], context)

                # 写入原始文件名作为来源表。
                row_data["来源表"] = os.path.basename(source_path)

                # 判断当前行是否至少包含一个业务字段。
                if any(not is_blank(row_data.get(field_name)) for field_name in ENTRY_DETAIL_DATA_FIELDS):

                    # 追加当前录入明细行。
                    result["行"].append(row_data)

                # 处理全部字段均为空的行。
                else:

                    # 增加跳过行计数。
                    result["跳过行数"] += 1

            # 关闭源工作簿。
            workbook.close()

            # 记录已经使用的源文件。
            result["来源文件"].append(source_path)

        # 将 Excel 读取异常记录为警告。
        except Exception as error:

            # 记录源文件读取失败。
            append_unique_message(result["警告"], f"源文件读取失败 {source_path}：{error}")

    # 读取采集结果行数。
    result["记录数量"] = len(result["行"])

    # 返回采集结果。
    return result


# 定义录入明细工作簿写入函数。
def build_entry_detail_workbook(company_name, rows, mapping_config, output_dir, collection_result=None):

    """
    将录入明细行写入单商业工作簿。

    :param company_name: 商业名称。
    :type company_name: str
    :param rows: 标准录入明细行列表。
    :type rows: list[dict]
    :param mapping_config: 商业映射配置。
    :type mapping_config: dict
    :param output_dir: 输出目录。
    :type output_dir: str
    :param collection_result: 可选的采集结果。
    :type collection_result: dict | None
    :return: 工作簿路径。
    :rtype: str
    """

    # 创建输出目录。
    os.makedirs(output_dir, exist_ok=True)

    # 构建商业录入明细文件名。
    output_path = os.path.join(output_dir, f"{company_name}_录入明细.xlsx")

    # 导入 Excel 工作簿对象。
    from openpyxl import Workbook

    # 创建录入明细工作簿。
    workbook = Workbook()

    # 读取录入明细工作表。
    detail_sheet = workbook.active

    # 设置录入明细工作表名称。
    detail_sheet.title = "录入明细"

    # 写入录入明细标准表头。
    detail_sheet.append(ENTRY_DETAIL_FIELDS)

    # 遍历标准录入明细行。
    for row_data in rows:

        # 按固定字段顺序写入行数据。
        detail_sheet.append([row_data.get(field_name, "") for field_name in ENTRY_DETAIL_FIELDS])

    # 创建字段映射工作表。
    mapping_sheet = workbook.create_sheet("字段映射")

    # 写入字段映射工作表内容。
    write_mapping_sheet(mapping_sheet, company_name, mapping_config, collection_result or {})

    # 创建处理说明工作表。
    note_sheet = workbook.create_sheet("处理说明")

    # 写入处理说明工作表内容。
    write_processing_sheet(note_sheet, company_name, collection_result or {})

    # 保存工作簿。
    workbook.save(output_path)

    # 关闭工作簿。
    workbook.close()

    # 应用项目统一 Excel 样式。
    apply_standard_excel_style(output_path)

    # 返回生成的工作簿路径。
    return output_path


# 定义录入明细生成总接口。
def generate_entry_detail_excel(company_name, source_files, output_dir, mapping_dir=None):

    """
    生成指定商业的聚合录入明细 Excel。

    :param company_name: 商业名称。
    :type company_name: str
    :param source_files: 原始 Excel 路径或路径描述列表。
    :type source_files: list[str | os.PathLike | dict]
    :param output_dir: 输出目录。
    :type output_dir: str
    :param mapping_dir: 可选的商业 JSON 配置目录。
    :type mapping_dir: str | os.PathLike | None
    :return: 录入明细生成结果。
    :rtype: dict
    """

    # 加载商业独立映射配置。
    mapping_config = load_entry_detail_mapping(company_name, mapping_dir)

    # 采集原始流向行。
    collection_result = collect_entry_detail_rows(company_name, source_files, mapping_config)

    # 初始化统一生成结果。
    result = {
        "输出文件": "",
        "记录数量": collection_result.get("记录数量", 0),
        "来源文件": collection_result.get("来源文件", []),
        "未匹配字段": collection_result.get("未匹配字段", []),
        "跳过行数": collection_result.get("跳过行数", 0),
        "警告": collection_result.get("警告", []),
        "配置错误": collection_result.get("配置错误", []),
        "字段映射": collection_result.get("字段映射", []),
    }

    # 判断当前商业映射不可用。
    if result["配置错误"] or not mapping_config.get("enabled", False):

        # 返回不生成文件的安全结果。
        return result

    # 生成录入明细工作簿。
    result["输出文件"] = build_entry_detail_workbook(company_name, collection_result["行"], mapping_config, output_dir, collection_result)

    # 返回录入明细生成结果。
    return result


# 定义字段映射工作表写入函数。
def write_mapping_sheet(worksheet, company_name, mapping_config, collection_result):

    """
    写入字段映射审计工作表。

    :param worksheet: Excel 工作表。
    :type worksheet: openpyxl.worksheet.worksheet.Worksheet
    :param company_name: 商业名称。
    :type company_name: str
    :param mapping_config: 商业映射配置。
    :type mapping_config: dict
    :param collection_result: 字段采集结果。
    :type collection_result: dict
    :return: None。
    :rtype: None
    """

    # 写入审计工作表标题。
    worksheet.append(["商业名称", company_name])

    # 写入映射配置文件路径。
    worksheet.append(["映射文件", mapping_config.get("_mapping_path", "")])

    # 写入空白分隔行。
    worksheet.append([])

    # 写入字段映射表头。
    worksheet.append(["数据类型", "标准字段", "配置定义", "匹配状态"])

    # 遍历数据类型字段配置。
    for data_type, data_type_config in mapping_config.get("data_types", {}).items():

        # 遍历当前数据类型标准字段。
        for field_name in ENTRY_DETAIL_DATA_FIELDS:

            # 读取当前字段配置。
            definition = data_type_config.get("fields", {}).get(field_name, {})

            # 读取字段来源名称。
            sources = resolve_definition_sources(definition)

            # 判断当前字段是否已经配置来源。
            status = "已配置" if sources or (isinstance(definition, dict) and definition.get("operation") in {"difference", "multiply"}) else "未配置"

            # 写入字段映射行。
            worksheet.append([data_type, field_name, json.dumps(definition, ensure_ascii=False), status])

    # 设置字段映射列宽。
    worksheet.column_dimensions["A"].width = 18

    # 设置标准字段列宽。
    worksheet.column_dimensions["B"].width = 24

    # 设置配置定义列宽。
    worksheet.column_dimensions["C"].width = 80

    # 设置匹配状态列宽。
    worksheet.column_dimensions["D"].width = 14


# 定义处理说明工作表写入函数。
def write_processing_sheet(worksheet, company_name, collection_result):

    """
    写入录入明细处理说明工作表。

    :param worksheet: Excel 工作表。
    :type worksheet: openpyxl.worksheet.worksheet.Worksheet
    :param company_name: 商业名称。
    :type company_name: str
    :param collection_result: 采集结果。
    :type collection_result: dict
    :return: None。
    :rtype: None
    """

    # 写入商业名称。
    worksheet.append(["商业名称", company_name])

    # 写入生成时间。
    worksheet.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])

    # 写入记录数量。
    worksheet.append(["记录数量", collection_result.get("记录数量", 0)])

    # 写入跳过行数。
    worksheet.append(["跳过行数", collection_result.get("跳过行数", 0)])

    # 写入来源文件标题。
    worksheet.append([])

    # 写入来源文件表头。
    worksheet.append(["来源文件"])

    # 写入全部来源文件。
    for source_path in collection_result.get("来源文件", []):

        # 写入来源文件名而不是绝对路径。
        worksheet.append([os.path.basename(source_path)])

    # 写入警告标题。
    worksheet.append([])

    # 写入警告表头。
    worksheet.append(["警告"])

    # 写入全部警告。
    for warning in collection_result.get("警告", []):

        # 写入警告文本。
        worksheet.append([warning])

    # 写入配置错误标题。
    worksheet.append([])

    # 写入配置错误表头。
    worksheet.append(["配置错误"])

    # 写入全部配置错误。
    for error in collection_result.get("配置错误", []):

        # 写入配置错误文本。
        worksheet.append([error])

    # 设置说明列宽。
    worksheet.column_dimensions["A"].width = 80
