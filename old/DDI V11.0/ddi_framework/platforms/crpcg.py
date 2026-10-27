# -*- coding: utf-8 -*-
"""
华润医药商业上游数据服务平台协议适配器。

本模块通过 RSA 密码加密、Cookie 会话和数据卡片接口查询华润湖南新特药有限公司的购进流向、配送流向和库存，并将标准字段与接口返回字段写入本地 Excel。

:return: 模块提供华润医药商业上游数据服务平台协议适配器。
:rtype: None
"""

# 导入 Base64 编解码模块。
import base64

# 导入日期时间处理模块。
from datetime import datetime

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入 SSL 上下文模块。
import ssl

# 导入 Cookie 处理器。
from urllib.request import HTTPCookieProcessor

# 导入 HTTPS 请求处理器。
from urllib.request import HTTPSHandler

# 导入请求客户端构造函数。
from urllib.request import build_opener

# 导入 Excel 工作簿对象。
from openpyxl import Workbook

# 导入 Excel 列标识转换函数。
from openpyxl.utils import get_column_letter

# 导入 PEM 公钥加载函数。
from cryptography.hazmat.primitives import serialization

# 导入 RSA 补齐规则。
from cryptography.hazmat.primitives.asymmetric import padding

# 导入 Excel 统一样式函数。
from ddi_framework import apply_standard_excel_style

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入公共日志打印器。
from ddi_framework import PrintHandler

# 导入平台适配器基类。
from ddi_framework import BasePlatform


# 定义数据类型、页面和数据卡片的协议映射。
CRPCG_DATA_TYPE_CONFIG = {
    "购进流向": {
        "page_id": "wcafd056bbc4844bbb4809e9",
        "card_id": "qf0b9cc2fc5a64db597d5ab3",
        "date_params": (("开始日期", "iee84e3cd28824a459a9647b", "lcc0ed3be07a5489289905bf"), ("结束日期", "cf5511e4a7c3447da91f7b2b", "ld09a993d61974f5faee4057")),
        "business_date_fields": ("data_date",),
    },
    "配送流向": {
        "page_id": "j5fad646cae7c4f13aa9ee05",
        "card_id": "qcfc843a9ad1348f68114aed",
        "date_params": (("开始日期", "iee84e3cd28824a459a9647b", "pfbab02a98b89479fb5f3490"), ("结束日期", "cf5511e4a7c3447da91f7b2b", "o58d5f165a56044f0bb9eefa")),
        "business_date_fields": ("data_date",),
    },
    "库存": {
        "page_id": "eae834dac0fef44d692cfaf1",
        "card_id": "n482635868237425f92c123e",
        "date_params": (("结束日期", "cf5511e4a7c3447da91f7b2b", "h08b725b353774731a1dbc1a"),),
        "business_date_fields": ("data_date",),
    },
}

# 定义标准字段及接口字段候选项。
CRPCG_STANDARD_FIELD_MAPPING = {
    "购进流向": (("标准数据类型", ()), ("标准业务日期", ("data_date",)), ("标准商品编码", ("产品编码", "product_id")), ("标准商品名称", ("product_name", "产品名称")), ("标准商品规格", ("product_spec",)), ("标准生产企业", ("manufacture_name",)), ("标准批号", ("lot_no",)), ("标准单位", ("product_unit",)), ("标准数量", ("qty",)), ("标准金额", ("total_amount",)), ("标准供应商名称", ("vendor_name",))),
    "配送流向": (("标准数据类型", ()), ("标准业务日期", ("data_date",)), ("标准商品编码", ("产品编码", "product_id")), ("标准商品名称", ("产品名称", "product_name")), ("标准商品规格", ("product_spec",)), ("标准生产企业", ("manufacture_name",)), ("标准批号", ("lot_no",)), ("标准单位", ("product_unit",)), ("标准数量", ("qty",)), ("标准金额", ("total_sale",)), ("标准客户名称", ("custom_name",))),
    "库存": (("标准数据类型", ()), ("标准业务日期", ("data_date",)), ("标准商品编码", ("产品编码", "product_id")), ("标准商品名称", ("product_name", "产品名称")), ("标准商品规格", ("product_spec",)), ("标准生产企业", ("manufacture_name",)), ("标准批号", ("lot_no",)), ("标准单位", ("product_unit",)), ("标准数量", ("qty",)), ("标准金额", ("total_amt",)), ("标准供应商名称", ("vendor_name",))),
}


# 定义华润医药商业上游数据服务平台适配器。
class CrpcgPlatform(BasePlatform):

    """
    华润医药商业上游数据服务平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 华润医药商业上游数据服务平台适配器实例。
    :rtype: CrpcgPlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化旧 TLS 兼容会话、账号配置和查询状态。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用平台基类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化公共日志打印器。
        self.ph = printer or PrintHandler()

        # 读取平台接口根地址。
        self.base_url = str(account.get("接口根地址", "https://lxdata.crpcg.com")).rstrip("/")

        # 读取平台租户标识。
        self.domain_id = str(account.get("租户", "guanbi"))

        # 读取平台登录账号。
        self.login_id = str(account.get("账号", ""))

        # 读取平台登录密码。
        self.password = str(account.get("密码", ""))

        # 读取稳定的单页查询数量。
        self.page_size = max(1, int(account.get("分页大小", 1000)))

        # 初始化带 Cookie 管理的协议客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 创建默认 TLS 校验上下文。
        ssl_context = ssl.create_default_context()

        # 启用目标旧服务器必需的 TLS 重新协商兼容选项。
        ssl_context.options |= ssl.OP_LEGACY_SERVER_CONNECT

        # 使用受限兼容上下文替换当前平台会话的 HTTPS 处理器。
        self.http.opener = build_opener(HTTPCookieProcessor(self.http.cookie_jar), HTTPSHandler(context=ssl_context))

        # 初始化登录令牌。
        self.token = ""

        # 初始化登录状态标识。
        self.logged_in = False

    # 定义平台登录方法。
    def login(self):

        """
        通过 RSA 密码加密与 Cookie 会话执行平台登录。

        :return: 不含敏感凭据的登录摘要。
        :rtype: dict
        """

        # 打印登录阶段分隔线。
        self.ph.section("华润医药商业上游数据服务平台登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 校验登录账号已配置。
        if not self.login_id:

            # 抛出账号缺失异常。
            raise RuntimeError("登录失败：未配置账号。")

        # 校验登录密码已配置。
        if not self.password:

            # 抛出密码缺失异常。
            raise RuntimeError("登录失败：未配置密码。")

        # 请求验证码策略以建立登录 Cookie 并确认无需人工验证码。
        captcha_data = self.post_json("/api/user/captcha", {"domain": self.domain_id, "account": self.login_id}, self.login_page_url(), raw_backend_response=True)

        # 读取验证码策略对象。
        captcha_response = captcha_data.get("response") or {}

        # 判断当前账号是否需要验证码。
        if captcha_response.get("needCaptcha"):

            # 抛出不支持的人工验证码异常。
            raise RuntimeError("登录失败：当前账号需要人工验证码，暂不能执行纯协议查询。")

        # 请求 RSA 登录公钥。
        key_data = self.get_json("/api/user/password/encrypt", self.login_page_url())

        # 读取 RSA 公钥文本。
        public_key = str(key_data.get("publicKey", ""))

        # 校验 RSA 公钥已返回。
        if not public_key:

            # 抛出 RSA 公钥缺失异常。
            raise RuntimeError("登录失败：密码加密接口未返回 RSA 公钥。")

        # 根据站点规则加密登录密码。
        encrypted_password = self.encrypt_password(public_key, self.password)

        # 构建登录请求数据。
        login_payload = {"domain": self.domain_id, "loginId": self.login_id, "password": encrypted_password}

        # 提交登录请求。
        login_data = self.post_json("/api/user/sign-in", login_payload, self.login_page_url())

        # 校验登录接口返回成功状态。
        if str(login_data.get("result", "")) != "ok":

            # 抛出不包含密码的登录失败异常。
            raise RuntimeError(f"登录失败：{login_data.get('message') or login_data.get('error') or '服务端未返回成功状态'}")

        # 保存会话令牌。
        self.token = str(login_data.get("uIdToken", ""))

        # 校验会话令牌已返回。
        if not self.token:

            # 抛出会话令牌缺失异常。
            raise RuntimeError("登录失败：服务端未返回会话令牌。")

        # 标记当前协议会话已登录。
        self.logged_in = True

        # 打印不含令牌正文的登录结果。
        self.ph.success(self.ph.join(self.ph.kv("登录状态", "成功", "green"), self.ph.kv("Cookie", len(list(self.http.cookie_jar)), "cyan"), self.ph.kv("会话令牌", "已获取", "green")))

        # 返回不含敏感值的登录摘要。
        return {"登录状态": "成功", "Cookie数量": len(list(self.http.cookie_jar)), "会话令牌": "已获取"}

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存公共任务运行时所需的 Cookie 文件。

        :param output_dir: Cookie 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建 Cookie 输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存当前协议会话 Cookie。
        self.http.save_cookie_file(cookie_path)

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义原始数据查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定数据类型的接口原始字段完整记录。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 接口原始记录列表。
        :rtype: list[dict]
        """

        # 查询当前数据类型的完整协议结果。
        query_result = self.query_platform_data(data_type, start_date, end_date)

        # 返回接口原始记录列表。
        return query_result["rows"]

    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定数据类型并生成标准字段加接口原始字段 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: Excel 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询当前数据类型的完整协议结果。
        query_result = self.query_platform_data(data_type, start_date, end_date)

        # 解析 Excel 输出目录。
        target_directory = output_dir or os.getcwd()

        # 创建 Excel 输出目录。
        os.makedirs(target_directory, exist_ok=True)

        # 构建公共标准 Excel 文件名。
        file_name = self.build_output_file_name(data_type, query_result["start_date"], query_result["end_date"])

        # 拼接 Excel 完整输出路径。
        output_path = os.path.join(target_directory, file_name)

        # 写入标准字段和接口原始字段。
        self.write_rows_to_excel(data_type, query_result["rows"], query_result["raw_headers"], output_path)

        # 打印 Excel 落地摘要。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan"), self.ph.kv("接口字段数", len(query_result["raw_headers"]), "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义协议导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        使用协议查询结果生成本地 Excel 文件。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: Excel 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 调用查询并落地 Excel 的统一流程。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义库存命名日期解析方法。
    def resolve_output_query_date(self, data_type, start_date=None, end_date=None):

        """
        为库存使用导出当天日期，为流向使用查询开始日期。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 文件命名使用的日期文本。
        :rtype: str
        """

        # 判断当前数据类型是否为库存。
        if data_type == "库存":

            # 返回库存导出当天日期。
            return datetime.now().strftime("%Y-%m-%d")

        # 返回流向查询开始日期。
        return start_date or end_date or datetime.now().strftime("%Y-%m-%d")

    # 定义平台数据查询方法。
    def query_platform_data(self, data_type, start_date, end_date):

        """
        建立登录态、按真实日期参数分页查询数据卡片并校验业务日期。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询记录、接口字段及实际日期范围。
        :rtype: dict
        """

        # 校验当前数据类型存在协议映射。
        if data_type not in CRPCG_DATA_TYPE_CONFIG:

            # 抛出不支持数据类型异常。
            raise ValueError(f"不支持的数据类型：{data_type}")

        # 确保当前协议会话已登录。
        if not self.logged_in:

            # 执行平台登录。
            self.login()

        # 规范化查询开始日期。
        normalized_start_date = self.normalize_date(start_date, "开始日期")

        # 规范化查询结束日期。
        normalized_end_date = self.normalize_date(end_date, "结束日期")

        # 校验查询日期范围顺序。
        if normalized_start_date > normalized_end_date:

            # 抛出日期范围异常。
            raise ValueError("查询开始日期不能晚于结束日期。")

        # 读取当前数据类型协议配置。
        type_config = CRPCG_DATA_TYPE_CONFIG[data_type]

        # 打印当前数据类型查询分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", normalized_start_date, "cyan"), self.ph.kv("结束日期", normalized_end_date, "cyan")))

        # 初始化累计业务记录集合。
        rows = []

        # 初始化接口字段顺序集合。
        raw_headers = []

        # 初始化接口分页偏移量。
        offset = 0

        # 初始化接口页码。
        page_index = 1

        # 循环读取全部接口分页数据。
        while True:

            # 构建当前页真实动态日期参数。
            dynamic_params = self.build_dynamic_params(type_config, normalized_start_date, normalized_end_date)

            # 构建当前页数据卡片请求体。
            payload = {"filters": [], "treeFilters": [], "dynamicParams": dynamic_params, "dynamicFieldFilters": [], "combinationFilters": [], "layerTreeFilters": [], "refresh": True, "view": "GRAPH", "isUniversalStructure": False, "offset": offset, "limit": self.page_size}

            # 拼接当前页数据卡片接口路径。
            card_path = f"/api/card/{type_config['card_id']}/data"

            # 提交当前页数据卡片请求。
            response_data = self.post_json(card_path, payload, self.page_url(type_config), raw_backend_response=True, page_url=self.page_url(type_config))

            # 校验数据卡片接口返回成功状态。
            if str(response_data.get("result", "")) != "ok":

                # 抛出数据卡片查询失败异常。
                raise RuntimeError(f"{data_type}查询失败：{response_data.get('message') or response_data.get('error') or '服务端未返回成功状态'}")

            # 读取当前页图表响应主体。
            chart_main = ((response_data.get("response") or {}).get("chartMain") or {})

            # 解析当前页接口字段和业务记录。
            page_headers, page_rows = self.parse_chart_rows(chart_main)

            # 合并当前页新增接口字段。
            self.merge_raw_headers(raw_headers, page_headers)

            # 追加当前页业务记录。
            rows.extend(page_rows)

            # 读取服务端是否仍有后续数据。
            has_more_data = bool(chart_main.get("hasMoreData"))

            # 打印当前页查询进度。
            self.ph.info(self.ph.join(self.ph.kv("页码", page_index, "cyan"), self.ph.kv("本页记录", len(page_rows), "cyan"), self.ph.kv("累计记录", len(rows), "cyan"), self.ph.kv("继续分页", "是" if has_more_data else "否", "cyan")))

            # 判断服务端已无后续数据。
            if not has_more_data:

                # 结束分页读取。
                break

            # 判断当前页未返回业务记录。
            if not page_rows:

                # 抛出分页异常避免无限循环。
                raise RuntimeError(f"{data_type}查询分页异常：服务端提示仍有后续数据但当前页无业务记录。")

            # 累加下一页接口偏移量。
            offset += len(page_rows)

            # 累加当前接口页码。
            page_index += 1

        # 校验接口返回的实际业务日期范围。
        self.validate_business_dates(data_type, rows, type_config["business_date_fields"], normalized_start_date, normalized_end_date)

        # 返回完整查询结果。
        return {"rows": rows, "raw_headers": raw_headers, "start_date": normalized_start_date, "end_date": normalized_end_date}

    # 定义动态日期参数构建方法。
    def build_dynamic_params(self, type_config, start_date, end_date):

        """
        根据页面元数据构建数据卡片的真实动态日期参数。

        :param type_config: 数据类型协议配置。
        :type type_config: dict
        :param start_date: 规范化开始日期。
        :type start_date: str
        :param end_date: 规范化结束日期。
        :type end_date: str
        :return: 数据卡片动态参数列表。
        :rtype: list[dict]
        """

        # 初始化动态日期参数集合。
        dynamic_params = []

        # 遍历当前页面定义的日期参数。
        for parameter_name, parameter_id, source_card_id in type_config["date_params"]:

            # 判断当前参数是否为开始日期。
            parameter_value = start_date if parameter_name == "开始日期" else end_date

            # 追加当前页面真实日期参数。
            dynamic_params.append({"dpId": parameter_id, "name": parameter_name, "valueType": "DATE", "defaultValue": parameter_value, "customize": False, "multiple": False, "optionValue": [""], "inheritParent": True, "sourceCdId": source_card_id})

        # 返回完整动态日期参数集合。
        return dynamic_params

    # 定义图表记录解析方法。
    def parse_chart_rows(self, chart_main):

        """
        将数据卡片透视表响应转换为字段完整的原始记录。

        :param chart_main: 数据卡片图表主体。
        :type chart_main: dict
        :return: 接口字段顺序与业务记录列表。
        :rtype: tuple[list[str], list[dict]]
        """

        # 读取行维度元数据。
        row_meta = ((chart_main.get("row") or {}).get("meta") or [])

        # 读取行维度数据。
        row_values = ((chart_main.get("row") or {}).get("values") or [])

        # 读取度量列分组数据。
        metric_groups = ((chart_main.get("column") or {}).get("values") or [])

        # 读取度量数值数据。
        metric_rows = chart_main.get("data") or []

        # 将行维度元数据转换为接口字段名。
        row_headers = [self.resolve_header_name(meta_item, index) for index, meta_item in enumerate(row_meta)]

        # 将度量分组元数据转换为接口字段名。
        metric_headers = [self.resolve_header_name((metric_group or [{}])[0], index + len(row_headers)) for index, metric_group in enumerate(metric_groups)]

        # 合并全部接口字段名。
        raw_headers = self.unique_headers(row_headers + metric_headers)

        # 初始化当前页业务记录集合。
        rows = []

        # 遍历当前页的行维度数据。
        for row_index, dimension_cells in enumerate(row_values):

            # 判断当前行是否为汇总行。
            if any(isinstance(cell, dict) and cell.get("isGrandtotal") for cell in dimension_cells):

                # 跳过汇总行。
                continue

            # 初始化当前记录对象。
            row_data = {}

            # 遍历行维度单元格。
            for column_index, header_name in enumerate(raw_headers[:len(row_headers)]):

                # 读取当前维度单元格对象。
                cell_data = dimension_cells[column_index] if column_index < len(dimension_cells) else {}

                # 写入当前维度字段值。
                row_data[header_name] = self.extract_cell_value(cell_data)

            # 读取当前行对应的度量单元格集合。
            metric_cells = metric_rows[row_index] if row_index < len(metric_rows) else []

            # 遍历度量字段。
            for metric_index, header_name in enumerate(raw_headers[len(row_headers):]):

                # 读取当前度量单元格对象。
                metric_data = metric_cells[metric_index] if metric_index < len(metric_cells) else {}

                # 写入当前度量字段值。
                row_data[header_name] = self.extract_cell_value(metric_data)

            # 追加当前业务记录。
            rows.append(row_data)

        # 返回接口字段和业务记录。
        return raw_headers, rows

    # 定义接口字段名解析方法。
    def resolve_header_name(self, meta_item, index):

        """
        从接口字段元数据解析可稳定保存的原始字段名。

        :param meta_item: 接口字段元数据。
        :type meta_item: dict
        :param index: 字段位置索引。
        :type index: int
        :return: 原始字段名。
        :rtype: str
        """

        # 读取接口原始字段名称。
        origin_name = str((meta_item or {}).get("originTitle") or "").strip()

        # 判断接口原始字段名称是否存在。
        if origin_name:

            # 返回接口原始字段名称。
            return origin_name

        # 读取接口展示字段标题。
        title_name = str((meta_item or {}).get("title") or "").strip()

        # 判断接口展示字段标题是否存在。
        if title_name:

            # 返回接口展示字段标题。
            return title_name

        # 返回兜底字段名称。
        return f"field_{index + 1}"

    # 定义唯一字段名构建方法。
    def unique_headers(self, headers):

        """
        为重复接口字段保留全部列并生成唯一字段名。

        :param headers: 原始字段名列表。
        :type headers: list[str]
        :return: 唯一字段名列表。
        :rtype: list[str]
        """

        # 初始化字段使用次数字典。
        name_counts = {}

        # 初始化唯一字段名列表。
        result = []

        # 遍历原始字段名。
        for header_name in headers:

            # 读取当前字段使用次数。
            current_count = name_counts.get(header_name, 0) + 1

            # 保存当前字段使用次数。
            name_counts[header_name] = current_count

            # 判断当前字段是否首次出现。
            if current_count == 1:

                # 追加首次出现的字段名。
                result.append(header_name)

                # 进入下一字段。
                continue

            # 追加带序号的重复字段名。
            result.append(f"{header_name}_{current_count}")

        # 返回唯一字段名列表。
        return result

    # 定义单元格值解析方法。
    def extract_cell_value(self, cell_data):

        """
        从接口单元格对象中提取原始业务值。

        :param cell_data: 接口单元格对象。
        :type cell_data: dict
        :return: 单元格业务值。
        :rtype: object
        """

        # 判断当前单元格是否为字典对象。
        if not isinstance(cell_data, dict):

            # 返回非字典单元格原始值。
            return cell_data

        # 判断当前单元格是否包含度量值。
        if "v" in cell_data:

            # 返回度量单元格数值。
            return cell_data.get("v")

        # 返回维度单元格标题值。
        return cell_data.get("title")

    # 定义接口字段合并方法。
    def merge_raw_headers(self, raw_headers, page_headers):

        """
        按首次出现顺序合并分页接口字段。

        :param raw_headers: 累计接口字段列表。
        :type raw_headers: list[str]
        :param page_headers: 当前页接口字段列表。
        :type page_headers: list[str]
        :return: None。
        :rtype: None
        """

        # 遍历当前页接口字段。
        for header_name in page_headers:

            # 判断当前字段是否已保存。
            if header_name in raw_headers:

                # 跳过已保存字段。
                continue

            # 追加当前页新字段。
            raw_headers.append(header_name)

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, data_type, rows, raw_headers, output_path):

        """
        写入标准字段、接口全部字段和统一 Excel 样式。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 接口原始记录列表。
        :type rows: list[dict]
        :param raw_headers: 接口原始字段列表。
        :type raw_headers: list[str]
        :param output_path: Excel 输出路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 创建新的 Excel 工作簿。
        workbook = Workbook()

        # 获取默认活动工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = data_type[:31]

        # 读取当前数据类型标准字段映射。
        standard_mapping = CRPCG_STANDARD_FIELD_MAPPING[data_type]

        # 构建标准字段标题列表。
        standard_headers = [header_name for header_name, _ in standard_mapping]

        # 构建接口字段标题列表。
        original_headers = [f"原始_{header_name}" for header_name in raw_headers]

        # 写入标准字段和接口字段表头。
        worksheet.append(standard_headers + original_headers)

        # 遍历接口原始记录。
        for row_data in rows:

            # 构建当前记录标准字段值。
            standard_values = [self.resolve_standard_value(data_type, header_name, candidates, row_data) for header_name, candidates in standard_mapping]

            # 构建当前记录接口原始字段值。
            original_values = [row_data.get(header_name, "") for header_name in raw_headers]

            # 写入当前完整记录行。
            worksheet.append(standard_values + original_values)

        # 遍历全部 Excel 列调整显示宽度。
        for column_index in range(1, worksheet.max_column + 1):

            # 初始化当前列最大文本长度。
            max_length = 0

            # 遍历当前列已写入单元格。
            for row_index in range(1, worksheet.max_row + 1):

                # 读取当前单元格文本。
                cell_text = str(worksheet.cell(row=row_index, column=column_index).value or "")

                # 更新当前列最大文本长度。
                max_length = max(max_length, len(cell_text))

            # 设置限制后的当前列宽度。
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 36)

        # 保存 Excel 工作簿。
        workbook.save(output_path)

        # 关闭 Excel 工作簿。
        workbook.close()

        # 应用项目统一 Excel 样式。
        apply_standard_excel_style(output_path)

    # 定义标准字段值解析方法。
    def resolve_standard_value(self, data_type, header_name, candidates, row_data):

        """
        从原始字段中解析单个标准字段值。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param header_name: 标准字段名称。
        :type header_name: str
        :param candidates: 原始字段候选项。
        :type candidates: tuple[str, ...]
        :param row_data: 接口原始记录。
        :type row_data: dict
        :return: 标准字段值。
        :rtype: object
        """

        # 判断当前字段是否为标准数据类型。
        if header_name == "标准数据类型":

            # 返回当前 DDI 数据类型。
            return data_type

        # 遍历当前标准字段候选项。
        for field_name in candidates:

            # 判断当前候选字段是否存在有效值。
            if row_data.get(field_name) not in (None, ""):

                # 返回当前候选字段值。
                return row_data.get(field_name)

        # 返回空标准字段值。
        return ""

    # 定义业务日期校验方法。
    def validate_business_dates(self, data_type, rows, date_fields, start_date, end_date):

        """
        校验接口返回记录的业务日期范围。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 接口原始记录列表。
        :type rows: list[dict]
        :param date_fields: 业务日期候选字段。
        :type date_fields: tuple[str, ...]
        :param start_date: 请求开始日期。
        :type start_date: str
        :param end_date: 请求结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 初始化可解析业务日期集合。
        business_dates = []

        # 遍历接口原始记录。
        for row_data in rows:

            # 遍历业务日期候选字段。
            for field_name in date_fields:

                # 解析当前候选业务日期。
                parsed_date = self.parse_date_value(row_data.get(field_name))

                # 判断当前候选业务日期可用。
                if parsed_date:

                    # 保存当前业务日期。
                    business_dates.append(parsed_date)

                    # 停止读取当前记录的其他候选字段。
                    break

        # 判断当前查询未返回业务记录。
        if not rows:

            # 打印空结果日期校验日志。
            self.ph.info(self.ph.join(self.ph.kv("业务日期范围", "空结果", "yellow"), self.ph.kv("日期校验", "无记录无需校验", "cyan")))

            # 结束日期校验。
            return

        # 判断当前业务日期字段无法解析。
        if not business_dates:

            # 抛出业务日期缺失异常。
            raise RuntimeError(f"{data_type}查询成功但无法从接口记录解析业务日期。")

        # 读取最小业务日期。
        minimum_date = min(business_dates)

        # 读取最大业务日期。
        maximum_date = max(business_dates)

        # 判断当前数据类型是否为库存。
        if data_type == "库存":

            # 校验库存快照日期等于查询结束日期。
            is_valid = minimum_date == end_date and maximum_date == end_date

        else:

            # 校验流向业务日期落在请求范围内。
            is_valid = start_date <= minimum_date <= end_date and start_date <= maximum_date <= end_date

        # 判断业务日期范围是否有效。
        if not is_valid:

            # 抛出业务日期范围异常。
            raise RuntimeError(f"{data_type}业务日期范围异常：{minimum_date} 至 {maximum_date}，请求范围为 {start_date} 至 {end_date}。")

        # 打印业务日期范围校验结果。
        self.ph.success(self.ph.join(self.ph.kv("业务日期范围", f"{minimum_date} 至 {maximum_date}", "green"), self.ph.kv("日期校验", "通过", "green")))

    # 定义日期文本规范化方法。
    def normalize_date(self, value, field_name):

        """
        将输入日期规范化为 YYYY-MM-DD 文本。

        :param value: 输入日期值。
        :type value: str | None
        :param field_name: 日期字段名称。
        :type field_name: str
        :return: 规范化日期文本。
        :rtype: str
        """

        # 解析默认日期文本。
        date_text = str(value or datetime.now().strftime("%Y-%m-%d")).strip()

        # 尝试解析标准日期格式。
        try:

            # 返回标准格式日期文本。
            return datetime.strptime(date_text[:10], "%Y-%m-%d").strftime("%Y-%m-%d")

        except ValueError as error:

            # 抛出日期格式异常。
            raise ValueError(f"{field_name}格式无效，应为 YYYY-MM-DD：{date_text}") from error

    # 定义接口日期值解析方法。
    def parse_date_value(self, value):

        """
        从接口日期值中提取 YYYY-MM-DD 日期文本。

        :param value: 接口日期值。
        :type value: object
        :return: 可解析日期文本或空值。
        :rtype: str
        """

        # 判断当前接口日期值是否为空。
        if value in (None, ""):

            # 返回空日期值。
            return ""

        # 提取当前接口日期文本前十位。
        date_text = str(value).strip()[:10]

        # 尝试解析当前接口日期文本。
        try:

            # 返回标准格式业务日期文本。
            return datetime.strptime(date_text, "%Y-%m-%d").strftime("%Y-%m-%d")

        except ValueError:

            # 返回不可解析日期空值。
            return ""

    # 定义页面地址构建方法。
    def page_url(self, type_config):

        """
        构建当前数据类型页面地址。

        :param type_config: 数据类型协议配置。
        :type type_config: dict
        :return: 页面完整地址。
        :rtype: str
        """

        # 返回当前数据类型页面完整地址。
        return f"{self.base_url}/page/{type_config['page_id']}"

    # 定义登录页面地址构建方法。
    def login_page_url(self):

        """
        构建平台登录页面地址。

        :return: 登录页面完整地址。
        :rtype: str
        """

        # 返回平台登录页面完整地址。
        return f"{self.base_url}/auth/index"

    # 定义 RSA 密码加密方法。
    def encrypt_password(self, public_key_text, password):

        """
        按平台 RSA-PKCS1v15 和双重 Base64 规则加密登录密码。

        :param public_key_text: 接口返回的 RSA 公钥文本。
        :type public_key_text: str
        :param password: 明文登录密码。
        :type password: str
        :return: 登录接口所需密码密文。
        :rtype: str
        """

        # 构建 PEM 格式公钥文本。
        pem_text = f"-----BEGIN PUBLIC KEY-----\n{public_key_text}\n-----END PUBLIC KEY-----\n"

        # 加载 RSA 公钥对象。
        public_key = serialization.load_pem_public_key(pem_text.encode("ascii"))

        # 按 PKCS1v15 规则加密密码字节。
        encrypted_bytes = public_key.encrypt(str(password).encode("utf-8"), padding.PKCS1v15())

        # 执行第一次 Base64 编码。
        first_base64 = base64.b64encode(encrypted_bytes)

        # 执行平台所需的第二次 Base64 编码。
        return base64.b64encode(first_base64).decode("ascii")

    # 定义 JSON GET 请求方法。
    def get_json(self, path, referer):

        """
        发送平台 JSON GET 请求并解析响应。

        :param path: 接口相对路径。
        :type path: str
        :param referer: 请求来源页面地址。
        :type referer: str
        :return: JSON 响应对象。
        :rtype: dict
        """

        # 构建当前请求头。
        headers = self.build_headers(referer)

        # 发送 JSON GET 请求。
        result = self.http.request("GET", f"{self.base_url}{path}", headers=headers, timeout=60)

        # 解析并返回 JSON 响应对象。
        return self.http.decode_json(result)

    # 定义 JSON POST 请求方法。
    def post_json(self, path, payload, referer, raw_backend_response=False, page_url=None):

        """
        发送平台 JSON POST 请求并解析响应。

        :param path: 接口相对路径。
        :type path: str
        :param payload: JSON 请求体。
        :type payload: dict
        :param referer: 请求来源页面地址。
        :type referer: str
        :param raw_backend_response: 是否请求原始后端响应。
        :type raw_backend_response: bool
        :param page_url: 数据页面地址。
        :type page_url: str | None
        :return: JSON 响应对象。
        :rtype: dict
        """

        # 构建当前请求头。
        headers = self.build_headers(referer, raw_backend_response=raw_backend_response, page_url=page_url)

        # 编码 JSON 请求体。
        request_body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 发送 JSON POST 请求。
        result = self.http.request("POST", f"{self.base_url}{path}", headers=headers, data=request_body, timeout=90)

        # 解析并返回 JSON 响应对象。
        return self.http.decode_json(result)

    # 定义平台请求头构建方法。
    def build_headers(self, referer, raw_backend_response=False, page_url=None):

        """
        构建平台登录和数据卡片接口所需请求头。

        :param referer: 请求来源页面地址。
        :type referer: str
        :param raw_backend_response: 是否请求原始后端响应。
        :type raw_backend_response: bool
        :param page_url: 数据页面地址。
        :type page_url: str | None
        :return: 平台请求头。
        :rtype: dict
        """

        # 构建基础 JSON 请求头。
        headers = self.http.json_headers(referer)

        # 设置浏览器可接受 JSON 响应头。
        headers["Accept"] = "application/json"

        # 设置禁用缓存请求头。
        headers["Cache-Control"] = "no-cache"

        # 判断是否需要原始后端响应。
        if raw_backend_response:

            # 写入原始后端响应请求头。
            headers["raw-backend-response"] = "TRUE"

        # 判断当前请求是否为数据页面请求。
        if page_url:

            # 写入租户编码请求头。
            headers["x-dom-id"] = base64.b64encode(self.domain_id.encode("utf-8")).decode("ascii")

            # 写入账号编码请求头。
            headers["user-id"] = base64.b64encode(self.login_id.encode("utf-8")).decode("ascii")

            # 判断当前会话令牌是否存在。
            if self.token:

                # 写入会话令牌请求头。
                headers["token"] = self.token

        # 返回完整平台请求头。
        return headers
