# -*- coding: utf-8 -*-
"""
国药控股安徽信息服务平台协议适配器。

本模块负责国药控股安徽信息服务平台的 WebForms 登录、验证码识别、认证 Cookie 维护和登录后首页校验。
"""

# 导入操作系统路径模块。
import os

# 导入警告控制模块。
import warnings

# 导入正则表达式模块。
import re

# 导入 ZIP 文件读取模块。
import zipfile

# 导入日期时间对象。
from datetime import datetime

# 导入 HTML 解析器。
from html.parser import HTMLParser

# 导入 URL 拼接函数。
from urllib.parse import urljoin

# 导入表单编码函数。
from urllib.parse import urlencode

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入验证码识别器。
from ddi_framework import CaptchaSolver

# 导入下载命名工具。
from ddi_framework import DownloadWatcher

# 导入日志打印器。
from ddi_framework import PrintHandler

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient


# 定义国控安徽登录表单解析器。
class GykgahLoginFormParser(HTMLParser):

    """
    国控安徽登录页表单解析器。

    :return: 国控安徽登录表单解析器实例。
    :rtype: GykgahLoginFormParser
    """

    # 定义解析器初始化方法。
    def __init__(self):

        """
        初始化登录表单解析器。

        :return: None。
        :rtype: None
        """

        # 调用父类初始化方法。
        super().__init__(convert_charrefs=True)

        # 初始化表单字段字典。
        self.fields = {}

        # 初始化验证码地址。
        self.captcha_url = ""

    # 定义开始标签处理方法。
    def handle_starttag(self, tag, attrs):

        """
        处理 HTML 开始标签。

        :param tag: 标签名称。
        :type tag: str
        :param attrs: 标签属性列表。
        :type attrs: list[tuple[str, str | None]]
        :return: None。
        :rtype: None
        """

        # 转换标签名称。
        tag_name = tag.lower()

        # 转换属性列表为字典。
        attr_map = {name: value or "" for name, value in attrs}

        # 处理输入控件。
        if tag_name == "input":

            # 读取输入控件名称。
            field_name = attr_map.get("name", "")

            # 判断输入控件名称是否存在。
            if field_name:

                # 保存输入控件默认值。
                self.fields[field_name] = attr_map.get("value", "")

        # 处理验证码图片。
        if tag_name == "img":

            # 读取图片地址。
            image_src = attr_map.get("src", "")

            # 判断图片地址是否为验证码地址。
            if "captcha" in image_src.lower():

                # 保存验证码地址。
                self.captcha_url = image_src


# 定义国控安徽信息服务平台协议适配器。
class GykgahPlatform(BasePlatform):

    """
    国控安徽信息服务平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 国控安徽信息服务平台适配器实例。
    :rtype: GykgahPlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化国控安徽信息服务平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "https://www.gykgah.com:9027/")

        # 保存平台根地址。
        self.base_url = account.get("接口根地址", "https://www.gykgah.com:9027")

        # 保存首页地址。
        self.index_url = account.get("首页地址", urljoin(self.base_url + "/", "Index.aspx"))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化验证码识别器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录状态。
        self.logged_in = False

        # 初始化登录结果。
        self.login_result = {}

    # 定义平台登录方法。
    def login(self):

        """
        执行国控安徽信息服务平台登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("国控安徽登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 执行验证码登录重试。
        self.login_result = self.login_with_captcha_retry()

        # 标记登录状态。
        self.logged_in = True

        # 返回登录结果。
        return self.login_result

    # 定义验证码登录重试方法。
    def login_with_captcha_retry(self, max_retries=8):

        """
        使用图片验证码执行登录重试。

        :param max_retries: 最大重试次数。
        :type max_retries: int
        :return: 登录结果。
        :rtype: dict
        """

        # 初始化最近一次登录错误。
        last_error = ""

        # 遍历验证码重试次数。
        for retry_index in range(max_retries):

            # 打印验证码尝试日志。
            self.ph.info(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{max_retries} 次", "cyan"))

            # 获取登录上下文。
            login_context = self.get_login_context()

            # 识别验证码文本。
            captcha_code = self.solve_captcha_code(login_context["验证码字节"])

            # 打印验证码识别日志。
            self.ph.info(self.ph.kv("图片验证码识别完成", f"长度={len(captcha_code)}", "cyan"))

            # 构建登录表单。
            payload = self.build_login_payload(login_context["表单字段"], captcha_code)

            # 提交登录表单。
            response_text = self.submit_login_payload(payload)

            # 读取登录响应状态。
            login_status = self.parse_login_status(response_text)

            # 判断登录是否成功。
            if login_status["登录成功"]:

                # 验证登录后首页。
                self.verify_index_page()

                # 读取密码状态。
                password_status = "需修改" if login_status["密码需修改"] else "正常"

                # 打印登录成功日志。
                self.ph.success(self.ph.join(self.ph.kv("正式登录成功", "认证 Cookie 已获取", "green"), self.ph.kv("密码状态", password_status, "yellow"), self.ph.kv("Cookie数", len(list(self.http.cookie_jar)), "cyan")))

                # 返回登录结果。
                return {"登录状态": "成功", "密码状态": password_status, "Cookie数量": len(list(self.http.cookie_jar)), "首页": self.index_url}

            # 保存本次登录错误。
            last_error = login_status["错误信息"]

            # 打印登录失败候选日志。
            self.ph.warning(self.ph.kv("登录校验未通过", last_error or "未知错误", "yellow"))

        # 抛出登录失败异常。
        raise RuntimeError(f"{self.company_name} 登录失败：{last_error}")

    # 定义登录上下文获取方法。
    def get_login_context(self):

        """
        获取登录页表单字段和验证码图片。

        :return: 登录上下文。
        :rtype: dict
        """

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 请求登录页。
        page_result = self.http.get(self.login_page_url, referer=self.base_url)

        # 解码登录页文本。
        page_text = self.decode_text(page_result)

        # 解析登录页表单。
        parser = self.parse_login_page(page_text)

        # 构建验证码地址。
        captcha_url = urljoin(self.login_page_url, parser.captcha_url)

        # 打印验证码请求日志。
        self.ph.info(self.ph.kv("请求图片验证码", self.extract_path(captcha_url), "cyan"))

        # 请求验证码图片。
        captcha_result = self.http.get(captcha_url, referer=self.login_page_url)

        # 返回登录上下文。
        return {"表单字段": parser.fields, "验证码字节": captcha_result.body}

    # 定义登录页解析方法。
    def parse_login_page(self, page_text):

        """
        解析登录页表单字段。

        :param page_text: 登录页 HTML 文本。
        :type page_text: str
        :return: 登录表单解析器。
        :rtype: GykgahLoginFormParser
        """

        # 初始化表单解析器。
        parser = GykgahLoginFormParser()

        # 输入登录页文本。
        parser.feed(page_text)

        # 校验页面状态字段。
        if "__VIEWSTATE" not in parser.fields:

            # 抛出页面状态异常。
            raise RuntimeError(f"{self.company_name} 登录页未返回 __VIEWSTATE")

        # 校验验证码地址。
        if not parser.captcha_url:

            # 抛出验证码地址异常。
            raise RuntimeError(f"{self.company_name} 登录页未返回验证码图片")

        # 返回表单解析器。
        return parser

    # 定义验证码识别方法。
    def solve_captcha_code(self, image_bytes):

        """
        识别登录图片验证码。

        :param image_bytes: 验证码图片字节。
        :type image_bytes: bytes
        :return: 验证码文本。
        :rtype: str
        """

        # 使用限定字符范围识别验证码。
        raw_text = self.captcha_solver.solve_raw_text_image_with_ranges(image_bytes, "0123456789abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ")

        # 清理非字母数字字符。
        captcha_code = re.sub(r"[^0-9A-Za-z]", "", raw_text or "")

        # 校验验证码长度。
        if len(captcha_code) < 4:

            # 抛出验证码识别异常。
            raise RuntimeError(f"{self.company_name} 验证码识别失败：{raw_text}")

        # 返回验证码文本。
        return captcha_code

    # 定义登录表单构建方法。
    def build_login_payload(self, form_fields, captcha_code):

        """
        构建登录表单字段。

        :param form_fields: 登录页表单字段。
        :type form_fields: dict
        :param captcha_code: 验证码文本。
        :type captcha_code: str
        :return: 登录表单字段。
        :rtype: dict
        """

        # 复制登录页表单字段。
        payload = dict(form_fields)

        # 写入登录账号。
        payload["UserId"] = self.clean_login_name()

        # 写入登录密码。
        payload["Password"] = self.clean_login_password()

        # 写入验证码文本。
        payload["Captcha"] = captcha_code

        # 写入登录按钮文本。
        payload["btnLogin"] = payload.get("btnLogin") or "登  陆"

        # 返回登录表单。
        return payload

    # 定义登录提交方法。
    def submit_login_payload(self, payload):

        """
        提交登录表单。

        :param payload: 登录表单字段。
        :type payload: dict
        :return: 登录响应文本。
        :rtype: str
        """

        # 打印登录提交日志。
        self.ph.info(self.ph.kv("提交正式登录请求", self.extract_path(self.login_page_url), "cyan"))

        # 构建请求头。
        headers = self.build_webforms_headers(self.login_page_url)

        # 编码请求体。
        data = urlencode(payload).encode("utf-8")

        # 提交登录表单。
        result = self.http.request("POST", self.login_page_url, headers=headers, data=data, timeout=60)

        # 返回登录响应文本。
        return self.decode_text(result)

    # 定义登录状态解析方法。
    def parse_login_status(self, response_text):

        """
        解析登录响应状态。

        :param response_text: 登录响应文本。
        :type response_text: str
        :return: 登录状态。
        :rtype: dict
        """

        # 判断认证 Cookie 是否存在。
        auth_cookie_exists = any(cookie.name == "Gykg.Web" for cookie in self.http.cookie_jar)

        # 判断密码是否需要修改。
        password_need_modify = "密码过期或复杂度不足" in response_text

        # 判断是否已经登录成功。
        if auth_cookie_exists:

            # 返回成功状态。
            return {"登录成功": True, "密码需修改": password_need_modify, "错误信息": ""}

        # 解析弹窗错误信息。
        error_message = self.extract_alert_message(response_text)

        # 返回失败状态。
        return {"登录成功": False, "密码需修改": password_need_modify, "错误信息": error_message}

    # 定义首页校验方法。
    def verify_index_page(self):

        """
        校验登录后首页是否可访问。

        :return: 首页 HTML 文本。
        :rtype: str
        """

        # 打印首页校验日志。
        self.ph.info(self.ph.kv("校验登录首页", self.extract_path(self.index_url), "cyan"))

        # 请求登录后首页。
        result = self.http.get(self.index_url, referer=self.login_page_url)

        # 解码首页文本。
        index_text = self.decode_text(result)

        # 校验首页关键内容。
        if "国控安徽信息平台" not in index_text or "FlowSystem" not in index_text:

            # 抛出首页校验异常。
            raise RuntimeError(f"{self.company_name} 登录后首页未返回有效菜单")

        # 返回首页文本。
        return index_text

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie。

        :param output_dir: Cookie 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建 Cookie 输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 打印 Cookie 保存日志。
        self.ph.success(self.ph.kv("登录 Cookie 已保存", self.ph.path(cookie_path), "green"))

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义平台导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        导出国控安徽信息服务平台数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: 导出文件路径。
        :rtype: str
        """

        # 返回平台导出的 Excel 文件。
        return self.download_flow_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义平台查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询国控安徽信息服务平台数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果列表。
        :rtype: list[dict]
        """

        # 下载平台 Excel 文件。
        excel_path = self.download_flow_excel(data_type, start_date=start_date, end_date=end_date, output_dir=None)

        # 返回 Excel 数据行。
        return self.read_rows_from_excel(excel_path)

    # 定义平台查询落地方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询国控安徽信息服务平台数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 返回平台导出的 Excel 文件。
        return self.download_flow_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义平台 Excel 下载方法。
    def download_flow_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        下载国控安徽信息服务平台 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 校验数据类型。
        self.validate_data_type(data_type)

        # 解析页面配置。
        page_config = self.resolve_page_config(data_type)

        # 构建业务页面地址。
        page_url = self.build_page_url(page_config)

        # 构建实际输出目录。
        save_dir = output_dir or "downloads"

        # 创建输出目录。
        os.makedirs(save_dir, exist_ok=True)

        # 打印数据类型查询分隔线。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 判断当前会话是否已登录。
        if not self.logged_in:

            # 执行平台登录。
            self.login()

        # 构建目标文件名。
        target_name = self.build_output_file_name(data_type, start_date, end_date)

        # 拼接目标文件路径。
        target_path = os.path.join(save_dir, target_name)

        # 打印导出页加载日志。
        self.ph.info(self.ph.join(self.ph.kv("加载导出页", page_config["页面名称"], "cyan"), self.ph.kv("端点", page_config["路径"], "yellow")))

        # 执行浏览器导出。
        self.download_with_browser(data_type, page_url, target_path, start_date, end_date)

        # 校验 Excel 文件结构。
        row_count = self.validate_xlsx_file(target_path)

        # 标准化 Excel 文件结构。
        self.normalize_xlsx_file(target_path)

        # 打印 Excel 文件生成日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(target_path), "green"), self.ph.kv("记录数", f"{row_count} 条", "cyan")))

        # 返回目标文件路径。
        return target_path

    # 定义浏览器导出方法。
    def download_with_browser(self, data_type, page_url, target_path, start_date, end_date):

        """
        使用浏览器执行 FineUI 导出并保存文件。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param page_url: 业务页面地址。
        :type page_url: str
        :param target_path: 目标文件路径。
        :type target_path: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 导入 Playwright 同步入口。
        from playwright.sync_api import sync_playwright

        # 打印浏览器导出日志。
        self.ph.info(self.ph.join(self.ph.kv("启动浏览器导出", data_type, "cyan"), self.ph.kv("FineUI状态", "__FSTATE", "yellow")))

        # 使用 Playwright 浏览器上下文。
        with sync_playwright() as playwright:

            # 启动 Chromium 浏览器。
            browser = self.launch_chromium_browser(playwright)

            # 捕获浏览器导出过程异常。
            try:

                # 创建浏览器上下文。
                context = browser.new_context(accept_downloads=True, viewport={"width": 1400, "height": 900})

                # 注入当前协议登录 Cookie。
                context.add_cookies(self.build_playwright_cookies())

                # 创建业务页面。
                page = context.new_page()

                # 打开业务页面。
                page.goto(page_url, wait_until="domcontentloaded", timeout=60000)

                # 等待 FineUI 页面渲染完成。
                self.wait_for_fineui_ready(page)

                # 设置页面查询日期。
                self.apply_browser_date_range(page, data_type, start_date, end_date)

                # 读取导出按钮文案。
                button_text = self.resolve_export_button_text(data_type)

                # 打印导出按钮日志。
                self.ph.info(self.ph.kv("触发导出按钮", button_text, "cyan"))

                # 删除已存在的目标文件。
                self.remove_existing_file(target_path)

                # 等待浏览器下载事件。
                with page.expect_download(timeout=120000) as download_info:

                    # 点击导出按钮。
                    page.get_by_text(button_text, exact=True).click(timeout=30000)

                # 读取下载对象。
                download = download_info.value

                # 保存下载文件。
                download.save_as(target_path)

                # 打印下载保存日志。
                self.ph.success(self.ph.kv("浏览器下载已保存", self.ph.path(target_path), "green"))

                # 关闭浏览器上下文。
                context.close()

            # 结束浏览器会话。
            finally:

                # 关闭浏览器。
                browser.close()

    # 定义 Chromium 启动方法。
    def launch_chromium_browser(self, playwright):

        """
        启动 Chromium 浏览器实例。

        :param playwright: Playwright 同步对象。
        :type playwright: Playwright
        :return: 浏览器实例。
        :rtype: Browser
        """

        # 读取 Chrome 可执行文件路径。
        chrome_path = self.account.get("Chrome路径", r"C:\Program Files\Google\Chrome\Application\chrome.exe")

        # 判断 Chrome 文件是否存在。
        if chrome_path and os.path.exists(chrome_path):

            # 使用本机 Chrome 启动浏览器。
            return playwright.chromium.launch(executable_path=chrome_path, headless=True, args=["--disable-gpu", "--no-sandbox"])

        # 使用 Playwright 默认 Chromium 启动浏览器。
        return playwright.chromium.launch(headless=True, args=["--disable-gpu", "--no-sandbox"])

    # 定义 Playwright Cookie 构建方法。
    def build_playwright_cookies(self):

        """
        构建 Playwright 浏览器 Cookie 列表。

        :return: 浏览器 Cookie 列表。
        :rtype: list[dict]
        """

        # 初始化 Cookie 列表。
        cookie_items = []

        # 遍历协议客户端 Cookie。
        for cookie_obj in self.http.cookie_jar:

            # 追加浏览器 Cookie。
            cookie_items.append({"name": cookie_obj.name, "value": cookie_obj.value, "domain": "www.gykgah.com", "path": cookie_obj.path or "/"})

        # 判断 Cookie 列表是否为空。
        if not cookie_items:

            # 抛出 Cookie 缺失异常。
            raise RuntimeError(f"{self.company_name} 浏览器导出缺少登录 Cookie")

        # 返回 Cookie 列表。
        return cookie_items

    # 定义 FineUI 页面等待方法。
    def wait_for_fineui_ready(self, page):

        """
        等待 FineUI 页面组件初始化完成。

        :param page: Playwright 页面对象。
        :type page: Page
        :return: None。
        :rtype: None
        """

        # 等待 FineUI 对象和导出按钮出现。
        page.wait_for_function("() => window.F && document.body && document.body.innerText.includes('导出')", timeout=60000)

        # 等待页面脚本稳定。
        page.wait_for_timeout(800)

        # 读取页面正文。
        body_text = page.locator("body").inner_text(timeout=30000)

        # 判断登录态是否异常。
        if "登录账号异常" in body_text or "请重新登录" in body_text:

            # 抛出登录态异常。
            raise RuntimeError(f"{self.company_name} 浏览器业务页提示登录异常")

    # 定义浏览器日期设置方法。
    def apply_browser_date_range(self, page, data_type, start_date, end_date):

        """
        设置浏览器页面日期区间。

        :param page: Playwright 页面对象。
        :type page: Page
        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: None。
        :rtype: None
        """

        # 判断库存页面是否需要日期。
        if "库存" in data_type:

            # 库存页面不设置日期。
            return None

        # 格式化开始日期。
        begin_text = self.format_page_date(start_date)

        # 格式化结束日期。
        end_text = self.format_page_date(end_date)

        # 执行 FineUI 日期控件赋值。
        page.evaluate(
            """
            ([startDate, endDate]) => {
                const begin = window.F && F('PanelView_GridView_Toolbar2_BeginDate');
                const end = window.F && F('PanelView_GridView_Toolbar2_EndDate');
                if (begin && begin.setText) begin.setText(startDate);
                if (begin && begin.setValue) begin.setValue(startDate);
                if (end && end.setText) end.setText(endDate);
                if (end && end.setValue) end.setValue(endDate);
                const beginInput = document.querySelector('[name="PanelView$GridView$Toolbar2$BeginDate"]');
                const endInput = document.querySelector('[name="PanelView$GridView$Toolbar2$EndDate"]');
                if (beginInput) beginInput.value = startDate;
                if (endInput) endInput.value = endDate;
            }
            """,
            [begin_text, end_text],
        )

        # 打印日期设置日志。
        self.ph.info(self.ph.join(self.ph.kv("写入页面日期", begin_text, "cyan"), self.ph.kv("结束日期", end_text, "cyan")))

        # 返回空结果。
        return None

    # 定义导出按钮文案解析方法。
    def resolve_export_button_text(self, data_type):

        """
        解析数据类型对应的导出按钮文案。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 导出按钮文案。
        :rtype: str
        """

        # 判断是否为库存页面。
        if "库存" in data_type:

            # 返回库存全量导出按钮文案。
            return "全部导出"

        # 返回流向日期导出按钮文案。
        return "按日期导出"

    # 定义 Excel 文件校验方法。
    def validate_xlsx_file(self, excel_path):

        """
        校验 Excel 文件结构并返回记录数。

        :param excel_path: Excel 文件路径。
        :type excel_path: str
        :return: 数据记录数。
        :rtype: int
        """

        # 判断文件是否存在。
        if not os.path.exists(excel_path):

            # 抛出文件缺失异常。
            raise RuntimeError(f"{self.company_name} 平台导出文件不存在：{excel_path}")

        # 校验 ZIP 文件结构。
        if not zipfile.is_zipfile(excel_path):

            # 抛出 Excel 结构异常。
            raise RuntimeError(f"{self.company_name} 平台未返回有效 xlsx 文件：{excel_path}")

        # 导入工作簿加载函数。
        from openpyxl import load_workbook

        # 忽略无默认样式提示。
        with warnings.catch_warnings():

            # 设置警告过滤策略。
            warnings.simplefilter("ignore")

            # 加载工作簿。
            workbook = load_workbook(excel_path, read_only=True, data_only=True)

            # 读取活动工作表。
            worksheet = workbook.active

            # 计算有效数据行数。
            row_count = self.count_worksheet_records(worksheet)

            # 关闭工作簿。
            workbook.close()

        # 返回记录数。
        return row_count

    # 定义工作表记录计数方法。
    def count_worksheet_records(self, worksheet):

        """
        统计工作表有效数据记录数。

        :param worksheet: 工作表对象。
        :type worksheet: Worksheet
        :return: 有效数据记录数。
        :rtype: int
        """

        # 初始化有效行数。
        row_count = 0

        # 遍历第二行之后的数据。
        for row_values in worksheet.iter_rows(min_row=2, values_only=True):

            # 判断当前行是否存在有效值。
            if any(value not in ("", None) for value in row_values):

                # 累加有效行数。
                row_count += 1

        # 返回有效行数。
        return row_count

    # 定义 Excel 文件标准化方法。
    def normalize_xlsx_file(self, excel_path):

        """
        标准化平台导出的 Excel 文件结构。

        :param excel_path: Excel 文件路径。
        :type excel_path: str
        :return: None。
        :rtype: None
        """

        # 导入工作簿加载函数。
        from openpyxl import load_workbook

        # 忽略平台文件缺少默认样式的提示。
        with warnings.catch_warnings():

            # 设置警告过滤策略。
            warnings.simplefilter("ignore")

            # 加载平台导出的工作簿。
            workbook = load_workbook(excel_path)

            # 重存为 openpyxl 标准结构。
            workbook.save(excel_path)

            # 关闭工作簿。
            workbook.close()

        # 返回空结果。
        return None

    # 定义 Excel 行读取方法。
    def read_rows_from_excel(self, excel_path):

        """
        读取 Excel 数据行。

        :param excel_path: Excel 文件路径。
        :type excel_path: str
        :return: 数据行列表。
        :rtype: list[dict]
        """

        # 导入工作簿加载函数。
        from openpyxl import load_workbook

        # 忽略无默认样式提示。
        with warnings.catch_warnings():

            # 设置警告过滤策略。
            warnings.simplefilter("ignore")

            # 加载工作簿。
            workbook = load_workbook(excel_path, read_only=True, data_only=True)

            # 读取活动工作表。
            worksheet = workbook.active

            # 读取全部行。
            raw_rows = list(worksheet.iter_rows(values_only=True))

            # 关闭工作簿。
            workbook.close()

        # 判断原始行是否为空。
        if not raw_rows:

            # 返回空列表。
            return []

        # 构建表头。
        headers = [str(value or f"字段{index}").strip() for index, value in enumerate(raw_rows[0], start=1)]

        # 初始化结果行列表。
        rows = []

        # 遍历原始数据行。
        for raw_row in raw_rows[1:]:

            # 构建数据行。
            row_data = {headers[index]: raw_row[index] if index < len(raw_row) else "" for index in range(len(headers))}

            # 判断数据行是否存在有效值。
            if any(value not in ("", None) for value in row_data.values()):

                # 追加数据行。
                rows.append(row_data)

        # 返回数据行。
        return rows

    # 定义页面配置解析方法。
    def resolve_page_config(self, data_type):

        """
        解析数据类型页面配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 页面配置。
        :rtype: dict
        """

        # 读取页面映射。
        page_map = self.account.get("页面映射", {})

        # 判断页面映射是否存在。
        if data_type not in page_map:

            # 抛出页面配置异常。
            raise ValueError(f"国控安徽信息服务平台暂未配置数据类型：{data_type}")

        # 返回页面配置。
        return page_map[data_type]

    # 定义页面地址构建方法。
    def build_page_url(self, page_config):

        """
        构建业务页面地址。

        :param page_config: 页面配置。
        :type page_config: dict
        :return: 业务页面地址。
        :rtype: str
        """

        # 读取页面路径。
        page_path = page_config.get("路径", "")

        # 判断页面路径是否为完整地址。
        if page_path.startswith("http://") or page_path.startswith("https://"):

            # 返回完整页面地址。
            return page_path

        # 返回拼接后的页面地址。
        return f"{self.base_url}{page_path}"

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验国控安徽信息服务平台数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 读取支持的数据类型列表。
        supported_types = self.account.get("数据类型列表", [])

        # 判断当前数据类型是否支持。
        if data_type not in supported_types:

            # 抛出数据类型异常。
            raise ValueError(f"国控安徽信息服务平台暂未接入数据类型：{data_type}")

    # 定义输出文件名构建方法。
    def build_output_file_name(self, data_type, start_date, end_date):

        """
        构建输出文件名。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 输出文件名。
        :rtype: str
        """

        # Delegate standard naming to the base platform.
        return super().build_output_file_name(data_type, start_date, end_date)

    # 定义页面日期格式化方法。
    def format_page_date(self, date_value):

        """
        格式化页面日期。

        :param date_value: 日期文本。
        :type date_value: str | None
        :return: 页面日期文本。
        :rtype: str
        """

        # 解析日期对象。
        date_obj = self.parse_date(date_value)

        # 返回页面日期格式。
        return f"{date_obj.year}/{date_obj.month:02d}/{date_obj.day:02d}"

    # 定义日期解析方法。
    def parse_date(self, date_value):

        """
        解析日期文本。

        :param date_value: 日期文本。
        :type date_value: str | None
        :return: 日期对象。
        :rtype: datetime
        """

        # 判断日期文本是否为空。
        if not date_value:

            # 返回当前日期。
            return datetime.now()

        # 遍历支持的日期格式。
        for date_format in ("%Y-%m-%d", "%Y/%m/%d"):

            # 捕获日期解析异常。
            try:

                # 返回解析后的日期。
                return datetime.strptime(str(date_value), date_format)

            # 处理日期解析异常。
            except ValueError:

                # 继续尝试下一种格式。
                continue

        # 抛出日期格式异常。
        raise ValueError(f"日期格式不支持：{date_value}")

    # 定义文件删除方法。
    def remove_existing_file(self, file_path):

        """
        删除已存在的目标文件。

        :param file_path: 文件路径。
        :type file_path: str
        :return: None。
        :rtype: None
        """

        # 判断目标文件是否存在。
        if os.path.exists(file_path):

            # 删除目标文件。
            os.remove(file_path)

    # 定义请求头构建方法。
    def build_webforms_headers(self, referer):

        """
        构建 WebForms 表单请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回浏览器表单请求头。
        return {"Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/x-www-form-urlencoded", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent()}

    # 定义文本解码方法。
    def decode_text(self, result):

        """
        解码 HTTP 响应文本。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 响应文本。
        :rtype: str
        """

        # 使用 UTF-8 解码响应体。
        return result.body.decode("utf-8", errors="replace")

    # 定义登录名清理方法。
    def clean_login_name(self):

        """
        读取登录账号。

        :return: 登录账号。
        :rtype: str
        """

        # 返回去除空白后的账号。
        return str(self.account.get("账号", "")).strip()

    # 定义登录密码清理方法。
    def clean_login_password(self):

        """
        读取登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回原始密码文本。
        return str(self.account.get("密码", ""))

    # 定义弹窗信息提取方法。
    def extract_alert_message(self, html_text):

        """
        提取页面弹窗信息。

        :param html_text: HTML 文本。
        :type html_text: str
        :return: 弹窗信息。
        :rtype: str
        """

        # 匹配 alert 文本。
        match_obj = re.search(r"alert\('([^']*)'\)", html_text)

        # 返回弹窗文本。
        return match_obj.group(1) if match_obj else ""

    # 定义路径提取方法。
    def extract_path(self, url):

        """
        提取 URL 路径用于日志展示。

        :param url: URL 地址。
        :type url: str
        :return: URL 路径。
        :rtype: str
        """

        # 移除平台根地址。
        if str(url).startswith(self.base_url):

            # 返回相对路径。
            return str(url)[len(self.base_url):] or "/"

        # 返回原始地址。
        return str(url)
