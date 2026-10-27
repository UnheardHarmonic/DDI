# -*- coding: utf-8 -*-
"""
广州医药股份有限公司供应链平台协议适配器。

本模块负责广州医药股份有限公司供应链平台的协议登录、认证令牌提取和本地会话保存。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期解析模块。
from datetime import datetime

# 导入 URL 解析函数。
from urllib.parse import urlparse

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义广州医药供应链平台协议适配器。
class GzyyPlatform(BasePlatform):

    """
    广州医药供应链平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 广州医药平台适配器实例。
    :rtype: GzyyPlatform
    """

    # 定义广州医药适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化广州医药平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "https://sip.gzmpc.com/user/login")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", "https://tccloud-gateway.gzmpc.com/sip-usercenter/v1/user/login")

        # 保存进货入库查询接口地址。
        self.purchase_query_url = account.get("进货入库接口", "https://tccloud-gateway.gzmpc.com/sip-dataflow/api/ebbms/su/suPage")

        # 保存销售出库查询接口地址。
        self.sale_query_url = account.get("销售出库接口", "https://tccloud-gateway.gzmpc.com/sip-dataflow/api/ebbms/sa/saleOutPage")

        # 保存即时库存查询接口地址。
        self.stock_query_url = account.get("即时库存接口", "https://tccloud-gateway.gzmpc.com/sip-dataflow/api/ebbms/immediate/page")

        # 保存平台页面根地址。
        self.page_base_url = self.build_origin(self.login_page_url)

        # 保存接口网关根地址。
        self.gateway_base_url = self.build_origin(self.login_api_url)

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.page_base_url)

        # 初始化登录令牌。
        self.login_token = ""

        # 初始化登录结果。
        self.login_result = {}

        # 保存默认分页大小。
        self.page_size = int(account.get("分页大小", 200))

    # 定义广州医药登录方法。
    def login(self):

        """
        执行广州医药协议登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("广州医药登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 请求登录页并初始化站点会话。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 构建登录请求体。
        payload = self.build_login_payload()

        # 打印正式登录请求日志。
        self.ph.info("提交正式登录请求: /sip-usercenter/v1/user/login")

        # 发送登录请求。
        result = self.post_login(payload)

        # 解析登录响应。
        login_result = self.http.decode_json(result)

        # 校验登录结果。
        if not self.is_login_success(login_result):

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 正式登录失败：{self.extract_login_message(login_result)}")

        # 保存登录响应。
        self.login_result = login_result

        # 提取登录令牌。
        self.login_token = self.extract_login_token(login_result)

        # 打印登录成功日志。
        self.ph.success(self.ph.kv("正式登录成功", self.extract_login_message(login_result), "green"))

        # 返回登录结果。
        return login_result

    # 定义登录请求体构建方法。
    def build_login_payload(self):

        """
        构建广州医药登录请求体。

        :return: 登录请求体。
        :rtype: dict
        """

        # 返回登录 JSON 字段。
        return {
            "account": self.clean_login_name(),
            "password": self.clean_login_password(),
        }

    # 定义登录请求提交方法。
    def post_login(self, payload):

        """
        提交广州医药登录请求。

        :param payload: 登录请求体。
        :type payload: dict
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 编码 JSON 请求体。
        request_body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建登录请求头。
        headers = self.build_login_headers()

        # 发送登录请求。
        return self.http.request("POST", self.login_api_url, headers=headers, data=request_body)

    # 定义登录请求头构建方法。
    def build_login_headers(self):

        """
        构建广州医药登录请求头。

        :return: 登录请求头。
        :rtype: dict
        """

        # 构建基础 JSON 请求头。
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            "Connection": "keep-alive",
            "Content-Type": "application/json;charset=UTF-8",
            "Origin": self.page_base_url,
            "Referer": self.login_page_url,
            "User-Agent": self.http.user_agent(),
        }

        # 读取可选的前端认证头。
        blade_auth = str(self.account.get("Blade-Auth", "")).strip()

        # 写入可选的前端认证头。
        if blade_auth:

            # 设置网关认证请求头。
            headers["Blade-Auth"] = blade_auth

        # 返回登录请求头。
        return headers

    # 定义登录成功判断方法。
    def is_login_success(self, login_result):

        """
        判断广州医药登录响应是否成功。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 是否登录成功。
        :rtype: bool
        """

        # 读取响应状态布尔值。
        status_value = login_result.get("status")

        # 读取响应编码。
        code_value = str(login_result.get("code", ""))

        # 判断响应状态和编码是否为成功。
        return bool(status_value) and code_value == "200"

    # 定义登录消息提取方法。
    def extract_login_message(self, login_result):

        """
        提取广州医药登录响应消息。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 登录响应消息。
        :rtype: str
        """

        # 返回接口消息或原始响应摘要。
        return str(login_result.get("msg") or login_result.get("message") or login_result)

    # 定义登录令牌提取方法。
    def extract_login_token(self, login_result):

        """
        提取广州医药登录响应中的认证令牌。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 认证令牌。
        :rtype: str
        """

        # 读取响应数据字段。
        data = login_result.get("data")

        # 处理数据字段为字符串令牌的响应。
        if isinstance(data, str):

            # 返回字符串令牌。
            return data

        # 处理数据字段为字典的响应。
        if isinstance(data, dict):

            # 遍历常见令牌字段名。
            for key in ["access_token", "accessToken", "token", "jwt", "Blade-Auth", "bladeAuth"]:

                # 读取当前字段值。
                token_value = data.get(key)

                # 返回非空令牌。
                if token_value:

                    # 返回令牌文本。
                    return str(token_value)

        # 遍历顶层常见令牌字段名。
        for key in ["access_token", "accessToken", "token", "jwt", "Blade-Auth", "bladeAuth"]:

            # 读取顶层字段值。
            token_value = login_result.get(key)

            # 返回非空令牌。
            if token_value:

                # 返回令牌文本。
                return str(token_value)

        # 返回空令牌。
        return ""

    # 定义登录账号清洗方法。
    def clean_login_name(self):

        """
        读取清洗后的登录账号。

        :return: 登录账号。
        :rtype: str
        """

        # 返回去除两端空白后的账号。
        return str(self.account["账号"]).strip()

    # 定义登录密码清洗方法。
    def clean_login_password(self):

        """
        读取清洗后的登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回去除两端空白后的密码。
        return str(self.account["密码"]).strip()

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie 与认证令牌。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 保存认证令牌文件。
        self.save_token_file(output_dir)

        # 打印 Cookie 保存日志。
        self.ph.success(self.ph.kv("登录 Cookie 已保存", self.ph.path(cookie_path), "green"))

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义令牌保存方法。
    def save_token_file(self, output_dir):

        """
        保存广州医药登录令牌到本地文件。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: 令牌文件路径。
        :rtype: str
        """

        # 构建令牌文件路径。
        token_path = os.path.join(output_dir, f"{self.company_name}_token.json")

        # 构建令牌文件内容。
        token_data = {"token": self.login_token, "login_result": self.login_result}

        # 写入令牌文件。
        with open(token_path, "w", encoding="utf-8") as file_obj:

            # 序列化令牌文件内容。
            json.dump(token_data, file_obj, ensure_ascii=False, indent=2)

        # 打印令牌保存日志。
        self.ph.success(self.ph.kv("登录令牌已保存", self.ph.path(token_path), "green"))

        # 返回令牌文件路径。
        return token_path

    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询广州医药流向数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 查询指定类型完整数据。
        query_result = self.query_gzyy_data(data_type, start_date, end_date)

        # 返回查询结果数据行。
        return query_result["rows"]

    # 定义流向导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        导出广州医药流向数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 导出目录。
        :type output_dir: str | None
        :return: 导出结果。
        :rtype: str
        """

        # 使用分页查询结果生成本地 Excel 文件。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询广州医药流向数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定类型完整数据。
        query_result = self.query_gzyy_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 读取当前数据类型的 Excel 列映射。
        columns = self.build_excel_columns(data_type)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], columns, output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{query_result['total']} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义广州医药数据查询方法。
    def query_gzyy_data(self, data_type, start_date, end_date):

        """
        查询广州医药指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 标准化查询结果。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 校验流向类数据的开始日期。
        if data_type in ["购进流向", "配送流向"] and not start_date:

            # 抛出查询开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验流向类数据的结束日期。
        if data_type in ["购进流向", "配送流向"] and not end_date:

            # 抛出查询结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 打印查询阶段标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date or "", "cyan"), self.ph.kv("结束日期", end_date or "", "cyan")))

        # 初始化当前页码。
        page_number = 1

        # 初始化累计数据行。
        all_rows = []

        # 初始化总记录数。
        total_count = None

        # 循环读取分页数据。
        while True:

            # 构建当前页查询请求体。
            payload = self.build_query_payload(data_type, page_number, self.page_size, start_date, end_date)

            # 打印分页查询日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"pageIndex={page_number}", "cyan"), self.ph.kv("pageSize", self.page_size, "cyan")))

            # 请求当前页数据。
            result = self.post_authorized_json(self.resolve_query_url(data_type), payload, self.resolve_referer(data_type))

            # 解析当前页响应。
            response_data = self.http.decode_json(result)

            # 校验当前页响应状态。
            self.validate_query_response(response_data, data_type)

            # 读取当前页数据块。
            data_block = self.extract_data_block(response_data)

            # 读取当前页数据行。
            page_rows = list(data_block.get("list") or data_block.get("records") or data_block.get("data") or [])

            # 读取响应总数。
            total_count = int(data_block.get("total") or len(page_rows))

            # 追加当前页数据行。
            all_rows.extend(page_rows)

            # 打印当前页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(page_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count, "cyan")))

            # 判断是否已经读取完整。
            if len(all_rows) >= total_count or not page_rows:

                # 结束分页循环。
                break

            # 递增页码。
            page_number += 1

        # 返回标准化查询结果。
        return {"rows": all_rows, "total": total_count or len(all_rows), "headers": self.build_excel_headers(data_type)}

    # 定义 Excel 表头构建方法。
    def build_excel_headers(self, data_type):

        """
        构建广州医药 Excel 中文表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: Excel 中文表头列表。
        :rtype: list[str]
        """

        # 返回列映射中的中文表头。
        return [column_name for column_name, _ in self.build_excel_columns(data_type)]

    # 定义 Excel 列映射构建方法。
    def build_excel_columns(self, data_type):

        """
        构建广州医药 Excel 列映射。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: Excel 中文表头与接口字段映射列表。
        :rtype: list[tuple[str, list[str]]]
        """

        # 返回配送流向列映射。
        if data_type == "配送流向":

            # 返回配送流向中文表头与接口字段。
            return [
                ("流向货主", ["ownerName"]),
                ("公司", ["companyOwnerName"]),
                ("单据类型", ["orderType", "flowType"]),
                ("销售日期", ["billDate", "lxDate"]),
                ("客户编码", ["customOwnedCode"]),
                ("客户名称", ["customOwnedName", "customMainName"]),
                ("货品编码", ["goodsOwnedId"]),
                ("货品名称", ["goodsMainName", "goodsOwnedName"]),
                ("货品规格", ["goodsMainSpec"]),
                ("生产厂商", ["goodsMainArea"]),
                ("数量", ["goodsqty"]),
                ("货品单位", ["goodsMainUnit"]),
                ("批号", ["lotNo"]),
                ("单价", ["realityPrice", "showPrice"]),
                ("金额", ["realityTotal", "totalLine"]),
                ("运输地址", ["shippingAddress"]),
                ("客户性质", ["office", "customerSegment"]),
                ("客户区域（省）", ["customProvince"]),
                ("客户区域（市）", ["customCity"]),
                ("货品状态*", ["goodsStatus"]),
                ("仓库名称", ["warehouse"]),
                ("失效日期", ["invalidDateStr"]),
                ("批准文号/注册证号-1", ["goodsMainApproveNo"]),
                ("批准文号/注册证号-2", ["goodsMainRgNumber", "approveRgNumber"]),
                ("处方属性", ["preseriptionFirstClass"]),
                ("客户平台ID", ["customMainId"]),
                ("货品平台编码", ["goodsMainMdmId"]),
                ("销售细单ID", ["saleOrderItemId"]),
                ("套装名称", ["bomName"]),
                ("单据删除标识", ["billDelTag"]),
                ("商品名", ["chinaApprovedDrugName"]),
                ("药材产地", ["medicineArea"]),
                ("联系人ID", ["lxrId"]),
                ("联系人", ["lxr"]),
                ("委托人ID", ["wtrId"]),
                ("委托人", ["wtr"]),
                ("颜色", ["color"]),
                ("味道", ["taste"]),
                ("包装", ["pack"]),
                ("供应商名称", ["supplierMainName"]),
                ("供应商MDM运营码", ["supplierOwnedId"]),
            ]

        # 返回购进流向列映射。
        if data_type == "购进流向":

            # 返回购进流向中文表头与接口字段。
            return [
                ("流向货主", ["ownerName"]),
                ("公司", ["companyOwnerName"]),
                ("单据类型", ["orderType"]),
                ("入库日期", ["billDate", "billDateStr"]),
                ("供应商编码", ["supplierOwnedId"]),
                ("供应商MDM运营码", ["supplierOwnedCode"]),
                ("供应商名称", ["supplierOwnedName", "sourceSupplierName", "supplierMainName"]),
                ("货品编码", ["goodsOwnedId"]),
                ("货品名称", ["goodsMainName"]),
                ("货品规格", ["goodsMainSpec"]),
                ("生产厂商", ["goodsMainArea"]),
                ("数量", ["goodsqty"]),
                ("货品单位", ["goodsMainUnit"]),
                ("批号", ["lotNo"]),
                ("失效日期", ["invalidDateStr"]),
                ("仓库名称", ["warehouse"]),
                ("批准文号/注册证号-1", ["goodsMainApproveNo"]),
                ("批准文号/注册证号-2", ["approveRgNumber", "goodsMainApproveNo"]),
                ("处方属性", ["preseriptionFirstClass"]),
                ("货品平台编码", ["goodsMainMdmId"]),
                ("供应商平台ID", ["supplierMainId"]),
                ("进货细单ID", ["sudocdtlId"]),
                ("套装名称", ["bomName"]),
                ("商品名", ["productName", "chinaApprovedDrugName"]),
                ("药材产地", ["medicineArea"]),
                ("颜色", ["color"]),
                ("味道", ["taste"]),
                ("包装", ["pack"]),
                ("联系人ID", ["lxrId"]),
                ("联系人", ["lxr"]),
                ("委托人ID", ["wtrId"]),
                ("委托人", ["wtr"]),
                ("金额", ["totalLine"]),
                ("生产日期", ["productionDateStr"]),
                ("单价", ["unitPrice"]),
            ]

        # 返回库存列映射。
        if data_type == "库存":

            # 返回库存中文表头与接口字段。
            return [
                ("流向货主", ["ownerName"]),
                ("库存日期", ["createdStr"]),
                ("货品编码", ["goodsOwnedId"]),
                ("货品名称", ["goodsName"]),
                ("货品规格", ["spec"]),
                ("生产厂商", ["factoryname"]),
                ("批号", ["lotNo"]),
                ("在库库存", ["stockQty"]),
                ("在途库存", ["intransitQty"]),
                ("数量", ["goodsqty"]),
                ("货品单位", ["unit"]),
                ("库存状态*", ["goodsStatus"]),
                ("仓库名称", ["warehouse"]),
                ("生产日期", ["productionDateStr"]),
                ("失效日期", ["invalidDateStr"]),
                ("批准文号", ["approvedocNo"]),
                ("注册证号", ["rgNumber"]),
                ("处方属性", ["preProperty"]),
                ("货源供应商", ["supplierMainName"]),
                ("委托人ID", ["wtrId"]),
                ("联系人ID", ["lxrId"]),
                ("委托人", ["wtr"]),
                ("联系人", ["lxr"]),
                ("货品MDM质量码", ["goodsMainMdmId"]),
                ("商品名", ["customizeGoodsName"]),
                ("药材产地", ["goodsPlace"]),
                ("颜色", ["color"]),
                ("味道", ["taste"]),
                ("包装", ["pack"]),
            ]

        # 抛出未接入类型异常。
        raise NotImplementedError(f"广州医药暂未接入数据类型：{data_type}")

    # 定义行字段读取方法。
    def read_row_value(self, row, field_names, column_name):

        """
        按字段优先级读取广州医药数据行。

        :param row: 查询结果数据行。
        :type row: dict
        :param field_names: 接口字段名列表。
        :type field_names: list[str]
        :param column_name: Excel 中文表头。
        :type column_name: str
        :return: Excel 单元格值。
        :rtype: object
        """

        # 遍历候选接口字段。
        for field_name in field_names:

            # 读取当前字段值。
            value = row.get(field_name)

            # 跳过空值字段。
            if value in [None, ""]:

                # 继续读取下一个候选字段。
                continue

            # 返回标准化后的字段值。
            return self.normalize_excel_value(column_name, value)

        # 返回空单元格值。
        return ""

    # 定义 Excel 值标准化方法。
    def normalize_excel_value(self, column_name, value):

        """
        标准化广州医药 Excel 单元格值。

        :param column_name: Excel 中文表头。
        :type column_name: str
        :param value: 原始接口字段值。
        :type value: object
        :return: 标准化后的单元格值。
        :rtype: object
        """

        # 标准化货品状态字段。
        if column_name in ["货品状态*", "库存状态*"]:

            # 返回页面导出使用的状态名称。
            return self.normalize_goods_status(value)

        # 标准化数值字段。
        if column_name in ["数量", "单价", "金额", "在库库存", "在途库存"]:

            # 返回数值型单元格。
            return self.normalize_number_value(value)

        # 返回原始字段值。
        return value

    # 定义货品状态标准化方法。
    def normalize_goods_status(self, value):

        """
        标准化广州医药货品状态名称。

        :param value: 原始货品状态。
        :type value: object
        :return: 标准化货品状态。
        :rtype: object
        """

        # 定义接口状态与页面状态映射。
        status_map = {"质量待定": "待检"}

        # 返回映射后的状态名称。
        return status_map.get(value, value)

    # 定义数值标准化方法。
    def normalize_number_value(self, value):

        """
        标准化广州医药数值字段。

        :param value: 原始数值字段。
        :type value: object
        :return: 数值字段。
        :rtype: int | float | object
        """

        # 保留原始数字类型。
        if isinstance(value, (int, float)):

            # 返回原始数字。
            return value

        # 捕获字符串数值转换异常。
        try:

            # 转换文本为浮点数。
            number_value = float(str(value).replace(",", ""))

            # 返回整数值。
            if number_value.is_integer():

                # 转换为整数类型。
                return int(number_value)

            # 返回浮点数值。
            return number_value

        # 处理非数值文本。
        except (TypeError, ValueError):

            # 返回原始字段值。
            return value

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验广州医药数据类型是否已接入。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 校验数据类型是否存在于接口映射。
        if data_type not in ["购进流向", "配送流向", "库存"]:

            # 抛出数据类型未接入异常。
            raise NotImplementedError(f"广州医药暂未接入数据类型：{data_type}")

    # 定义查询请求体构建方法。
    def build_query_payload(self, data_type, page_number, page_size, start_date, end_date):

        """
        构建广州医药分页查询请求体。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param page_number: 当前页码。
        :type page_number: int
        :param page_size: 每页记录数。
        :type page_size: int
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 分页查询请求体。
        :rtype: dict
        """

        # 读取数据类型对应的菜单配置。
        menu_config = self.resolve_menu_config(data_type)

        # 构建基础分页请求体。
        payload = {
            "current": page_number,
            "pageIndex": page_number,
            "pageSize": page_size,
            "size": page_size,
            "isOpenSub": 0,
            "menuId": menu_config["query_menu_id"],
            "pMenuId": menu_config["page_menu_id"],
            "viewer": "supplier",
            "paramList": [],
        }

        # 判断当前数据类型是否需要日期范围。
        if data_type in ["购进流向", "配送流向"]:

            # 写入业务日期范围筛选条件。
            payload["paramList"].append(self.build_date_param("bill_date", start_date, end_date))

        # 返回分页请求体。
        return payload

    # 定义日期条件构建方法。
    def build_date_param(self, field_name, start_date, end_date):

        """
        构建广州医药日期范围查询条件。

        :param field_name: 接口日期字段名。
        :type field_name: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 日期范围查询条件。
        :rtype: dict
        """

        # 返回日期范围条件。
        return {
            "condition": "等于",
            "paramField": field_name,
            "paramType": "datetime",
            "paramValue": f"{start_date} 00:00:00,{end_date} 23:59:59",
        }

    # 定义菜单配置解析方法。
    def resolve_menu_config(self, data_type):

        """
        解析广州医药数据类型对应的菜单权限配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 菜单权限配置。
        :rtype: dict
        """

        # 配置内部数据类型与页面菜单 ID 的映射。
        menu_map = {
            "购进流向": {"page_menu_id": "104", "query_menu_id": "134"},
            "配送流向": {"page_menu_id": "105", "query_menu_id": "136"},
            "库存": {"page_menu_id": "175", "query_menu_id": "181"},
        }

        # 返回当前数据类型对应的菜单配置。
        return menu_map[data_type]

    # 定义查询接口解析方法。
    def resolve_query_url(self, data_type):

        """
        解析广州医药数据类型对应的查询接口。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 查询接口地址。
        :rtype: str
        """

        # 返回购进流向查询接口。
        if data_type == "购进流向":

            # 返回进货入库查询接口地址。
            return self.purchase_query_url

        # 返回配送流向查询接口。
        if data_type == "配送流向":

            # 返回销售出库查询接口地址。
            return self.sale_query_url

        # 返回库存查询接口地址。
        return self.stock_query_url

    # 定义来源页面解析方法。
    def resolve_referer(self, data_type):

        """
        解析广州医药数据类型对应的来源页面。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 来源页面地址。
        :rtype: str
        """

        # 返回购进流向来源页面。
        if data_type == "购进流向":

            # 返回进货入库页面地址。
            return self.page_base_url + "/flowPlatform/dataQuery/goodsIn/_1779850140728"

        # 返回配送流向来源页面。
        if data_type == "配送流向":

            # 返回销售出库页面地址。
            return self.page_base_url + "/flowPlatform/dataQuery/pageSaleOutWholesale/_1779850140728"

        # 返回库存页面地址。
        return self.page_base_url + "/flowPlatform/dataQuery/immediate/_1779850140728"

    # 定义授权 JSON 请求方法。
    def post_authorized_json(self, url, payload, referer):

        """
        发送带广州医药认证令牌的 JSON POST 请求。

        :param url: 请求地址。
        :type url: str
        :param payload: 请求体。
        :type payload: dict
        :param referer: 来源页面地址。
        :type referer: str
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 编码 JSON 请求体。
        request_body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 构建授权请求头。
        headers = self.build_authorized_json_headers(referer)

        # 发送授权请求。
        return self.http.request("POST", url, headers=headers, data=request_body)

    # 定义授权请求头构建方法。
    def build_authorized_json_headers(self, referer):

        """
        构建广州医药授权 JSON 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 授权请求头。
        :rtype: dict
        """

        # 返回授权请求头。
        return {
            "Accept": "application/json",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            "Blade-Auth": self.login_token,
            "Connection": "keep-alive",
            "Content-Type": "application/json;charset=UTF-8",
            "Origin": self.page_base_url,
            "Referer": referer,
            "User-Agent": self.http.user_agent(),
        }

    # 定义查询响应校验方法。
    def validate_query_response(self, response_data, data_type):

        """
        校验广州医药查询响应状态。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 读取响应状态。
        status_value = response_data.get("status")

        # 读取响应编码。
        code_value = response_data.get("code")

        # 判断响应状态是否为成功。
        if status_value not in [0, True, 200] and code_value not in [0, 200]:

            # 抛出查询失败异常。
            raise RuntimeError(f"{self.company_name} {data_type}查询失败：{response_data.get('msg') or response_data.get('message') or response_data}")

    # 定义查询数据块提取方法。
    def extract_data_block(self, response_data):

        """
        提取广州医药查询响应中的数据块。

        :param response_data: 查询响应数据。
        :type response_data: dict
        :return: 查询数据块。
        :rtype: dict
        """

        # 读取响应数据字段。
        data_block = response_data.get("data")

        # 返回字典型数据块。
        if isinstance(data_block, dict):

            # 返回数据块。
            return data_block

        # 返回空数据块。
        return {"list": [], "total": 0}

    # 定义表头推断方法。
    def infer_headers(self, rows):

        """
        根据广州医药数据行推断 Excel 表头。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :return: Excel 表头。
        :rtype: list[str]
        """

        # 空数据返回空表头。
        if not rows:

            # 返回空列表。
            return []

        # 初始化表头列表。
        headers = []

        # 遍历全部数据行。
        for row in rows:

            # 遍历当前数据行字段。
            for key in row.keys():

                # 追加尚未出现的字段。
                if key not in headers:

                    # 保存字段名称。
                    headers.append(key)

        # 返回表头列表。
        return headers

    # 定义输出文件名构建方法。
    def build_output_file_name(self, data_type, start_date, end_date):

        """
        构建广州医药输出文件名。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 输出文件名。
        :rtype: str
        """

        # Delegate standard naming to the base platform.
        return super().build_output_file_name(data_type, start_date, end_date)

    # 定义流向月份标签构建方法。
    def build_month_label(self, start_date):

        """
        构建广州医药流向文件月份标签。

        :param start_date: 查询开始日期。
        :type start_date: str
        :return: 月份标签。
        :rtype: str
        """

        # 解析开始日期。
        date_value = datetime.strptime(start_date, "%Y-%m-%d")

        # 返回 yy.m月 标签。
        return f"{date_value.year % 100}.{date_value.month}月"

    # 定义库存日期标签构建方法。
    def build_stock_date_label(self, end_date):

        """
        构建广州医药库存文件日期标签。

        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 库存日期标签。
        :rtype: str
        """

        # 读取导出当天日期。
        date_value = datetime.now()

        # 返回 yy.m.d 标签。
        return f"{date_value.year % 100}.{date_value.month}.{date_value.day}"

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, columns, output_path, sheet_name):

        """
        写入广州医药查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param columns: Excel 中文表头与接口字段映射列表。
        :type columns: list[tuple[str, list[str]]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 提取中文表头列表。
        headers = [column_name for column_name, _ in columns]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按列映射顺序写入当前数据行。
            worksheet.append([self.read_row_value(row, field_names, column_name) for column_name, field_names in columns])

        # 遍历所有导出列设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 捕获文件保存异常。
        try:

            # 保存 Excel 文件。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

    # 定义同名文件删除方法。
    def remove_existing_file(self, output_path):

        """
        删除广州医药已有同名输出文件。

        :param output_path: 输出文件路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 判断同名文件是否存在。
        if not os.path.exists(output_path):

            # 文件不存在时直接返回。
            return

        # 捕获删除文件时的占用异常。
        try:

            # 删除已有同名文件。
            os.remove(output_path)

        # 处理文件被占用导致的删除异常。
        except PermissionError as error:

            # 抛出带业务路径的文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

    # 定义站点根地址构建方法。
    def build_origin(self, url):

        """
        根据完整地址构建站点根地址。

        :param url: 完整地址。
        :type url: str
        :return: 站点根地址。
        :rtype: str
        """

        # 解析 URL 地址。
        parsed_url = urlparse(url)

        # 返回协议与主机组成的根地址。
        return f"{parsed_url.scheme}://{parsed_url.netloc}"
