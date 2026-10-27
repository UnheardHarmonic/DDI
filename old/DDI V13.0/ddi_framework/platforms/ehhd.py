# -*- coding: utf-8 -*-
"""
恩华和达经典 ASP 查询系统协议适配器。

本模块负责恩华和达查询系统的协议登录、模式页面访问、HTML 表格解析和本地 Excel 生成。
"""

# 导入 HTML 实体解码函数。
from html import unescape

# 导入 HTML 标准解析器。
from html.parser import HTMLParser

# 导入操作系统路径模块。
import os

# 导入正则表达式模块。
import re

# 导入日期时间解析对象。
from datetime import datetime

# 导入 Cookie 对象。
from http.cookiejar import Cookie

# 导入 URL 解析函数。
from urllib.parse import urljoin

# 导入 URL 解析对象。
from urllib.parse import urlparse

# 导入 URL 表单编码函数。
from urllib.parse import urlencode

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient


# 定义恩华和达系统页面表格解析器。
class EhhdTableParser(HTMLParser):

    """
    恩华和达系统 HTML 表格解析器。

    :return: HTML 表格解析器实例。
    :rtype: EhhdTableParser
    """

    # 定义解析器初始化方法。
    def __init__(self):

        """
        初始化 HTML 表格解析器。

        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(convert_charrefs=True)

        # 初始化表格集合。
        self.tables = []

        # 初始化当前表格。
        self.current_table = None

        # 初始化当前行。
        self.current_row = None

        # 初始化当前单元格。
        self.current_cell = None

        # 初始化表格嵌套深度。
        self.table_depth = 0

    # 定义开始标签处理方法。
    def handle_starttag(self, tag, attrs):

        """
        处理 HTML 开始标签。

        :param tag: 标签名称。
        :type tag: str
        :param attrs: 标签属性列表。
        :type attrs: list[tuple[str, str | None]]
        :return: None。
        :rtype: None
        """

        # 转换标签名称为小写。
        tag_name = tag.lower()

        # 判断当前标签是否为表格。
        if tag_name == "table":

            # 增加当前表格深度。
            self.table_depth += 1

            # 判断当前是否进入顶层表格。
            if self.table_depth == 1:

                # 初始化当前表格。
                self.current_table = []

        # 判断当前标签是否为表格行。
        if tag_name == "tr" and self.current_table is not None:

            # 初始化当前行缓存。
            self.current_row = []

        # 判断当前标签是否为表格单元格。
        if tag_name in ("td", "th") and self.current_row is not None:

            # 初始化当前单元格缓存。
            self.current_cell = []

        # 判断当前标签是否为换行标签。
        if tag_name == "br" and self.current_cell is not None:

            # 追加换行占位符。
            self.current_cell.append("\n")

    # 定义结束标签处理方法。
    def handle_endtag(self, tag):

        """
        处理 HTML 结束标签。

        :param tag: 标签名称。
        :type tag: str
        :return: None。
        :rtype: None
        """

        # 转换标签名称为小写。
        tag_name = tag.lower()

        # 判断当前结束标签是否为表格单元格。
        if tag_name in ("td", "th"):

            # 刷新当前单元格内容。
            self.flush_current_cell()

        # 判断当前结束标签是否为表格行。
        if tag_name == "tr":

            # 刷新当前行内容。
            self.flush_current_row()

        # 判断当前结束标签是否为表格。
        if tag_name == "table" and self.table_depth:

            # 判断当前是否离开顶层表格。
            if self.table_depth == 1:

                # 刷新当前表格内容。
                self.flush_current_table()

            # 减少当前表格深度。
            self.table_depth -= 1

    # 定义文本处理方法。
    def handle_data(self, data):

        """
        处理 HTML 文本。

        :param data: 文本内容。
        :type data: str
        :return: None。
        :rtype: None
        """

        # 判断当前是否正在采集单元格文本。
        if self.current_cell is not None:

            # 追加当前文本内容。
            self.current_cell.append(data)

    # 定义单元格文本清理方法。
    def clean_cell_text(self, value):

        """
        清理表格单元格文本。

        :param value: 原始单元格文本。
        :type value: object
        :return: 清理后的单元格文本。
        :rtype: str
        """

        # 解码 HTML 实体。
        decoded_text = unescape(str(value))

        # 替换不间断空格。
        normalized_text = decoded_text.replace("\xa0", " ")

        # 压缩连续空白字符。
        compact_text = " ".join(normalized_text.split())

        # 返回清理后的文本。
        return compact_text.strip()

    # 定义当前单元格刷新方法。
    def flush_current_cell(self):

        """
        刷新当前单元格缓存。

        :return: None。
        :rtype: None
        """

        # 判断当前单元格或当前行是否不可用。
        if self.current_cell is None or self.current_row is None:

            # 当前无可刷新的单元格时直接返回。
            return

        # 拼接当前单元格文本。
        cell_text = "".join(self.current_cell)

        # 清理当前单元格文本。
        clean_text = self.clean_cell_text(cell_text)

        # 保存当前单元格文本。
        self.current_row.append(clean_text)

        # 清理当前单元格缓存。
        self.current_cell = None

    # 定义当前表格行刷新方法。
    def flush_current_row(self):

        """
        刷新当前表格行缓存。

        :return: None。
        :rtype: None
        """

        # 刷新当前行中未关闭的单元格。
        self.flush_current_cell()

        # 判断当前行或当前表格是否不可用。
        if self.current_row is None or self.current_table is None:

            # 当前无可刷新的表格行时直接返回。
            return

        # 判断当前行是否存在有效文本。
        if any(cell for cell in self.current_row):

            # 保存当前表格行。
            self.current_table.append(self.current_row)

        # 清理当前行缓存。
        self.current_row = None

    # 定义当前表格刷新方法。
    def flush_current_table(self):

        """
        刷新当前表格缓存。

        :return: None。
        :rtype: None
        """

        # 刷新当前表格中未关闭的表格行。
        self.flush_current_row()

        # 判断当前表格是否不可用。
        if self.current_table is None:

            # 当前无可刷新的表格时直接返回。
            return

        # 判断当前表格是否包含有效行。
        if self.current_table:

            # 保存当前表格。
            self.tables.append(self.current_table)

        # 清理当前表格缓存。
        self.current_table = None

    # 定义解析结束收尾方法。
    def finalize(self):

        """
        对残缺 HTML 进行收尾处理。

        :return: None。
        :rtype: None
        """

        # 刷新当前未闭合的顶层表格。
        self.flush_current_table()

        # 清理表格深度状态。
        self.table_depth = 0


# 定义恩华和达经典 ASP 查询系统协议适配器。
class EhhdPlatform(BasePlatform):

    """
    恩华和达经典 ASP 查询系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 恩华和达系统协议适配器实例。
    :rtype: EhhdPlatform
    """

    # 定义平台初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化恩华和达系统协议适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "http://58.218.188.74:88/Login.asp")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", self.login_page_url)

        # 保存首页地址。
        self.home_page_url = account.get("首页", "http://58.218.188.74:88/default.asp")

        # 保存平台根地址。
        self.base_url = account.get("接口根地址", self.build_origin(self.login_page_url))

        # 保存表单编码。
        self.form_encoding = account.get("表单编码", "gbk")

        # 保存响应编码。
        self.response_encoding = account.get("响应编码", "gbk")

        # 保存最大查询天数。
        self.max_query_days = int(account.get("最大查询天数", 90))

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化登录状态。
        self.logged_in = False

    # 定义平台登录方法。
    def login(self):

        """
        执行恩华和达系统协议登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("恩华和达登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.join(self.ph.kv("访问登录页", self.login_page_url, "cyan"), self.ph.kv("账号", self.clean_login_name(), "cyan")))

        # 访问登录页以建立站点会话 Cookie。
        self.http.get(self.login_page_url, referer=self.base_url)

        # 构建登录请求体。
        payload = self.build_login_payload()

        # 构建登录请求头。
        headers = self.build_form_headers(self.login_page_url)

        # 编码登录请求体。
        data = self.encode_form_payload(payload)

        # 发送登录请求。
        result = self.http.request("POST", self.login_api_url, headers=headers, data=data, timeout=60)

        # 解码登录响应文本。
        response_text = self.decode_text(result)

        # 校验登录响应文本。
        self.validate_login_response(response_text)

        # 标记当前登录状态。
        self.logged_in = True

        # 打印登录成功日志。
        self.ph.success(self.ph.join(self.ph.kv("登录状态", "成功", "green"), self.ph.kv("Cookie", len(list(self.http.cookie_jar)), "cyan")))

        # 返回登录结果。
        return {"登录状态": "成功", "Cookie数量": len(list(self.http.cookie_jar))}

    # 定义登录请求体构建方法。
    def build_login_payload(self):

        """
        构建恩华和达系统登录请求体。

        :return: 登录请求体。
        :rtype: dict
        """

        # 返回登录表单字段。
        return {"username": self.clean_login_name(), "password": self.clean_login_password()}

    # 定义登录响应校验方法。
    def validate_login_response(self, response_text):

        """
        校验恩华和达系统登录响应。

        :param response_text: 登录响应文本。
        :type response_text: str
        :return: None。
        :rtype: None
        """

        # 判断响应文本是否包含登录失败提示。
        if "账号或密码错误" in response_text:

            # 抛出登录失败异常。
            raise RuntimeError(f"{self.company_name} 登录失败：账号或密码错误")

        # 判断响应文本是否包含欢迎语。
        if "欢迎您：" not in response_text:

            # 抛出登录异常。
            raise RuntimeError(f"{self.company_name} 登录后未返回有效系统首页")

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 打印 Cookie 保存日志。
        self.ph.success(self.ph.kv("登录 Cookie 已保存", self.ph.path(cookie_path), "green"))

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义平台导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        导出指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 使用查询结果生成本地 Excel。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义平台查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询指定类型业务数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果行。
        :rtype: list[dict]
        """

        # 查询恩华和达系统指定类型数据。
        query_result = self.query_ehhd_data(data_type, start_date, end_date)

        # 返回查询结果行。
        return query_result["rows"]

    # 定义查询并生成 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询指定类型业务数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询恩华和达系统指定类型数据。
        query_result = self.query_ehhd_data(data_type, start_date, end_date)

        # 读取默认输出目录。
        output_dir = output_dir or os.getcwd()

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建输出文件名。
        file_name = self.build_output_file_name(data_type, query_result["start_date"], query_result["end_date"])

        # 构建输出文件路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan"), self.ph.kv("字段数", len(query_result["headers"]), "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义主查询方法。
    def query_ehhd_data(self, data_type, start_date, end_date):

        """
        查询恩华和达系统指定类型数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果、表头和日期范围。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 确保当前会话已经登录。
        self.ensure_login()

        # 解析查询日期范围。
        start_date, end_date = self.resolve_query_dates(start_date, end_date)

        # 读取当前数据类型页面配置。
        page_config = self.resolve_page_config(data_type)

        # 判断当前数据类型是否需要校验日期跨度。
        if page_config.get("需要日期", True):

            # 校验查询日期范围。
            self.validate_query_range(start_date, end_date)

        # 打印查询阶段标题。
        self.print_query_section(data_type, start_date, end_date)

        # 构建查询页面地址。
        query_url = self.build_page_url(page_config["查询路径"])

        # 构建查询请求体。
        payload = self.build_query_payload(page_config, start_date, end_date)

        # 构建查询请求头。
        headers = self.build_form_headers(query_url)

        # 编码查询请求体。
        data = self.encode_form_payload(payload)

        # 发送查询请求。
        result = self.http.request("POST", query_url, headers=headers, data=data, timeout=180)

        # 解码查询响应文本。
        html_text = self.decode_text(result)

        # 校验查询响应文本。
        self.validate_query_response(data_type, html_text)

        # 解析页面中的全部表格。
        tables = self.parse_tables(html_text)

        # 选择结果数据表格。
        data_table = self.select_data_table(data_type, tables)

        # 构建表头与数据行。
        headers_row, rows = self.build_rows_from_table(data_table)

        # 校验结果日期范围。
        self.validate_result_date_range(data_type, rows, start_date, end_date)

        # 打印查询结果日志。
        self.ph.success(self.ph.join(self.ph.kv("结果表格", page_config["查询路径"], "green"), self.ph.kv("记录数", f"{len(rows)} 条", "cyan"), self.ph.kv("字段数", len(headers_row), "cyan")))

        # 返回查询结果。
        return {"headers": headers_row, "rows": rows, "start_date": start_date, "end_date": end_date}

    # 定义页面配置解析方法。
    def resolve_page_config(self, data_type):

        """
        解析指定数据类型的页面配置。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: 页面配置。
        :rtype: dict
        """

        # 返回当前数据类型的页面配置。
        return dict(self.account.get("页面映射", {}).get(data_type, {}))

    # 定义页面 URL 构建方法。
    def build_page_url(self, page_path):

        """
        构建页面访问地址。

        :param page_path: 页面相对路径。
        :type page_path: str
        :return: 页面访问地址。
        :rtype: str
        """

        # 返回拼接后的页面地址。
        return urljoin(f"{self.base_url}/", str(page_path).lstrip("/"))

    # 定义查询请求体构建方法。
    def build_query_payload(self, page_config, start_date, end_date):

        """
        构建页面查询请求体。

        :param page_config: 页面配置。
        :type page_config: dict
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 查询请求体。
        :rtype: dict
        """

        # 初始化查询请求体。
        payload = {
            page_config.get("商品编号字段", "sp"): "",
            page_config.get("厂家字段", "sc"): "",
        }

        # 判断当前页面是否需要日期参数。
        if page_config.get("需要日期", True):

            # 写入开始日期参数。
            payload[page_config.get("开始日期字段", "n1")] = start_date

            # 写入结束日期参数。
            payload[page_config.get("结束日期字段", "n2")] = end_date

        # 返回查询请求体。
        return payload

    # 定义表单编码方法。
    def encode_form_payload(self, payload):

        """
        使用指定编码构建表单请求体。

        :param payload: 表单字段字典。
        :type payload: dict
        :return: 编码后的表单字节串。
        :rtype: bytes
        """

        # 按指定编码构建 URL 编码文本。
        encoded_text = urlencode(payload, encoding=self.form_encoding, errors="strict")

        # 返回 ASCII 字节串。
        return encoded_text.encode("ascii")

    # 定义 HTML 表格解析方法。
    def parse_tables(self, html_text):

        """
        解析页面中的全部顶层 HTML 表格。

        :param html_text: 页面 HTML 文本。
        :type html_text: str
        :return: 表格集合。
        :rtype: list[list[list[str]]]
        """

        # 创建 HTML 表格解析器。
        parser = EhhdTableParser()

        # 解析当前页面 HTML。
        parser.feed(str(html_text or ""))

        # 收尾处理残缺 HTML。
        parser.finalize()

        # 返回解析后的表格集合。
        return parser.tables

    # 定义结果表格选择方法。
    def select_data_table(self, data_type, tables):

        """
        从页面表格集合中选择结果数据表。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param tables: 页面表格集合。
        :type tables: list[list[list[str]]]
        :return: 结果数据表。
        :rtype: list[list[str]]
        """

        # 判断页面表格集合是否为空。
        if not tables:

            # 抛出结果表格缺失异常。
            raise RuntimeError(f"{self.company_name} {data_type} 页面未返回可解析的表格")

        # 初始化候选表格列表。
        candidate_tables = [table for table in tables if len(table) >= 2]

        # 判断候选表格列表是否为空。
        if not candidate_tables:

            # 抛出结果表格缺失异常。
            raise RuntimeError(f"{self.company_name} {data_type} 页面未返回有效结果表格")

        # 返回行数最多的候选表格。
        return max(candidate_tables, key=lambda item: (len(item), max((len(row) for row in item), default=0)))

    # 定义表头与数据行构建方法。
    def build_rows_from_table(self, table_rows):

        """
        根据表格行构建表头与数据行。

        :param table_rows: 表格行集合。
        :type table_rows: list[list[str]]
        :return: 表头与数据行。
        :rtype: tuple[list[str], list[dict]]
        """

        # 判断表格行是否为空。
        if not table_rows:

            # 返回空表头和空数据。
            return [], []

        # 定位结果表头所在行。
        header_index = self.locate_header_row_index(table_rows)

        # 读取原始表头行。
        raw_headers = table_rows[header_index]

        # 标准化结果表头。
        headers = self.normalize_headers(raw_headers)

        # 初始化结果数据行集合。
        rows = []

        # 遍历表头之后的数据行。
        for raw_row in table_rows[header_index + 1:]:

            # 判断当前行是否为空行。
            if not any(str(cell).strip() for cell in raw_row):

                # 空行直接进入下一项。
                continue

            # 按表头顺序重建当前数据行。
            row = {header: raw_row[index] if index < len(raw_row) else "" for index, header in enumerate(headers)}

            # 保存当前数据行。
            rows.append(row)

        # 返回表头与数据行。
        return headers, rows

    # 定义结果表头定位方法。
    def locate_header_row_index(self, table_rows):

        """
        定位结果表头所在的行索引。

        :param table_rows: 表格行集合。
        :type table_rows: list[list[str]]
        :return: 表头行索引。
        :rtype: int
        """

        # 读取当前表格最大列数。
        max_columns = max((len(row) for row in table_rows), default=0)

        # 遍历全部表格行。
        for index, row in enumerate(table_rows):

            # 判断当前行列数是否小于最大列数。
            if len(row) < max_columns:

                # 当前行不是结果表头时进入下一项。
                continue

            # 统计当前行非空单元格数量。
            filled_cells = sum(1 for cell in row if str(cell).strip())

            # 判断当前行是否具备结果表头特征。
            if filled_cells >= 3:

                # 返回当前表头行索引。
                return index

        # 返回默认首行索引。
        return 0

    # 定义表头标准化方法。
    def normalize_headers(self, headers):

        """
        标准化表头并处理重名字段。

        :param headers: 原始表头列表。
        :type headers: list[str]
        :return: 标准化表头列表。
        :rtype: list[str]
        """

        # 初始化字段出现次数字典。
        name_counts = {}

        # 初始化标准化表头列表。
        normalized_headers = []

        # 遍历原始表头。
        for index, header in enumerate(headers, start=1):

            # 读取清洗后的字段名。
            field_name = str(header or "").strip() or f"字段{index}"

            # 递增当前字段出现次数。
            name_counts[field_name] = name_counts.get(field_name, 0) + 1

            # 判断当前字段是否首次出现。
            if name_counts[field_name] == 1:

                # 保存当前字段名。
                normalized_headers.append(field_name)

                # 进入下一个字段。
                continue

            # 保存带序号的字段名。
            normalized_headers.append(f"{field_name}_{name_counts[field_name]}")

        # 返回标准化表头列表。
        return normalized_headers

    # 定义登录状态确保方法。
    def ensure_login(self):

        """
        确保当前适配器已经完成登录。

        :return: None。
        :rtype: None
        """

        # 判断当前是否已经登录。
        if self.logged_in:

            # 已登录时直接返回。
            return

        # 执行登录流程。
        self.login()

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验当前数据类型是否已接入。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 判断当前数据类型是否存在页面映射。
        if data_type in self.account.get("页面映射", {}):

            # 已接入时直接返回。
            return

        # 抛出数据类型未接入异常。
        raise ValueError(f"{self.company_name} 暂不支持数据类型：{data_type}")

    # 定义日期范围校验方法。
    def validate_query_range(self, start_date, end_date):

        """
        校验查询日期范围。

        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 解析开始日期。
        start_value = datetime.strptime(start_date, "%Y-%m-%d")

        # 解析结束日期。
        end_value = datetime.strptime(end_date, "%Y-%m-%d")

        # 判断开始日期是否晚于结束日期。
        if start_value > end_value:

            # 抛出日期范围异常。
            raise ValueError(f"{self.company_name} 查询开始日期不能晚于结束日期：{start_date} > {end_date}")

        # 判断查询跨度是否超过平台限制。
        if (end_value - start_value).days > self.max_query_days:

            # 抛出日期跨度异常。
            raise ValueError(f"{self.company_name} 单次查询范围不能超过 {self.max_query_days} 天")

    # 定义查询日期解析方法。
    def resolve_query_dates(self, start_date, end_date):

        """
        解析查询日期范围。

        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询开始日期与结束日期。
        :rtype: tuple[str, str]
        """

        # 读取当天日期文本。
        today_text = datetime.now().strftime("%Y-%m-%d")

        # 返回解析后的查询日期范围。
        return start_date or today_text, end_date or today_text

    # 定义查询标题打印方法。
    def print_query_section(self, data_type, start_date, end_date):

        """
        打印查询阶段标题。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 打印查询阶段标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date, "cyan"), self.ph.kv("结束日期", end_date, "cyan")))

    # 定义查询响应校验方法。
    def validate_query_response(self, data_type, html_text):

        """
        校验查询页面响应文本。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param html_text: 页面 HTML 文本。
        :type html_text: str
        :return: None。
        :rtype: None
        """

        # 判断页面是否跳回登录页。
        if "网络超时或您还没有登陆" in html_text or "window.location.href='Login.asp'" in html_text:

            # 抛出登录态失效异常。
            raise RuntimeError(f"{self.company_name} {data_type} 查询时登录态失效")

        # 判断页面是否包含欢迎语。
        if "欢迎您：" not in html_text:

            # 抛出系统页面异常。
            raise RuntimeError(f"{self.company_name} {data_type} 查询未返回有效系统页面")

    # 定义结果日期范围校验方法。
    def validate_result_date_range(self, data_type, rows, start_date, end_date):

        """
        校验流向结果中的业务日期范围。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果行。
        :type rows: list[dict]
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 判断当前数据类型是否为库存。
        if data_type == "库存":

            # 库存数据不校验日期范围。
            return

        # 判断结果行是否为空。
        if not rows:

            # 空结果直接返回。
            return

        # 初始化结果日期集合。
        result_dates = []

        # 遍历全部结果行。
        for row in rows:

            # 读取当前行日期文本。
            date_text = str(row.get("日期", "")).strip()

            # 判断当前行日期文本是否为空。
            if not date_text:

                # 空日期直接进入下一行。
                continue

            try:

                # 解析当前行日期。
                result_dates.append(datetime.strptime(date_text, "%Y%m%d"))

            except ValueError:

                # 非标准日期时直接进入下一项。
                continue

        # 判断结果日期集合是否为空。
        if not result_dates:

            # 无有效日期时直接返回。
            return

        # 解析查询开始日期。
        start_value = datetime.strptime(start_date, "%Y-%m-%d")

        # 解析查询结束日期。
        end_value = datetime.strptime(end_date, "%Y-%m-%d")

        # 读取结果日期最小值。
        min_date = min(result_dates)

        # 读取结果日期最大值。
        max_date = max(result_dates)

        # 判断结果日期是否超出查询范围。
        if min_date < start_value or max_date > end_value:

            # 抛出结果日期越界异常。
            raise RuntimeError(f"{self.company_name} {data_type} 结果日期超出查询范围：{min_date:%Y-%m-%d} ~ {max_date:%Y-%m-%d}")

        # 打印结果日期范围日志。
        self.ph.info(self.ph.join(self.ph.kv("结果日期最小值", min_date.strftime("%Y-%m-%d"), "cyan"), self.ph.kv("结果日期最大值", max_date.strftime("%Y-%m-%d"), "cyan")))

    # 定义文本解码方法。
    def decode_text(self, result):

        """
        按平台编码解码响应文本。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 解码后的文本。
        :rtype: str
        """

        # 使用平台编码解码响应体。
        return result.body.decode(self.response_encoding, errors="replace")

    # 定义 HTML 请求头构建方法。
    def build_html_headers(self, referer):

        """
        构建 HTML 请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回 HTML 请求头。
        return {"Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Referer": referer, "User-Agent": self.http.user_agent()}

    # 定义表单请求头构建方法。
    def build_form_headers(self, referer):

        """
        构建表单请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回表单请求头。
        return {"Accept": "*/*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Content-Type": "application/x-www-form-urlencoded", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent()}

    # 定义输出文件名构建方法。
    def build_output_file_name(self, data_type, start_date, end_date):

        """
        构建输出文件名。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 输出文件名。
        :rtype: str
        """

        # Delegate standard naming to the base platform.
        return super().build_output_file_name(data_type, start_date, end_date)

    # 定义月份标签构建方法。
    def build_month_label(self, start_date):

        """
        构建流向文件月份标签。

        :param start_date: 查询开始日期。
        :type start_date: str
        :return: 月份标签。
        :rtype: str
        """

        # 解析查询开始日期。
        date_value = datetime.strptime(start_date, "%Y-%m-%d")

        # 返回 yy.m月 标签。
        return f"{date_value.year % 100}.{date_value.month}月"

    # 定义库存日期标签构建方法。
    def build_stock_date_label(self):

        """
        构建库存文件日期标签。

        :return: 日期标签。
        :rtype: str
        """

        # 读取当前日期。
        date_value = datetime.now()

        # 返回 yy.m.d 标签。
        return f"{date_value.year % 100}.{date_value.month}.{date_value.day}"

    # 定义 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头列表。
        :type headers: list[str]
        :param output_path: Excel 输出路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append(headers)

        # 遍历结果数据行。
        for row in rows:

            # 按表头顺序写入当前行。
            worksheet.append([row.get(header, "") for header in headers])

        # 遍历表头列设置列宽。
        for column_index, header_name in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 尝试保存 Excel 文件。
        try:

            # 保存当前工作簿。
            workbook.save(output_path)

        except PermissionError as error:

            # 抛出文件占用异常。
            raise PermissionError(f"无法覆盖文件，请先关闭已打开的 Excel 文件：{output_path}") from error

    # 定义已有文件删除方法。
    def remove_existing_file(self, output_path):

        """
        删除已有同名输出文件。

        :param output_path: 输出文件路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 判断输出文件是否存在。
        if not os.path.exists(output_path):

            # 文件不存在时直接返回。
            return

        # 删除已有输出文件。
        os.remove(output_path)

    # 定义登录账号读取方法。
    def clean_login_name(self):

        """
        读取并清洗登录账号。

        :return: 登录账号。
        :rtype: str
        """

        # 返回去除首尾空白的账号。
        return str(self.account.get("账号", "")).strip()

    # 定义登录密码读取方法。
    def clean_login_password(self):

        """
        读取并清洗登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回原始密码字符串。
        return str(self.account.get("密码", ""))

    # 定义站点源地址构建方法。
    def build_origin(self, url):

        """
        构建站点源地址。

        :param url: 页面地址。
        :type url: str
        :return: 站点源地址。
        :rtype: str
        """

        # 解析当前 URL。
        parsed_url = urlparse(url)

        # 读取端口片段。
        port_part = f":{parsed_url.port}" if parsed_url.port else ""

        # 返回站点源地址。
        return f"{parsed_url.scheme}://{parsed_url.hostname}{port_part}"

    # 定义 Cookie 写入方法。
    def add_cookie(self, name, value, path="/"):

        """
        向当前会话写入 Cookie。

        :param name: Cookie 名称。
        :type name: str
        :param value: Cookie 值。
        :type value: str
        :param path: Cookie 路径。
        :type path: str
        :return: None。
        :rtype: None
        """

        # 解析接口主机域名。
        domain = urlparse(self.base_url).hostname or ""

        # 创建 Cookie 对象。
        cookie = Cookie(0, name, str(value), None, False, domain, True, False, path, True, False, None, True, None, None, {}, False)

        # 写入当前 Cookie 容器。
        self.http.cookie_jar.set_cookie(cookie)
