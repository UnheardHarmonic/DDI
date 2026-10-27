# -*- coding: utf-8 -*-
"""
四川本草堂药业有限公司协同办公系统协议适配器。

本模块通过图片验证码登录、Authorization 会话令牌和分页 JSON 接口采集购进流向、配送流向与库存数据。

:return: 模块提供四川本草堂药业平台适配器。
:rtype: None
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入日期时间处理对象。
from datetime import datetime

# 导入 URL 查询字符串编码函数。
from urllib.parse import urlencode

# 导入 Excel 工作簿对象。
from openpyxl import Workbook

# 导入 Excel 列标识转换函数。
from openpyxl.utils import get_column_letter

# 导入平台适配器基类。
from ddi_framework import BasePlatform

# 导入验证码识别器。
from ddi_framework import CaptchaSolver

# 导入 Excel 统一样式函数。
from ddi_framework import apply_standard_excel_style

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入日志打印器。
from ddi_framework import PrintHandler


# 定义数据类型到平台接口的协议映射。
DATA_TYPE_CONFIG = {
    "购进流向": {"endpoint": "/api/flowimportsu/index", "date_field": "cofirmdate"},
    "配送流向": {"endpoint": "/api/flowimportsa/index", "date_field": "cofirmdate"},
    "库存": {"endpoint": "/api/flowimportst/index", "date_field": ""},
}

# 定义标准字段与原始字段的前置映射。
STANDARD_FIELD_MAP = {
    "购进流向": [
        ("购进日期", "cofirmdate"),
        ("供货商ID", "supplyid"),
        ("入库单号", "sudocid"),
        ("商品名称", "goodsname"),
        ("规格", "goodstype"),
        ("生产厂家", "factoryname"),
        ("批号", "lotno"),
        ("数量", "qty"),
        ("单价", "price"),
        ("客户名称", "companyname"),
        ("有效期", "invaliddate"),
    ],
    "配送流向": [
        ("配送日期", "cofirmdate"),
        ("客户ID", "customid"),
        ("销售单号", "erpid"),
        ("商品名称", "goodsname"),
        ("规格", "goodstype"),
        ("生产厂家", "factoryname"),
        ("批号", "lotno"),
        ("数量", "qty"),
        ("单价", "price"),
        ("客户名称", "customname"),
        ("有效期", "invaliddate"),
    ],
    "库存": [
        ("库存日期", "createTime"),
        ("商品名称", "goodsname"),
        ("规格", "goodstype"),
        ("生产厂家", "factoryname"),
        ("批号", "loton"),
        ("库存数量", "qty"),
        ("批次价", "unitprice"),
        ("供货单位", "companyname"),
        ("有效期", "invaliddate"),
    ],
}

# 定义用户界面数据类型别名。
DATA_TYPE_ALIASES = {"销售数据": "配送流向", "购进数据": "购进流向", "库存数据": "库存"}


# 定义四川本草堂药业协议适配器。
class ScbctPlatform(BasePlatform):
    """
    四川本草堂药业有限公司协同办公系统协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 四川本草堂药业平台适配器实例。
    :rtype: ScbctPlatform
    """

    # 定义平台适配器初始化方法。
    def __init__(self, company_name, account, printer=None):
        """
        初始化 HTTP 会话、验证码处理器和平台接口地址。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 初始化平台基类。
        super().__init__(company_name, account, printer=printer)

        # 初始化公共日志打印器。
        self.ph = printer or PrintHandler()

        # 读取平台接口根地址。
        self.base_url = str(account.get("接口根地址", "https://flow.scbct.com")).rstrip("/")

        # 读取登录页面地址。
        self.login_page_url = str(account.get("登录页", f"{self.base_url}/login"))

        # 读取验证码接口地址。
        self.captcha_api_url = str(account.get("验证码接口", f"{self.base_url}/api/login/captcha"))

        # 读取登录接口地址。
        self.login_api_url = str(account.get("登录接口", f"{self.base_url}/api/login/login"))

        # 读取单页查询数量并限制在平台允许范围内。
        self.page_size = min(int(account.get("分页大小", 400)), 400)

        # 读取验证码重试次数。
        self.captcha_retries = int(account.get("验证码重试次数", 8))

        # 初始化带 Cookie 管理的协议客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化本地验证码识别器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录令牌。
        self.login_token = ""

        # 初始化登录状态。
        self.logged_in = False

    # 定义登录方法。
    def login(self):
        """
        获取图片验证码并建立 Authorization 登录会话。

        :return: 不包含令牌正文的登录摘要。
        :rtype: dict
        """

        # 打印登录阶段分隔线。
        self.ph.section("四川本草堂药业协同办公系统登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 读取登录账号。
        username = str(self.account.get("账号", "")).strip()

        # 读取登录密码。
        password = str(self.account.get("密码", ""))

        # 校验账号密码配置完整。
        if not username or not password:
            raise RuntimeError("登录失败：未配置账号或密码。")

        # 保存最近一次登录失败原因。
        last_error = ""

        # 执行验证码识别和登录重试。
        for attempt in range(1, self.captcha_retries + 1):
            try:
                # 获取验证码上下文。
                captcha = self._get_captcha()

                # 解码验证码图片。
                image_bytes = self.captcha_solver.decode_base64_image(captcha["captcha"])

                # 识别验证码文本。
                raw_code = self.captcha_solver.solve_raw_text_image_with_ranges(image_bytes, "0123456789abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ")

                # 清洗验证码文本。
                captcha_code = "".join(char for char in str(raw_code or "") if char.isalnum())

                # 校验验证码识别结果长度。
                if len(captcha_code) < 4:
                    raise RuntimeError("图片验证码识别结果长度不足。")

                # 构建登录请求体。
                payload = {"username": username, "password": password, "captcha": captcha_code, "key": captcha["key"], "remember": True}

                # 提交登录请求。
                result = self.http.post_json(self.login_api_url, payload, self.login_page_url)

                # 解析登录响应。
                response_data = self._decode_json(result.body, "登录")

                # 判断登录接口返回状态。
                if str(response_data.get("code", "")) == "0":
                    # 提取登录响应中的访问令牌。
                    self.login_token = str((response_data.get("data") or {}).get("access_token") or "").strip()

                    # 校验访问令牌已返回。
                    if not self.login_token:
                        raise RuntimeError("登录成功但服务端未返回 access_token。")

                    # 标记当前会话已经登录。
                    self.logged_in = True

                    # 打印脱敏登录成功日志。
                    self.ph.success(self.ph.join(self.ph.kv("登录状态", "成功", "green"), self.ph.kv("验证码尝试", attempt, "cyan"), self.ph.kv("令牌长度", len(self.login_token), "cyan")))

                    # 返回脱敏登录摘要。
                    return {"登录状态": "成功", "验证码尝试": attempt, "令牌长度": len(self.login_token)}

                # 保存服务端失败消息。
                last_error = str(response_data.get("msg", "登录校验未通过"))

            except Exception as error:
                # 保存当前轮次异常消息。
                last_error = str(error)

            # 打印重试提示。
            self.ph.warning(self.ph.join(self.ph.kv("登录重试", f"第 {attempt}/{self.captcha_retries} 次", "yellow"), self.ph.kv("原因", last_error, "yellow")))

        # 抛出登录失败异常。
        raise RuntimeError(f"{self.company_name} 登录失败：{last_error}")

    # 定义验证码获取方法。
    def _get_captcha(self):
        """
        请求当前登录验证码及其一次性 key。

        :return: 验证码上下文。
        :rtype: dict
        """

        # 构建验证码请求头。
        headers = self._build_headers(self.login_page_url)

        # 请求验证码接口。
        result = self.http.request("GET", self.captcha_api_url, headers=headers)

        # 解析验证码响应。
        response_data = self._decode_json(result.body, "验证码")

        # 读取验证码数据对象。
        payload = response_data.get("data") or {}

        # 校验验证码字段完整。
        if str(response_data.get("code", "")) != "0" or not payload.get("captcha") or not payload.get("key"):
            raise RuntimeError(f"验证码接口失败：{response_data.get('msg', '响应字段不完整')}")

        # 返回验证码上下文。
        return {"captcha": str(payload["captcha"]), "key": str(payload["key"])}

    # 定义通用请求头构建方法。
    def _build_headers(self, referer):
        """
        构建本草堂 JSON 接口请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 请求头字典。
        :rtype: dict
        """

        # 返回浏览器协议请求头。
        return {"Accept": "application/json, text/plain, */*", "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8", "Connection": "keep-alive", "Origin": self.base_url, "Referer": referer, "User-Agent": self.http.user_agent(), "X-Requested-With": "XMLHttpRequest"}

    # 定义带令牌查询请求头构建方法。
    def _build_query_headers(self, referer):
        """
        构建带 Authorization 令牌的查询请求头。

        :param referer: 来源页面地址。
        :type referer: str
        :return: 查询请求头字典。
        :rtype: dict
        """

        # 构建基础查询请求头。
        headers = self._build_headers(referer)

        # 写入平台要求的原样访问令牌。
        headers["Authorization"] = self.login_token

        # 返回查询请求头。
        return headers

    # 定义原始 JSON 响应解析方法。
    def _decode_json(self, body, stage_name):
        """
        将接口响应字节解析为字典。

        :param body: JSON 响应字节。
        :type body: bytes
        :param stage_name: 协议阶段名称。
        :type stage_name: str
        :return: JSON 对象。
        :rtype: dict
        """

        # 解码响应文本。
        text = body.decode("utf-8-sig", errors="replace").strip()

        # 解析 JSON 文本。
        try:
            return json.loads(text)
        except json.JSONDecodeError as error:
            raise RuntimeError(f"{stage_name}失败：响应不是有效 JSON。") from error

    # 定义数据类型规范化方法。
    def _normalize_data_type(self, data_type):
        """
        将界面数据类型别名转换为平台标准数据类型。

        :param data_type: 输入数据类型名称。
        :type data_type: str
        :return: 平台标准数据类型名称。
        :rtype: str
        """

        # 读取别名映射结果。
        normalized = DATA_TYPE_ALIASES.get(str(data_type).strip(), str(data_type).strip())

        # 校验平台数据类型。
        if normalized not in DATA_TYPE_CONFIG:
            raise ValueError(f"不支持的数据类型：{data_type}")

        # 返回平台标准数据类型。
        return normalized

    # 定义日期规范化方法。
    def _normalize_date(self, value, field_name):
        """
        将输入日期规范化为 YYYY-MM-DD。

        :param value: 输入日期文本。
        :type value: str | None
        :param field_name: 日期字段名称。
        :type field_name: str
        :return: 规范化日期文本。
        :rtype: str
        """

        # 使用当天日期填充空值。
        text = str(value or "").strip() or datetime.now().strftime("%Y-%m-%d")

        # 校验日期格式。
        try:
            parsed = datetime.strptime(text, "%Y-%m-%d")
        except ValueError as error:
            raise ValueError(f"{field_name}格式必须为 YYYY-MM-DD：{text}") from error

        # 返回标准日期文本。
        return parsed.strftime("%Y-%m-%d")

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):
        """
        保存当前协议会话 Cookie。

        :param output_dir: Cookie 保存目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建 Cookie 保存目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存 Cookie 内容。
        self.http.save_cookie_file(path)

        # 返回 Cookie 文件路径。
        return path

    # 定义原始数据查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):
        """
        查询指定数据类型的全部原始记录。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 原始记录列表。
        :rtype: list[dict]
        """

        # 查询并返回原始记录。
        return self._query_platform_data(data_type, start_date, end_date)["rows"]

    # 定义协议查询并落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):
        """
        查询指定数据类型并生成标准字段加全部原始字段 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: Excel 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询当前数据类型的全部分页数据。
        result = self._query_platform_data(data_type, start_date, end_date)

        # 创建 Excel 输出目录。
        target_dir = output_dir or os.getcwd()
        os.makedirs(target_dir, exist_ok=True)

        # 构建标准输出文件名。
        file_name = self.build_output_file_name(result["data_type"], result["start_date"], result["end_date"])

        # 拼接 Excel 输出路径。
        output_path = os.path.join(target_dir, file_name)

        # 写入标准字段和全部原始字段。
        self._write_excel(result["rows"], result["raw_headers"], result["data_type"], output_path)

        # 打印 Excel 生成摘要。
        self.ph.success(self.ph.join(self.ph.kv(f"{result['data_type']} Excel已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", len(result["rows"]), "cyan"), self.ph.kv("原始字段数", len(result["raw_headers"]), "cyan")))

        # 返回 Excel 输出路径。
        return output_path

    # 定义导出接口兼容方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):
        """
        使用协议查询结果生成本地 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: Excel 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 复用协议查询落地流程。
        return self.query_flow_to_excel(data_type, start_date, end_date, output_dir)

    # 定义输出文件日期解析方法。
    def resolve_output_query_date(self, data_type, start_date=None, end_date=None):
        """
        为库存使用导出当天日期，为流向使用查询开始日期。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 文件命名日期。
        :rtype: str
        """

        # 库存文件名使用导出当天日期。
        if self._normalize_data_type(data_type) == "库存":
            return datetime.now().strftime("%Y-%m-%d")

        # 流向文件名使用查询开始日期。
        return start_date or end_date or datetime.now().strftime("%Y-%m-%d")

    # 定义分页查询方法。
    def _query_platform_data(self, data_type, start_date, end_date):
        """
        登录并分页读取平台 JSON 数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询记录、原始字段和日期信息。
        :rtype: dict
        """

        # 规范化数据类型和日期范围。
        normalized_type = self._normalize_data_type(data_type)
        normalized_start = self._normalize_date(start_date, "开始日期")
        normalized_end = self._normalize_date(end_date, "结束日期")

        # 校验日期范围顺序。
        if normalized_start > normalized_end:
            raise ValueError("查询开始日期不能晚于结束日期。")

        # 确保查询前已经登录。
        if not self.logged_in:
            self.login()

        # 打印数据类型查询分隔线。
        self.ph.section(f"{normalized_type}查询", self.ph.join(self.ph.kv("开始日期", normalized_start, "cyan"), self.ph.kv("结束日期", normalized_end, "cyan")))

        # 读取当前数据类型接口配置。
        type_config = DATA_TYPE_CONFIG[normalized_type]

        # 初始化累计记录。
        rows = []

        # 初始化原始字段顺序。
        raw_headers = []

        # 初始化页码。
        page = 1

        # 初始化总记录数。
        total = None

        # 循环读取分页数据。
        while True:
            # 构建平台查询参数。
            params = {"page": page, "limit": self.page_size}

            # 仅向流向接口传递日期参数。
            if normalized_type != "库存":
                params.update({"startDate": normalized_start, "endDate": normalized_end})

            # 拼接查询接口地址。
            request_url = f"{self.base_url}{type_config['endpoint']}?{urlencode(params)}"

            # 提交分页查询请求。
            response = self.http.request("GET", request_url, headers=self._build_query_headers(self.base_url), timeout=90)

            # 解析分页响应。
            response_data = self._decode_json(response.body, f"{normalized_type}查询")

            # 校验业务响应状态。
            if str(response_data.get("code", "")) != "0":
                raise RuntimeError(f"{normalized_type}查询失败：{response_data.get('msg', '未知错误')}")

            # 读取分页数据对象。
            payload = response_data.get("data") or {}

            # 校验分页数据结构。
            if not isinstance(payload, dict):
                raise RuntimeError(f"{normalized_type}查询失败：data 字段不是对象。")

            # 读取当前页记录。
            page_rows = payload.get("records") or []

            # 校验记录数组结构。
            if not isinstance(page_rows, list):
                raise RuntimeError(f"{normalized_type}查询失败：records 字段不是数组。")

            # 读取服务端总记录数。
            if total is None and payload.get("total") is not None:
                total = int(payload.get("total") or 0)

            # 累积当前页原始字段顺序。
            for row in page_rows:
                if isinstance(row, dict):
                    for field_name in row:
                        if field_name not in raw_headers:
                            raw_headers.append(field_name)

            # 累积当前页有效记录。
            rows.extend(row for row in page_rows if isinstance(row, dict))

            # 打印分页进度。
            self.ph.info(self.ph.join(self.ph.kv("页码", page, "cyan"), self.ph.kv("本页记录", len(page_rows), "cyan"), self.ph.kv("累计记录", len(rows), "cyan"), self.ph.kv("服务端总数", total if total is not None else "未返回", "cyan")))

            # 判断已读取服务端总数。
            if total is not None and len(rows) >= total:
                break

            # 判断当前页不足一页。
            if len(page_rows) < self.page_size:
                break

            # 进入下一页。
            page += 1

        # 校验流向业务日期范围。
        self._validate_dates(normalized_type, rows, normalized_start, normalized_end)

        # 返回完整查询结果。
        return {"data_type": normalized_type, "rows": rows, "raw_headers": raw_headers, "start_date": normalized_start, "end_date": normalized_end}

    # 定义业务日期范围校验方法。
    def _validate_dates(self, data_type, rows, start_date, end_date):
        """
        校验流向记录的业务日期处于请求范围内。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 原始记录列表。
        :type rows: list[dict]
        :param start_date: 请求开始日期。
        :type start_date: str
        :param end_date: 请求结束日期。
        :type end_date: str
        :return: None。
        :rtype: None
        """

        # 库存是当前快照，不执行流向日期校验。
        if data_type == "库存":
            self.ph.info(self.ph.kv("库存日期校验", "当前快照接口无日期筛选", "yellow"))
            return

        # 提取业务日期的日期部分。
        date_values = [str(row.get("cofirmdate", ""))[:10] for row in rows if str(row.get("cofirmdate", "")).strip()]

        # 处理空结果日期校验。
        if not date_values:
            self.ph.info(self.ph.kv("业务日期校验", "无返回记录", "yellow"))
            return

        # 读取返回数据日期边界。
        minimum_date = min(date_values)
        maximum_date = max(date_values)

        # 校验返回日期没有超出请求范围。
        if minimum_date < start_date or maximum_date > end_date:
            raise RuntimeError(f"{data_type}查询日期越界：实际 {minimum_date} 至 {maximum_date}，请求 {start_date} 至 {end_date}。")

        # 打印日期校验结果。
        self.ph.info(self.ph.join(self.ph.kv("业务日期最小值", minimum_date, "cyan"), self.ph.kv("业务日期最大值", maximum_date, "cyan")))

    # 定义 Excel 输出方法。
    def _write_excel(self, rows, raw_headers, data_type, output_path):
        """
        写入标准字段和接口返回的全部原始字段。

        :param rows: 原始记录列表。
        :type rows: list[dict]
        :param raw_headers: 原始字段顺序。
        :type raw_headers: list[str]
        :param data_type: 数据类型名称。
        :type data_type: str
        :param output_path: Excel 输出路径。
        :type output_path: str
        :return: None。
        :rtype: None
        """

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 获取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = data_type[:31]

        # 构建标准字段表头。
        standard_headers = [title for title, _ in STANDARD_FIELD_MAP[data_type]]

        # 构建全部原始字段表头。
        raw_titles = [f"原始_{field_name}" for field_name in raw_headers]

        # 写入完整表头。
        worksheet.append(standard_headers + raw_titles)

        # 写入每一条业务记录。
        for row in rows:
            # 写入标准字段值与原始字段值。
            values = [self._excel_value(row.get(field_name, "")) for _, field_name in STANDARD_FIELD_MAP[data_type]]
            values.extend(self._excel_value(row.get(field_name, "")) for field_name in raw_headers)
            worksheet.append(values)

        # 调整全部列宽。
        for column_index in range(1, worksheet.max_column + 1):
            # 计算当前列最大文本长度。
            width = max(len(str(worksheet.cell(row_index, column_index).value or "")) for row_index in range(1, worksheet.max_row + 1))
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(width + 2, 12), 36)

        # 保存工作簿。
        workbook.save(output_path)

        # 关闭工作簿。
        workbook.close()

        # 应用统一 Excel 样式。
        apply_standard_excel_style(output_path)

    # 定义 Excel 值转换方法。
    def _excel_value(self, value):
        """
        将嵌套字段转换为 Excel 可写值。

        :param value: 原始字段值。
        :type value: object
        :return: Excel 可写值。
        :rtype: object
        """

        # 将列表或字典保留为 JSON 文本。
        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

        # 返回普通字段值。
        return value
