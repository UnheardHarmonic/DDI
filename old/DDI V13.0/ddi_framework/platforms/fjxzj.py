# -*- coding: utf-8 -*-
"""
福建新紫金医药供应商平台协议适配器。

本模块负责福建新紫金医药供应商平台的验证码获取、登录加密、协议登录、令牌保存和 Cookie 保存。
"""

# 导入 JSON 编解码模块。
import json

# 导入操作系统路径模块。
import os

# 导入随机数模块。
import random

# 导入时间模块。
import time

# 导入 UUID 生成模块。
import uuid

# 导入日期时间解析对象。
from datetime import datetime

# 导入 URL 解析函数。
from urllib.parse import urlencode, urlparse

# 导入验证码处理器。
from ddi_framework import CaptchaSolver

# 导入协议请求客户端。
from ddi_framework import ProtocolHttpClient

# 导入平台基类。
from ddi_framework import BasePlatform

# 导入日志打印器。
from ddi_framework import PrintHandler

# 导入通用逆向摘要与加密函数。
from ddi_framework import md5_hex, sha256_hex, triple_des_cbc_pkcs7_base64


# 定义福建新紫金接口字段标题映射。
FIELD_TITLE_MAPPING = {
    "goodsno": "商品编码",
    "goodsCode": "商品编码",
    "goodsname": "商品名称",
    "goodstype": "规格",
    "goodsunit": "单位",
    "supName": "供应商名称",
    "supCode": "供应商账号",
    "factoryname": "生产企业",
    "djRq": "单据日期",
    "djLx": "单据类型",
    "proQty": "进出量",
    "deliveryPrice": "配送价格",
    "mdBrId": "门店编码",
    "upFlow": "上游",
    "buyer": "采购商",
    "gsName": "采购商",
    "gsId": "采购商",
    "sppqSupId": "供应商账号",
    "areaName": "区域",
    "ifShow": "显示标记",
    "index": "序号",
    "mdBrName": "门店简称",
    "mdBrQc": "门店全称",
    "proPh": "批号",
}


# 定义福建新紫金医药供应商平台协议适配器。
class FjxzjPlatform(BasePlatform):

    """
    福建新紫金医药供应商平台协议适配器。

    :param company_name: 商业公司名称。
    :type company_name: str
    :param account: 商业账号配置。
    :type account: dict
    :param printer: 日志打印器。
    :type printer: PrintHandler | None
    :return: 福建新紫金医药平台适配器实例。
    :rtype: FjxzjPlatform
    """

    # 定义福建新紫金医药适配器初始化方法。
    def __init__(self, company_name, account, printer=None):

        """
        初始化福建新紫金医药平台适配器。

        :param company_name: 商业公司名称。
        :type company_name: str
        :param account: 商业账号配置。
        :type account: dict
        :param printer: 日志打印器。
        :type printer: PrintHandler | None
        :return: None。
        :rtype: None
        """

        # 调用父类初始化逻辑。
        super().__init__(company_name, account, printer=printer)

        # 初始化日志打印器。
        self.ph = printer or PrintHandler()

        # 保存登录页地址。
        self.login_page_url = account.get("登录页", "https://provider.kbjcn.com/user/login")

        # 保存验证码接口地址。
        self.captcha_api_url = account.get("验证码接口", "https://provider.kbjcn.com/provider-admin/api/account/captcha")

        # 保存登录接口地址。
        self.login_api_url = account.get("登录接口", "https://provider.kbjcn.com/provider-admin/api/account/login")

        # 保存查询接口地址。
        self.query_api_url = account.get("查询接口", "https://provider.kbjcn.com/provider-admin/api/goods_detail/list")

        # 保存平台页面根地址。
        self.base_url = self.build_origin(self.login_page_url)

        # 保存业务页面地址。
        self.goods_detail_page_url = f"{self.base_url}/goods/goodsDetail"

        # 初始化协议请求客户端。
        self.http = ProtocolHttpClient(self.base_url)

        # 初始化验证码处理器。
        self.captcha_solver = CaptchaSolver()

        # 初始化登录响应。
        self.login_result = {}

        # 初始化登录令牌。
        self.login_token = ""

        # 保存默认分页大小。
        self.page_size = int(account.get("分页大小", 100))

    # 定义福建新紫金医药登录方法。
    def login(self):

        """
        执行福建新紫金医药协议登录。

        :return: 登录结果。
        :rtype: dict
        """

        # 打印登录阶段标题。
        self.ph.section("福建新紫金医药登录", self.ph.kv("商业", self.company_name, "cyan"))

        # 打印登录页访问日志。
        self.ph.info(self.ph.kv("访问登录页", self.login_page_url, "cyan"))

        # 访问登录页以初始化站点 Cookie。
        self.http.get(self.login_page_url)

        # 打印登录页访问完成日志。
        self.ph.success(self.ph.kv("登录页访问完成，当前 Cookie 数", len(list(self.http.cookie_jar)), "green"))

        # 执行带验证码重试的登录流程。
        return self.login_with_captcha_retry()

    # 定义带验证码重试的登录方法。
    def login_with_captcha_retry(self, max_retries=8):

        """
        使用验证码重试机制执行登录。

        :param max_retries: 最大验证码刷新次数。
        :type max_retries: int
        :return: 登录结果。
        :rtype: dict
        """

        # 初始化最后一次失败原因。
        last_error = ""

        # 循环获取验证码并提交登录。
        for retry_index in range(max_retries):

            # 打印验证码尝试日志。
            self.ph.info(self.ph.kv("图片验证码尝试", f"第 {retry_index + 1}/{max_retries} 次", "cyan"))

            # 捕获当前验证码登录异常。
            try:

                # 获取验证码上下文。
                captcha_context = self.get_captcha_context()

                # 识别验证码答案。
                captcha_code = self.solve_captcha_code(captcha_context["image_bytes"])

                # 打印验证码识别完成日志。
                self.ph.info(self.ph.kv("图片验证码识别完成", f"长度={len(captcha_code)}", "cyan"))

                # 构建登录请求体。
                payload = self.build_login_payload(captcha_context["key"], captcha_code)

                # 打印正式登录请求日志。
                self.ph.info("提交正式登录请求: /provider-admin/api/account/login")

                # 提交登录请求。
                result = self.post_login(payload)

                # 解析登录响应。
                login_result = self.decode_login_result(result)

                # 判断登录响应是否成功。
                if self.is_login_success(login_result):

                    # 保存登录响应。
                    self.login_result = login_result

                    # 保存登录令牌。
                    self.login_token = self.extract_login_token(login_result)

                    # 打印登录成功日志。
                    self.ph.success(self.ph.kv("正式登录成功", self.extract_login_message(login_result), "green"))

                    # 返回登录响应。
                    return login_result

                # 记录登录失败原因。
                last_error = self.extract_login_message(login_result)

                # 打印验证码或登录失败日志。
                self.ph.warning("登录校验未通过，准备刷新验证码")

            # 处理验证码识别或登录提交异常。
            except Exception as error:

                # 记录当前异常原因。
                last_error = str(error)

                # 打印验证码处理失败日志。
                self.ph.warning("图片验证码识别或登录校验未通过，准备刷新验证码")

            # 最后一轮失败后跳过等待。
            if retry_index >= max_retries - 1:

                # 进入失败处理。
                continue

            # 计算失败后的等待秒数。
            delay_seconds = random.uniform(2.2, 4.4)

            # 打印等待日志。
            self.ph.wait(self.ph.kv("验证码失败后等待，预计", f"{delay_seconds:.1f} 秒", "cyan"))

            # 等待后再刷新验证码。
            time.sleep(delay_seconds)

        # 抛出登录失败异常。
        raise RuntimeError(f"{self.company_name} 登录失败：{last_error}")

    # 定义验证码上下文获取方法。
    def get_captcha_context(self):

        """
        获取登录验证码图片与验证码会话编号。

        :return: 验证码上下文字典。
        :rtype: dict
        """

        # 生成验证码会话编号。
        captcha_key = str(uuid.uuid4())

        # 构建验证码请求参数。
        params = {"captchaKey": captcha_key, "time": int(time.time() * 1000)}

        # 拼接验证码请求地址。
        request_url = f"{self.captcha_api_url}?{urlencode(params)}"

        # 打印验证码请求日志。
        self.ph.info("请求图片验证码: /provider-admin/api/account/captcha")

        # 构建验证码请求头。
        headers = self.build_captcha_headers()

        # 请求验证码图片。
        result = self.http.request("GET", request_url, headers=headers)

        # 返回验证码上下文。
        return {"key": captcha_key, "image_bytes": result.body}

    # 定义验证码请求头构建方法。
    def build_captcha_headers(self):

        """
        构建验证码请求头。

        :return: 验证码请求头。
        :rtype: dict
        """

        # 返回验证码接口请求头。
        return {
            "Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            "Connection": "keep-alive",
            "Referer": self.login_page_url,
            "User-Agent": self.http.user_agent(),
        }

    # 定义验证码答案识别方法。
    def solve_captcha_code(self, image_bytes):

        """
        识别验证码答案。

        :param image_bytes: 验证码图片字节流。
        :type image_bytes: bytes
        :return: 验证码答案。
        :rtype: str
        """

        # 识别验证码原始文本。
        raw_text = self.captcha_solver.solve_raw_text_image(image_bytes)

        # 清洗验证码文本。
        cleaned_text = str(raw_text or "").strip().lower()

        # 校验验证码文本是否为空。
        if not cleaned_text:

            # 抛出验证码识别异常。
            raise RuntimeError("图片验证码识别为空")

        # 返回清洗后的验证码文本。
        return cleaned_text

    # 定义登录请求体构建方法。
    def build_login_payload(self, captcha_key, captcha_code):

        """
        构建福建新紫金医药登录请求体。

        :param captcha_key: 验证码会话编号。
        :type captcha_key: str
        :param captcha_code: 验证码答案。
        :type captcha_code: str
        :return: 登录请求体。
        :rtype: dict
        """

        # 计算验证码 MD5 摘要。
        captcha_md5 = md5_hex(str(captcha_code).lower())

        # 计算验证码提交摘要。
        captcha_sha256 = sha256_hex(captcha_md5)

        # 计算登录密码 MD5 摘要。
        password_md5 = md5_hex(self.clean_login_password())

        # 使用验证码 MD5 摘要加密密码摘要。
        encrypted_password = triple_des_cbc_pkcs7_base64(password_md5, captcha_md5, "01234567")

        # 返回登录请求体。
        return {
            "account": self.clean_login_name(),
            "captcha": captcha_sha256,
            "captchaKey": captcha_key,
            "password": encrypted_password,
        }

    # 定义登录请求提交方法。
    def post_login(self, payload):

        """
        提交福建新紫金医药登录请求。

        :param payload: 登录请求体。
        :type payload: dict
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 构建登录请求头。
        headers = self.build_login_headers()

        # 序列化登录请求体。
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        # 发送登录请求。
        return self.http.request("POST", self.login_api_url, headers=headers, data=body)

    # 定义登录请求头构建方法。
    def build_login_headers(self):

        """
        构建登录请求头。

        :return: 登录请求头。
        :rtype: dict
        """

        # 返回登录接口请求头。
        return {
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            "Connection": "keep-alive",
            "Content-Type": "application/json;charset=UTF-8",
            "Origin": self.base_url,
            "Referer": self.login_page_url,
            "User-Agent": self.http.user_agent(),
        }

    # 定义登录响应解析方法。
    def decode_login_result(self, result):

        """
        解析登录响应。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 登录响应数据。
        :rtype: dict
        """

        # 解码登录响应文本。
        response_text = result.body.decode("utf-8", errors="replace").strip()

        # 处理空登录响应。
        if not response_text:

            # 返回空响应登录结果。
            return {"success": False, "message": "登录接口返回空响应"}

        # 尝试解析 JSON 响应。
        try:

            # 返回 JSON 响应数据。
            return json.loads(response_text)

        # 处理非 JSON 响应。
        except json.JSONDecodeError:

            # 返回原始响应文本。
            return {"success": False, "message": response_text}

    # 定义登录成功判断方法。
    def is_login_success(self, login_result):

        """
        判断登录响应是否成功。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 是否登录成功。
        :rtype: bool
        """

        # 判断登录令牌是否存在。
        if self.extract_login_token(login_result):

            # 返回登录成功状态。
            return True

        # 判断成功状态码是否存在。
        if str(login_result.get("code", "")) in ["0", "200"] and str(login_result.get("success", "")).lower() != "false":

            # 返回登录成功状态。
            return True

        # 返回登录失败状态。
        return False

    # 定义登录响应消息提取方法。
    def extract_login_message(self, login_result):

        """
        提取登录响应消息。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 登录响应消息。
        :rtype: str
        """

        # 返回响应消息字段。
        if login_result.get("msg") or login_result.get("message"):

            # 返回响应消息文本。
            return str(login_result.get("msg") or login_result.get("message"))

        # 返回令牌登录成功消息。
        if self.extract_login_token(login_result):

            # 返回令牌成功摘要。
            return "token"

        # 返回原始响应摘要。
        return str(login_result)

    # 定义登录令牌提取方法。
    def extract_login_token(self, login_result):

        """
        提取登录响应中的访问令牌。

        :param login_result: 登录响应数据。
        :type login_result: dict
        :return: 访问令牌。
        :rtype: str
        """

        # 遍历顶层常见令牌字段。
        for key in ["access_token", "accessToken", "token", "jwt"]:

            # 读取当前令牌字段值。
            token_value = login_result.get(key)

            # 返回非空令牌。
            if token_value:

                # 返回令牌文本。
                return str(token_value)

        # 读取响应数据字段。
        data = login_result.get("data")

        # 处理嵌套字典令牌字段。
        if isinstance(data, dict):

            # 遍历嵌套常见令牌字段。
            for key in ["access_token", "accessToken", "token", "jwt"]:

                # 读取当前令牌字段值。
                token_value = data.get(key)

                # 返回非空令牌。
                if token_value:

                    # 返回令牌文本。
                    return str(token_value)

        # 读取响应结果字段。
        result_value = login_result.get("result")

        # 返回字符串形式的登录凭证。
        if isinstance(result_value, str) and result_value.strip():

            # 返回登录凭证文本。
            return result_value.strip()

        # 处理嵌套字典形式的响应结果字段。
        if isinstance(result_value, dict):

            # 遍历嵌套常见令牌字段。
            for key in ["access_token", "accessToken", "token", "jwt"]:

                # 读取当前令牌字段值。
                token_value = result_value.get(key)

                # 返回非空令牌。
                if token_value:

                    # 返回令牌文本。
                    return str(token_value)

        # 返回空令牌。
        return ""

    # 定义登录账号清洗方法。
    def clean_login_name(self):

        """
        读取清洗后的登录账号。

        :return: 登录账号。
        :rtype: str
        """

        # 返回去除两端空白后的账号。
        return str(self.account["账号"]).strip()

    # 定义登录密码清洗方法。
    def clean_login_password(self):

        """
        读取清洗后的登录密码。

        :return: 登录密码。
        :rtype: str
        """

        # 返回去除两端空白后的密码。
        return str(self.account["密码"]).strip()

    # 定义站点根地址构建方法。
    def build_origin(self, url):

        """
        根据页面地址构建站点根地址。

        :param url: 页面地址。
        :type url: str
        :return: 站点根地址。
        :rtype: str
        """

        # 解析页面地址。
        parsed_url = urlparse(url)

        # 返回协议与域名组成的站点根地址。
        return f"{parsed_url.scheme}://{parsed_url.netloc}"

    # 定义 Cookie 保存方法。
    def save_cookie(self, output_dir):

        """
        保存当前登录 Cookie 与认证令牌。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: Cookie 文件路径。
        :rtype: str
        """

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Cookie 文件路径。
        cookie_path = os.path.join(output_dir, f"{self.company_name}_cookie.json")

        # 保存 Cookie 文件。
        self.http.save_cookie_file(cookie_path)

        # 保存认证令牌文件。
        self.save_token_file(output_dir)

        # 打印 Cookie 保存日志。
        self.ph.success(self.ph.kv("登录 Cookie 已保存", self.ph.path(cookie_path), "green"))

        # 返回 Cookie 文件路径。
        return cookie_path

    # 定义令牌保存方法。
    def save_token_file(self, output_dir):

        """
        保存登录令牌到本地文件。

        :param output_dir: 输出目录。
        :type output_dir: str
        :return: 令牌文件路径。
        :rtype: str
        """

        # 构建令牌文件路径。
        token_path = os.path.join(output_dir, f"{self.company_name}_token.json")

        # 构建令牌文件内容。
        token_data = {"token": self.login_token, "login_result": self.login_result}

        # 写入令牌文件。
        with open(token_path, "w", encoding="utf-8") as file_obj:

            # 序列化令牌文件内容。
            json.dump(token_data, file_obj, ensure_ascii=False, indent=2)

        # 打印令牌保存日志。
        self.ph.success(self.ph.kv("登录令牌已保存", self.ph.path(token_path), "green"))

        # 返回令牌文件路径。
        return token_path

    # 定义流向查询方法。
    def query_flow(self, data_type, start_date=None, end_date=None):

        """
        查询福建新紫金医药流向数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果。
        :rtype: list[dict]
        """

        # 查询指定数据类型完整数据。
        query_result = self.query_fjxzj_data(data_type, start_date, end_date)

        # 返回查询结果数据行。
        return query_result["rows"]

    # 定义流向导出方法。
    def export_flow(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        导出福建新紫金医药流向数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: 导出文件路径。
        :rtype: str
        """

        # 使用分页查询结果生成本地 Excel 文件。
        return self.query_flow_to_excel(data_type, start_date=start_date, end_date=end_date, output_dir=output_dir)

    # 定义查询落地 Excel 方法。
    def query_flow_to_excel(self, data_type, start_date=None, end_date=None, output_dir=None):

        """
        查询福建新紫金医药流向数据并生成 Excel。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :param output_dir: 输出目录。
        :type output_dir: str | None
        :return: Excel 文件路径。
        :rtype: str
        """

        # 查询指定业务数据。
        query_result = self.query_fjxzj_data(data_type, start_date, end_date)

        # 使用调用方指定目录或默认下载目录。
        output_dir = output_dir or os.path.join(os.getcwd(), "downloads")

        # 创建输出目录。
        os.makedirs(output_dir, exist_ok=True)

        # 构建 Excel 文件名。
        file_name = self.build_output_file_name(data_type, start_date, end_date)

        # 构建 Excel 输出路径。
        output_path = os.path.join(output_dir, file_name)

        # 写入 Excel 文件。
        self.write_rows_to_excel(query_result["rows"], query_result["headers"], output_path, data_type)

        # 打印 Excel 保存日志。
        self.ph.success(self.ph.join(self.ph.kv(f"{data_type} Excel 已生成", self.ph.path(output_path), "green"), self.ph.kv("记录数", f"{len(query_result['rows'])} 条", "cyan")))

        # 返回 Excel 文件路径。
        return output_path

    # 定义福建新紫金医药数据查询方法。
    def query_fjxzj_data(self, data_type, start_date, end_date):

        """
        查询福建新紫金医药指定类型完整数据。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 查询结果与表头。
        :rtype: dict
        """

        # 校验数据类型是否已接入。
        self.validate_data_type(data_type)

        # 校验查询开始日期。
        if not start_date:

            # 抛出查询开始日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 start_date")

        # 校验查询结束日期。
        if not end_date:

            # 抛出查询结束日期缺失异常。
            raise ValueError(f"查询{data_type}必须提供 end_date")

        # 打印查询阶段标题。
        self.ph.section(f"{data_type}查询", self.ph.join(self.ph.kv("开始日期", start_date, "cyan"), self.ph.kv("结束日期", end_date, "cyan")))

        # 初始化累计数据行。
        all_rows = []

        # 遍历数据类型对应的平台单据类型。
        for platform_type in self.platform_types_for_data_type(data_type):

            # 查询当前平台单据类型全部分页。
            type_rows = self.query_platform_type_rows(data_type, platform_type, start_date, end_date)

            # 合并当前平台单据类型数据行。
            all_rows.extend(type_rows)

        # 构建当前数据类型 Excel 表头。
        headers = self.build_output_headers(data_type, all_rows)

        # 打印查询完成日志。
        self.ph.success(self.ph.join(self.ph.kv("查询解析完成", f"{len(all_rows)} 条", "green"), self.ph.kv("表头数", len(headers), "cyan")))

        # 返回标准查询结果。
        return {"rows": all_rows, "headers": headers}

    # 定义平台单据类型分页查询方法。
    def query_platform_type_rows(self, data_type, platform_type, start_date, end_date):

        """
        查询指定平台单据类型的全部分页数据。

        :param data_type: DDI 数据类型名称。
        :type data_type: str
        :param platform_type: 平台单据类型名称。
        :type platform_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :return: 平台单据类型数据行。
        :rtype: list[dict]
        """

        # 初始化当前页码。
        current_page = 1

        # 初始化当前单据类型累计数据行。
        all_rows = []

        # 初始化当前单据类型总记录数。
        total_count = None

        # 按页循环请求业务数据。
        while True:

            # 构建当前页查询参数。
            params = self.build_query_params(platform_type, start_date, end_date, current_page)

            # 打印分页请求日志。
            self.ph.info(self.ph.join(self.ph.kv(f"请求{data_type}分页", f"djLx={platform_type}", "cyan"), self.ph.kv("pageNo", current_page, "yellow"), self.ph.kv("pageSize", params["pageSize"], "yellow")))

            # 发送当前页查询请求。
            result = self.get_query_page(params)

            # 解析当前页查询响应。
            page_result = self.decode_query_page(result)

            # 读取当前页数据行。
            page_rows = page_result["rows"]

            # 保存总记录数。
            total_count = page_result["total"]

            # 追加当前页数据行。
            all_rows.extend(page_rows)

            # 打印当前页解析日志。
            self.ph.success(self.ph.join(self.ph.kv("当前页解析完成", f"{len(page_rows)} 条", "green"), self.ph.kv("累计", f"{len(all_rows)} 条", "cyan"), self.ph.kv("总数", total_count, "yellow")))

            # 判断是否已经完成全部分页。
            if len(all_rows) >= int(total_count or 0) or not page_rows:

                # 结束分页循环。
                break

            # 递增页码。
            current_page += 1

        # 返回当前平台单据类型数据行。
        return all_rows

    # 定义查询参数构建方法。
    def build_query_params(self, platform_type, start_date, end_date, current_page):

        """
        构建流向明细分页查询参数。

        :param platform_type: 平台单据类型名称。
        :type platform_type: str
        :param start_date: 查询开始日期。
        :type start_date: str
        :param end_date: 查询结束日期。
        :type end_date: str
        :param current_page: 当前页码。
        :type current_page: int
        :return: 分页查询参数。
        :rtype: dict
        """

        # 返回平台查询参数。
        return {
            "goodsno": "",
            "startDate": f"{start_date} 00:00:00",
            "endDate": f"{end_date} 23:59:59",
            "gsId": str(self.account.get("采购商编码", "1")),
            "current": current_page,
            "pageSize": self.page_size,
            "djLx": platform_type,
            "time": [start_date, end_date],
            "searchType": "submit",
            "pageNo": current_page,
        }

    # 定义查询请求提交方法。
    def get_query_page(self, params):

        """
        请求流向明细分页接口。

        :param params: 分页查询参数。
        :type params: dict
        :return: HTTP 响应结果。
        :rtype: HttpResult
        """

        # 构建授权查询请求头。
        headers = self.build_authorized_headers()

        # 拼接查询参数后的请求地址。
        request_url = f"{self.query_api_url}?{urlencode(params, doseq=True)}"

        # 发送授权查询请求。
        return self.http.request("GET", request_url, headers=headers, timeout=120)

    # 定义查询响应解析方法。
    def decode_query_page(self, result):

        """
        解析分页查询响应。

        :param result: HTTP 响应结果。
        :type result: HttpResult
        :return: 分页查询结果。
        :rtype: dict
        """

        # 解析 JSON 响应。
        response_data = self.http.decode_json(result)

        # 校验响应状态。
        if str(response_data.get("code", "")) not in ["0", "200"]:

            # 抛出查询失败异常。
            raise RuntimeError(response_data.get("msg") or response_data)

        # 读取响应结果字段。
        result_data = response_data.get("result") or {}

        # 读取分页数据行。
        rows = result_data.get("records") or []

        # 读取分页总数。
        total = result_data.get("total", len(rows))

        # 返回分页查询结果。
        return {"rows": rows, "total": int(total or 0)}

    # 定义授权请求头构建方法。
    def build_authorized_headers(self):

        """
        构建授权业务请求头。

        :return: 授权业务请求头。
        :rtype: dict
        """

        # 返回授权业务请求头。
        return {
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
            "Connection": "keep-alive",
            "Referer": self.goods_detail_page_url,
            "Token": self.login_token,
            "User-Agent": self.http.user_agent(),
            "X-Requested-With": "XMLHttpRequest",
        }

    # 定义数据类型到平台单据类型映射方法。
    def platform_types_for_data_type(self, data_type):

        """
        读取 DDI 数据类型对应的平台单据类型列表。

        :param data_type: DDI 数据类型名称。
        :type data_type: str
        :return: 平台单据类型列表。
        :rtype: list[str]
        """

        # 定义 DDI 数据类型与平台单据类型映射关系。
        type_map = {
            "购进流向": ["采购入库", "采购退厂"],
            "配送流向": ["销售", "退货"],
            "报损流向": ["报损", "报溢"],
        }

        # 返回当前数据类型对应的平台单据类型列表。
        return type_map[data_type]

    # 定义数据类型校验方法。
    def validate_data_type(self, data_type):

        """
        校验福建新紫金医药数据类型。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: None。
        :rtype: None
        """

        # 校验数据类型是否存在。
        if data_type not in ["配送流向", "购进流向", "报损流向"]:

            # 抛出数据类型未接入异常。
            raise NotImplementedError(f"当前福建新紫金医药暂未接入: {data_type}")

    # 定义 Excel 表头构建方法。
    def build_excel_headers(self, data_type):

        """
        构建业务 Excel 表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :return: Excel 表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 返回流向明细表头。
        return [("商品编码", "goodsno"), ("商品名称", "goodsname"), ("规格", "goodstype"), ("单位", "goodsunit"), ("供应商名称", "supName"), ("生产企业", "factoryname"), ("单据日期", "djRq"), ("单据类型", "djLx"), ("进出量", "proQty"), ("配送价格", "deliveryPrice"), ("门店编码", "mdBrId"), ("上游", "upFlow"), ("门店简称", "mdBrName"), ("门店全称", "mdBrQc"), ("批号", "proPh")]

    # 定义输出表头构建方法。
    def build_output_headers(self, data_type, rows):

        """
        构建固定业务字段与接口原始字段合并后的输出表头。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :return: 输出表头配置。
        :rtype: list[tuple[str, str]]
        """

        # 复制当前业务固定表头。
        headers = list(self.build_excel_headers(data_type))

        # 读取已存在字段集合。
        existing_fields = {field_name for _, field_name in headers}

        # 遍历查询结果数据行。
        for row in rows:

            # 跳过非字典数据行。
            if not isinstance(row, dict):

                # 进入下一条数据行。
                continue

            # 遍历当前数据行字段。
            for field_name in row.keys():

                # 读取字段名称文本。
                field_key = str(field_name or "").strip()

                # 跳过空字段名称。
                if not field_key:

                    # 进入下一项字段。
                    continue

                # 跳过已经输出的字段。
                if field_key in existing_fields:

                    # 进入下一项字段。
                    continue

                # 添加字段到去重集合。
                existing_fields.add(field_key)

                # 追加原始接口字段。
                headers.append((FIELD_TITLE_MAPPING.get(field_key, field_key), field_key))

        # 返回合并后的输出表头。
        return headers

    # 定义输出文件名构建方法。
    def build_output_file_name(self, data_type, start_date, end_date):

        """
        构建输出文件名。

        :param data_type: 数据类型名称。
        :type data_type: str
        :param start_date: 查询开始日期。
        :type start_date: str | None
        :param end_date: 查询结束日期。
        :type end_date: str | None
        :return: 输出文件名。
        :rtype: str
        """

        # Delegate standard naming to the base platform.
        return super().build_output_file_name(data_type, start_date, end_date)

    # 定义流向月份标签构建方法。
    def build_month_label(self, start_date):

        """
        构建流向文件月份标签。

        :param start_date: 查询开始日期。
        :type start_date: str
        :return: 月份标签。
        :rtype: str
        """

        # 解析开始日期。
        date_value = datetime.strptime(start_date, "%Y-%m-%d")

        # 返回 yy.m月 标签。
        return f"{date_value.year % 100}.{date_value.month}月"

    # 定义通用 Excel 写入方法。
    def write_rows_to_excel(self, rows, headers, output_path, sheet_name):

        """
        写入查询结果到 Excel。

        :param rows: 查询结果数据行。
        :type rows: list[dict]
        :param headers: 表头配置。
        :type headers: list[tuple[str, str]]
        :param output_path: 输出文件路径。
        :type output_path: str
        :param sheet_name: 工作表名称。
        :type sheet_name: str
        :return: None。
        :rtype: None
        """

        # 导入 Excel 工作簿对象。
        from openpyxl import Workbook

        # 导入列宽工具函数。
        from openpyxl.utils import get_column_letter

        # 创建 Excel 工作簿。
        workbook = Workbook()

        # 读取默认工作表。
        worksheet = workbook.active

        # 设置工作表名称。
        worksheet.title = sheet_name[:31]

        # 写入表头行。
        worksheet.append([header_name for header_name, _ in headers])

        # 遍历查询结果并写入数据行。
        for row in rows:

            # 按表头顺序组装一行数据。
            worksheet.append([row.get(field_name, "") if isinstance(row, dict) else "" for _, field_name in headers])

        # 遍历所有导出列设置列宽。
        for column_index, header_item in enumerate(headers, start=1):

            # 读取当前列字母。
            column_letter = get_column_letter(column_index)

            # 读取当前表头名称。
            header_name = header_item[0]

            # 设置当前列宽。
            worksheet.column_dimensions[column_letter].width = max(12, min(34, len(str(header_name)) + 10))

        # 冻结表头行。
        worksheet.freeze_panes = "A2"

        # 启用自动筛选。
        worksheet.auto_filter.ref = worksheet.dimensions

        # 创建输出目录。
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        # 删除已有同名文件。
        self.remove_existing_file(output_path)

        # 保存 Excel 文件。
        try:

            # 将工作簿保存到目标路径。
            workbook.save(output_path)

        # 处理文件被占用导致的保存异常。
        except PermissionError as error:

            # 抛出带路径提示的保存异常。
            raise PermissionError(f"Excel 文件被占用，无法写入：{output_path}") from error

    # 定义本地文件删除方法。
    def remove_existing_file(self, file_path):

        """
        删除已存在的本地文件。

        :param file_path: 文件路径。
        :type file_path: str
        :return: None。
        :rtype: None
        """

        # 判断目标文件是否存在。
        if os.path.exists(file_path):

            # 删除目标文件。
            os.remove(file_path)
