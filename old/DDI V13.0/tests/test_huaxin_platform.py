# -*- coding: utf-8 -*-
"""
四川华信 FineReport 协议适配器测试。

本模块使用固定 FineReport 坐标单元格验证原始字段解析、标准字段写入和流向日期范围约束。

:return: 模块提供四川华信平台协议适配器测试。
:rtype: None
"""

# 导入临时目录模块。
import tempfile

# 导入单元测试模块。
import unittest

# 导入 Excel 工作簿读取函数。
from openpyxl import load_workbook

# 导入四川华信平台适配器。
from ddi_framework.platforms.huaxin import HuaxinPlatform


# 定义四川华信平台适配器测试类。
class HuaxinPlatformTests(unittest.TestCase):

    """
    验证四川华信 FineReport 原始字段解析和 Excel 落地能力。

    :return: None。
    :rtype: None
    """

    # 定义 FineReport 表格解析与 Excel 写入测试方法。
    def test_parser_preserves_raw_headers_and_rows_in_excel(self):

        """
        验证坐标单元格会解析为完整原始字段并写入 Excel。

        :return: None。
        :rtype: None
        """

        # 构建模拟 FineReport 页面 HTML。
        html_text = "<script>FR._p.reportTotalPage=1;</script><td id=\"A5&#45;0&#45;1\">购进日期</td><td id=\"B5&#45;0&#45;1\">货号</td><td id=\"C5&#45;0&#45;1\">名称</td><td id=\"D5&#45;0&#45;1\">数量</td><td id=\"A7&#45;0&#45;1\">2026-06-03</td><td id=\"B7&#45;0&#45;1\">10001</td><td id=\"C7&#45;0&#45;1\">测试药品</td><td id=\"D7&#45;0&#45;1\">5</td>"

        # 构建最小平台账号配置。
        account = {"平台": "scxh"}

        # 创建待测平台适配器。
        platform = HuaxinPlatform("测试商业", account)

        # 解析模拟 FineReport 页面内容。
        headers, rows = platform.parse_report_content(html_text)

        # 断言原始表头完整保留。
        self.assertEqual(headers, ["购进日期", "货号", "名称", "数量"])

        # 断言坐标单元格记录完整解析。
        self.assertEqual(rows, [{"购进日期": "2026-06-03", "货号": "10001", "名称": "测试药品", "数量": "5"}])

        # 创建临时 Excel 输出目录。
        with tempfile.TemporaryDirectory() as output_dir:

            # 构建临时 Excel 文件路径。
            output_path = f"{output_dir}/huaxin.xlsx"

            # 写入测试 Excel 文件。
            platform.write_rows_to_excel(rows, headers, "购进流向", output_path)

            # 读取生成的 Excel 工作簿。
            workbook = load_workbook(output_path, read_only=True)

            # 读取当前工作表。
            worksheet = workbook.active

            # 读取 Excel 表头。
            excel_headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]

            # 读取 Excel 首行记录。
            excel_row = [cell.value for cell in next(worksheet.iter_rows(min_row=2, max_row=2))]

            # 关闭 Excel 工作簿。
            workbook.close()

        # 断言 Excel 包含原始字段列。
        self.assertIn("原始_货号", excel_headers)

        # 断言标准货品 ID 列保留页面货号。
        self.assertEqual(excel_row[excel_headers.index("货品id")], "10001")

        # 断言原始货号列同步保留页面值。
        self.assertEqual(excel_row[excel_headers.index("原始_货号")], "10001")

    # 定义流向跨月查询校验测试方法。
    def test_flow_query_range_rejects_cross_month(self):

        """
        验证流向报表不能跨越 FineReport 单月参数范围。

        :return: None。
        :rtype: None
        """

        # 创建待测平台适配器。
        platform = HuaxinPlatform("测试商业", {"平台": "scxh"})

        # 断言跨月流向请求会抛出明确异常。
        with self.assertRaisesRegex(ValueError, "同一自然月"):

            # 校验跨月购进流向日期范围。
            platform.validate_query_range("购进流向", "2026-06-30", "2026-07-01")
